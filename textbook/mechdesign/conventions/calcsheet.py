"""设计计算书（第 11 轮第 4 步）：已知 → 求 → 计算 → 结论，每一步写明公式、代入、结果、单位和出处。

书里的算例程序用它记下计算过程，同一份记录既交给正文（bookout），又能导出 Word 计算书——工程任务单的
计算书模板就是这样生成的（学生交的计算书与书中算例同一格式）。

    from calcsheet import CalcSheet
    cs = CalcSheet("SH-301 输出轴强度校核", item="SH-301", rev="B")
    T = cs.given("T", 350, "N·m", "输出轴额定转矩", src="任务单 TS-33-1")
    Ft = cs.step("F_t", "F_t = 2T/d", 2 * T * 1e3 / 210, "N", "齿轮圆周力")
    cs.check("S", 2.1, ">=", 1.5, "疲劳安全系数", src="[S] 取 1.5（濮良贵表 15-x，待核对）")
    cs.docx("SH-301_计算书_B.docx")
"""
from __future__ import annotations

import datetime
import operator
import os

OPS = {">=": operator.ge, ">": operator.gt, "<=": operator.le, "<": operator.lt}
EN = os.environ.get("WQ_LANG", "zh") == "en"


def _fmt(v, digits=4) -> str:
    if isinstance(v, str):
        return v
    if v == 0:
        return "0"
    a = abs(v)
    if a >= 1e5 or a < 1e-3:
        return f"{v:.{digits - 1}e}"
    s = f"{v:.{digits}g}"
    return s


class CalcSheet:
    def __init__(self, title: str, item: str = "", rev: str = "", author: str = "", task: str = ""):
        self.title, self.item, self.rev, self.author, self.task = title, item, rev, author, task
        self.rows: list[dict] = []
        self.finds: list[tuple[str, str]] = []
        self.checks: list[dict] = []
        self.notes: list[str] = []
        self.sources: dict[str, None] = {}

    # -------------------------------------------------------------- recording
    def given(self, sym: str, value, unit: str, what: str, src: str = ""):
        self.rows.append({"kind": "given", "sym": sym, "value": value, "unit": unit, "what": what, "src": src})
        if src:
            self.sources[src] = None
        return value

    def find(self, sym: str, what: str):
        self.finds.append((sym, what))

    def step(self, sym: str, formula: str, value, unit: str, what: str = "", src: str = "", subst: str = ""):
        self.rows.append({"kind": "step", "sym": sym, "formula": formula, "subst": subst, "value": value, "unit": unit,
                          "what": what, "src": src})
        if src:
            self.sources[src] = None
        return value

    def table_value(self, sym: str, value, unit: str, what: str, src: str):
        """查表（或按标准公式）得到的系数：出处必填。"""
        if not src:
            raise ValueError(f"{sym}：查表得到的数值必须写出处")
        return self.step(sym, "查表", value, unit, what, src)

    def check(self, sym: str, value, op: str, limit, what: str, src: str = "") -> bool:
        ok = bool(OPS[op](value, limit))
        self.checks.append({"sym": sym, "value": value, "op": op, "limit": limit, "what": what, "src": src, "ok": ok})
        if src:
            self.sources[src] = None
        return ok

    def note(self, text: str):
        self.notes.append(text)

    @property
    def ok(self) -> bool:
        return all(c["ok"] for c in self.checks)

    # -------------------------------------------------------------- output
    def conclusion(self) -> str:
        if not self.checks:
            return ""
        bad = [c for c in self.checks if not c["ok"]]
        if EN:
            return "All checks pass." if not bad else "Not acceptable: " + "; ".join(c["what"] for c in bad) + "."
        return "全部校核通过。" if not bad else "不满足要求：" + "；".join(c["what"] for c in bad) + "。"

    def markdown(self) -> str:
        """供实验指导书或网页预览用的简明表格。"""
        L = [f"**{self.title}**", "", "| 项 | 公式或出处 | 结果 |", "|---|---|---|"]
        for r in self.rows:
            how = r.get("formula", "已知") if r["kind"] == "step" else "已知"
            src = f"（{r['src']}）" if r.get("src") else ""
            L.append(f"| {r['what'] or r['sym']} | {how}{src} | {r['sym']} = {_fmt(r['value'])} {r['unit']} |")
        for c in self.checks:
            L.append(f"| 校核：{c['what']} | {c['sym']} {c['op']} {_fmt(c['limit'])} | {'✓' if c['ok'] else '✗'} {_fmt(c['value'])} |")
        return "\n".join(L)

    def docx(self, path: str, blank: bool = False) -> str:
        """导出 Word 计算书。blank=True 时只留公式、出处和空白结果栏，作为任务单发给学生的模板。"""
        from docx import Document
        from docx.shared import Pt

        T = (lambda zh, en: en) if EN else (lambda zh, en: zh)
        doc = Document()
        st = doc.styles["Normal"]
        st.font.name = "Noto Sans CJK SC"
        st.font.size = Pt(10.5)
        doc.add_heading(self.title, level=1)
        meta = doc.add_table(rows=2, cols=4)
        meta.style = "Table Grid"
        cells = [T("零件号", "Part no."), self.item, T("版本", "Revision"), self.rev,
                 T("任务单", "Task sheet"), self.task, T("编制 / 日期", "By / date"),
                 f"{self.author or '　　　'} / {datetime.date.today().isoformat() if not blank else '　　　'}"]
        for i, c in enumerate(cells):
            meta.cell(i // 4, i % 4).text = str(c)

        doc.add_heading(T("一、已知", "1. Given"), level=2)
        given = [r for r in self.rows if r["kind"] == "given"]
        t = doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        for i, h in enumerate([T("符号", "Symbol"), T("含义", "Meaning"), T("数值", "Value"), T("来源", "Source")]):
            t.rows[0].cells[i].text = h
        for r in given:
            c = t.add_row().cells
            c[0].text, c[1].text, c[2].text, c[3].text = r["sym"], r["what"], f"{_fmt(r['value'])} {r['unit']}", r["src"]

        if self.finds:
            doc.add_heading(T("二、求", "2. Find"), level=2)
            for sym, what in self.finds:
                doc.add_paragraph(f"{sym}：{what}" if not EN else f"{sym}: {what}", style="List Bullet")

        doc.add_heading(T("三、计算", "3. Calculation"), level=2)
        t = doc.add_table(rows=1, cols=5)
        t.style = "Table Grid"
        for i, h in enumerate([T("项目", "Item"), T("公式", "Formula"), T("代入", "Substitution"), T("结果", "Result"), T("出处", "Source")]):
            t.rows[0].cells[i].text = h
        for r in (x for x in self.rows if x["kind"] == "step"):
            c = t.add_row().cells
            c[0].text = r["what"] or r["sym"]
            c[1].text = r["formula"]
            c[2].text = "" if blank else r["subst"]
            c[3].text = f"{r['sym']} = " + ("" if blank else f"{_fmt(r['value'])} {r['unit']}")
            c[4].text = r["src"]

        doc.add_heading(T("四、校核与结论", "4. Checks and conclusion"), level=2)
        t = doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        for i, h in enumerate([T("校核项", "Check"), T("要求", "Requirement"), T("计算值", "Value"), T("结论", "Result")]):
            t.rows[0].cells[i].text = h
        for ck in self.checks:
            c = t.add_row().cells
            c[0].text = ck["what"]
            c[1].text = f"{ck['sym']} {ck['op']} {_fmt(ck['limit'])}" + (f"（{ck['src']}）" if ck["src"] else "")
            c[2].text = "" if blank else _fmt(ck["value"])
            c[3].text = "" if blank else (T("合格", "Pass") if ck["ok"] else T("不合格", "Fail"))
        doc.add_paragraph("" if blank else self.conclusion())
        for n in self.notes:
            doc.add_paragraph(n)
        if self.sources:
            doc.add_heading(T("五、引用的标准与资料", "5. Standards and references used"), level=2)
            for s in self.sources:
                doc.add_paragraph(s, style="List Number")
        doc.add_paragraph(T("审核：　　　　　　批准：　　　　　　", "Checked by:　　　　　　Approved by:　　　　　　"))
        doc.save(path)
        return path
