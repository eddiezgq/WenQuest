"""云端 GPU 实验的教材写法（人工智能第 14 轮附 4.3 节）：lab.yaml + lab.py 的检查、转成笔记本、打包，指导书与报告模板。"""
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import gpulab  # noqa: E402
import labdocs  # noqa: E402

YAML = """标题: ["实验 4.1 第一个核函数", "Lab 4.1 The first kernel"]
目的: ["编译运行向量加法，测出有效带宽。", "Compile and run vector addition; measure the effective bandwidth."]
显卡: basic
结果:
  - {名: t_ms, 说明: ["核函数耗时", "kernel time"], 单位: ms}
  - {名: bw_GBs, 说明: ["有效带宽", "effective bandwidth"], 单位: GB/s}
文件: [vadd.cu]
实际问题: {标题: ["检测相机", "Inspection camera"], 内容: ["每帧八百万像素。", "Eight million pixels a frame."]}
原理: [["带宽 = 字节 / 时间。", "Bandwidth = bytes / time."]]
步骤: [["运行各格。", "Run the cells."]]
数据表:
  - 标题: ["表 1  测量", "Table 1  Measurements"]
    表头: ["量", "值"]
    行: [["耗时（ms）", "{{t_ms}}"], ["带宽（GB/s）", "{{bw_GBs}}"]]
注意: [["机器按秒计费。", "Billed by the second."]]
"""
NB = """# 实验 4.1 的笔记本（构建时转成 .ipynb）
# %% [markdown]
# # 实验 4.1
# 先检查环境。
# %%
import sys; sys.path.insert(0, "..")
import wqgpu
wqgpu.check(("gpu", "nvcc"))
# %%
!nvcc -O3 -o vadd vadd.cu
out = wqgpu.sh("./vadd")
t_ms, bw = 1.0, 800.0
# %%
wqgpu.submit(t_ms=t_ms, bw_GBs=bw)
"""


def make(tmp: Path, yaml_text=YAML, nb=NB) -> Path:
    d = tmp / "lab4_1"
    d.mkdir(parents=True, exist_ok=True)
    (d / "lab.yaml").write_text(yaml_text, encoding="utf-8")
    (d / "lab.py").write_text(nb, encoding="utf-8")
    (d / "vadd.cu").write_text("__global__ void add() {}\n", encoding="utf-8")
    return d


def test_a_good_lab_packs_with_its_documents(tmp_path):
    d = make(tmp_path)
    g, cs, bad = gpulab.load(d)
    assert bad == []
    assert [c["type"] for c in cs] == ["markdown", "code", "code", "code"] and cs[0]["source"].startswith("# 实验 4.1")
    out = tmp_path / "build" / "gpulab"
    idx = gpulab.pack(out, [("4.1", d, g, cs)])
    assert idx["labs"]["4.1"] == {"title": g["标题"], "tier": "basic", "notebook": "lab4_1.ipynb", "files": ["lab4_1/vadd.cu"],
                                  "results": ["t_ms", "bw_GBs"]}
    nb = json.loads((out / "lab4_1.ipynb").read_text(encoding="utf-8"))
    assert nb["nbformat"] == 4 and nb["cells"][2]["source"][0].startswith("!nvcc")
    assert (out / "lab4_1" / "vadd.cu").exists() and "def submit" in (out / "wqgpu.py").read_text(encoding="utf-8")
    m = gpulab.meta(g)
    labdocs.guide(m, g, out / "lab4_1-guide.docx", "4.1", ("书 4.1 节", "Book, Section 4.1"))
    labdocs.report(m, g, out / "lab4_1-report.docx", "4.1", ("书 4.1 节", "Book, Section 4.1"))
    from docx import Document
    guide = "\n".join(p.text for p in Document(str(out / "lab4_1-guide.docx")).paragraphs)
    assert "打开 GPU 实验" in guide and "检测相机" in guide
    rep = Document(str(out / "lab4_1-report.docx"))
    assert any("{{t_ms}}" in c.text for t in rep.tables for row in t.rows for c in row.cells)


def test_problems_are_reported(tmp_path):
    y = YAML.replace("显卡: basic", "显卡: fast").replace("{{bw_GBs}}", "{{bw}}")
    nb = NB.replace('wqgpu.check(("gpu", "nvcc"))', "print(1)").replace("wqgpu.submit(t_ms=t_ms, bw_GBs=bw)", "wqgpu.submit(t_ms=t_ms)\nx = (")
    _, _, bad = gpulab.load(make(tmp_path, y, nb))
    text = "\n".join(bad)
    assert "显卡" in text and "没有登记的结果 bw" in text and "wqgpu.check" in text
    assert "没有回传 bw_GBs" in text and "语法错误" in text
    _, _, bad = gpulab.load(tmp_path / "nothing")
    assert bad and "没有 GPU 实验文件夹" in bad[0]


def test_measurement_records(tmp_path):
    import gpurun
    run = tmp_path / "ch04" / "code" / "runs" / "run4_1"
    run.mkdir(parents=True)
    (run / "run.yaml").write_text("显卡: basic\n", encoding="utf-8")
    (run / "run.py").write_text("print(1)\n", encoding="utf-8")
    vals, probs = gpurun.values(tmp_path, {4}, {4: "draft"})
    assert vals["run4_1"] == {gpurun.PENDING: True} and probs[0][0] == "warning"
    _, probs = gpurun.values(tmp_path, {4}, {4: "ai"})
    assert probs[0][0] == "error"                                     # past draft, it must be measured
    rec = {"fingerprint": gpurun.fingerprint(run), "at": "2026-10-03T00:00:00Z", "env": {"gpu": "NVIDIA A10", "driver": "570", "nvcc": "12.8"},
           "values": {"t_ms": 0.5}}
    (run.parent / "run4_1.json").write_text(json.dumps(rec), encoding="utf-8")
    vals, probs = gpurun.values(tmp_path, {4}, {4: "ai"})
    assert probs == [] and vals["run4_1"]["t_ms"] == 0.5 and vals["run4_1"]["_gpu"] == "NVIDIA A10" and vals["run4_1"]["_date"] == "2026-10-03"
    (run / "run.py").write_text("print(2)\n", encoding="utf-8")       # changed after the measurement
    vals, probs = gpurun.values(tmp_path, {4}, {4: "ai"})
    assert vals["run4_1"].get(gpurun.PENDING) and "改过" in probs[0][2]
