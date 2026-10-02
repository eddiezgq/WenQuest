"""第 13 轮第 4 步：数字化工艺数据表（自洽、公式、抽查值）与工艺计算书（完整、出处、Word）。"""
import math
import sys
from pathlib import Path

import pytest
import yaml

import stdtab

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "mfgtech" / "conventions"))
import mfgcalc  # noqa: E402

STD = Path(__file__).resolve().parents[2] / "mfgtech" / "std"


def load(name):
    return yaml.safe_load((STD / f"{name}.yaml").read_text(encoding="utf-8"))


def test_every_table_is_well_formed():
    files = sorted(STD.glob("*.yaml"))
    assert len(files) >= 8
    for p in files:
        assert stdtab.check(p) == [], p.name


def test_it_grades_follow_the_standard_tolerance_factor():
    """IT5–IT18 = k·i, i = 0.45·D^(1/3) + 0.001·D (D: geometric mean of the range); the standard rounds, so above 3 mm
    every value lies within 10 % of the formula, and the 30–50 mm row has the well-known values."""
    t = stdtab.table("it_grades")
    k = {5: 7, 6: 10, 7: 16, 8: 25, 9: 40, 10: 64, 11: 100, 12: 160, 13: 250, 14: 400, 15: 640, 16: 1000, 17: 1600, 18: 2500}
    for r in t.rows:
        if r["size_over_mm"] < 3:
            continue
        D = math.sqrt(r["size_over_mm"] * r["size_to_mm"])
        i = 0.45 * D ** (1 / 3) + 0.001 * D
        for g, m in k.items():
            assert r[f"IT{g}_um"] == pytest.approx(m * i, rel=0.10), (r["key"], g)
        vals = [r[f"IT{g}_um"] for g in range(1, 19)]
        assert vals == sorted(vals), r["key"]
    r = t.find(size_over_mm__lt=35, size_to_mm__ge=35)
    assert (r["IT6_um"], r["IT7_um"], r["IT8_um"], r["IT11_um"]) == (16, 25, 39, 160)


def test_casting_tolerance_grows_with_size_and_grade():
    t = stdtab.table("casting_dctg")
    cols = [c for c in t.columns if c.startswith("DCTG")]
    prev = None
    for r in t.rows:
        row = [r[c] for c in cols if r.get(c) is not None]
        assert row == sorted(row), r["key"]
        if prev is not None:
            for c in cols:
                if r.get(c) is not None and prev.get(c) is not None:
                    assert r[c] >= prev[c], (r["key"], c)
        prev = r


def test_kienzle_rows_agree_with_their_own_power_law():
    """kc(h) = kc1.1·h^(−mc): the h = 0.1 and 0.4 mm columns of the main source agree with kc1.1 and mc."""
    t = stdtab.table("kienzle")
    n = 0
    for r in t.rows:
        for h, col in ((0.1, "kc_h01_MPa"), (0.4, "kc_h04_MPa")):
            if r.get(col):
                assert r[col] == pytest.approx(r["kc11_MPa"] * h ** (-r["mc"]), rel=0.02), (r["key"], col)
                n += 1
        assert 0.05 < r["mc"] < 0.6 and 300 < r["kc11_MPa"] < 4000, r["key"]
    assert n >= 20


def test_taylor_exponents_are_in_the_textbook_ranges():
    t = stdtab.table("taylor")
    for r in t.rows:
        for c in ("m", "m_min", "m_max"):
            if r.get(c) is not None:
                assert 0.05 <= r[c] <= 0.9, (r["key"], c)


def test_insert_designation_covers_the_common_codes():
    t = stdtab.table("iso1832")
    codes = {(str(r.get("position")), str(r.get("code"))) for r in t.rows}
    for pos, code in (("1", "C"), ("1", "D"), ("1", "S"), ("1", "T"), ("1", "V"), ("1", "W"), ("2", "N"), ("3", "M")):
        assert (pos, code) in codes, (pos, code)


def test_a_calculation_sheet_cites_its_tables_and_exports_to_word(tmp_path):
    """The process calculation sheet is the shared CalcSheet (《机械设计》); a value looked up in a table carries its
    standard and data source into the sheet, and missing() lists what an incomplete sheet lacks."""
    cs = mfgcalc.CalcSheet("SH-301 轴承位精车工序公差", item="SH-301", author="测试")
    cs.given("d", 35, "mm", "轴承位直径")
    cs.find("T_3", "精车工序尺寸公差")
    it7 = mfgcalc.lookup(cs, "IT7", "it_grades", dict(size_over_mm__lt=35, size_to_mm__ge=35), "IT7_um", "μm", "精车公差等级 IT7")
    cs.step("T_3", "T_3 = IT7", it7 / 1000, "mm", "精车工序公差")
    cs.check("Z_4", 0.3, ">", it7 / 1000, "磨削余量大于精车公差")
    assert it7 == 25 and mfgcalc.missing(cs) == [] and cs.ok
    assert any("GB/T 1800.1-2020" in s for s in cs.sources)
    p = cs.docx(str(tmp_path / "c.docx"))
    from docx import Document
    d = Document(p)
    text = "\n".join([x.text for x in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
    for w in ("已知", "计算", "校核", "GB/T 1800.1-2020", "IT7"):
        assert w in text, w
    bad = mfgcalc.CalcSheet("缺项的计算书")
    bad.given("d", 35, "mm", "直径")
    assert set(mfgcalc.missing(bad)) == {"缺少“求”", "缺少计算步骤", "缺少校核"}
