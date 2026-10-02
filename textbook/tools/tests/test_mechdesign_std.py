"""《机械设计》数字化标准表与计算书（第 11 轮第 4 步）：查表结果与数字工厂的检验模板一致，计算书能导出 Word。"""
import sys
from pathlib import Path

import pytest

CONV = Path(__file__).resolve().parents[2] / "mechdesign" / "conventions"
sys.path.insert(0, str(CONV))

import calcsheet  # noqa: E402
import mdstd  # noqa: E402


def test_fits_agree_with_the_factory_inspection_templates():
    # factory/工厂设计.md 零件检验-输出轴：轴承位 Ø35 k6 35.002–35.018，齿轮位 Ø40 k6 40.002–40.018，键槽宽 12 N9 11.957–12
    assert mdstd.fit_limits(35, "k6") == (35.002, 35.018)
    assert mdstd.fit_limits(40, "k6") == (40.002, 40.018)
    assert mdstd.fit_limits(12, "N9") == (11.957, 12.0)
    assert mdstd.fit_limits(10, "N9") == (9.964, 10.0)          # 中间轴键槽 10 N9
    assert mdstd.fit_limits(35, "H7") == (35.0, 35.025)         # 大齿轮 z72 孔 Ø35 H7


def test_tables_are_monotonic_and_complete():
    t = mdstd.table("iso286_it")
    for row in t["rows"]:
        assert row[2:] == sorted(row[2:]), row               # IT5 < IT6 < ... < IT11
    for a, b in zip(t["rows"], t["rows"][1:]):
        assert a[1] == b[0] and all(x <= y for x, y in zip(a[2:], b[2:]))
    keys = mdstd.table("gbt1095_key")["rows"]
    assert all(a[1] == b[0] for a, b in zip(keys, keys[1:]))
    assert mdstd.key_for(40)["b"] == 12 and mdstd.key_for(30)["b"] == 8   # SH-301：12×8、8×7


def test_materials_come_from_the_factory_library():
    m = mdstd.material("45-QT")
    assert (m["sigma_b"], m["sigma_s"], m["sigma_1"], m["tau_1"]) == (640, 355, 275, 155)
    assert "GB/T 699" in m["src"]


def test_fatigue_factors_are_in_the_usual_ranges():
    assert 0.85 < mdstd.surface_factor("ground", 640) < 0.95
    assert 0.75 < mdstd.surface_factor("machined", 640) < 0.85
    assert 0.80 < mdstd.size_factor(40) < 0.87
    q = mdstd.notch_sensitivity(1.0, 640)
    assert 0.6 < q < 0.85
    with pytest.raises(ValueError):
        mdstd.size_factor(400)


def test_calc_sheet_records_and_exports(tmp_path):
    cs = calcsheet.CalcSheet("测试", item="SH-301", rev="B")
    T = cs.given("T", 350, "N·m", "转矩", src="任务单")
    Ft = cs.step("F_t", "F_t = 2T/d", 2 * T * 1e3 / 210, "N", "圆周力")
    with pytest.raises(ValueError):
        cs.table_value("K_A", 1.25, "", "使用系数", src="")
    assert cs.check("S", 2.0, ">=", 1.5, "安全系数")
    assert not cs.check("S2", 1.2, ">=", 1.5, "另一安全系数")
    assert not cs.ok and "另一安全系数" in cs.conclusion()
    p = cs.docx(str(tmp_path / "a.docx"))
    from docx import Document
    text = "\n".join(c.text for t in Document(p).tables for r in t.rows for c in r.cells)
    assert "3333" in text and "不合格" in text
    blank = cs.docx(str(tmp_path / "b.docx"), blank=True)
    text = "\n".join(c.text for t in Document(blank).tables for r in t.rows for c in r.cells)
    assert "3333" not in text and "F_t = 2T/d" in text
    assert abs(Ft - 3333.33) < 0.01
