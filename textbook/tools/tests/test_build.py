"""Textbook build and checks (第 7 轮 2.3): the real section 4.1 passes; every kind of mistake is stopped."""
import shutil
import sys
from pathlib import Path

import pytest

TOOLS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TOOLS))
import build  # noqa: E402

ROBOTICS = TOOLS.parent / "robotics"


@pytest.fixture
def book(tmp_path):
    """A copy of the robotics book in a temporary folder, to break on purpose."""
    root = tmp_path / "textbook" / "robotics"
    shutil.copytree(ROBOTICS, root, ignore=shutil.ignore_patterns("__pycache__"))
    return root


CH4 = {"4.0", "4.end", *(f"4.{k}" for k in range(1, 9))}


def run(root, **kw):
    kw.setdefault("only", CH4)          # the tests break chapter 4 on purpose; building it alone keeps them quick
    return build.build("robotics", root=root, **kw)


def kinds(rep):
    return {(p.kind, p.text.split("：")[0]) for p in rep.errors}


def edit(root, old, new, path="ch04/04-1.md"):
    p = root / path
    t = p.read_text(encoding="utf-8")
    assert old in t, old
    p.write_text(t.replace(old, new, 1), encoding="utf-8")


def test_the_real_book_passes_and_is_in_line_with_the_outline(tmp_path):
    out = tmp_path / "out"
    rep = build.build("robotics", out_dir=out)
    assert not rep.errors, "\n".join(map(str, rep.errors))
    frag = (out / "web" / "4.1.html").read_text(encoding="utf-8")
    assert "<svg" in frag and "x_p = 0.123205" in frag and "{{" not in frag   # numbers in formulas are SVG glyphs
    assert "WQMATH" not in frag and "WQTOKEN" not in frag
    idx = (out / "web" / "index.json").read_text(encoding="utf-8")
    assert '"no": 61' in idx


def test_a_number_typed_by_hand_is_stopped(book):
    edit(book, "{{ex4_1_1.x_p}}\\ \\mathrm{m}, \\\\", "0.12322\\ \\mathrm{m}, \\\\")
    rep = run(book)
    assert any(p.kind == "手写数字" and "0.12322" in p.text for p in rep.errors)


def test_placeholders_programs_formulas_numbering_references_terms(book):
    edit(book, "{{ex4_1_1.r}}", "{{ex4_1_1.radius}}")                       # not handed over by the program
    edit(book, "\\tag{4.1.3}", "\\tag{4.1.9}")                              # numbering out of order
    edit(book, "式 (4.1.5) 的 2×2", "式 (4.1.12) 的 2×2")                   # reference to nothing
    edit(book, "R(\\theta) = \\begin{pmatrix} \\cos", "R(\\theta = \\frac{ \\begin{pmatrix} \\cos")  # broken formula
    edit(book, "**主动转动**", "**主动旋转术**")                              # term not in the glossary
    edit(book, "第 5 章作为定理", "第 75 章作为定理")                          # no such chapter
    rep = run(book)
    text = "\n".join(map(str, rep.errors))
    assert "ex4_1_1.radius" in text
    assert "公式编号应从 1 起连续" in text
    assert "式 4.1.12 不存在" in text
    assert any(p.kind == "公式" for p in rep.errors)
    assert "主动旋转术" in text
    assert "第 75 章不存在" in text


def test_a_failing_program_stops_the_build(book):
    p = book / "ch04" / "code" / "ex4_1_1.py"
    for q in (p, p.parent / "en" / p.name):          # the English copy must stay the same code
        q.write_text(q.read_text(encoding="utf-8").replace("theta = math.radians(30)", "theta = math.radians(31)"), encoding="utf-8")
    edit(book, "R(30^\\circ) = \\begin{pmatrix}", "R(30^\\circ) = \\begin{pmatrix}")   # text unchanged
    rep = run(book)
    # the program's own checks still hold at 31°, so it runs; the numbers in the book follow it automatically
    assert not [e for e in rep.errors if e.kind == "程序"]
    frag = (book.parent / "build" / "robotics" / "web" / "4.1.html").read_text(encoding="utf-8")
    assert "x_p = 0.123205" not in frag and "x_p = 0.11" in frag
    p.write_text(p.read_text(encoding="utf-8") + "\nraise SystemExit('核对失败')\n", encoding="utf-8")
    rep = run(book)
    assert any(e.kind == "程序" for e in rep.errors)


def test_section_files_must_match_the_outline(book):
    (book / "ch04" / "04-9.md").write_text("---\nid: \"4.99\"\ntitle: 不存在的节\n---\n正文\n", encoding="utf-8")
    rep = run(book, only=CH4 | {"4.99"})
    assert any("4.99 不在提纲里" in p.text for p in rep.errors)


def test_significant_digits_keep_trailing_zeros():
    assert build.sig(0.1866025403784) == "0.18660"
    assert build.sig(26.565051177) == "26.565"
    assert build.sig(0.05) == "0.050000"


def test_an_animation_needs_its_static_figure(book):
    """动画代替不了示意图，两个都要 (Eddie 2026-10-01)."""
    edit(book, "::: 动画 4.1.1\nsrc: a4_1_1\n图: 4.1.1\n", "::: 动画 4.1.1\nsrc: a4_1_1\n")
    rep = run(book)
    assert any("动画 4.1.1 没有配示意图" in p.text for p in rep.errors)
    frag = (book.parent / "build" / "robotics" / "web" / "4.1.html").read_text(encoding="utf-8")
    assert "data:image/svg+xml;base64" in frag and "图 4.1.2" in frag      # the figures are in the page itself


def test_animations_and_labs_are_packaged_and_checked(book):
    """Scenes must exist and pass the animation service's own check; labs must exist and pass the lab kit's check."""
    rep = run(book)
    out = book.parent / "build" / "robotics"
    assert (out / "anim" / "a4_1_1.py").exists() and (out / "lab" / "ch04.html").exists()
    p = book / "ch04" / "anim" / "a4_1_1.py"
    p.write_text(p.read_text(encoding="utf-8").replace("import math", "import math\nimport os"), encoding="utf-8")
    q = book / "ch04" / "lab" / "lab4_1.js"
    q.write_text(q.read_text(encoding="utf-8") + "\nfetch('x');\n", encoding="utf-8")
    edit(book, "src: a4_1_2\n", "src: no_such_scene\n")
    rep = run(book)
    text = "\n".join(map(str, rep.errors))
    assert "动画 4.1.1：场景程序不能通过渲染服务的检查" in text
    assert "动画 4.1.2：没有场景程序" in text
    assert "实验 4.1：not allowed: network access (fetch)" in text


def test_every_lab_has_a_guide_and_a_report_template(book):
    """第 7 轮第 4 步补充: lab/NAME.yaml → 实验指导书 and 实验报告模板 (Word); a missing or broken file stops the build."""
    import docx
    rep = run(book)
    assert not rep.errors
    lab = book.parent / "build" / "robotics" / "lab"
    for k in range(1, 9):
        for kind in ("guide", "report"):
            assert (lab / f"lab4_{k}-{kind}.docx").exists()
    text = "\n".join(p.text for p in docx.Document(str(lab / "lab4_3-guide.docx")).paragraphs)
    assert "实验 4.3 找出看不见的转轴" in text and "四、实验步骤" in text and "让一次转动与目标姿态相差不到 1°" in text
    rtext = "\n".join(p.text for p in docx.Document(str(lab / "lab4_3-report.docx")).paragraphs)
    assert "九、AI 使用声明" in rtext and "为什么转角为 180° 时" in rtext
    (book / "ch04" / "lab" / "lab4_2.yaml").unlink()
    y = book / "ch04" / "lab" / "lab4_3.yaml"
    y.write_text(y.read_text(encoding="utf-8").replace('["(0.6, 0.8, 0)", "", "", "", ""]', '["(0.6, 0.8, 0)", "", ""]'), encoding="utf-8")
    text = "\n".join(map(str, run(book).errors))
    assert "实验 4.2：没有实验说明文件 lab4_2.yaml" in text
    assert "实验 4.3：lab4_3.yaml“表 2  转了半圈”有一行 3 格，表头是 5 格" in text


def test_the_english_edition_is_checked_against_the_chinese(book):
    """第 8 轮：NN-M.en.md keeps the Chinese structure, uses English glossary terms, and lists a program copy that differs
    only in comments; every figure is drawn again with English labels."""
    rep = run(book)
    assert not rep.errors, "\n".join(map(str, rep.errors))
    out = book.parent / "build" / "robotics"
    en = (out / "web" / "en" / "4.1.html").read_text(encoding="utf-8")
    assert "lang='en'" in en and "Figure 4.1.1" in en and "{{" not in en
    assert "x_p = 0.123205" in en and "Example 4.1.1" in en and "运行结果" not in en and "Output" in en
    assert (out / "figs" / "en" / "fig4_1_1.svg").exists() and (out / "lab" / "lab4_1-guide.en.docx").exists()
    idx = (out / "web" / "index.json").read_text(encoding="utf-8")
    assert '"title_en": "Rotation in the Plane and the Rotation Matrix"' in idx and '"status": "first"' in idx
    edit(book, "\\tag{4.1.3}", "", "ch04/04-1.en.md")                                  # an equation number lost
    edit(book, "**passive rotation**", "passive rotation", "ch04/04-1.en.md")         # a term no longer marked
    edit(book, "**reflection**", "**mirror flip**", "ch04/04-1.en.md")                # not a glossary term
    edit(book, "theta = math.radians(30)", "theta = math.radians(31)", "ch04/code/en/ex4_1_1.py")   # different code
    text = "\n".join(map(str, run(book).errors))
    assert "公式编号与中文版不一致：英文版缺少 4.1.3" in text
    assert "“被动转动”，英文版没有加粗对应的“passive rotation”" in text
    assert "“mirror flip”不是术语表里的英文名" in text
    assert "code/en/ex4_1_1.py 与中文版的代码不同" in text


def test_the_same_term_has_the_same_english_name_in_every_book(book):
    """第 9 轮 2.3: a Chinese term shared by two books should share an English name (brackets, dashes ignored);
    a difference is a 提醒 (one book never stops another's build); "含义不同" in 备注 marks a deliberate one."""
    other = book.parent / "physics" / "conventions"
    other.mkdir(parents=True)
    (other / "术语表.csv").write_text("首次出现章,中文,English,备注\n1,旋转矩阵,rotation matrix (SO(3)),\n1,四元数,hypercomplex number,\n"
                                    "1,右手定则,right–hand rule,\n1,刚体运动,shift,含义不同\n", encoding="utf-8")
    rep = run(book, only={"4.1"})
    bad = [p.text for p in rep.problems if p.kind == "术语" and "不一致" in p.text]
    assert len(bad) == 1 and "四元数" in bad[0], bad
    assert not [p for p in rep.errors if "不一致" in p.text]


def test_the_index_lists_each_chapters_animations_and_labs(book, tmp_path):
    """第 9 轮 2.5: index.json carries the 互动资源 of every chapter, in reading order, with the section they are in."""
    out = tmp_path / "out"
    rep = build.build("robotics", root=book, out_dir=out, only={"4.1"})
    assert not rep.errors, "\n".join(map(str, rep.errors))
    import json
    res = json.loads((out / "web" / "index.json").read_text(encoding="utf-8"))["resources"]["4"]
    kinds_nums = [(r["kind"], r["num"]) for r in res]
    assert ("anim", "4.1.1") in kinds_nums and ("lab", "4.1") in kinds_nums
    assert all(r["sec"] == "4.1" and r["title"] for r in res)


def test_a_law_is_a_numbered_item(book):
    """第 9 轮: **定律 x.y.z** is numbered like definitions and theorems, and can be referred to."""
    edit(book, "**定义 4.1.1（平面旋转矩阵）**", "**定律 4.1.1（试验）** 一条定律。\n\n**定义 4.1.1（平面旋转矩阵）**")
    edit(book, "### 4.1.4 旋转矩阵", "### 4.1.4 旋转矩阵\n\n由定律 4.1.1 与定律 4.1.2 可知。")
    rep = run(book, only={"4.1"})
    bad = [p.text for p in rep.errors if p.kind == "引用"]
    assert bad == ["定律 4.1.2 不存在"], bad
