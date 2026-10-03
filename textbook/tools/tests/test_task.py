"""工程任务单（第 13 轮）：说明文件的检查、任务单与评分量规的 Word、空白计算书。"""
from pathlib import Path

import yaml

import tasksheet

GOOD = {
    "编号": "TS-1-1",
    "标题": ["读懂一条工艺路线", "Read a routing"],
    "角色": ["工艺员", "process engineer"],
    "背景": ["工厂接到 50 件订单。", "The factory has an order for 50 pieces."],
    "输入": [{"名称": ["零件图", "Part drawing"], "来源": ["设计与工艺页", "Design page"]}],
    "交付物": [{"名称": ["工艺过程卡", "Process route sheet"], "验收": ["AI 评审无“必须改”意见", "No must-fix comment"]}],
    "工位": ["ERP", "PLM"],
    "步骤": [["打开 ERPNext", "Open ERPNext"]],
    "评审要点": [["工序顺序先粗后精", "Roughing before finishing"]],
    "评分": [{"项": ["工艺正确", "Correct plan"], "分": 70, "标准": ["无错误", "No errors"]},
             {"项": ["文件规范", "Tidy documents"], "分": 30, "标准": ["格式符合 JB/T 9165", "Format per JB/T 9165"]}],
    "学时": 2,
    "计算书": {"程序": "code/calc.py", "函数": "sheet"},
}


def write(tmp_path, data):
    ch = tmp_path / "ch01"
    (ch / "task").mkdir(parents=True)
    (ch / "code").mkdir()
    (ch / "code" / "calc.py").write_text(
        "from mfgcalc import CalcSheet\n"
        "def sheet():\n"
        "    cs = CalcSheet('余量计算', item='SH-301')\n"
        "    cs.given('d', 35, 'mm', '直径')\n"
        "    cs.step('Z', 'Z = d1 - d2', 0.3, 'mm', '磨削余量', subst='35.3 - 35.0')\n"
        "    cs.check('Z', 0.3, '>', 0.039, '余量大于精车公差')\n"
        "    return cs\n", encoding="utf-8")
    p = ch / "task" / "ts1_1.yaml"
    p.write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    return ch, p


def test_a_good_task_sheet_makes_three_word_files(tmp_path):
    ch, p = write(tmp_path, GOOD)
    t, bad = tasksheet.load(p)
    assert bad == []
    where = ("《机械制造技术》1.1 节", "Mechanical Manufacturing Technology, Section 1.1")
    out = tmp_path / "out"
    from docx import Document

    def text(f):
        d = Document(f)
        return "\n".join([x.text for x in d.paragraphs] + [c.text for tb in d.tables for r in tb.rows for c in r.cells])
    a = text(tasksheet.task_doc(t, out / "a.docx", "1.1", where))
    for w in ("工程任务单 TS-1-1", "交付物与验收标准", "工艺过程卡", "AI 评审", "〔ERP〕", "Rubric" if False else "评分"):
        assert w in a, w
    b = text(tasksheet.rubric_doc(t, out / "b.docx", "1.1", where))
    assert "满分标准" in b and "合计 Total" in b and "格式符合 JB/T 9165" in b
    conv = Path(__file__).resolve().parents[2] / "mfgtech" / "conventions"
    c = text(tasksheet.blank_calc(t, ch, out / "c.docx", conv))
    assert "Z = d1 - d2" in c and "35.3 - 35.0" not in c and "0.3 mm" not in c      # 空白：公式在，代入和结果留空


def test_problems_are_reported(tmp_path):
    bad_data = dict(GOOD)
    bad_data["工位"] = ["ERP", "XYZ"]
    bad_data["评分"] = [{"项": ["工艺", "Plan"], "分": 60, "标准": ["对", "Right"]}]
    bad_data["计算书"] = {"程序": "code/none.py", "函数": "sheet"}
    del bad_data["背景"]
    _, p = write(tmp_path, bad_data)
    _, bad = tasksheet.load(p)
    text = " ".join(bad)
    assert "缺少“背景”" in text and "XYZ" in text and "合计 60 分" in text and "none.py 不存在" in text


def test_deliverable_kinds_for_the_platform(tmp_path):
    """第 11 轮 2.7（5）：交付物的类型、格式、零件、必列。"""
    import copy
    d = copy.deepcopy(GOOD)
    d["交付物"] = [
        dict(GOOD["交付物"][0], 类型="文件", 格式=["docx", "pdf"]),
        dict(GOOD["交付物"][0], 类型="设计发布", 零件="SH-301", 轴承=["6207"]),
        dict(GOOD["交付物"][0], 类型="更改单", 必列=["GR-302", "KEY|键"]),
    ]
    _, p = write(tmp_path, d)
    assert tasksheet.load(p)[1] == []
    d["交付物"] = [dict(GOOD["交付物"][0], 类型="照片"), dict(GOOD["交付物"][0], 类型="分析"),
                  dict(GOOD["交付物"][0], 格式=["exe"]), dict(GOOD["交付物"][0], 必列=["GR-302"])]
    (p.parent / "ts1_1.yaml").write_text(yaml.safe_dump(d, allow_unicode=True), encoding="utf-8")
    bad = " ".join(tasksheet.load(p)[1])
    assert "类型 照片" in bad and "要写“零件”" in bad and "格式" in bad and "只用于类型“更改单”" in bad


def test_parts_library_links():
    """第 11 轮：[[零件:编号/规格]] 链到零件库，编号、规格都要在零件库里"""
    import build
    rep = build.Report()

    class Sec:
        id, lang = "33.2", "zh"
    h = build.part_html(Sec, "A-BRG-DG", "6207", rep)
    assert "深沟球轴承 6207" in h and "ref=A-BRG-DG/6207" in h and not rep.problems
    build.part_html(Sec, "A-BRG-DG", "9999", rep)
    build.part_html(Sec, "X-NOPE", None, rep)
    assert len(rep.problems) == 2
