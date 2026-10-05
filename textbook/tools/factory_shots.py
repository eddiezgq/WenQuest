# -*- coding: utf-8 -*-
"""数字工厂界面截图（第 13 轮 8.8，R5）：在本机起一套教学用数字工厂，灌测试数据，用 Playwright 以教师、学生两种账号截图。

    python3 textbook/tools/factory_shots.py mfgtech            # 全部
    python3 textbook/tools/factory_shots.py mfgtech --only 51   # 只截第 51 章的

要本机有 PostgreSQL（默认 postgresql://postgres:wq@localhost:5432，可用 --dsn 或环境变量 WQ_SHOTS_DSN 改），程序新建一个临时库
wq_shots（每次重建）；不接 ERPNext、MQTT。输出 textbook/<书>/shots/<名称>.png 和 shots/清单.yaml（页面、账号、数据来源、日期），
书中由各章的插图程序把截图包成插图。测试数据全部来自仓库：参照工艺规程 std/SH-301_process.yaml，以及按任务单 TS-51-1～53-1 要求
“故意改坏”的三个变体（改法见 VARIANTS）；任务单内容取自 textbook/<书>/ch*/task/*.yaml。账号用“教师”“学生甲”，不出现真实人名。
"""
from __future__ import annotations

import argparse
import base64
import copy
import datetime as dt
import hashlib
import hmac
import json
import os
import subprocess
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

import yaml

REPO = Path(__file__).resolve().parents[2]
DIGITAL = REPO / "factory" / "digital"
SECRET = "shots-secret"
TASK_KEY = "shots-key"
PORT = 8137
BASE = f"http://127.0.0.1:{PORT}"
STUDENT = {"name": "学生甲", "uid": "S001", "teacher": False, "role": "engineer", "mode": "teach"}


# ---------------------------------------------------------------- 故意改坏的工艺规程（与任务单的要求一致）
def _op(plan, seq):
    return next(o for o in plan["operations"] if o["seq"] == seq)


def v_locate(plan):
    """TS-51-1：把定位外圆的公差 Td 改为 0.3 mm（相当于在不控制公差的粗车外圆上定位）"""
    _op(plan, 50)["locate_check"]["Td"] = 0.3
    return "工序 50 的 locate_check 中 Td 由 0.039 改为 0.3（TS-51-1 的“故意改坏”）"


def v_fixture(plan):
    """TS-52-1：工序 50 写入自己的夹具数据，拧紧扭矩改为 1 N·m（hand_force 10、wrench 100）"""
    fx = yaml.safe_load((DIGITAL / "std" / "JJ-301-50_fixture.yaml").read_text(encoding="utf-8"))
    fx["clamping"]["hand_force"] = 10
    _op(plan, 50)["fixture"] = fx
    return "工序 50 写入夹具数据（同 JJ-301-50），hand_force 由 50 改为 10，即拧紧扭矩 1 N·m（TS-52-1 的“故意改坏”）"


def v_chart(plan):
    """TS-53-1：跟踪图第 40.3 刀（工序 40 总长）的公差改为 ±0.3"""
    c = next(c for c in plan["tolerance_chart"]["cuts"] if abs(c["seq"] - 40.3) < 1e-9)
    c["es"], c["ei"] = 0.3, -0.3
    return "tolerance_chart 第 40.3 刀（工序 40 总长）的公差由 ±0.2 改为 ±0.3（TS-53-1 的“故意改坏”）"


VARIANTS = {"locate": v_locate, "fixture": v_fixture, "chart": v_chart}


# ---------------------------------------------------------------- 数字工厂
def _sign(payload):          # 与 hub/app.py 的 _sign 相同（同一个 WQ_SECRET）
    raw = base64.urlsafe_b64encode(json.dumps(payload, ensure_ascii=False).encode()).decode()
    return raw + "." + hmac.new(SECRET.encode(), raw.encode(), hashlib.sha256).hexdigest()[:32]


def call(method, path, body=None, token=None, key=None):
    h = {"content-type": "application/json"}
    if token:
        h["x-wq-token"] = token
    if key:
        h["authorization"] = "Bearer " + key
    req = urllib.request.Request(BASE + path, method=method, headers=h,
                                 data=None if body is None else json.dumps(body, ensure_ascii=False).encode())
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            t = r.read().decode()
            return json.loads(t) if t and r.headers.get_content_type() == "application/json" else t
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"{method} {path} → {e.code}: {e.read().decode()[:300]}") from None


def upload(path, token, filename, data):
    """multipart 上传一个文件（任务单的“上传文件”交付物）"""
    bnd = "----wqshots" + hashlib.md5(data).hexdigest()
    body = (f"--{bnd}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
            "Content-Type: application/octet-stream\r\n\r\n").encode() + data + f"\r\n--{bnd}--\r\n".encode()
    req = urllib.request.Request(BASE + path, method="POST", data=body,
                                 headers={"x-wq-token": token, "content-type": f"multipart/form-data; boundary={bnd}"})
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.loads(r.read().decode())


def sample_docx(title, paras):
    """测试用的交付物（Word）：标题加几段说明，注明是截图用的示例"""
    import io
    import docx
    d = docx.Document()
    d.add_heading(title, 1)
    for t in paras:
        d.add_paragraph(t)
    d.add_paragraph("（数字工厂界面截图用的示例文件，不是答案。）")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def fresh_db(dsn):
    import psycopg
    with psycopg.connect(dsn.rsplit("/", 1)[0] + "/postgres", autocommit=True) as c:
        c.execute("drop database if exists wq_shots with (force)")
        c.execute("create database wq_shots")
    return dsn.rsplit("/", 1)[0] + "/wq_shots"


def start_hub(db):
    env = dict(os.environ, WQ_DB=db, WQ_SECRET=SECRET, WQ_FACTORY_TASK_KEY=TASK_KEY, WQ_AUTO_SCENARIO="0", WQ_MQTT_TIMEOUT="1",
               WQ_AUTH="local", WQ_MODE="teach")
    log = open(REPO / "textbook" / "build" / "factory_shots_hub.log", "w")
    p = subprocess.Popen([sys.executable, "-m", "uvicorn", "hub.app:app", "--host", "127.0.0.1", "--port", str(PORT)],
                         cwd=DIGITAL, env=env, stdout=log, stderr=subprocess.STDOUT)
    for _ in range(90):
        time.sleep(1)
        if p.poll() is not None:
            raise RuntimeError("数字工厂没有起来，见 textbook/build/factory_shots_hub.log")
        try:
            call("POST", "/api/login", {"name": "教师", "role": "engineer", "mode": "teach"})
            return p
        except Exception:  # noqa: BLE001
            continue
    p.terminate()
    raise RuntimeError("数字工厂 90 秒内没有响应")


def seed(book):
    """灌测试数据：学生提交参照工艺规程和三个变体，下达任务单 TS-50-1～53-1，学生在任务单里选自己的工艺规程。"""
    teacher = call("POST", "/api/login", {"name": "教师", "role": "engineer", "mode": "teach"})["token"]
    student = _sign(STUDENT)
    ref = call("GET", "/api/process/reference/SH-301", token=student)["yaml"]
    subs, notes = {}, {}
    s = call("POST", "/api/process/check", {"yaml": ref, "submit": True, "note": "按参照工艺规程提交"}, token=student)
    subs["ref"] = s["id"]
    for k, fn in VARIANTS.items():
        plan = yaml.safe_load(ref)
        notes[k] = fn(plan)
        s = call("POST", "/api/process/check", {"yaml": yaml.safe_dump(plan, allow_unicode=True, sort_keys=False), "submit": True,
                                                "note": notes[k]}, token=student)
        subs[k] = s["id"]
    tasks = {}
    root = REPO / "textbook" / book
    for ch, code, link in ((50, "TS-50-1", None), (51, "TS-51-1", "ref"), (52, "TS-52-1", "ref"), (53, "TS-53-1", "ref")):
        spec = yaml.safe_load((root / f"ch{ch}" / "task" / f"ts{ch}_1.yaml").read_text(encoding="utf-8"))
        spec.update(book=book, no=f"{ch}.1")
        due = (dt.datetime.now(dt.timezone.utc) + dt.timedelta(days=10)).isoformat()
        tid = call("POST", "/api/tasks/course/issue", {"spec": spec, "course_id": 101, "course_name": "机械制造技术（示例班）",
                                                       "due": due, "issued_by": "教师"}, key=TASK_KEY)["id"]
        tasks[ch] = tid
        if link:
            i = next(n for n, x in enumerate(spec["交付物"]) if x.get("类型") == "工艺规程")
            call("POST", f"/api/tasks/{tid}/deliverables/{i}/link", {"process": subs[link]}, token=student)
    # 第 52 章截老师审阅页：学生把其余交付物也交上（示例 Word 文件），再提交审批
    spec52 = yaml.safe_load((root / "ch52" / "task" / "ts52_1.yaml").read_text(encoding="utf-8"))
    for i, x in enumerate(spec52["交付物"]):
        if x.get("类型", "文件") == "文件":
            upload(f"/api/tasks/{tasks[52]}/deliverables/{i}/file", student, f"TS-52-1_{i + 1}.docx",
                   sample_docx(x["名称"][0], [x["验收"][0]]))
    call("POST", f"/api/tasks/{tasks[52]}/submit", {}, token=student)
    rev = call("GET", f"/api/tasks/{tasks[52]}/submissions", token=teacher)["submissions"][0]["id"]
    return {"teacher": teacher, "student": student, "subs": subs, "notes": notes, "tasks": tasks, "review52": rev}


# ---------------------------------------------------------------- 截图
def shots_list(d):
    s, t = d["subs"], d["tasks"]
    return [  # (名称, 章, 账号, 页面, 操作, 说明)
        ("shot50_1", 50, "student", "/process", "precheck_ref", "“工艺规程与 AI 评审”页：下载参照工艺规程、上传 AI 预检、提交记录"),
        ("shot50_2", 50, "student", f"/process/{s['ref']}", None, "按参照工艺规程提交后，AI 工艺评审员的意见与建议分"),
        ("shot50_3", 50, "student", f"/tasks/{t[50]}", None, "工程任务单 TS-50-1 在数字工厂的学生页：交付物、验收标准、进度"),
        ("shot51_1", 51, "student", f"/process/{s['locate']}", None, d["notes"]["locate"]),
        ("shot51_2", 51, "student", f"/tasks/{t[51]}", "precheck_task", "任务单 TS-51-1 的学生页：选定自己的工艺规程提交，点“AI 预检”"),
        ("shot52_1", 52, "student", f"/process/{s['fixture']}", None, d["notes"]["fixture"]),
        ("shot52_2", 52, "teacher", f"/tasks/review/{d['review52']}", None, "任务单 TS-52-1 的老师审阅页：交付物、AI 评审员的意见、评分"),
        ("shot53_1", 53, "student", f"/process/{s['chart']}", None, d["notes"]["chart"]),
        ("shot53_2", 53, "student", f"/tasks/{t[53]}", "precheck_task", "任务单 TS-53-1 的学生页：选定工艺规程提交，AI 预检"),
    ]


def take(book, d, only=None):
    from playwright.sync_api import sync_playwright
    out = REPO / "textbook" / book / "shots"
    out.mkdir(exist_ok=True)
    users = {"teacher": {"token": d["teacher"], "user": {"name": "教师", "role": "engineer", "mode": "teach", "teacher": True}},
             "student": {"token": d["student"], "user": STUDENT}}
    ref = call("GET", "/api/process/reference/SH-301", token=d["student"])["yaml"]
    done = []
    with sync_playwright() as pw:
        b = pw.chromium.launch()
        for name, ch, who, path, act, note in shots_list(d):
            if only and str(ch) not in only:
                continue
            ctx = b.new_context(viewport={"width": 1440, "height": 900}, device_scale_factor=2, locale="zh-CN")
            ctx.add_init_script(f"localStorage.setItem('wq-df-session', {json.dumps(json.dumps(users[who], ensure_ascii=False))})")
            pg = ctx.new_page()
            pg.goto(BASE + path, wait_until="networkidle")
            pg.wait_for_timeout(800)
            if act == "precheck_ref":       # 上传参照工艺规程，点“AI 预检（不提交）”
                pg.set_input_files("input[type=file]", files=[{"name": "SH-301_process.yaml", "mimeType": "text/yaml", "buffer": ref.encode()}])
                pg.wait_for_timeout(300)
                pg.get_by_role("button", name="AI 预检（不提交）").click()
                pg.wait_for_timeout(1500)
            elif act == "precheck_task":
                pg.get_by_role("button", name="AI 预检").click()
                pg.wait_for_timeout(1500)
            cards = pg.locator(".page section.card")
            n = min(cards.count(), 2 if path.startswith("/tasks") else 9)     # 任务单页只截前两块（任务、交付物），后面是步骤和量规
            boxes = [cards.nth(i).bounding_box() for i in range(n)]
            top = pg.locator(".page .page-title").first.bounding_box() or boxes[0]
            x0 = min(b["x"] for b in boxes) - 8
            clip = {"x": x0, "y": top["y"] - 8, "width": max(b["x"] + b["width"] for b in boxes) + 8 - x0,
                    "height": max(b["y"] + b["height"] for b in boxes) + 8 - top["y"] + 8}
            pg.screenshot(path=str(out / f"{name}.png"), clip=clip, full_page=True)
            done.append({"名称": name, "章": ch, "账号": "教师" if who == "teacher" else "学生甲", "页面": path,
                         "说明": note, "数据": "测试数据（factory_shots.py 灌入，见程序说明）"})
            ctx.close()
        b.close()
    man = out / "清单.yaml"
    old = yaml.safe_load(man.read_text(encoding="utf-8")) if man.exists() else {}
    items = {x["名称"]: x for x in (old or {}).get("截图", [])}
    items.update({x["名称"]: x for x in done})
    commit = subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=REPO, capture_output=True, text=True).stdout.strip()
    man.write_text("# 数字工厂界面截图清单（由 textbook/tools/factory_shots.py 生成，不要手改）\n" + yaml.safe_dump(
        {"日期": dt.date.today().isoformat(), "数字工厂版本": commit, "截图": sorted(items.values(), key=lambda x: x["名称"])},
        allow_unicode=True, sort_keys=False, width=200), encoding="utf-8")
    return done


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("book")
    ap.add_argument("--only", nargs="*", help="只截这几章（如 51 52）")
    ap.add_argument("--dsn", default=os.environ.get("WQ_SHOTS_DSN", "postgresql://postgres:wq@localhost:5432/postgres"))
    a = ap.parse_args(argv)
    (REPO / "textbook" / "build").mkdir(exist_ok=True)
    hub = start_hub(fresh_db(a.dsn))
    try:
        d = seed(a.book)
        for x in take(a.book, d, a.only):
            print("截图", x["名称"], x["页面"])
    finally:
        hub.terminate()
        hub.wait(20)


if __name__ == "__main__":
    main()
