"""Engineering calculation sheets (设计 / 工艺计算书, 第 11、13 轮): 已知 → 求 → 计算 → 结论, every looked-up value
with its standard and table, exported to Word.

The same sheet is used three ways: a worked example in the book (the program builds the sheet and hands its numbers
to the text with ``out``), a lab or task where the student fills in the inputs, and the factory's process-document
editor, which attaches the sheet to the process plan it reviews.

    from calcsheet import Sheet
    s = Sheet("SH-301 φ35k6 轴承位工序尺寸", part="SH-301", author="张三")
    s.given("d", "轴承位直径", 35, "mm")
    s.find("工序尺寸 A_i 及公差")
    it7 = s.lookup("IT7", "it_grades", dict(size_over_mm__lt=35, size_to_mm__ge=35), "IT7_um", "μm")
    s.step("精车工序公差取 IT7", "T_3 = IT7", it7 / 1000, "mm")
    s.conclude("精车工序尺寸 φ35.3₋₀.₀₂₅ mm")
    s.docx("calc.docx")
"""
from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass, field
from pathlib import Path

import stdtab


@dataclass
class Item:
    kind: str                       # given / find / lookup / step / check / conclude / note
    name: str = ""
    text: str = ""
    formula: str = ""
    value: object = None
    unit: str = ""
    cite: str = ""
    ok: bool | None = None


def _fmt(v, digits: int = 4) -> str:
    if isinstance(v, bool) or v is None:
        return "" if v is None else str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v == 0:
            return "0"
        return f"{v:.{digits}g}"
    return str(v)


@dataclass
class Sheet:
    title: str
    part: str = ""
    author: str = ""
    task: str = ""
    lang: str = "zh"
    items: list[Item] = field(default_factory=list)

    # ------------------------------------------------------------ building
    def given(self, name: str, text: str, value, unit: str = "", cite: str = ""):
        self.items.append(Item("given", name, text, value=value, unit=unit, cite=cite))
        return value

    def find(self, text: str):
        self.items.append(Item("find", text=text))

    def lookup(self, name: str, table: str, where: dict, col: str, unit: str = "", text: str = ""):
        """Look a value up in a digitized standard table; the sheet records the standard and table it came from."""
        t = stdtab.table(table)
        row = t.find(**where) if where else None
        v = t.value(row, col)
        self.items.append(Item("lookup", name, text or f"{t.columns[col]['zh']}（{row.get('key')}）",
                               value=v, unit=unit, cite=t.cite(row, self.lang)))
        return v

    def step(self, text: str, formula: str, value, unit: str = "", cite: str = ""):
        self.items.append(Item("step", text=text, formula=formula, value=value, unit=unit, cite=cite))
        return value

    def check(self, text: str, ok: bool, formula: str = ""):
        """A verification (校核): the sheet marks it passed or failed, and failed checks are listed in the conclusion."""
        self.items.append(Item("check", text=text, formula=formula, ok=bool(ok)))
        return ok

    def note(self, text: str):
        self.items.append(Item("note", text=text))

    def conclude(self, text: str):
        self.items.append(Item("conclude", text=text))

    # ------------------------------------------------------------ reading back
    @property
    def failed(self) -> list[str]:
        return [i.text for i in self.items if i.kind == "check" and i.ok is False]

    @property
    def citations(self) -> list[str]:
        seen: list[str] = []
        for i in self.items:
            if i.cite and i.cite not in seen:
                seen.append(i.cite)
        return seen

    def to_dict(self) -> dict:
        return {"title": self.title, "part": self.part, "author": self.author, "task": self.task,
                "items": [i.__dict__ for i in self.items], "failed": self.failed, "citations": self.citations}

    def missing(self) -> list[str]:
        """What a complete sheet lacks (used by the AI reviewer and the tests): 已知, 求, 计算, 结论; every 查表 cited."""
        L = {"zh": ("缺少“已知”", "缺少“求”", "缺少计算步骤", "缺少“结论”", "查表值没有注明出处"),
             "en": ("no givens", "nothing to find", "no calculation steps", "no conclusion", "a looked-up value has no source")}[self.lang]
        kinds = {i.kind for i in self.items}
        out = []
        for k, msg in zip(("given", "find", "step", "conclude"), L):
            if k not in kinds and not (k == "step" and "lookup" in kinds):
                out.append(msg)
        if any(i.kind == "lookup" and not i.cite for i in self.items):
            out.append(L[4])
        return out

    # ------------------------------------------------------------ Word
    def docx(self, path: str | Path) -> Path:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor

        zh = self.lang == "zh"
        H = (("工艺计算书", "已知", "求", "计算", "校核", "结论", "数据出处", "零件", "编制", "日期", "任务单", "项目", "符号", "数值", "单位",
              "出处", "通过", "不通过", "未通过的校核")
             if zh else ("Calculation Sheet", "Given", "Find", "Calculation", "Checks", "Conclusion", "Sources", "Part", "Prepared by",
                         "Date", "Task", "Item", "Symbol", "Value", "Unit", "Source", "pass", "FAIL", "Failed checks"))
        doc = Document()
        st = doc.styles["Normal"]
        st.font.name = "Noto Sans CJK SC"
        st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        h = doc.add_heading(H[0], level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(self.title, level=1)
        meta = doc.add_table(rows=1, cols=4)
        meta.style = "Table Grid"
        c = meta.rows[0].cells
        c[0].text = f"{H[7]}：{self.part}" if zh else f"{H[7]}: {self.part}"
        c[1].text = f"{H[8]}：{self.author}" if zh else f"{H[8]}: {self.author}"
        c[2].text = f"{H[10]}：{self.task}" if zh else f"{H[10]}: {self.task}"
        c[3].text = f"{H[9]}：{_dt.date.today().isoformat()}" if zh else f"{H[9]}: {_dt.date.today().isoformat()}"

        def table_of(items, cols):
            t = doc.add_table(rows=1, cols=len(cols))
            t.style = "Table Grid"
            for j, name in enumerate(cols):
                t.rows[0].cells[j].text = name
            for it in items:
                r = t.add_row().cells
                vals = {"given": [it.text, it.name, _fmt(it.value), it.unit, it.cite],
                        "lookup": [it.text, it.name, _fmt(it.value), it.unit, it.cite]}[it.kind]
                for j, v in enumerate(vals):
                    r[j].text = v
            return t

        givens = [i for i in self.items if i.kind == "given"]
        if givens:
            doc.add_heading(H[1], level=2)
            table_of(givens, [H[11], H[12], H[13], H[14], H[15]])
        finds = [i for i in self.items if i.kind == "find"]
        if finds:
            doc.add_heading(H[2], level=2)
            for i in finds:
                doc.add_paragraph(i.text, style="List Bullet")
        doc.add_heading(H[3], level=2)
        n = 0
        for it in self.items:
            if it.kind in ("step", "lookup"):
                n += 1
                p = doc.add_paragraph()
                p.add_run(f"{n}. {it.text}").bold = True
                line = doc.add_paragraph()
                if it.formula:
                    line.add_run(f"{it.formula} = ")
                elif it.name:
                    line.add_run(f"{it.name} = ")
                line.add_run(f"{_fmt(it.value)} {it.unit}".strip())
                if it.cite:
                    r = line.add_run(f"    ［{it.cite}］")
                    r.font.color.rgb = RGBColor(0x55, 0x5f, 0x66)
            elif it.kind == "note":
                q = doc.add_paragraph(it.text)
                q.runs[0].italic = True
        checks = [i for i in self.items if i.kind == "check"]
        if checks:
            doc.add_heading(H[4], level=2)
            for it in checks:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(it.text + (f"：{it.formula}" if it.formula and zh else (f": {it.formula}" if it.formula else "")))
                r = p.add_run(f"  〔{H[16] if it.ok else H[17]}〕")
                r.bold = True
                r.font.color.rgb = RGBColor(0x1a, 0x7f, 0x37) if it.ok else RGBColor(0xc0, 0x1c, 0x28)
        doc.add_heading(H[5], level=2)
        for it in self.items:
            if it.kind == "conclude":
                doc.add_paragraph(it.text)
        if self.failed:
            p = doc.add_paragraph()
            r = p.add_run(f"{H[18]}：" + "；".join(self.failed) if zh else f"{H[18]}: " + "; ".join(self.failed))
            r.bold = True
            r.font.color.rgb = RGBColor(0xc0, 0x1c, 0x28)
        if self.citations:
            doc.add_heading(H[6], level=2)
            for c_ in self.citations:
                doc.add_paragraph(c_, style="List Number")
        path = Path(path)
        doc.save(path)
        return path
