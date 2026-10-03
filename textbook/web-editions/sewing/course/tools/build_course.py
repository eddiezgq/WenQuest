#!/usr/bin/env python3
"""Render the course pack of 《缝纫机设计与制造》 (第 16 轮).

    python3 course/tools/build_course.py <out>        # CI: textbook/build/sewing/course

Reads course/lessons/chNN.json (format in ../FORMAT.md), the book's chapter fragments (src/<lang>/chNN.html),
its figures (img/) and labs (src/<lang>/labs/), and writes everything the learning platform publishes:

    <out>/manifest.json          what to publish, lesson by lesson (read by the gateway, course_pack.py)
    <out>/files/chNN/...         slides (.pptx), lab guides, report templates and lesson plans (.docx),
                                 labs (.html, Chinese and English), figures of the lecture notes (.webp)

Every text in the manifest is a [zh, en] pair; the gateway turns pairs into Moodle multilang markup.
"""
from __future__ import annotations

import hashlib
import html
import io
import json
import logging
import re
import shutil
import sys
import warnings
from pathlib import Path

HERE = Path(__file__).resolve().parent
COURSE = HERE.parent
BOOK = COURSE.parent
sys.path.insert(0, str(BOOK))
from book import CH, PARTS  # noqa: E402

BATCHES = [  # Eddie 第 16 轮 Q1：一门课 31 讲，分三批上线
    {"no": 1, "lessons": list(range(1, 9)), "exams": []},
    {"no": 2, "lessons": list(range(9, 23)), "exams": ["midterm"]},
    {"no": 3, "lessons": list(range(23, 32)), "exams": ["final"]},
]
MIDTERM_UNITS = 4          # 期中：第 1–4 单元（第 1–12 讲）
EXAM_SECTION = 33          # 考试放在最后一节（第 1 节课程说明，第 2–32 节为第 1–31 讲）
E = html.escape


def esc(s: str) -> str:
    return E(s, quote=False)


def pair(zh: str, en: str) -> list[str]:
    return [zh, en]


# --- the book ------------------------------------------------------------------------------------------

def unit_of(no: int) -> tuple[int, tuple]:
    for i, p in enumerate(PARTS):
        if no in p[3]:
            return i + 1, p
    raise KeyError(no)


def unit_name(no: int) -> list[str]:
    u, p = unit_of(no)
    return pair(f"{p[0]} {p[1]}", p[2])


RES = {lang: {Path(r["href"]).stem: r for r in json.loads((BOOK / f"res_{lang}.json").read_text(encoding="utf-8"))}
       for lang in ("zh", "en")}


def figure_file(lang: str, png_name: str) -> Path | None:
    stem = Path(png_name).stem
    if lang == "en" and (BOOK / "img" / "en" / f"{stem}.webp").exists():
        return BOOK / "img" / "en" / f"{stem}.webp"
    p = BOOK / "img" / f"{stem}.webp"
    return p if p.exists() else None


def figure_upload_name(lang: str, png_name: str) -> str:
    p = figure_file(lang, png_name)
    stem = Path(png_name).stem
    return f"en-{stem}.webp" if p and p.parent.name == "en" else f"{stem}.webp"


_DETAILS = re.compile(r"<details[^>]*>\s*<summary>(.*?)</summary>(.*?)</details>", re.S)


def split_fragment(lang: str, no: int) -> tuple[str, str, str]:
    """(body, exercises <ol>…</ol>, tail sections such as references) of a chapter fragment."""
    s = (BOOK / "src" / lang / f"ch{no:02d}.html").read_text(encoding="utf-8")
    h3 = [(m.start(), m.end(), re.sub(r"<[^>]+>", "", m.group(1))) for m in re.finditer(r"<h3[^>]*>(.*?)</h3>", s, re.S)]
    ex = [x for x in h3 if "习题" in x[2] or "xercise" in x[2]]
    if len(ex) != 1:
        raise SystemExit(f"ch{no:02d} {lang}: expected one exercises heading, found {len(ex)}")
    ex_start, ex_end = ex[0][0], ex[0][1]
    later = [x for x in h3 if x[0] > ex_start]
    ex_html = s[ex_end:later[0][0]] if later else s[ex_end:]
    tail = ""
    for i, x in enumerate(later):
        if "配套资源" in x[2] or "ompanion" in x[2]:
            continue
        end = later[i + 1][0] if i + 1 < len(later) else len(s)
        tail += s[x[0]:end]
    return s[:ex_start], ex_html, tail


def clean_html(lang: str, frag: str, labs_here: dict[str, str]) -> tuple[str, list[str]]:
    """The book's HTML made fit for a Moodle page: no lab widgets or buttons, figures as <div class="wq-fig">
    with their file referenced as @@PLUGINFILE@@/<name>, links inside the book reduced to text.
    Returns the HTML and the figure files (book-relative png names) it uses."""
    s = frag
    s = re.sub(r'<p class="eyebrow">.*?</p>', "", s, flags=re.S)
    s = re.sub(r"<h1>.*?</h1>", "", s, flags=re.S)
    s = re.sub(r'<div class="seam"[^>]*></div>', "", s)
    s = re.sub(r'<div class="chips">.*?</div>', "", s, flags=re.S)

    def lab_box(m: re.Match) -> str:
        sec = m.group(0)
        title = re.sub(r"<[^>]+>", "", (re.search(r"<h4[^>]*>(.*?)</h4>", sec, re.S) or [None, ""])[1]).strip()
        desc = re.search(r"</div><p>(.*?)</p>", sec, re.S)
        where = labs_here.get(m.group(1), "")
        note = where or ("（在它所属那一讲的活动里打开）" if lang == "zh" else "(open it in the lesson it belongs to)")
        return (f"<blockquote><p><strong>▶ {esc(title)}</strong></p>"
                f"{'<p>' + desc.group(1) + '</p>' if desc else ''}<p><em>{note}</em></p></blockquote>")

    s = re.sub(r'<section class="lab" id="lab-([a-z0-9-]+)"[^>]*>.*?</section>', lab_box, s, flags=re.S)
    s = re.sub(r"<button[^>]*>.*?</button>", "", s, flags=re.S)
    s = _DETAILS.sub(lambda m: f"<p><strong>{m.group(1)}</strong></p>{m.group(2)}", s)
    figs: list[str] = []

    def fig(m: re.Match) -> str:
        block = m.group(0)
        img = re.search(r'<img src="img/([^"]+)"(?:[^>]*alt="([^"]*)")?', block)
        cap = re.search(r"<figcaption>(.*?)</figcaption>", block, re.S)
        if not img:
            return block
        name = img.group(1)
        if not figure_file(lang, name):
            raise SystemExit(f"{lang}: figure {name} missing")
        figs.append(name)
        up = figure_upload_name(lang, name)
        return (f'<div class="wq-fig"><img src="@@PLUGINFILE@@/{up}" alt="{img.group(2) or ""}">'
                + (f'<p class="wq-cap">{cap.group(1)}</p>' if cap else "") + "</div>")

    s = re.sub(r"<figure[^>]*>.*?</figure>", fig, s, flags=re.S)
    if re.search(r'src="img/', s):
        raise SystemExit(f"{lang}: an image outside a figure")
    # links: keep the web, drop the ones inside the book (anchors, chapters, labs)
    s = re.sub(r'<a\b[^>]*href="(?!https?://)[^"]*"[^>]*>(.*?)</a>', r"\1", s, flags=re.S)
    s = re.sub(r'<a\b[^>]*class="zoom"[^>]*>(.*?)</a>', r"\1", s, flags=re.S)
    s = re.sub(r"<(/?)(aside|section)\b[^>]*>", lambda m: f"<{m.group(1)}div>", s)
    s = re.sub(r'\s(?:loading|data-src|aria-hidden|aria-label)="[^"]*"', "", s)
    return s.strip(), figs


def exercises(lang: str, ex_html: str) -> list[tuple[str, str]]:
    """[(question html, answer html or "")] from the exercises list."""
    ol = re.search(r"<ol>(.*)</ol>", ex_html, re.S)
    items = re.findall(r"<li>(.*?)</li>", ol.group(1), re.S) if ol else []
    out = []
    for it in items:
        ans = re.search(r'<details class="ans">\s*<summary>.*?</summary>(.*?)</details>', it, re.S)
        q = re.sub(r"<details.*?</details>", "", it, flags=re.S).strip()
        q = re.sub(r'<a\b[^>]*href="(?!https?://)[^"]*"[^>]*>(.*?)</a>', r"\1", q, flags=re.S)
        out.append((q, ans.group(1).strip() if ans else ""))
    return out


# --- HTML pieces -------------------------------------------------------------------------------------

L = {
    "zh": dict(lesson=lambda n: f"第 {n} 讲", hours=lambda h: f"建议 {h} 学时", goals="学习目标", problem="车间问题",
               question="问题", robot="机器人联系", life="生活中的例子", route="本讲路线",
               route_text="车间问题 → 概念（下面的讲义正文）→ 动画与虚拟实验 → 建模求解（回到车间问题）。讲义正文即教材第 {n} 章。",
               back="回到车间问题：解答", summary="本讲小结", labs="本讲的动画与虚拟实验",
               labs_note="都在本讲的活动里，中文版、英文版各一个，点开即可在浏览器里操作。", ex_intro="教材第 {n} 章习题。带“计算”的题请写出过程；参考答案在测验之后由老师公布。",
               ans_none="开放题：按讲义要点作答，老师批阅。", ans="参考答案", here="（在本讲的“{kind}”里打开）",
               tasks="作业题", deliverable="提交要求", rubric="评分要点", quiz_intro="本讲测验，{k} 题，自动评分，可做 3 次，取最高分。",
               unit=lambda u: f"第 {u} 单元"),
    "en": dict(lesson=lambda n: f"Lesson {n}", hours=lambda h: f"{h} class hours", goals="Learning goals", problem="Shop-floor problem",
               question="Question", robot="Robot link", life="An everyday example", route="Route of this lesson",
               route_text="Shop-floor problem → concepts (the notes below) → animations and virtual labs → modelling and solving (back to the problem). The notes are chapter {n} of the textbook.",
               back="Back to the shop-floor problem: the solution", summary="Summary", labs="Animations and virtual labs of this lesson",
               labs_note="They are activities of this lesson, each in a Chinese and an English version; open one to run it in your browser.",
               ex_intro="Exercises of chapter {n} of the textbook. Show your working for calculations; the answer key is released by the teacher.",
               ans_none="Open question: answer from the notes; the teacher marks it.", ans="Answer", here="(open it under “{kind}” in this lesson)",
               tasks="Tasks", deliverable="What to hand in", rubric="Marking", quiz_intro="Lesson quiz: {k} questions, marked automatically; 3 attempts, the best one counts.",
               unit=lambda u: f"Unit {u}"),
}


def lab_kind(lang: str, lab_id: str) -> str:
    return RES[lang][lab_id]["kind"]


def lesson_intro(lang: str, d: dict) -> str:
    i = 0 if lang == "zh" else 1
    t = L[lang]
    no = d["no"]
    u, p = unit_of(no)
    unit = (p[0] + " " + p[1]) if lang == "zh" else p[2]
    out = [f"<p><em>{esc(t['hours'](d['hours']))} · {esc(unit)}</em></p>",
           f"<h3>{t['goals']}</h3><ol>" + "".join(f"<li>{esc(g[i])}</li>" for g in d["goals"]) + "</ol>",
           f"<h3>{t['problem']}：{esc(d['problem']['title'][i])}</h3>" if lang == "zh"
           else f"<h3>{t['problem']}: {esc(d['problem']['title'][i])}</h3>",
           f"<p>{esc(d['problem']['text'][i])}</p>",
           f"<p><strong>{t['question']}{'：' if lang == 'zh' else ': '}</strong>{esc(d['problem']['question'][i])}</p>",
           f"<h3>{t['robot']}{'：' if lang == 'zh' else ': '}{esc(d['robot']['title'][i])}</h3><p>{esc(d['robot']['text'][i])}</p>",
           f"<h3>{t['life']}{'：' if lang == 'zh' else ': '}{esc(d['life']['title'][i])}</h3><p>{esc(d['life']['text'][i])}</p>",
           f"<h3>{t['route']}</h3><p>{esc(t['route_text'].format(n=no))}</p><hr>"]
    return "".join(out)


def lesson_outro(lang: str, d: dict) -> str:
    i = 0 if lang == "zh" else 1
    t = L[lang]
    labs = "".join(f"<li><strong>{esc(lab_kind(lang, lb['id']))} {esc(lb['no'])}</strong> {esc(lb['title'][i])}</li>"
                   for lb in d["labs"])
    return (f"<hr><h3>{t['back']}</h3><p>{esc(d['problem']['answer'][i])}</p>"
            f"<h3>{t['summary']}</h3><ul>" + "".join(f"<li>{esc(x[i])}</li>" for x in d["summary"]) + "</ul>"
            f"<h3>{t['labs']}</h3><ul>{labs}</ul><p>{t['labs_note']}</p>")


def q_html(text: str) -> str:
    return "".join(f"<p>{esc(p)}</p>" for p in text.split("\n") if p.strip())


def question(q: dict, mark: float) -> dict:
    out = {"type": q["type"], "text": [q_html(q["text"][0]), q_html(q["text"][1])],
           "feedback": [q_html(q["feedback"][0]), q_html(q["feedback"][1])], "mark": mark, "answers": [], "correct": True}
    if q["type"] in ("single", "multiple"):
        out["answers"] = [{"text": [esc(a["text"][0]), esc(a["text"][1])], "fraction": 1 if a["correct"] else 0, "tolerance": 0}
                          for a in q["answers"]]
    elif q["type"] == "truefalse":
        out["correct"] = bool(q["correct"])
    elif q["type"] == "numerical":
        unit = (q.get("unit") or "").strip()
        if unit:
            out["text"][0] += f"<p><em>（只填数字，单位：{esc(unit)}）</em></p>"
            out["text"][1] += f"<p><em>(Enter the number only, in {esc(unit)}.)</em></p>"
        out["answers"] = [{"text": [repr(float(q["answer"])).rstrip("0").rstrip(".")] * 2, "fraction": 1,
                           "tolerance": float(q["tolerance"])}]
    return out


# --- files: slides, documents, labs ---------------------------------------------------------------------

def render_formula(tex: str, dest: Path) -> bool:
    """LaTeX -> PNG with matplotlib's mathtext (a subset of LaTeX). False if it cannot be drawn."""
    logging.getLogger("matplotlib").setLevel(logging.ERROR)
    warnings.filterwarnings("ignore")
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    cjk = next((f.name for f in font_manager.fontManager.ttflist
                if any(k in f.name for k in ("CJK SC", "CJK JP", "WenQuanYi", "Noto Sans CJK"))), None)
    if not cjk:
        raise SystemExit("no Chinese font for the formulas (CI: apt-get install fonts-noto-cjk)")
    plt.rcParams.update({"mathtext.fontset": "custom", "mathtext.rm": cjk, "mathtext.it": "DejaVu Serif:italic",
                         "mathtext.bf": f"{cjk}:bold", "mathtext.sf": cjk})
    f = re.sub(r"\\[bB]igg?[lr]?(?![a-zA-Z])", "", tex)
    f = (f.replace(r"\tfrac", r"\frac").replace(r"\iff", r"\Leftrightarrow").replace(r"\Longrightarrow", r"\Rightarrow")
         .replace(r"\le ", r"\leq ").replace(r"\ge ", r"\geq "))
    f = re.sub(r"\\le(?![a-zA-Z])", r"\\leq", f)
    f = re.sub(r"\\ge(?![a-zA-Z])", r"\\geq", f)
    f = re.sub(r"\\rm\s*([A-Za-z]+)", r"\\mathrm{\1}", f)
    f = re.sub(r"\\sqrt\s+([A-Za-z0-9])", r"\\sqrt{\1}", f)
    cases = re.search(r"\\begin\{cases\}(.*?)\\end\{cases\}", f, re.S)
    if cases:
        rows = [r.replace("&", r",\ ") for r in cases.group(1).split(r"\\")]
        f = f[:cases.start()] + r"\left\{" + r";\quad ".join(x.strip() for x in rows) + r"\right." + f[cases.end():]
    fig = plt.figure(figsize=(0.01, 0.01))
    try:
        fig.text(0, 0, f"${f}$", fontsize=24, color="#1f2933")
        fig.savefig(dest, dpi=200, transparent=True, bbox_inches="tight", pad_inches=0.05)
        return True
    except Exception:  # noqa: BLE001 - shown as text instead
        return False
    finally:
        plt.close(fig)


def png_of(src: Path, dest: Path, width: int = 1600) -> Path:
    from PIL import Image
    im = Image.open(src)
    if im.width > width:
        im = im.resize((width, round(im.height * width / im.width)), Image.LANCZOS)
    if im.mode in ("RGBA", "LA", "P"):
        bg = Image.new("RGB", im.size, "white")
        im = im.convert("RGBA")
        bg.paste(im, mask=im.split()[-1])
        im = bg
    im.convert("RGB").save(dest, "JPEG", quality=88)
    return dest


NAVY, ORANGE, INK, GREY, LIGHT = "1F3A5F", "D9822B", "1F2933", "5F6B7A", "EEF2F6"
ZH_FONT, EN_FONT = "Microsoft YaHei", "Calibri"


class Deck:
    def __init__(self) -> None:
        from pptx import Presentation
        from pptx.util import Inches
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.blank = self.prs.slide_layouts[6]

    @staticmethod
    def _run(p, text: str, size: float, color: str = INK, bold: bool = False, italic: bool = False):
        from pptx.dml.color import RGBColor
        from pptx.oxml.ns import qn
        from pptx.util import Pt
        r = p.add_run()
        r.text = text
        f = r.font
        f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, EN_FONT
        f.color.rgb = RGBColor.from_string(color)
        rpr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = rpr.makeelement(qn(tag), {})
                rpr.append(el)
            el.set("typeface", ZH_FONT)
        return r

    def _box(self, slide, x, y, w, h):
        from pptx.util import Inches
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        from pptx.util import Pt
        tf.margin_left = tf.margin_right = Pt(4)
        return tf

    def _rect(self, slide, x, y, w, h, color):
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(color)
        s.line.fill.background()
        return s

    def slide(self, title: list[str], notes: list[str] | None = None, footer: str = ""):
        s = self.prs.slides.add_slide(self.blank)
        self._rect(s, 0, 0, 13.333, 1.15, NAVY)
        self._rect(s, 0, 1.15, 13.333, 0.06, ORANGE)
        tf = self._box(s, 0.5, 0.12, 12.3, 0.95)
        p = tf.paragraphs[0]
        self._run(p, title[0], 26, "FFFFFF", bold=True)
        p2 = tf.add_paragraph()
        self._run(p2, title[1], 15, "C9D6E3")
        if footer:
            ft = self._box(s, 0.5, 7.05, 12.3, 0.35)
            self._run(ft.paragraphs[0], footer, 10, GREY)
        if notes:
            s.notes_slide.notes_text_frame.text = "\n\n".join(x for x in notes if x)
        return s

    def bullets(self, s, items: list[list[str]], x: float, y: float, w: float, h: float, numbered: bool = False):
        from pptx.util import Pt
        chars = sum(len(a) + len(b) * 0.5 for a, b in items) * (12.3 / w)
        zs, es = (22, 15) if chars < 260 else (20, 14) if chars < 380 else (18, 13) if chars < 520 else (16, 12) if chars < 700 else (14, 11)
        tf = self._box(s, x, y, w, h)
        first = True
        for k, (zh, en) in enumerate(items):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_before = Pt(10 if k else 0)
            self._run(p, (f"{k + 1}. " if numbered else "■ "), zs * 0.8 if not numbered else zs, ORANGE, bold=True)
            self._run(p, zh, zs, INK)
            if en:
                q = tf.add_paragraph()
                self._run(q, ("     " if numbered else "     ") + en, es, GREY, italic=True)
        return tf

    def picture(self, s, path: Path, x: float, y: float, w: float, h: float):
        from PIL import Image
        from pptx.util import Inches
        with Image.open(path) as im:
            iw, ih = im.size
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        s.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))

    def save(self, path: Path) -> None:
        self.prs.save(str(path))


def build_deck(d: dict, out: Path, tmp: Path) -> None:
    from pptx.util import Pt
    no = d["no"]
    foot = f"《缝纫机设计与制造》 Sewing Machine Design and Manufacturing · 第 {no} 讲 Lesson {no} · 问渠机器人学院 WenQuest Robotics Academy"
    dk = Deck()
    # title
    s = dk.prs.slides.add_slide(dk.blank)
    dk._rect(s, 0, 0, 13.333, 7.5, NAVY)
    dk._rect(s, 0.8, 4.55, 4.0, 0.07, ORANGE)
    tf = dk._box(s, 0.8, 1.6, 11.8, 3.0)
    dk._run(tf.paragraphs[0], f"第 {no} 讲 · Lesson {no}", 20, "F2B880", bold=True)
    p = tf.add_paragraph(); p.space_before = Pt(10)
    dk._run(p, d["title"][0], 44, "FFFFFF", bold=True)
    p = tf.add_paragraph()
    dk._run(p, d["title"][1], 26, "C9D6E3")
    tf = dk._box(s, 0.8, 4.85, 11.8, 1.6)
    u = unit_name(no)
    dk._run(tf.paragraphs[0], f"{u[0]}  ·  {u[1]}", 16, "C9D6E3")
    p = tf.add_paragraph(); p.space_before = Pt(8)
    dk._run(p, "《缝纫机设计与制造》 Sewing Machine Design and Manufacturing · 问渠机器人学院 WenQuest Robotics Academy", 14, "9FB3C8")
    s.notes_slide.notes_text_frame.text = f"第 {no} 讲 {d['title'][0]}\n\nLesson {no}: {d['title'][1]}"
    # goals
    s = dk.slide(pair("学习目标", "Learning goals"), footer=foot)
    dk.bullets(s, d["goals"], 0.6, 1.5, 12.1, 5.4, numbered=True)
    # problem
    pr = d["problem"]
    s = dk.slide(pair("车间问题：" + pr["title"][0], "Shop-floor problem: " + pr["title"][1]), footer=foot,
                 notes=[pr["text"][0] + "\n" + pr["question"][0], pr["text"][1] + "\n" + pr["question"][1]])
    dk.bullets(s, [pr["text"], [("问题：" + pr["question"][0]), ("Question: " + pr["question"][1])]], 0.6, 1.5, 12.1, 5.4)
    # robot & everyday life
    s = dk.slide(pair("机器人联系与生活中的例子", "Robot link and an everyday example"), footer=foot,
                 notes=[d["robot"]["text"][0] + "\n" + d["life"]["text"][0], d["robot"]["text"][1] + "\n" + d["life"]["text"][1]])
    dk.bullets(s, [["机器人联系：" + d["robot"]["title"][0] + "。" + d["robot"]["text"][0],
                    "Robot link: " + d["robot"]["title"][1] + ". " + d["robot"]["text"][1]],
                   ["生活中的例子：" + d["life"]["title"][0] + "。" + d["life"]["text"][0],
                    "Everyday example: " + d["life"]["title"][1] + ". " + d["life"]["text"][1]]], 0.6, 1.5, 12.1, 5.4)
    # content
    for k, sl in enumerate(d["slides"]):
        s = dk.slide(sl["title"], notes=sl["narration"], footer=foot)
        fig = figure_file("zh", sl["figure"]) if sl.get("figure") else None
        fpng = None
        if sl.get("formula"):
            fpng = tmp / f"f{no}-{k}.png"
            if not render_formula(sl["formula"], fpng):
                fpng = None
        w = 6.9 if fig else 12.1
        th = 4.25 if (fpng or (sl.get("formula") and not fpng)) else 5.4
        dk.bullets(s, sl["points"], 0.6, 1.45, w, th)
        if sl.get("formula"):
            if fpng:
                dk.picture(s, fpng, 0.6, 5.75, w, 1.2)
            else:
                tf = dk._box(s, 0.6, 5.8, w, 1.1)
                dk._run(tf.paragraphs[0], sl["formula"], 14, NAVY)
        if fig:
            jpg = png_of(fig, tmp / f"s{no}-{k}.jpg")
            dk.picture(s, jpg, 7.75, 1.5, 5.2, 5.4)
    # labs
    s = dk.slide(pair("动画与虚拟实验", "Animations and virtual labs"), footer=foot)
    dk.bullets(s, [[f"{lab_kind('zh', lb['id'])} {lb['no']}　{lb['title'][0]}：{RES['zh'][lb['id']]['desc']}",
                    f"{lab_kind('en', lb['id'])} {lb['no']}  {lb['title'][1]}: {RES['en'][lb['id']]['desc']}"] for lb in d["labs"]],
               0.6, 1.5, 12.1, 5.4)
    # solution
    s = dk.slide(pair("回到车间问题：解答", "Back to the shop-floor problem: the solution"), footer=foot,
                 notes=pr["answer"])
    dk.bullets(s, [pr["answer"]], 0.6, 1.5, 12.1, 5.4)
    # summary & after class
    s = dk.slide(pair("本讲小结", "Summary"), footer=foot, notes=["\n".join(x[0] for x in d["summary"]), "\n".join(x[1] for x in d["summary"])])
    dk.bullets(s, d["summary"], 0.6, 1.5, 12.1, 5.4, numbered=True)
    s = dk.slide(pair("课后", "After class"), footer=foot)
    dk.bullets(s, [["作业：" + d["assignment"]["title"][0], "Assignment: " + d["assignment"]["title"][1]],
                   [f"测验：本讲测验 {len(d['quiz'])} 题，自动评分", f"Quiz: {len(d['quiz'])} questions, marked automatically"],
                   ["练习：教材本章习题", "Practice: the exercises of this chapter of the textbook"],
                   ["虚拟实验：按实验指导书完成并提交实验报告", "Virtual labs: follow the lab guide and hand in the report"]],
               0.6, 1.5, 12.1, 5.4)
    dk.save(out)


class Doc:
    def __init__(self) -> None:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = "Calibri", Pt(10.5)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
        for name in ("Title", "Heading 1", "Heading 2"):
            h = self.doc.styles[name]
            h.font.name = "Calibri"
            h.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")

    def _font(self, run, size=None, bold=None, italic=None, color=None):
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体" if not bold else "黑体")
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic
        if color:
            run.font.color.rgb = RGBColor.from_string(color)

    def title(self, zh: str, en: str, sub: str = "") -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._font(p.add_run(zh), 18, True, color=NAVY)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._font(p.add_run(en), 12, False, True, GREY)
        if sub:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._font(p.add_run(sub), 9.5, False, False, GREY)

    def heading(self, zh: str, en: str) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = __import__("docx.shared", fromlist=["Pt"]).Pt(10)
        self._font(p.add_run(zh), 13, True, color=NAVY)
        self._font(p.add_run("  " + en), 11, False, True, GREY)

    def bi(self, items: list[list[str]], numbered: bool = True) -> None:
        for k, (zh, en) in enumerate(items):
            p = self.doc.add_paragraph()
            self._font(p.add_run(f"{k + 1}. " if numbered else "• "), bold=True)
            self._font(p.add_run(zh))
            if en:
                q = self.doc.add_paragraph()
                q.paragraph_format.left_indent = __import__("docx.shared", fromlist=["Cm"]).Cm(0.6)
                self._font(q.add_run(en), 9.5, italic=True, color=GREY)

    def text(self, zh: str, en: str = "") -> None:
        p = self.doc.add_paragraph()
        self._font(p.add_run(zh))
        if en:
            q = self.doc.add_paragraph()
            self._font(q.add_run(en), 9.5, italic=True, color=GREY)

    def lines(self, n: int) -> None:
        for _ in range(n):
            p = self.doc.add_paragraph("_" * 86)
            self._font(p.runs[0], 10, color="B0B8C1")

    def table(self, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
        from docx.shared import Cm
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = "Table Grid"
        for j, h in enumerate(headers):
            c = t.rows[0].cells[j]
            c.text = ""
            self._font(c.paragraphs[0].add_run(h), 9.5, True)
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                c = t.rows[i + 1].cells[j]
                c.text = ""
                self._font(c.paragraphs[0].add_run(v), 9.5)
        if widths:
            for row in t.rows:
                for j, w in enumerate(widths):
                    row.cells[j].width = Cm(w)

    def save(self, path: Path) -> None:
        self.doc.save(str(path))


def build_guide(d: dict, lb: dict, out: Path) -> None:
    no = lb["no"]
    k = (lab_kind("zh", lb["id"]), lab_kind("en", lb["id"]))
    doc = Doc()
    doc.title(f"{k[0]} {no}　{lb['title'][0]} · 指导书", f"{k[1]} {no}: {lb['title'][1]} · Lab guide",
              f"《缝纫机设计与制造》第 {d['no']} 讲 {d['title'][0]} · Lesson {d['no']}: {d['title'][1]}")
    doc.heading("一、实验目的", "Objectives"); doc.bi(lb["goal"])
    doc.heading("二、实验原理", "Theory"); doc.bi(lb["theory"])
    doc.heading("三、实验环境", "Setup")
    doc.text(f"在问渠学习平台本讲的活动里打开“{k[0]} {no}”（中文版或英文版），用电脑浏览器操作；也可在“教材”栏目的网页版教材第 {d['no']} 章里运行。",
             f"Open “{k[1]} {no}” (Chinese or English version) among the activities of this lesson on WenQuest, in a desktop browser; "
             f"it also runs inside chapter {d['no']} of the web edition of the textbook.")
    doc.heading("四、实验步骤", "Procedure"); doc.bi(lb["steps"])
    doc.heading("五、数据记录", "Data record")
    rec = lb["record"]
    doc.text(rec["caption"][0], rec["caption"][1])
    doc.table([f"{h[0]}\n{h[1]}" for h in rec["headers"]], [[""] * len(rec["headers"]) for _ in range(int(rec["rows"]))])
    doc.heading("六、思考题", "Questions"); doc.bi(lb["questions"])
    doc.heading("七、实验报告", "Report")
    doc.text("按“实验报告模板”填写，连同截图一起提交到本讲作业。", "Fill in the report template and hand it in, with screenshots, to this lesson's assignment.")
    doc.save(out)


def build_report(d: dict, lb: dict, out: Path) -> None:
    no = lb["no"]
    k = (lab_kind("zh", lb["id"]), lab_kind("en", lb["id"]))
    doc = Doc()
    doc.title(f"{k[0]} {no}　{lb['title'][0]} · 实验报告", f"{k[1]} {no}: {lb['title'][1]} · Lab report",
              f"《缝纫机设计与制造》第 {d['no']} 讲 · Lesson {d['no']}")
    doc.table(["姓名 Name", "", "学号 Student no.", ""], [["班级 Class", "", "日期 Date", ""]], [3.2, 5, 3.2, 5])
    doc.heading("一、实验目的", "Objectives"); doc.bi(lb["goal"])
    doc.heading("二、实验原理（用自己的话简述）", "Theory (in your own words)"); doc.lines(5)
    doc.heading("三、实验步骤与现象（附截图）", "Procedure and observations (with screenshots)"); doc.lines(7)
    doc.heading("四、数据记录", "Data record")
    rec = lb["record"]
    doc.text(rec["caption"][0], rec["caption"][1])
    doc.table([f"{h[0]}\n{h[1]}" for h in rec["headers"]], [[""] * len(rec["headers"]) for _ in range(int(rec["rows"]))])
    doc.heading("五、数据处理与结果", "Analysis and results"); doc.lines(6)
    doc.heading("六、思考题解答", "Answers to the questions")
    for j, q in enumerate(lb["questions"]):
        doc.text(f"{j + 1}. {q[0]}", q[1]); doc.lines(3)
    doc.heading("七、结论与体会", "Conclusions"); doc.lines(4)
    doc.save(out)


def build_plan(d: dict, out: Path) -> None:
    pl = d["plan"]
    doc = Doc()
    doc.title(f"第 {d['no']} 讲　{d['title'][0]} · 教案", f"Lesson {d['no']}: {d['title'][1]} · Lesson plan",
              f"{unit_name(d['no'])[0]} · {unit_name(d['no'])[1]} · 建议 {d['hours']} 学时 / {d['hours']} class hours")
    doc.heading("一、教学目标", "Goals"); doc.bi(d["goals"])
    doc.heading("二、重点", "Key points"); doc.text(*pl["key"])
    doc.heading("三、难点", "Difficult points"); doc.text(*pl["difficult"])
    doc.heading("四、导入：车间问题", "Opening: the shop-floor problem")
    doc.text(d["problem"]["title"][0], d["problem"]["title"][1])
    doc.text(d["problem"]["text"][0] + " " + d["problem"]["question"][0], d["problem"]["text"][1] + " " + d["problem"]["question"][1])
    doc.text("解答：" + d["problem"]["answer"][0], "Solution: " + d["problem"]["answer"][1])
    doc.heading("五、机器人联系与生活例子", "Robot link and everyday example")
    doc.text(d["robot"]["title"][0] + "：" + d["robot"]["text"][0], d["robot"]["title"][1] + ": " + d["robot"]["text"][1])
    doc.text(d["life"]["title"][0] + "：" + d["life"]["text"][0], d["life"]["title"][1] + ": " + d["life"]["text"][1])
    doc.heading("六、教学过程", "Teaching process")
    doc.table(["环节 Phase", "分钟 Min", "内容 Content"],
              [[f"{x['phase'][0]}\n{x['phase'][1]}", str(x["minutes"]), f"{x['content'][0]}\n{x['content'][1]}"] for x in pl["process"]],
              [3.6, 1.6, 11.4])
    doc.heading("七、动画与虚拟实验", "Animations and virtual labs")
    doc.bi([[f"{lab_kind('zh', lb['id'])} {lb['no']} {lb['title'][0]}", f"{lab_kind('en', lb['id'])} {lb['no']} {lb['title'][1]}"]
            for lb in d["labs"]], numbered=False)
    doc.heading("八、作业", "Homework"); doc.text(*pl["homework"])
    doc.heading("九、小结", "Summary"); doc.bi(d["summary"])
    doc.save(out)


def lab_copy(lang: str, lab_id: str, dest: Path) -> None:
    s = (BOOK / "src" / lang / "labs" / f"{lab_id}.html").read_text(encoding="utf-8")
    s = re.sub(r'<a href="\.\./index\.html"[^>]*>[^<]*</a>', "", s)
    if "../" in s:
        raise SystemExit(f"lab {lang}/{lab_id} still points outside itself")
    dest.write_text(s, encoding="utf-8")


# --- the manifest -------------------------------------------------------------------------------------

def lesson_entry(d: dict, out: Path, tmp: Path) -> dict:
    no = d["no"]
    rel = f"files/ch{no:02d}"
    fdir = out / rel
    fdir.mkdir(parents=True, exist_ok=True)
    title = d["title"]
    acts: list[dict] = []

    # 讲义 (lecture notes): the lesson's opening + chapter body + the solution and summary
    labs_here = {lb["id"]: L["zh"]["here"].format(kind=lab_kind("zh", lb["id"])) for lb in d["labs"]}
    labs_here_en = {lb["id"]: L["en"]["here"].format(kind=lab_kind("en", lb["id"])) for lb in d["labs"]}
    pages, files, ex = [], [], {}
    for lang, here in (("zh", labs_here), ("en", labs_here_en)):
        body, ex_html, tail = split_fragment(lang, no)
        b, figs = clean_html(lang, body, here)
        t, figs2 = clean_html(lang, tail, here)
        pages.append(lesson_intro(lang, d) + b + t + lesson_outro(lang, d))
        for name in figs + figs2:
            up = figure_upload_name(lang, name)
            if not (fdir / up).exists():
                shutil.copy(figure_file(lang, name), fdir / up)
            if f"{rel}/{up}" not in files:
                files.append(f"{rel}/{up}")
        ex[lang] = exercises(lang, ex_html)
    acts.append({"type": "page", "name": pair(f"讲义：第 {no} 讲 {title[0]}", f"Notes: Lesson {no} {title[1]}"),
                 "content": pages, "files": files, "visible": 1})

    # 动画 first (the benchmark order: notes, video, animation, slides ...)
    def lab_acts(kinds: set[str]) -> list[dict]:
        res = []
        for lb in d["labs"]:
            kz, ke = lab_kind("zh", lb["id"]), lab_kind("en", lb["id"])
            if (kz == "动画") != ("动画" in kinds):
                continue
            for lang, lab_name in (("zh", pair(f"{kz} {lb['no']} {lb['title'][0]}（中文版）", f"{ke} {lb['no']} {lb['title'][1]} (Chinese)")),
                                   ("en", pair(f"{kz} {lb['no']} {lb['title'][0]}（英文版）", f"{ke} {lb['no']} {lb['title'][1]} (English)"))):
                fname = f"{kz}{lb['no']}-{lb['title'][0]}（中文版）.html" if lang == "zh" else \
                    f"{ke.replace(' ', '-')}-{lb['no']}-{re.sub(r'[^A-Za-z0-9]+', '-', lb['title'][1]).strip('-')}-English.html"
                lab_copy(lang, lb["id"], fdir / fname)
                res.append({"type": "resource", "name": lab_name, "file": f"{rel}/{fname}", "visible": 1})
        return res

    acts += lab_acts({"动画"})
    deck = f"第{no}讲-{title[0]}-课件（中英对照）.pptx"
    build_deck(d, fdir / deck, tmp)
    acts.append({"type": "resource", "name": pair(f"课件：第 {no} 讲 {title[0]}（中英对照）", f"Slides: Lesson {no} {title[1]} (bilingual)"),
                 "file": f"{rel}/{deck}", "visible": 1})

    # 练习 and its answer key (teachers only)
    def ex_page(lang: str, answers: bool) -> str:
        t = L[lang]
        items = []
        for q, a in ex[lang]:
            body = q
            if answers:
                body += f"<p><strong>{t['ans']}{'：' if lang == 'zh' else ': '}</strong>{a or t['ans_none']}</p>"
            items.append(f"<li>{body}</li>")
        return f"<p>{t['ex_intro'].format(n=no)}</p><ol>{''.join(items)}</ol>"

    if ex["zh"]:
        acts.append({"type": "page", "name": pair(f"练习：第 {no} 讲 {title[0]}", f"Practice: Lesson {no} {title[1]}"),
                     "content": [ex_page("zh", False), ex_page("en", False)], "files": [], "visible": 1})
    # lab guides and report templates, then the labs themselves
    for lb in d["labs"]:
        kz, ke = lab_kind("zh", lb["id"]), lab_kind("en", lb["id"])
        g = f"{kz}{lb['no']}-{lb['title'][0]}-指导书.docx"
        build_guide(d, lb, fdir / g)
        acts.append({"type": "resource", "name": pair(f"实验指导书：{kz} {lb['no']} {lb['title'][0]}", f"Lab guide: {ke} {lb['no']} {lb['title'][1]}"),
                     "file": f"{rel}/{g}", "visible": 1})
    for lb in d["labs"]:
        kz, ke = lab_kind("zh", lb["id"]), lab_kind("en", lb["id"])
        r = f"{kz}{lb['no']}-{lb['title'][0]}-实验报告模板.docx"
        build_report(d, lb, fdir / r)
        acts.append({"type": "resource", "name": pair(f"实验报告模板：{kz} {lb['no']} {lb['title'][0]}", f"Report template: {ke} {lb['no']} {lb['title'][1]}"),
                     "file": f"{rel}/{r}", "visible": 1})
    acts += lab_acts({"虚拟实验", "三维模型"})

    # 作业 and 测验
    a = d["assignment"]

    def assign_intro(lang: str) -> str:
        i = 0 if lang == "zh" else 1
        t = L[lang]
        return (f"<h4>{t['tasks']}</h4><ol>" + "".join(f"<li>{esc(x[i])}</li>" for x in a["tasks"]) + "</ol>"
                f"<h4>{t['deliverable']}</h4><p>{esc(a['deliverable'][i])}</p>"
                f"<h4>{t['rubric']}</h4><ul>" + "".join(f"<li>{esc(x[i])}</li>" for x in a["rubric"]) + "</ul>")

    acts.append({"type": "assign", "name": pair(f"作业：{a['title'][0]}", f"Assignment: {a['title'][1]}"),
                 "intro": [assign_intro("zh"), assign_intro("en")], "grade": 100, "visible": 1})
    acts.append({"type": "quiz", "name": pair(f"测验：第 {no} 讲 {title[0]}", f"Quiz: Lesson {no} {title[1]}"),
                 "intro": [f"<p>{L['zh']['quiz_intro'].format(k=len(d['quiz']))}</p>", f"<p>{L['en']['quiz_intro'].format(k=len(d['quiz']))}</p>"],
                 "timelimit": 0, "attempts": 3, "grade": 10, "showanswers": "immediately", "visible": 1,
                 "questions": [question(q, 1) for q in d["quiz"]]})

    # teachers only: answers to the practice, lesson plan
    if ex["zh"]:
        acts.append({"type": "page", "name": pair(f"练习参考答案：第 {no} 讲 {title[0]}", f"Answer key: Lesson {no} {title[1]}"),
                     "content": [ex_page("zh", True), ex_page("en", True)], "files": [], "visible": 0})
    plan = f"第{no}讲-{title[0]}-教案.docx"
    build_plan(d, fdir / plan)
    acts.append({"type": "resource", "name": pair(f"教案：第 {no} 讲 {title[0]}", f"Lesson plan: Lesson {no} {title[1]}"),
                 "file": f"{rel}/{plan}", "visible": 0})

    u, p = unit_of(no)
    summary = [f"<p>{esc(p[0])} {esc(p[1])} · 建议 {d['hours']} 学时。{esc(CH[no][2])}。</p>",
               f"<p>{esc(p[2])} · {d['hours']} class hours. {esc(CH[no][3])}.</p>"]
    return {"no": no, "section": no + 1, "name": pair(f"第 {no} 讲 {title[0]}", f"Lesson {no}: {title[1]}"),
            "summary": summary, "activities": acts}


def exam_entry(kind: str, lessons: dict[int, dict]) -> dict:
    """期中：第 1–4 单元每讲 4 题；期末：第 1–4 单元每讲 1 题（期中没用的那题）+ 第 5–8 单元每讲 2 题。"""
    mid = [n for n in sorted(lessons) if unit_of(n)[0] <= MIDTERM_UNITS]
    qs = []
    if kind == "midterm":
        for n in mid:
            qs += lessons[n]["exam"][:4]
        name = pair("期中考试（第 1–4 单元）", "Midterm exam (units 1–4)")
        scope = pair(f"范围：第 1–{max(mid)} 讲", f"Covers lessons 1–{max(mid)}")
    else:
        for n in sorted(lessons):
            qs += lessons[n]["exam"][4:5] if n in mid else lessons[n]["exam"][:2]
        name = pair("期末考试（全课）", "Final exam (the whole course)")
        scope = pair("范围：第 1–31 讲", "Covers lessons 1–31")
    mark = round(100 / len(qs), 4)
    intro = [f"<p>{scope[0]}。共 {len(qs)} 题，限时 120 分钟，只能作答一次，交卷后可看答案与解析。</p>",
             f"<p>{scope[1]}. {len(qs)} questions, 120 minutes, one attempt; answers and explanations are shown after you submit.</p>"]
    return {"id": kind, "type": "quiz", "name": name, "intro": intro, "timelimit": 7200, "attempts": 1, "grade": 100,
            "showanswers": "immediately", "visible": 1, "questions": [question(q, mark) for q in qs]}


def course_info(lessons: dict[int, dict]) -> dict:
    rows_zh, rows_en = [], []
    for i, p in enumerate(PARTS):
        ls_zh = "；".join(f"第 {n} 讲 {CH[n][0]}" for n in p[3])
        ls_en = "; ".join(f"{n} {CH[n][1]}" for n in p[3])
        rows_zh.append(f"<tr><td>第 {i + 1} 单元 {esc(p[1])}</td><td>{esc(ls_zh)}</td></tr>")
        rows_en.append(f"<tr><td>Unit {i + 1}: {esc(p[2].split(' · ')[-1])}</td><td>{esc(ls_en)}</td></tr>")
    hours = sum(d["hours"] for d in lessons.values())
    zh = ("<p>《缝纫机设计与制造》是问渠机器人学院“智能制造装备”方向的课程，以同名教材（问渠学习平台“教材”栏目，网页版）为蓝本，"
          "从一针线迹讲起：平缝、包缝、特种机的机构，电气控制，电控原型机的搭建、自动化测试与迭代，主要零件的制造工艺，"
          "直到数字工厂与服装生产线。</p>"
          f"<p>全课 8 个单元、31 讲，建议共 {hours} 学时。每讲按“车间实际问题 → 概念 → 动画与虚拟实验 → 建模求解”展开，"
          "并配一个机器人联系和一个生活例子。</p>"
          "<h4>每讲包含</h4><ul><li>讲义（教材本章 + 本讲导入、解答与小结）</li><li>动画、虚拟实验或三维模型（中文版、英文版）</li>"
          "<li>中英对照课件（PPT，备注里有讲稿）</li><li>练习（教材习题）</li><li>实验指导书与实验报告模板（Word）</li>"
          "<li>作业 1 份、测验 1 份（自动评分）</li></ul>"
          "<h4>单元与各讲</h4><table><tr><th>单元</th><th>各讲</th></tr>" + "".join(rows_zh) + "</table>"
          "<h4>考核（建议）</h4><ul><li>每讲作业（含实验报告）30%</li><li>每讲测验 20%</li><li>期中考试（第 1–4 单元）20%</li><li>期末考试（全课）30%</li></ul>"
          "<h4>学习建议</h4><p>先读讲义里的车间问题，带着问题读正文；每讲至少动手做一个虚拟实验，把数据填进报告模板；"
          "测验可以做 3 次，错题看解析后再做。讲解视频随后陆续补进各讲。</p>")
    en = ("<p><em>Sewing Machine Design and Manufacturing</em> is a course of WenQuest Robotics Academy (intelligent manufacturing equipment). "
          "It follows the textbook of the same name (web edition, under Textbooks on WenQuest) from a single stitch to the whole factory: "
          "the mechanisms of lockstitch, overlock and special machines, electrical control, building, testing and iterating control prototypes, "
          "the manufacture of the key parts, and the digital factory and garment lines.</p>"
          f"<p>8 units, 31 lessons, about {hours} class hours. Every lesson runs shop-floor problem → concepts → animations and virtual labs → "
          "modelling and solving, with a robot link and an everyday example.</p>"
          "<h4>Every lesson has</h4><ul><li>Notes (the textbook chapter plus the lesson's opening, solution and summary)</li>"
          "<li>Animations, virtual labs or 3D models (Chinese and English)</li><li>Bilingual slides (PPT, with the script in the notes)</li>"
          "<li>Practice (the textbook exercises)</li><li>Lab guides and report templates (Word)</li><li>One assignment and one quiz (marked automatically)</li></ul>"
          "<h4>Units and lessons</h4><table><tr><th>Unit</th><th>Lessons</th></tr>" + "".join(rows_en) + "</table>"
          "<h4>Assessment (suggested)</h4><ul><li>Assignments with lab reports 30%</li><li>Lesson quizzes 20%</li><li>Midterm (units 1–4) 20%</li><li>Final (whole course) 30%</li></ul>"
          "<h4>How to study</h4><p>Read the shop-floor problem first and keep it in mind while you read; run at least one virtual lab per lesson and "
          "fill in the report template; each quiz allows 3 attempts. Lecture videos will be added to the lessons later.</p>")
    return {"name": pair("课程说明", "Course information"),
            "activities": [{"type": "forum", "name": pair("课程讨论区", "Course discussion"),
                            "intro": ["<p>课程问题、学习心得都可以在这里讨论。</p>", "<p>Ask questions and share what you learn here.</p>"]},
                           {"type": "page", "name": pair("课程说明与学习指南", "About this course and how to study"),
                            "content": [zh, en], "files": [], "visible": 1}]}


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    out = Path(sys.argv[1]).resolve()
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    tmp = out / ".tmp"
    tmp.mkdir()
    lessons = {int(p.stem[2:]): json.loads(p.read_text(encoding="utf-8")) for p in sorted((COURSE / "lessons").glob("ch*.json"))}
    if sorted(lessons) != list(range(1, 32)):
        raise SystemExit(f"lessons missing: {sorted(set(range(1, 32)) - set(lessons))}")
    entries = []
    for n in sorted(lessons):
        entries.append(lesson_entry(lessons[n], out, tmp))
        print(f"ch{n:02d}: {len(entries[-1]['activities'])} activities")
    shutil.rmtree(tmp)
    for e in entries:   # Moodle names: at most 255 characters with the multilang markup (about 80)
        for a in [e] + e["activities"]:
            if len(a["name"][0]) + len(a["name"][1]) + 80 > 250:
                raise SystemExit(f"name too long: {a['name']}")
    exams = [exam_entry("midterm", lessons), exam_entry("final", lessons)]
    src_hash = hashlib.sha256()
    for p in sorted((COURSE / "lessons").glob("*.json")) + sorted((BOOK / "src").rglob("*.html")):
        src_hash.update(p.read_bytes())
    manifest = {
        "format": 1, "book": "sewing", "version": src_hash.hexdigest()[:12],
        "course": {"fullname": pair("缝纫机设计与制造", "Sewing Machine Design and Manufacturing"),
                   "shortname": "SMDM",
                   "summary": ["<p>从一针线迹到整座数字工厂：缝纫机的机构、电控、原型机搭建与自动化测试、零件制造和服装生产线。"
                               "8 个单元 31 讲，每讲有讲义、中英课件、动画与虚拟实验、实验指导书、作业和测验。</p>",
                               "<p>From a single stitch to the digital factory: mechanisms, electrical control, building and testing control "
                               "prototypes, manufacturing the parts, and garment lines. 8 units, 31 lessons, each with notes, bilingual slides, "
                               "animations and virtual labs, lab guides, an assignment and a quiz.</p>"],
                   "blurb": "从一针线迹到数字工厂：缝纫机的机构、电控、原型机与自动化测试、零件制造与智能工厂。31 讲，中英双语，配虚拟实验。",
                   "info": course_info(lessons)},
        "batches": [{"no": b["no"], "name": pair(f"第 {b['no']} 批：第 {b['lessons'][0]}–{b['lessons'][-1]} 讲",
                                                    f"Batch {b['no']}: lessons {b['lessons'][0]}–{b['lessons'][-1]}"),
                     "lessons": b["lessons"], "exams": b["exams"]} for b in BATCHES],
        "exam_section": {"section": EXAM_SECTION, "name": pair("考试", "Examinations")},
        "lessons": entries,
        "exams": exams,
    }
    (out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    size = sum(p.stat().st_size for p in out.rglob("*") if p.is_file())
    print(f"manifest: {len(entries)} lessons, {sum(len(e['activities']) for e in entries)} activities, "
          f"exams {[len(e['questions']) for e in exams]} questions; {size / 1e6:.1f} MB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
