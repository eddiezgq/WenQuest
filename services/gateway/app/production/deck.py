"""Lesson slides in the chapter-1 benchmark layout (build_ch1_deck_bi.js, ported to python-pptx).

One deck per lesson, Chinese–English:
  cover · robot problem · concept · animation · virtual lab · back to the problem (5 steps +
  everyday example) · worked example · summary
The animation slide embeds the lesson's video once it is rendered (A2); the lab slide shows the
lab screenshot once the lab exists (A3) and always links to it ("#lab-2-1"), which the WenQuest
presenter turns into "go to lab 2.1".
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY, ICE, AMBER, INK, MUTED, WHITE, NIGHT, PALE = "1E2761", "EEF3FA", "F2B705", "1F2D3A", "5B6B75", "FFFFFF", "0F1419", "FFF6DA"
CN, EN = "Microsoft YaHei", "Calibri"
W, H = 13.333, 7.5
LAB_URL = "https://wenquestrobotics.com/lab/#lab-{ch}-{sec}"


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def _run(par, text, size, color=INK, bold=False, font=CN, link=None):
    r = par.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.name = Pt(size), bold, font
    f.color.rgb = rgb(color)
    rpr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", CN)
    if link:
        r.hyperlink.address = link
    return r


def box(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def text(slide, x, y, w, h, s, size, color=INK, bold=False, font=CN, align=None, anchor=MSO_ANCHOR.TOP):
    tf = box(slide, x, y, w, h, anchor)
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    _run(p, s, size, color, bold, font)
    return tf


def bi(slide, x, y, w, h, pair, zh=16, en=11, color=INK, en_color=MUTED, bold=False, anchor=MSO_ANCHOR.TOP):
    tf = box(slide, x, y, w, h, anchor)
    _run(tf.paragraphs[0], pair[0], zh, color, bold)
    p = tf.add_paragraph()
    _run(p, pair[1], en, en_color, False, EN)
    return tf


def bi_list(slide, x, y, w, h, items, zh=16, en=11, color=INK, en_color=MUTED, numbered=False):
    tf = box(slide, x, y, w, h)
    first = True
    for i, (a, b) in enumerate(items, 1):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        _run(p, (f"{i}. " if numbered else "• ") + a, zh, color)
        q = tf.add_paragraph()
        q.space_after = Pt(8)
        _run(q, ("    " if numbered else "   ") + b, en, en_color, False, EN)
    return tf


def title(slide, pair, y=0.62, color=INK, en_color=MUTED):
    text(slide, 0.7, y, 12, 0.6, pair[0], 26, color, True)
    text(slide, 0.7, y + 0.62, 12, 0.35, pair[1], 14, en_color, False, EN)


def tag(slide, s, color=AMBER):
    text(slide, 0.7, 0.3, 9, 0.3, s, 12, color, True)


def shape(slide, kind, x, y, w, h, fill, line=None):
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = rgb(fill)
    if line:
        sp.line.color.rgb = rgb(line)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = min(0.12, 0.12 / max(w, h) * 4)  # small corners, like the benchmark
    return sp


def background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def placeholder(slide, x, y, w, h, label: list[str], lines: list[list[str]], dark=False):
    """Where the animation or lab picture goes until it has been made."""
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, "1A2430" if dark else "F7F9FB", "3A4A5A" if dark else "C9D3DD")
    tf = box(slide, x + 0.4, y + 0.35, w - 0.8, h - 0.7)
    _run(tf.paragraphs[0], label[0], 15, "F2F5F8" if dark else INK, True)
    q = tf.add_paragraph()
    _run(q, label[1], 11, "9FB3D1" if dark else MUTED, False, EN)
    for a, b in lines[:6]:
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        _run(p, "· " + a, 13, "E6EDF7" if dark else INK)
        p2 = tf.add_paragraph()
        _run(p2, "  " + b, 10, "9FB3D1" if dark else MUTED, False, EN)


def build(spec: dict, out: Path, *, no: str, course: list[str], chapter: list[str],
          video: Path | None = None, poster: Path | None = None, lab_image: Path | None = None) -> Path:
    """Write the lesson deck. `no` is the lesson number such as '2.1'."""
    ch, _, sec = no.partition(".")
    lab_url = LAB_URL.format(ch=ch or "1", sec=sec or "1")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    s = spec

    # cover
    sl = prs.slides.add_slide(blank)
    background(sl, NAVY)
    text(sl, 0.8, 0.7, 11, 0.4, f"{course[0]}  ·  {chapter[0]}", 14, AMBER, True)
    text(sl, 0.8, 1.05, 11, 0.35, f"{course[1]}  ·  {chapter[1]}", 11, "9FB3D1", False, EN)
    text(sl, 0.8, 2.5, 11.5, 1.1, f"{no}  {s['title'][0]}", 40, WHITE, True)
    text(sl, 0.8, 3.6, 11.5, 0.6, s["title"][1], 22, "C9D6EA", False, EN)
    bi(sl, 0.8, 4.7, 11.5, 0.9, [f"本节目标：{s['goal'][0]}", f"Goal: {s['goal'][1]}"], 16, 12, "E6EDF7", "9FB3D1")
    shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, 6.3, 11.7, 0.6, "283A7A")
    text(sl, 1.0, 6.3, 11.4, 0.6, "每节课：实际问题 → 概念 → 动画 → 虚拟实验 → 建模求解    Each lesson: problem → concept → animation → lab → modelling",
         12, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    notes(sl, "本节从一个机器人实际问题出发，最后用本节的模型解决它。\nThis lesson starts from a robot problem and ends by solving it with the lesson's model.")

    # A. the robot problem
    p = s["problem"]
    sl = prs.slides.add_slide(blank)
    background(sl, PALE)
    tag(sl, f"{no}  机器人问题  Robot problem")
    title(sl, p["title"])
    bi(sl, 0.7, 1.8, 5.6, 2.4, p["text"], 16, 12)
    text(sl, 0.7, 4.35, 5, 0.35, "已知  Given", 13, MUTED, True)
    bi_list(sl, 0.7, 4.75, 5.6, 2.3, p["given"], 14, 10.5)
    iw, ih = 6.3, 6.3 * 10 / 16
    if lab_image and lab_image.exists():
        sl.shapes.add_picture(str(lab_image), Inches(6.6), Inches(1.8), Inches(iw), Inches(ih))
    else:
        placeholder(sl, 6.6, 1.8, iw, ih, ["机器人场景", "Robot scene"], [s["lab"]["robot_scene"]])
    text(sl, 6.6, 1.8 + ih + 0.15, iw, 0.6, "本节结束时，我们用建立的模型回答这个问题。  We will answer it with this lesson’s model at the end.", 11, MUTED)
    notes(sl, "先让学生猜一猜、讨论 2 分钟，不急于给答案。本节最后回到这一页。\nLet students guess and discuss for two minutes; return to this problem at the end.")

    # B. concept
    c = s["concept"]
    sl = prs.slides.add_slide(blank)
    background(sl, WHITE)
    tag(sl, f"{no}  概念  Concept")
    title(sl, c["title"])
    bi_list(sl, 0.7, 1.8, 11.9, 3.4, c["points"])
    shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 5.35, 11.9, 1.4, ICE)
    bi(sl, 1.0, 5.4, 11.4, 1.3, c["formula"], 18, 13, NAVY, MUTED, True, MSO_ANCHOR.MIDDLE)
    notes(sl, "讲清公式的含义，再进入动画。\nExplain what each formula means before the animation.")

    # C. animation
    a = s["animation"]
    sl = prs.slides.add_slide(blank)
    background(sl, NIGHT)
    tag(sl, f"{no}  动画  Animation")
    bi(sl, 0.7, 0.6, 12, 0.85, a["question"], 19, 13, WHITE, "9FB3D1", True)
    vw = 9.3
    vh = vw * 9 / 16
    vx, vy = (W - vw) / 2, 1.6
    if video and video.exists():
        sl.shapes.add_movie(str(video), Inches(vx), Inches(vy), Inches(vw), Inches(vh),
                            poster_frame_image=str(poster) if poster and poster.exists() else None, mime_type="video/mp4")
        text(sl, vx, vy + vh + 0.08, vw, 0.3, "点击画面播放 · 中英字幕 · Click to play · bilingual captions", 11, "9FB3D1", align=PP_ALIGN.CENTER)
    else:
        placeholder(sl, vx, vy, vw, vh, ["动画分镜（动画渲染完成后替换）", "Storyboard (replaced by the rendered animation)"], a["beats"], dark=True)
    notes(sl, "播放动画，在关键画面暂停提问。\nPlay the clip and pause at the key frame to ask the question on the slide.")

    # D. virtual lab
    lb = s["lab"]
    sl = prs.slides.add_slide(blank)
    background(sl, ICE)
    tag(sl, f"{no}  虚拟实验  Virtual lab")
    title(sl, ["动手试一试：" + s["title"][0], "Try it: " + s["title"][1]])
    iw2, ih2 = 7.3, 7.3 * 10 / 16
    if lab_image and lab_image.exists():
        pic = sl.shapes.add_picture(str(lab_image), Inches(0.7), Inches(1.8), Inches(iw2), Inches(ih2))
        pic.click_action.hyperlink.address = lab_url
    else:
        placeholder(sl, 0.7, 1.8, iw2, ih2, lb["title"], [lb["robot_scene"], lb["life_scene"]] + lb["params"][:3])
    text(sl, 8.4, 1.8, 4.3, 0.35, "任务  Tasks", 14, MUTED, True)
    bi_list(sl, 8.4, 2.2, 4.3, 3.0, lb["tasks"], 14, 10.5, numbered=True)
    btn = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 8.4, 5.35, 4.3, 0.65, NAVY)
    btn.click_action.hyperlink.address = lab_url
    tf = box(sl, 8.4, 5.35, 4.3, 0.65, MSO_ANCHOR.MIDDLE)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    _run(tf.paragraphs[0], "打开虚拟实验  Open lab ▶", 15, WHITE, True)  # the button shape carries the link
    text(sl, 8.4, 6.1, 4.3, 0.5, "课后按实验指导书完成实验报告 · Write the lab report using the lab guide", 11, MUTED, align=PP_ALIGN.CENTER)
    notes(sl, "课堂投屏演示一遍机器人场景，其余任务课后完成，并按实验指导书写实验报告。\nDemonstrate the robot scene in class; students finish the tasks and the lab report after class.")

    # E. back to the problem: the five steps
    m = s["model"]
    sl = prs.slides.add_slide(blank)
    background(sl, WHITE)
    tag(sl, f"{no}  回到实际问题  Back to the problem")
    title(sl, ["建模求解：" + p["title"][0], "Modelling: " + p["title"][1]])
    steps = [("① 问题", "Problem", p["text"]), ("② 建模", "Model", m["assume"]), ("③ 求解", "Solve", m["solve"]),
             ("④ 检验", "Check", m["check"]), ("⑤ 修正", "Improve", m["improve"])]
    tbl = sl.shapes.add_table(len(steps) + 1, 2, Inches(0.7), Inches(1.75), Inches(8.4), Inches(4.9)).table
    tbl.columns[0].width, tbl.columns[1].width = Inches(1.35), Inches(7.05)
    for j, head in enumerate(("步骤 Step", "内容 What we do")):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(ICE)
        cell.text_frame.paragraphs[0].text = ""
        _run(cell.text_frame.paragraphs[0], head, 13, INK, True)
    for i, (zh, en, pr) in enumerate(steps, 1):
        c0, c1 = tbl.cell(i, 0), tbl.cell(i, 1)
        for c in (c0, c1):
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(WHITE)
        c0.text_frame.paragraphs[0].text = ""
        _run(c0.text_frame.paragraphs[0], zh, 13, INK, True)
        _run(c0.text_frame.add_paragraph(), en, 10, MUTED, False, EN)
        c1.text_frame.paragraphs[0].text = ""
        c1.text_frame.word_wrap = True
        _run(c1.text_frame.paragraphs[0], pr[0], 12.5, INK)
        _run(c1.text_frame.add_paragraph(), pr[1], 10, MUTED, False, EN)
    shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 9.4, 1.75, 3.3, 3.6, PALE)
    tf = box(sl, 9.6, 1.9, 2.9, 3.3)
    _run(tf.paragraphs[0], "生活中的例子", 15, INK, True)
    _run(tf.add_paragraph(), "Everyday example", 11, MUTED, False, EN)
    q = tf.add_paragraph()
    q.space_before = Pt(8)
    _run(q, s["everyday"]["text"][0], 13, INK)
    _run(tf.add_paragraph(), s["everyday"]["text"][1], 10, MUTED, False, EN)
    text(sl, 9.4, 5.5, 3.3, 1.4, "工程和科学都靠建立模型来描述世界：先简化，再求解，再用实验检验，最后找出模型的局限。\n"
         "Engineering and science describe the world through models: simplify, solve, test, then find where the model breaks.", 10.5, MUTED)
    notes(sl, "对照五步讲解，强调第⑤步：模型在哪里失效。实验报告第六部分按这五步写。\n"
              "Walk through the five steps and stress step ⑤. Section 6 of the lab report follows the same steps.")

    # worked example
    e = s["example"]
    sl = prs.slides.add_slide(blank)
    background(sl, ICE)
    tag(sl, f"{no}  例题  Worked example")
    bi(sl, 0.7, 0.75, 12, 1.4, e["question"], 20, 13, INK, MUTED, True)
    shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 2.4, 12, 4.3, WHITE, "D5DEE8")
    bi_list(sl, 1.0, 2.65, 11.4, 3.9, e["steps"] + ([e["answer"]] if e["answer"][0] else []), 17, 12, numbered=True)
    notes(sl, "先让学生独立做 3 分钟，再展开解答。\nGive students three minutes, then work through it.")

    # summary
    sl = prs.slides.add_slide(blank)
    background(sl, NAVY)
    text(sl, 0.8, 0.5, 8, 0.7, "本节小结", 30, WHITE, True)
    text(sl, 0.8, 1.15, 8, 0.4, "Summary", 16, "9FB3D1", False, EN)
    bi_list(sl, 0.8, 1.8, 11.8, 3.0, s["summary"], 17, 11.5, "E6EDF7", "9FB3D1")
    if s["self_check"]:
        text(sl, 0.8, 4.9, 11.8, 0.35, "想一想  Check yourself", 14, AMBER, True)
        bi_list(sl, 0.8, 5.3, 11.8, 1.9, s["self_check"][:3], 13, 10.5, "E6EDF7", "9FB3D1", numbered=True)
    notes(sl, "回顾本节要点，提醒练习和实验报告。\nReview the lesson and remind students of the practice set and the lab report.")

    # sources: library entries used (问渠零件与机器人库) with their licences
    if spec.get("sources"):
        sl = prs.slides.add_slide(blank)
        background(sl, ICE)
        tag(sl, f"{no}  素材来源  Sources")
        lines = [[f"{x['name'].get('zh', '')}（{x['id']} · {x['version']}）",
                  f"{x['name'].get('en', '')} — {x['license']}; {x['attribution']}"] for x in spec["sources"][:10]]
        bi_list(sl, 0.8, 1.0, 11.8, 5.8, lines, 13, 10)
        notes(sl, "本课用到的模型、图纸来自问渠零件与机器人库。\nModels and drawings in this lesson come from the WenQuest parts and robot library.")

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
