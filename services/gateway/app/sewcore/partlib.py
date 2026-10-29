# -*- coding: utf-8 -*-
"""本单位零件库的类别、字段、接进设计的方式和造型程序（第 49 轮，细则第七节 2）。

网关（sewing_parts.py）、网页（GET /api/v1/sewing/parts/cats）、Excel 模板、AI 整理样本都按这张表走。
造型：有造型程序的类别按字段生成三维（build123d 出 STEP），和导出的机构零件是同一套做法；机壳、其他要自己上传 STEP。
"""
import io
import math

# 字段：(键, 中文, 单位, 最小, 最大, 必填)
CATS = {
    "motor": dict(zh="伺服 / 主轴电机", fields=[
        ("power_w", "额定功率", "W", 10, 20000, True),
        ("rated_nm", "额定转矩", "N·m", 0.01, 200, True),
        ("peak_nm", "峰值转矩", "N·m", 0.01, 600, True),
        ("rated_rpm", "额定转速", "r/min", 100, 20000, True),
        ("max_rpm", "最高转速", "r/min", 100, 20000, True),
        ("inertia", "转子惯量", "×10⁻⁴ kg·m²", 0.001, 500, True),
        ("flange", "法兰边长", "mm", 20, 300, False),
        ("length", "机身长", "mm", 20, 600, False),
        ("shaft_d", "轴径", "mm", 3, 60, False)],
        uses=[("pattern", "X 轴电机", {"TxR": "rated_nm", "TxP": "peak_nm", "Jx": "inertia", "Nmot": "rated_rpm"}),
              ("pattern", "Y 轴电机", {"TyR": "rated_nm", "TyP": "peak_nm", "Jy": "inertia", "Nmot": "rated_rpm"})],
        cae="电机选型的候选表里一起比", shape=True),
    "needle": dict(zh="机针", fields=[
        ("nm", "针号 Nm", "", 40, 250, True),
        ("blade_d", "针杆直径", "mm", 0.4, 2.5, True),
        ("shank_d", "针柄直径", "mm", 1.0, 4.0, False),
        ("length", "全长", "mm", 20, 60, False)],
        uses=[("lockstitch", "机针", {"d": "blade_d"})], shape=True),
    "needlebar": dict(zh="针杆组件", fields=[
        ("mass", "质量", "kg", 0.01, 1.0, True),
        ("d", "针杆直径", "mm", 3, 20, False),
        ("length", "长度", "mm", 40, 300, False)],
        uses=[("lockstitch", "针杆组件", {"mr": "mass"}), ("overlock", "针杆组件", {"mr": "mass"})], shape=True),
    "bushing": dict(zh="针杆轴套", fields=[
        ("load", "允许载荷", "N", 50, 5000, True),
        ("d_in", "内径", "mm", 3, 30, False),
        ("d_out", "外径", "mm", 5, 50, False),
        ("length", "长度", "mm", 5, 80, False)],
        uses=[("lockstitch", "针杆轴套", {"F_lim": "load"})], shape=True),
    "belt": dict(zh="同步带", fields=[
        ("EA", "拉伸刚度 EA", "kN", 10, 2000, True),
        ("width", "带宽", "mm", 3, 100, False),
        ("pitch", "节距", "mm", 1, 20, False)],
        uses=[("pattern", "同步带", {"EA": "EA"})], shape=False),
    "pulley": dict(zh="同步带轮", fields=[
        ("r", "节圆半径", "mm", 3, 150, True),
        ("teeth", "齿数", "", 8, 200, False),
        ("width", "带宽", "mm", 3, 100, False),
        ("bore", "孔径", "mm", 2, 60, False)],
        uses=[("pattern", "X 轴带轮", {"rx": "r"}), ("pattern", "Y 轴带轮", {"ry": "r"}), ("embroid", "绷架带轮", {"rf": "r"})], shape=True),
    "link": dict(zh="曲柄 / 连杆（两孔板）", fields=[
        ("L", "孔距", "mm", 5, 200, True),
        ("d1", "大孔径", "mm", 2, 60, False),
        ("d2", "小孔径", "mm", 1, 60, False),
        ("R1", "大端外圆半径", "mm", 3, 80, False),
        ("R2", "小端外圆半径", "mm", 2, 80, False),
        ("t", "厚度", "mm", 1, 40, False)],
        uses=[("lockstitch", "刺布曲柄", {"r": "L"}), ("lockstitch", "针杆连杆", {"l": "L"}),
              ("overlock", "引针曲柄", {"r": "L"}), ("overlock", "引针连杆", {"l": "L"})], shape=True),
    "frame": dict(zh="机壳 / 绷架 / 横梁", fields=[
        ("mass", "质量", "kg", 0.1, 2000, False),
        ("material", "材料", "", None, None, False)],
        uses=[], cae="“机壳固有频率”可以直接选这里的件", shape=False, need_step=True),
    "other": dict(zh="其他", fields=[
        ("spec", "规格", "", None, None, False),
        ("material", "材料", "", None, None, False)],
        uses=[], shape=False),
}
ORDER = list(CATS)
TEXT_FIELDS = {"material", "spec"}


def cats_public():
    out = []
    for k in ORDER:
        c = CATS[k]
        out.append(dict(key=k, zh=c["zh"], shape=c.get("shape", False), need_step=c.get("need_step", False), cae=c.get("cae", ""),
                        fields=[dict(key=f[0], zh=f[1], unit=f[2], min=f[3], max=f[4], required=f[5], text=f[0] in TEXT_FIELDS)
                                for f in c["fields"]],
                        uses=[dict(machine=m, zh=zh, map=mp) for m, zh, mp in c["uses"]]))
    return out


def check_fields(cat, fields):
    """返回 (整理好的字段, 问题列表)。数值字段转成数、查范围；必填不能空。"""
    c = CATS.get(cat)
    if not c:
        return {}, ["没有这个类别"]
    out, errs = {}, []
    for key, zh, unit, lo, hi, req in c["fields"]:
        v = (fields or {}).get(key)
        if v is None or (isinstance(v, str) and not v.strip()):
            if req:
                errs.append(f"{zh}没填")
            continue
        if key in TEXT_FIELDS:
            out[key] = str(v).strip()[:120]
            continue
        try:
            x = float(str(v).replace(",", "").strip())
        except ValueError:
            errs.append(f"{zh}不是数：{v}")
            continue
        if not math.isfinite(x) or (lo is not None and not lo <= x <= hi):
            errs.append(f"{zh} {x:g} {unit} 不在 {lo:g}–{hi:g} 内")
            continue
        out[key] = int(x) if x == int(x) and abs(x) < 1e12 else x
    return out, errs


def param_values(cat, fields, machine):
    """这个零件接进某机型的哪些参数：[{label, params}]"""
    c = CATS.get(cat) or {}
    res = []
    for m, zh, mp in c.get("uses", []):
        if m != machine:
            continue
        ps = {pk: fields[fk] for pk, fk in mp.items() if fk in fields}
        if len(ps) == len(mp):
            res.append(dict(label=zh, params=ps))
    return res


def cae_motor(part):
    """本单位电机 → CAE 电机表的一行（单位与数字工厂电机表一致）"""
    f = part["fields"]
    return dict(name=f"本单位 · {part['code']} {part['name']}", power_w=f["power_w"], rated_nm=f["rated_nm"], peak_nm=f["peak_nm"],
                max_rpm=f["max_rpm"], inertia_kgm2=f["inertia"] * 1e-4, org=True, part_id=part["id"])


# ---------- 造型 ----------
def can_shape(cat, fields):
    c = CATS.get(cat) or {}
    return bool(c.get("shape")) and not (cat == "motor" and not fields.get("flange"))


def to_step(cat, fields, label="part"):
    """按字段生成三维，返回 STEP 字节；服务器没装 build123d 或这类没有造型时返回 None"""
    try:
        import build123d as bd
    except Exception:
        return None
    from .cad import _bd_solid
    f = fields
    Z = (0, 0, 1)

    def cyl(r, h, z0=0.0, d=Z, x0=0.0):
        return bd.Solid.make_cylinder(r, h, bd.Plane(origin=(x0, 0, z0), z_dir=d))

    s = None
    if cat == "motor" and f.get("flange"):
        a = float(f["flange"])
        L = float(f.get("length") or a * 1.5)
        sd = float(f.get("shaft_d") or max(5.0, a / 6))
        s = bd.Solid.make_box(a, a, 8, bd.Plane(origin=(-a / 2, -a / 2, 0)))
        s = s + bd.Solid.make_box(a * 0.9, a * 0.9, L, bd.Plane(origin=(-a * 0.45, -a * 0.45, -L)))
        s = s + cyl(sd / 2, sd * 3, 8)
    elif cat == "needle":
        L = float(f.get("length") or 38)
        sh = float(f.get("shank_d") or 1.6)
        bl = float(f["blade_d"])
        s = cyl(sh / 2, L * 0.4, L * 0.6) + cyl(bl / 2, L * 0.6 - bl, bl) + bd.Solid.make_cone(0.05, bl / 2, bl, bd.Plane(origin=(0, 0, 0), z_dir=Z))
    elif cat == "needlebar":
        s = cyl(float(f.get("d") or 7) / 2, float(f.get("length") or 110))
    elif cat == "bushing":
        di = float(f.get("d_in") or 7)
        do = float(f.get("d_out") or di + 5)
        L = float(f.get("length") or 20)
        if do <= di:
            return None
        s = cyl(do / 2, L) - cyl(di / 2, L + 2, -1)
    elif cat == "pulley":
        r = float(f["r"])
        w = float(f.get("width") or 10)
        bore = float(f.get("bore") or max(3.0, r * 0.4))
        s = cyl(r, w) + cyl(r + 1.5, 1.2, -1.2) + cyl(r + 1.5, 1.2, w)
        if bore / 2 < r:
            s = s - cyl(bore / 2, w + 6, -3)
    elif cat == "link":
        L = float(f["L"])
        d1 = float(f.get("d1") or max(4.0, L * 0.25))
        d2 = float(f.get("d2") or max(3.0, d1 * 0.6))
        R1 = float(f.get("R1") or d1 * 0.9)
        R2 = float(f.get("R2") or d2 * 0.9)
        if R1 <= d1 / 2 or R2 <= d2 / 2:
            return None
        s = _bd_solid(bd, dict(type="plate", circles=[[0, 0, R1], [L, 0, R2]], t=float(f.get("t") or 6),
                               holes=[[0, 0, d1 / 2], [L, 0, d2 / 2]],
                               frame=dict(o=(0, 0, 0), u=(1, 0, 0), v=(0, 1, 0), n=(0, 0, 1))))
    if s is None:
        return None
    import os
    import tempfile
    s.label = label
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "part.step")
        bd.export_step(s, path)
        with open(path, "rb") as fh:
            return fh.read()


# ---------- Excel 模板与读入 ----------
HEAD = [("code", "本单位编号", True), ("name", "名称", True), ("erp", "ERP 物料号", False)]


def template_xlsx(cat):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    c = CATS[cat]
    wb = Workbook()
    ws = wb.active
    ws.title = c["zh"].replace("/", "、")[:30]
    cols = [(k, zh, "", req) for k, zh, req in HEAD] + [(k, zh, u, req) for k, zh, u, _lo, _hi, req in c["fields"]]
    for i, (k, zh, u, req) in enumerate(cols, 1):
        cell = ws.cell(1, i, f"{zh}{'（' + u + '）' if u else ''}{' *' if req else ''}")
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EAF2FA")
        ws.cell(2, i, k).font = Font(color="888888", size=8)
        ws.column_dimensions[cell.column_letter].width = max(12, len(cell.value) * 2)
    ws.cell(3, 1, "（从第 3 行起，一行一个规格；第 2 行的灰字是字段键，别删。带 * 的必填）").font = Font(color="888888")
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def read_xlsx(cat, data):
    """读模板填好的表：返回 [{row, code, name, erp, fields, errors}]。按第 2 行的字段键认列，列顺序可以变。"""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    keys = [str(x).strip() if x is not None else "" for x in rows[1]]
    known = {k for k, _z, _r in HEAD} | {f[0] for f in CATS[cat]["fields"]}
    if not ({"code", "name"} <= set(keys)) or not (set(keys) - {""}) <= known:
        raise ValueError("表头不对：请用本类别的模板，第 2 行的字段键别改")
    out = []
    for i, r in enumerate(rows[2:], 3):
        d = {keys[j]: v for j, v in enumerate(r) if j < len(keys) and keys[j]}
        if not any(v not in (None, "") for v in d.values()):
            continue
        if str(d.get("code") or "").startswith("（"):
            continue
        code, name = str(d.pop("code", "") or "").strip(), str(d.pop("name", "") or "").strip()
        erp = str(d.pop("erp", "") or "").strip()
        fields, errs = check_fields(cat, d)
        if not code:
            errs.insert(0, "本单位编号没填")
        if not name:
            errs.insert(0, "名称没填")
        out.append(dict(row=i, code=code[:60], name=name[:80], erp=erp[:60], fields=fields, errors=errs))
        if len(out) > 2000:
            break
    seen = {}
    for x in out:
        if x["code"] in seen and x["code"]:
            x["errors"].append(f"编号和第 {seen[x['code']]} 行重复")
        seen.setdefault(x["code"], x["row"])
    return out


def erp_xlsx(parts):
    """ERPNext 物料导入表（数据导入工具：物料 Item）：物料号、名称、物料组、单位、说明"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Item"
    ws.append(["Item Code", "Item Name", "Item Group", "Stock UOM", "Is Stock Item", "Description"])
    for p in parts:
        c = CATS.get(p["cat"], {})
        spec = "；".join(f"{zh} {p['fields'][k]}{(' ' + u) if u else ''}" for k, zh, u, *_ in c.get("fields", []) if k in p["fields"])
        ws.append([p["erp"] or p["code"], p["name"], "Sub Assemblies", "Nos", 1,
                   f"本单位编号 {p['code']}；{c.get('zh', '')}；{spec}"])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()
