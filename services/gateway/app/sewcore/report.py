# -*- coding: utf-8 -*-
"""出件（第 39 轮）：设计计算书（Word）、调机卡（Word）、零件参数表（Excel）。只在网关里用（要 python-docx、openpyxl、matplotlib）。"""
import io
import math
from datetime import date

from .design import CHAIN
from .machines import get

STATE = {"ok": ("合格", "Pass", "2E9E6A"), "warn": ("注意", "Check", "E8913A"), "bad": ("不合格", "Fail", "C0392B")}
FONT = "Noto Sans CJK SC"


def _t(zh, en, lang):
    return en if lang == "en" else zh


def _choice(M, k, v, lang):
    for c in M.CHOICES.get(k, []):
        if c[0] == v:
            return c[2] if lang == "en" else c[1]
    return v


def _fmt(v):
    if isinstance(v, float):
        return f"{v:g}" if abs(v) >= 1e-3 or v == 0 else f"{v:.4f}"
    return str(v)


# ---------- 图 ----------
def _mpl_cjk():
    """matplotlib 不一定认 .ttc 里的中文字体：按 fc-match 找到文件后手动登记"""
    import subprocess
    from matplotlib import font_manager
    try:
        path = subprocess.run(["fc-match", "-f", "%{file}", FONT + ":style=Regular"], capture_output=True, text=True, timeout=5).stdout.strip()
        if path:
            font_manager.fontManager.addfont(path)
            return font_manager.FontProperties(fname=path).get_name()
    except Exception:
        pass
    return FONT


def figures(R, lang="zh"):
    """计算结果里的 charts（各机型通用）加一张时序图"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.family"] = [_mpl_cjk(), "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    out = []

    def save(fig, title):
        b = io.BytesIO()
        fig.tight_layout()
        fig.savefig(b, format="png", dpi=150)
        plt.close(fig)
        out.append((title, b.getvalue()))
    no = 0
    for c in R.get("charts", []):
        no += 1
        if c["kind"] == "xy":
            fig, ax = plt.subplots(figsize=(5, 3.4))
            sr = c["series"][0]
            if c.get("open"):
                ax.plot(c["x"], sr["y"], color=sr["color"], marker=".", ms=3)
            else:
                ax.plot(c["x"] + c["x"][:1], sr["y"] + sr["y"][:1], color=sr["color"])
            ax.axhline(0, color="#8a9aa8", lw=0.8)
            ax.set_aspect("equal", adjustable="datalim")
        else:
            fig, ax = plt.subplots(figsize=(7, 3))
            for sr in c["series"]:
                ax.plot(c["x"], sr["y"], color=sr["color"], label=sr[lang] or None)
            if c.get("ref") is not None:
                ax.axhline(c["ref"], color="#c0392b", lw=0.8, ls="--")
            if c.get("zero"):
                ax.axhline(0, color="#8a9aa8", lw=0.8)
            if c.get("xr"):
                ax.set_xlim(*c["xr"])
                ax.set_xlabel(c["xl"][lang])
            else:
                ax.set_xlim(0, 360)
                ax.set_xticks(range(0, 361, 45))
                ax.set_xlabel(_t("主轴转角 °", "Main-shaft angle, °", lang))
            if len(c["series"]) > 1:
                ax.legend()
        ax.grid(alpha=0.3)
        save(fig, _t(f"图 {no} {c['zh']}", f"Fig. {no} {c['en']}", lang))
    rows = R.get("timing", [])
    if rows:
        no += 1
        fig, ax = plt.subplots(figsize=(7, 0.6 + 0.45 * len(rows)))
        for i, t in enumerate(rows):
            a, b = t["a"], t["b"]
            segs = [(a, b - a)] if b >= a else [(a, 360 - a), (0, b)]
            if a == b:
                segs = [(a - 1, 2)]
            ax.broken_barh(segs, (i - 0.35, 0.7), color=["#2d6fb3", "#7d5ba6", "#c0392b", "#2e9e6a", "#e8913a"][i % 5])
            if "mark" in t:
                ax.plot([t["mark"]] * 2, [i - 0.45, i + 0.45], color="k", lw=1.2)
        ax.set_yticks(range(len(rows)))
        ax.set_yticklabels([t["en" if lang == "en" else "zh"] for t in rows], fontsize=8)
        ax.set_xlim(0, 360)
        ax.set_xticks(range(0, 361, 30))
        ax.grid(axis="x", alpha=0.3)
        if R.get("timing_x") == "pass":            # 横机（第 54 轮）：没有主轴转角，横轴是一趟的时间
            ax.set_xticklabels([f"{round(x / 3.6)}%" for x in range(0, 361, 30)])
            ax.set_xlabel(_t("机头走一趟的时间（%）", "Time of one carriage pass, %", lang))
        else:
            ax.set_xlabel(_t("主轴转角 °", "Main-shaft angle, °", lang))
        save(fig, _t(f"图 {no} {'机头一趟的时间分配' if R.get('timing_x') == 'pass' else '全机运动时序'}",
                     f"Fig. {no} {'Time split of one carriage pass' if R.get('timing_x') == 'pass' else 'Machine timing diagram'}", lang))
    return out


# ---------- Word ----------
def _doc():
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
    d = Document()
    st = d.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return d


def _shade(cell, hexcolor):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def _table(d, head, rows, widths=None):
    t = d.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for i, h in enumerate(head):
        t.rows[0].cells[i].text = h
        _shade(t.rows[0].cells[i], "EAF2FA")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = _fmt(v)
    return t


NUM = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]


def ebox_section(d, head, E, lang):
    """第 62 轮：电控箱的热与流体（教材第 18 章）——参数、结果、校核、改进方案"""
    from .ebox import CHOICES as EC, GROUPS as EG
    q = E["params"]
    head("电控箱：热、电容寿命、开关时间", "Control box: heat, capacitor life, switching time")
    d.add_paragraph(_t("模型见教材第 18 章（功率级损耗、热阻网络、间歇缝纫负载、电解电容十度法则、开关时间与共模电流的折中），"
                       "电流取本机的设计（主轴转矩或伺服轴转矩 ÷ 转矩常数）；判据是教材的示意值，按本单位的器件资料改。",
                       "Models follow chapter 18 (power-stage loss, thermal network, intermittent sewing load, the 10-degree rule for "
                       "electrolytic capacitors, switching time vs common-mode current); currents come from this design (spindle or servo-axis "
                       "torque ÷ torque constant); the criteria are the textbook's illustrative values — replace them with your parts' data.", lang))
    _table(d, [_t("驱动", "Drive", lang), _t("运行电流 A", "Run current A", lang), _t("峰值电流 A", "Peak current A", lang),
               _t("平均损耗 W", "Average loss W", lang), _t("结温峰值 °C", "Peak junction °C", lang)],
           [[_t(x["zh"], x["en"], lang), x["Irun"], x["Ipk"], x["Pavg"], x["Tjpk"]] for x in E["drives"]])
    d.add_paragraph("")
    _table(d, [_t("项目", "Item", lang), _t("数值", "Value", lang), _t("单位", "Unit", lang)],
           [[_t(r["zh"], r["en"], lang), r["value"], r["unit"]] for r in E["results"]])
    d.add_paragraph("")
    _table(d, [_t("校核项", "Check", lang), _t("数值", "Value", lang), _t("判据", "Criterion", lang), _t("结论", "Result", lang)],
           [[c["name"][lang], c["value"][lang], c["limit"][lang], STATE[c["state"]][1 if lang == "en" else 0]] for c in E["checks"]])
    v = E["variants"]
    d.add_paragraph(_t(v["zh"], v["en"], lang))
    _table(d, [h[1] if lang == "en" else h[0] for i, h in enumerate(v["cols"]) if i != 1],
           [[r[1] if lang == "en" else r[0]] + r[2:] for r in v["rows"]])
    rows = []
    for gk, gzh, gen, items in EG:
        for k, zh, en, unit, *_ in items:
            val = q[k]
            if k in EC:
                val = next((c[2] if lang == "en" else c[1] for c in EC[k] if c[0] == val), val)
            rows.append([_t(gzh, gen, lang), _t(zh, en, lang), val, unit])
    d.add_paragraph(_t("电控箱参数", "Control-box parameters", lang))
    _table(d, [_t("部分", "Group", lang), _t("参数", "Parameter", lang), _t("取值", "Value", lang), _t("单位", "Unit", lang)], rows)


def calc_report(params=None, lang="zh", project="", company="", machine="lockstitch", cae=None, desc=None, ebox=None):
    from docx.shared import Cm, RGBColor
    M = get(machine)
    R = M.analyze(params)
    p = R["params"]
    d = _doc()
    sec = [0]

    def head(zh, en):
        sec[0] += 1
        d.add_heading(_t(f"{NUM[sec[0] - 1]}、{zh}", f"{sec[0]} {en}", lang), 1)
    d.add_heading(_t(f"{M.NAME[0]}设计计算书", f"{M.NAME[1]} Machine Design Calculation", lang), 0)
    d.add_paragraph(_t(f"项目：{project or '未命名'}　　单位：{company or '—'}　　日期：{date.today().isoformat()}　　计算核心版本：{R['version']}",
                       f"Project: {project or 'Untitled'}    Company: {company or '—'}    Date: {date.today().isoformat()}    Core version: {R['version']}", lang))
    sch = (("机头质量、传动、电机、摇床伺服和全成型对比用的是示例值", "carriage mass, drive, motor, racking servo and the whole-garment comparison use example values")
           if machine == "flatknit" else ("线环、送布、张力、面料等为教材中的示意模型", "loop, feed, tension and fabric models are the textbook's schematic models"))
    d.add_paragraph(_t(f"本计算书由问渠缝纫机设计软件生成。模型与判据见教材《缝纫机设计与制造》{M.LECTURES[0]}；"
                       f"其中{sch[0]}，结论用于设计比较和调机，不代替样机试验。",
                       "Generated by the WenQuest sewing-machine design software. Models and criteria follow the textbook "
                       f"Sewing Machine Design and Manufacturing, {M.LECTURES[1]}; {sch[1]} — use the results to compare designs and to set up "
                       "the machine, not in place of prototype tests.", lang))
    bad = [c for c in R["checks"] if c["state"] == "bad"]
    warn = [c for c in R["checks"] if c["state"] == "warn"]
    head("结论", "Conclusion")
    if not bad and not warn:
        d.add_paragraph(_t(f"全部 {len(R['checks'])} 项校核合格。", f"All {len(R['checks'])} checks pass.", lang))
    else:
        d.add_paragraph(_t(f"{len(R['checks'])} 项校核：不合格 {len(bad)} 项，需注意 {len(warn)} 项。",
                           f"{len(R['checks'])} checks: {len(bad)} fail, {len(warn)} need attention.", lang))
        for c in bad + warn:
            d.add_paragraph(f"{c['name'][lang]}：{c['value'][lang]}（{_t('判据', 'criterion', lang)} {c['limit'][lang]}）{c['advice'][lang]}", style="List Bullet")
    if "trial" in R:
        cn = R["trial"]["counts"]
        d.add_paragraph(_t(f"虚拟试缝 30 针：成缝 {cn['ok']}，跳针 {cn['skip']}，碰针 {cn['hit']}，划线 {cn['cut']}，断线 {cn['break']}。",
                           f"Virtual test seam, 30 stitches: good {cn['ok']}, skipped {cn['skip']}, needle hit {cn['hit']}, split {cn['cut']}, broken {cn['break']}.", lang))
    if desc:                                    # 第 49 轮：AI 助手起草的设计说明（数字已由 aicheck 与本组结果核对）
        from .aicheck import verify
        keep, _ = verify(desc, R, M.NAME[0])
        if keep:
            head("设计说明", "Design description")
            d.add_paragraph("".join(keep))
            d.add_paragraph(_t("（AI 起草，文中数字已由程序与本计算书逐个核对。）",
                               "(Drafted by AI; every number has been checked by the software against this calculation sheet.)", lang))
    head("设计参数", "Design parameters")
    rows = []
    for gk, gzh, gen, items in M.GROUPS:
        for k, zh, en, unit, *_ in items:
            v = p[k]
            rows.append([_t(gzh, gen, lang), _t(zh, en, lang), _choice(M, k, v, lang) if k in M.CHOICES else v, unit])
    _table(d, [_t("部分", "Group", lang), _t("参数", "Parameter", lang), _t("取值", "Value", lang), _t("单位", "Unit", lang)], rows)
    head("主要计算结果", "Main results")
    _table(d, [_t("项目", "Item", lang), _t("数值", "Value", lang), _t("单位", "Unit", lang)],
           [[_t(r["zh"], r["en"], lang), r["value"], r["unit"]] for r in R["results"]])
    head("校核", "Checks")
    t = _table(d, [_t("讲次", "Lecture", lang), _t("校核项", "Check", lang), _t("数值", "Value", lang), _t("判据", "Criterion", lang),
                   _t("结论", "Result", lang)],
               [[c["lecture"], c["name"][lang], c["value"][lang], c["limit"][lang], STATE[c["state"]][1 if lang == "en" else 0]] for c in R["checks"]])
    for i, c in enumerate(R["checks"]):
        cell = t.rows[i + 1].cells[4]
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor.from_string(STATE[c["state"]][2])
            run.font.bold = True
    head("运动曲线与时序", "Motion curves and timing")
    for title, png in figures(R, lang):
        d.add_picture(io.BytesIO(png), width=Cm(15))
        d.add_paragraph(title)
    head("零件参数", "Part parameters")
    rows = []
    for part in R["parts"]:
        for zh, en, v, unit in part["dims"]:
            rows.append([_t(part["zh"], part["en"], lang), _t(zh, en, lang), v, unit])
    _table(d, [_t("零件", "Part", lang), _t("尺寸", "Dimension", lang), _t("数值", "Value", lang), _t("单位", "Unit", lang)], rows)
    for tb in R.get("tables", []):
        head(tb["zh"], tb["en"])
        if tb["key"] == "plan":
            d.add_paragraph(_t(f"订单：每件缝长 {p['Lcm']:g} cm × {p['Q']:g} 件，损耗 {p['W'] * 100:g}%，每个线团 {p['cone']:g} m。",
                               f"Order: {p['Lcm']:g} cm of seam × {p['Q']:g} pieces, {p['W'] * 100:g}% waste, {p['cone']:g} m per cone.", lang))
        _table(d, [_t(a, b, lang) for a, b in tb["head"]], [[_t(*c, lang) if isinstance(c, (list, tuple)) else c for c in r] for r in tb["rows"]])
    if "chain" in R:
        head("制造：梭尖间隙尺寸链与公差建议", "Manufacturing: hook-clearance chain and tolerances")
        CH = R["chain"]
        d.add_paragraph(_t(f"梭尖间隙 0.07 ± 0.03 mm 由八个环决定。极值法 ±{CH['worst']:.3f} mm；统计法 σ = {CH['sigma']:.4f} mm，"
                           f"完全互换装配合格率 {CH['y_int'] * 100:.1f}%。装配时调整旋梭位置吸收四个环后 σ = {CH['sigma_adj']:.4f} mm，"
                           f"合格率 {CH['y_adj'] * 100:.2f}%；用户换针后 σ = {CH['sigma_swap']:.4f} mm，合格率 {CH['y_swap'] * 100:.1f}%。"
                           "建议：机壳与针杆套、旋梭轴孔按下表公差加工，装配时必须调整梭尖间隙，不要靠零件精度保证。",
                           f"The 0.07 ± 0.03 mm hook clearance depends on eight links. Worst case ±{CH['worst']:.3f} mm; statistical σ = {CH['sigma']:.4f} mm, "
                           f"yield with full interchangeability {CH['y_int'] * 100:.1f}%. Adjusting the hook at assembly absorbs four links: σ = {CH['sigma_adj']:.4f} mm, "
                           f"yield {CH['y_adj'] * 100:.2f}%; after a needle change σ = {CH['sigma_swap']:.4f} mm, yield {CH['y_swap'] * 100:.1f}%. "
                           "Recommendation: machine the parts to the tolerances below and always set the clearance at assembly.", lang))
        _table(d, [_t("环", "Link", lang), _t("公差 ±mm", "Tolerance ±mm", lang), _t("装配调整吸收", "Absorbed by adjustment", lang),
                   _t("换针时变化", "Changes with needle", lang)],
               [[_t(c[0], c[1], lang), c[2], "✓" if c[3] else "", "✓" if c[4] else ""] for c in CHAIN])
    if ebox is not None:
        from . import ebox as EBX
        ebox_section(d, head, EBX.analyze(ebox, machine, p), lang)
    if cae:
        head("CAE 分析结果（问渠数字工厂）", "CAE results (WenQuest digital factory)")
        d.add_paragraph(_t("以下是同一组设计参数在问渠数字工厂计算服务里算出的结果：有限元强度与疲劳（Gmsh + CalculiX + pyLife）、"
                           "机壳模态、机构动力学（MuJoCo）、电控箱散热（OpenFOAM 流体 + CalculiX 温度场）。详细报告在设计页“CAE 分析”里下载。",
                           "Results computed for this same set of design parameters by the WenQuest digital-factory computing service: "
                           "FEA strength and fatigue (Gmsh + CalculiX + pyLife), frame modes, mechanism dynamics (MuJoCo). "
                           "Full reports can be downloaded from the CAE tab of the design page.", lang))
        rows = []
        for x in cae:
            for c in (x["result"] or {}).get("checks", []):
                rows.append([x["title"], _t(c["zh"], c["en"], lang), c["value"], c["limit"], STATE[c["state"]][1 if lang == "en" else 0]])
        if rows:
            _table(d, [_t("分析", "Analysis", lang), _t("项目", "Item", lang), _t("数值", "Value", lang), _t("判据", "Criterion", lang),
                       _t("结论", "Result", lang)], rows)
    head("装配调机顺序", "Assembly adjustment sequence")
    for i, s_ in enumerate(R["card"], 1):
        d.add_paragraph(f"{i}. {_t(s_['zh'], s_['en'], lang)}：{_t(s_['do_zh'], s_['do_en'], lang)}")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def card_docx(params=None, lang="zh", project="", machine="lockstitch"):
    from docx.shared import Pt
    M = get(machine)
    p = M.clean(params)
    d = _doc()
    d.add_heading(_t(f"{M.NAME[0]}调机卡", f"{M.NAME[1]} Adjustment Card", lang), 0)
    basic = [x for x in M.GROUPS[0][3]]
    desc = "　".join(f"{_t(zh, en, lang)} {_choice(M, k, p[k], lang) if k in M.CHOICES else _fmt(p[k])}{(' ' + u) if u and k not in M.CHOICES else ''}"
                    for k, zh, en, u, *_ in basic)
    d.add_paragraph(_t(f"项目：{project or '未命名'}　", f"Project: {project or 'Untitled'}  ", lang) + desc)
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, h in enumerate([_t("顺序", "Step", lang), _t("项目", "Item", lang), _t("怎么调、调到多少", "How and to what", lang), _t("确认", "Done", lang)]):
        t.rows[0].cells[i].text = h
        _shade(t.rows[0].cells[i], "EAF2FA")
    for i, s in enumerate(M.adjust_card(p), 1):
        c = t.add_row().cells
        c[0].text, c[1].text, c[2].text, c[3].text = str(i), _t(s["zh"], s["en"], lang), _t(s["do_zh"], s["do_en"], lang), "□"
    d.add_paragraph()
    d.add_paragraph(_t("试缝：用本卡的面料与缝线缝 30 cm，检查跳针、断线、浮线、起皱、布边是否平整、针距。有问题按设计软件校核页的提示回到对应步骤。",
                       "Test seam: sew 30 cm with this fabric and thread; check for skips, breaks, loops, pucker, a flat edge and stitch length. "
                       "If anything is wrong, go back to the step named on the software's check page.", lang)).runs[0].font.size = Pt(9.5)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def params_xlsx(params=None, lang="zh", project="", machine="lockstitch", ebox=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    M = get(machine)
    R = M.analyze(params, trial=False)
    p = R["params"]
    wb = Workbook()
    ws = wb.active
    ws.title = _t("设计参数", "Parameters", lang)
    ws.append([_t("键", "Key", lang), _t("部分", "Group", lang), _t("参数", "Parameter", lang), _t("取值", "Value", lang), _t("单位", "Unit", lang)])
    for gk, gzh, gen, items in M.GROUPS:
        for k, zh, en, unit, *_ in items:
            ws.append([k, _t(gzh, gen, lang), _t(zh, en, lang), p[k], unit])
    ws2 = wb.create_sheet(_t("零件参数", "Parts", lang))
    ws2.append([_t("零件", "Part", lang), _t("尺寸", "Dimension", lang), _t("数值", "Value", lang), _t("单位", "Unit", lang)])
    for part in R["parts"]:
        for zh, en, v, unit in part["dims"]:
            ws2.append([_t(part["zh"], part["en"], lang), _t(zh, en, lang), v, unit])
    ws3 = wb.create_sheet(_t("校核", "Checks", lang))
    ws3.append([_t("讲次", "Lecture", lang), _t("校核项", "Check", lang), _t("数值", "Value", lang), _t("判据", "Criterion", lang), _t("结论", "Result", lang)])
    for c in R["checks"]:
        ws3.append([c["lecture"], c["name"][lang], c["value"][lang], c["limit"][lang], STATE[c["state"]][1 if lang == "en" else 0]])
        ws3.cell(ws3.max_row, 5).font = Font(bold=True, color=STATE[c["state"]][2])
    ws4 = wb.create_sheet(_t("运动曲线", "Curves", lang))
    lines = [c for c in R.get("charts", []) if c["kind"] == "line" and not c.get("xr")]
    if lines:
        cols = [(c, sr) for c in lines for sr in c["series"]]
        ws4.append([_t("主轴转角 °", "Shaft angle °", lang)] + [f"{_t(c['zh'], c['en'], lang)} · {sr[lang]}" for c, sr in cols])
        for i, x in enumerate(lines[0]["x"]):
            ws4.append([x] + [sr["y"][i] for _, sr in cols])
    sheets = [ws, ws2, ws3, ws4]
    if not lines:                                  # 横机等没有主轴转角曲线的机型：不留空页
        wb.remove(ws4)
        sheets.remove(ws4)
    for c in [c for c in R.get("charts", []) if c["kind"] == "line" and c.get("xr")]:      # 横轴不是转角的曲线（例如每针转速）：各占一页，全部数据
        w = wb.create_sheet(_t(c["zh"], c["en"], lang).split("（")[0].split(" (")[0][:30])
        w.append([c["xl"][lang]] + [sr[lang] for sr in c["series"]])
        for i, x in enumerate(c["x"]):
            w.append([x] + [sr["y"][i] for sr in c["series"]])
        sheets.append(w)
    for tb in R.get("tables", []):
        w = wb.create_sheet(_t(tb["zh"], tb["en"], lang).split("（")[0].split(" (")[0][:30])
        w.append([_t(a, b, lang) for a, b in tb["head"]])
        for r in tb["rows"]:
            w.append([_t(*c, lang) if isinstance(c, (list, tuple)) else c for c in r])
        sheets.append(w)
    if ebox is not None:                            # 第 62 轮：电控箱一页（参数、结果、校核）
        from . import ebox as EBX
        from .ebox import GROUPS as EG
        E = EBX.analyze(ebox, machine, p)
        w = wb.create_sheet(_t("电控箱", "Control box", lang))
        w.append([_t("键", "Key", lang), _t("部分", "Group", lang), _t("项目", "Item", lang), _t("数值", "Value", lang), _t("单位", "Unit", lang)])
        for gk, gzh, gen, items in EG:
            for k, zh, en, unit, *_ in items:
                w.append([k, _t(gzh, gen, lang), _t(zh, en, lang), E["params"][k], unit])
        for r in E["results"]:
            w.append(["", _t("结果", "Result", lang), _t(r["zh"], r["en"], lang), r["value"], r["unit"]])
        for c in E["checks"]:
            w.append([c["key"], _t("校核", "Check", lang), c["name"][lang], c["value"][lang], f"{c['limit'][lang]} · {STATE[c['state']][1 if lang == 'en' else 0]}"])
        sheets.append(w)
    for w in sheets:
        for cell in w[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EAF2FA")
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = max(10, min(40, max(len(str(c.value or "")) for c in col) * 1.6))
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()
