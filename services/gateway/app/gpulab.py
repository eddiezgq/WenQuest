"""Cloud GPU labs (《人工智能》第 14 轮附：云端 GPU 实验环境).

A student opens a GPU lab of a textbook: the gateway checks the student's GPU-hour quota, rents one GPU on a platform
(AutoDL in China, Lambda in the US; a provider adapter each), waits for its JupyterLab, uploads the lab notebook, the
`wqgpu` helper and a per-session key through JupyterLab's file interface, and hands the student the JupyterLab link.
The notebook's last cell sends the measured results back, signed with the session key; they fill the student's
lab report. A loop stops machines left idle (JupyterLab's last activity) or running too long, and keeps a ledger.

Quota (Q3: paid by the school, no personal top-up): a student's hours = the largest grant among the courses they are
enrolled in (set by the course's teacher) + an administrator's personal grant. One machine at a time per student.

The textbook build packs the GPU labs of a book into <WQ_TEXTBOOK_DIR>/<book>/gpulab/: index.json
{"labs": {"4.1": {"title": [zh, en], "tier": "basic", "notebook": "lab4_1.ipynb", "files": [...], "results": [...]}}}
and the files themselves.
"""


import asyncio
import hashlib
import hmac
import json
import logging
import re
import secrets
import sqlite3
import threading
import time
from pathlib import Path
from typing import Annotated, Any
from urllib.parse import urlsplit

from fastapi import Depends, Request, Response
from fastapi.responses import HTMLResponse
from html import escape as html_escape
from pydantic import BaseModel, Field

from .moodle import EngineError
from .session import Session

log = logging.getLogger("wenquest.gpulab")

BOOK = re.compile(r"^[a-z][a-z0-9_-]{0,40}$")
LAB = re.compile(r"^\d{1,3}\.\d{1,3}$")
NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_]{0,40}$")
TIERS = ("basic", "hopper", "profile")
LIVE = ("starting", "provisioning", "ready")
LINK_TTL = 12 * 3600
PLACE = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]{0,40})\s*\}\}")


# ---------------------------------------------------------------- providers

class Provider:
    """One GPU platform. Every method may raise EngineError("gpu_provider", ...)."""
    name = "base"

    async def launch(self, tier: str, label: str) -> str: ...
    async def status(self, pid: str) -> dict: ...          # {state: starting|running|stopped|failed, jupyter_url, jupyter_token, price_hour}
    async def stop(self, pid: str) -> None: ...


def _pairs(raw: str) -> dict[str, str]:
    """"basic=v-48g:base-image-x,hopper=h800:base-image-y" → {"basic": "v-48g:base-image-x", ...}"""
    out = {}
    for part in (raw or "").split(","):
        if "=" in part:
            k, v = part.split("=", 1)
            out[k.strip()] = v.strip()
    return out


async def _req(http, method: str, url: str, provider: str, **kw) -> Any:
    try:
        r = await http.request(method, url, timeout=30, **kw)
    except Exception as e:  # noqa: BLE001
        raise EngineError("gpu_provider", f"{provider} unreachable: {type(e).__name__}", 503) from e
    if r.status_code >= 400:
        raise EngineError("gpu_provider", f"{provider} {r.status_code}: {r.text[:200]}", 502)
    try:
        return r.json()
    except ValueError as e:
        raise EngineError("gpu_provider", f"{provider}: not JSON", 502) from e


class AutoDL(Provider):
    """AutoDL 容器实例 Pro (https://www.autodl.com/docs/instance_pro_api/): create → status / snapshot → power_off → release.
    Responses are {"code": "Success", "data": ...}."""
    name = "autodl"

    def __init__(self, http, token: str, url: str, specs: str, cuda: int):
        self.http, self.token, self.url, self.specs, self.cuda = http, token, url.rstrip("/"), _pairs(specs), cuda

    async def _call(self, method: str, path: str, **kw) -> Any:
        if not self.token:
            raise EngineError("gpu_unavailable", "AutoDL is not configured", 503)
        data = await _req(self.http, method, self.url + path, "AutoDL", headers={"Authorization": self.token}, **kw)
        if data.get("code") != "Success":
            raise EngineError("gpu_provider", f"AutoDL: {data.get('msg') or data.get('code')}", 502)
        return data.get("data")

    async def launch(self, tier: str, label: str) -> str:
        spec = self.specs.get(tier)
        if not spec or ":" not in spec:
            raise EngineError("gpu_unavailable", f"no AutoDL machine for {tier}", 503)
        gpu, image = spec.split(":", 1)
        data = await self._call("POST", "/api/v1/dev/instance/pro/create", json={
            "req_gpu_amount": 1, "gpu_spec_uuid": gpu, "image_uuid": image, "cuda_v_from": self.cuda,
            "expand_system_disk_by_gb": 0, "instance_name": label})
        return str(data)

    async def status(self, pid: str) -> dict:
        st = await self._call("GET", "/api/v1/dev/instance/pro/status", params={"instance_uuid": pid})
        state = {"running": "running", "shutdown": "stopped", "released": "stopped"}.get(str(st), "starting")
        if str(st) in ("failed", "create_failed"):
            state = "failed"
        out = {"state": state}
        if state == "running":
            snap = await self._call("GET", "/api/v1/dev/instance/pro/snapshot", params={"instance_uuid": pid})
            dom = str(snap.get("jupyter_domain") or "")
            out.update(jupyter_url=dom if dom.startswith("http") else f"https://{dom}" if dom else "",
                       jupyter_token=snap.get("jupyter_token") or "", price_hour=float(snap.get("payg_price") or 0))
        return out

    async def stop(self, pid: str) -> None:
        try:
            await self._call("POST", "/api/v1/dev/instance/pro/power_off", json={"instance_uuid": pid})
        except EngineError as e:
            log.warning("AutoDL power_off %s: %s", pid, e)
        for _ in range(10):                                   # release only after it is off
            st = await self._call("GET", "/api/v1/dev/instance/pro/status", params={"instance_uuid": pid})
            if str(st) in ("shutdown", "released"):
                break
            await asyncio.sleep(3)
        await self._call("POST", "/api/v1/dev/instance/pro/release", json={"instance_uuid": pid})


class Lambda(Provider):
    """Lambda Cloud API v1: POST /instance-operations/launch, GET /instances/{id}, POST /instance-operations/terminate."""
    name = "lambda"

    def __init__(self, http, key: str, url: str, types: str, region: str, ssh_key: str, prices: str):
        self.http, self.key, self.url = http, key, url.rstrip("/")
        self.types, self.region, self.ssh_key = _pairs(types), region, ssh_key
        self.prices = {k: float(v) for k, v in _pairs(prices).items() if re.fullmatch(r"[\d.]+", v)}

    async def _call(self, method: str, path: str, **kw) -> Any:
        if not self.key:
            raise EngineError("gpu_unavailable", "Lambda is not configured", 503)
        data = await _req(self.http, method, self.url + path, "Lambda", headers={"Authorization": f"Bearer {self.key}"}, **kw)
        return data.get("data")

    async def launch(self, tier: str, label: str) -> str:
        kind = self.types.get(tier)
        if not kind:
            raise EngineError("gpu_unavailable", f"no Lambda machine for {tier}", 503)
        data = await self._call("POST", "/instance-operations/launch", json={
            "region_name": self.region, "instance_type_name": kind, "ssh_key_names": [self.ssh_key] if self.ssh_key else [],
            "name": label[:60], "quantity": 1})
        return str((data.get("instance_ids") or [""])[0])

    async def status(self, pid: str) -> dict:
        inst = await self._call("GET", f"/instances/{pid}")
        st = inst.get("status")
        state = {"active": "running", "booting": "starting", "terminated": "stopped", "terminating": "stopped",
                 "unhealthy": "failed"}.get(st, "starting")
        out = {"state": state}
        if state == "running":
            kind = (inst.get("instance_type") or {}).get("name", "")
            url = str(inst.get("jupyter_url") or "")
            out.update(jupyter_url=url.split("?", 1)[0].rstrip("/"), jupyter_token=inst.get("jupyter_token") or "",
                       price_hour=self.prices.get(kind, 0.0))
        return out

    async def stop(self, pid: str) -> None:
        await self._call("POST", "/instance-operations/terminate", json={"instance_ids": [pid]})


class Mock(Provider):
    """For tests and the local environment: a machine that is up after `delay` status calls."""
    name = "mock"

    def __init__(self, jupyter_url: str = "http://jupyter.mock", delay: int = 1):
        self.machines: dict[str, dict] = {}
        self.jupyter_url, self.delay = jupyter_url, delay

    async def launch(self, tier: str, label: str) -> str:
        pid = f"mock-{len(self.machines) + 1}"
        self.machines[pid] = {"tier": tier, "label": label, "calls": 0, "on": True}
        return pid

    async def status(self, pid: str) -> dict:
        m = self.machines.get(pid)
        if not m or not m["on"]:
            return {"state": "stopped"}
        m["calls"] += 1
        if m["calls"] <= self.delay:
            return {"state": "starting"}
        return {"state": "running", "jupyter_url": self.jupyter_url, "jupyter_token": f"tok-{pid}", "price_hour": 2.68}

    async def stop(self, pid: str) -> None:
        if pid in self.machines:
            self.machines[pid]["on"] = False


def provider_of(settings, http) -> Provider:
    name = settings.gpu_provider
    if name == "autodl":
        return AutoDL(http, settings.gpu_autodl_token, settings.gpu_autodl_url, settings.gpu_autodl_specs, settings.gpu_autodl_cuda)
    if name == "lambda":
        return Lambda(http, settings.gpu_lambda_key, settings.gpu_lambda_url, settings.gpu_lambda_types, settings.gpu_lambda_region,
                      settings.gpu_lambda_ssh_key, settings.gpu_lambda_prices)
    if name == "mock":
        return Mock()
    raise EngineError("gpu_unavailable", "GPU labs are not set up on this platform", 503)


# ---------------------------------------------------------------- store

class Store:
    def __init__(self, folder: str):
        Path(folder).mkdir(parents=True, exist_ok=True)
        self.path = str(Path(folder) / "gpulab.db")
        self.lock = threading.Lock()
        self._db()

    def _db(self):
        c = sqlite3.connect(self.path)
        c.row_factory = sqlite3.Row
        with c:
            c.execute("create table if not exists course_grant (course_id integer primary key, hours real not null, set_by integer, at integer)")
            c.execute("create table if not exists user_grant (user_id integer primary key, hours real not null, set_by integer, at integer)")
            c.execute("create table if not exists session (id text primary key, user_id integer not null, book text, lab text, tier text,"
                      " provider text, pid text, state text, key text, jupyter_url text, jupyter_token text, price_hour real,"
                      " started_at integer, ready_at integer, ended_at integer, last_active integer, end_reason text)")
            c.execute("create table if not exists result (session_id text, user_id integer, book text, lab text, at integer,"
                      " vals text, env text)")
        return c

    def grant_course(self, course_id: int, hours: float, by: int) -> None:
        with self.lock, self._db() as c:
            c.execute("insert into course_grant values (?, ?, ?, ?) on conflict (course_id) do update set hours=excluded.hours,"
                      " set_by=excluded.set_by, at=excluded.at", (course_id, hours, by, int(time.time())))

    def grant_user(self, user_id: int, hours: float, by: int) -> None:
        with self.lock, self._db() as c:
            c.execute("insert into user_grant values (?, ?, ?, ?) on conflict (user_id) do update set hours=excluded.hours,"
                      " set_by=excluded.set_by, at=excluded.at", (user_id, hours, by, int(time.time())))

    def course_hours(self, course_ids: list[int]) -> float:
        if not course_ids:
            return 0.0
        with self._db() as c:
            q = ",".join("?" * len(course_ids))
            r = c.execute(f"select max(hours) from course_grant where course_id in ({q})", course_ids).fetchone()
            return float(r[0] or 0)

    def course_grant(self, course_id: int) -> float:
        with self._db() as c:
            r = c.execute("select hours from course_grant where course_id=?", (course_id,)).fetchone()
            return float(r[0]) if r else 0.0

    def user_hours(self, user_id: int) -> float:
        with self._db() as c:
            r = c.execute("select hours from user_grant where user_id=?", (user_id,)).fetchone()
            return float(r[0]) if r else 0.0

    def used_hours(self, user_id: int, now: int | None = None) -> float:
        now = now or int(time.time())
        with self._db() as c:
            rows = c.execute("select started_at, ended_at from session where user_id=?", (user_id,)).fetchall()
        return sum(((r["ended_at"] or now) - r["started_at"]) for r in rows) / 3600

    def new(self, **row: Any) -> None:
        with self.lock, self._db() as c:
            c.execute("insert into session (id, user_id, book, lab, tier, provider, pid, state, key, started_at, last_active)"
                      " values (:id, :user_id, :book, :lab, :tier, :provider, :pid, :state, :key, :started_at, :started_at)", row)

    def update(self, sid: str, **cols: Any) -> None:
        if not cols:
            return
        with self.lock, self._db() as c:
            c.execute("update session set " + ", ".join(f"{k}=?" for k in cols) + " where id=?", (*cols.values(), sid))

    def get(self, sid: str) -> dict | None:
        with self._db() as c:
            r = c.execute("select * from session where id=?", (sid,)).fetchone()
            return dict(r) if r else None

    def live(self, user_id: int | None = None) -> list[dict]:
        with self._db() as c:
            q = "select * from session where state in ('starting','provisioning','ready')"
            rows = c.execute(q + (" and user_id=?" if user_id is not None else ""), (() if user_id is None else (user_id,))).fetchall()
            return [dict(r) for r in rows]

    def of_users(self, user_ids: list[int]) -> list[dict]:
        if not user_ids:
            return []
        with self._db() as c:
            q = ",".join("?" * len(user_ids))
            return [dict(r) for r in c.execute(f"select * from session where user_id in ({q}) order by started_at", user_ids)]

    def all(self, limit: int = 500) -> list[dict]:
        with self._db() as c:
            return [dict(r) for r in c.execute("select * from session order by started_at desc limit ?", (limit,))]

    def add_result(self, sess: dict, vals: dict, env: dict) -> None:
        with self.lock, self._db() as c:
            c.execute("insert into result values (?, ?, ?, ?, ?, ?, ?)", (sess["id"], sess["user_id"], sess["book"], sess["lab"],
                      int(time.time()), json.dumps(vals, ensure_ascii=False), json.dumps(env, ensure_ascii=False)))

    def results(self, user_id: int, book: str, lab: str) -> list[dict]:
        with self._db() as c:
            rows = c.execute("select * from result where user_id=? and book=? and lab=? order by at", (user_id, book, lab)).fetchall()
        return [dict(r, vals=json.loads(r["vals"]), env=json.loads(r["env"])) for r in rows]


def cost(s: dict, now: int | None = None) -> float:
    end = s.get("ended_at") or now or int(time.time())
    return round((end - s["started_at"]) / 3600 * float(s.get("price_hour") or 0), 2)


def fmt(v: Any) -> str:
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, int) or float(v).is_integer() and abs(v) < 1e9:
        return str(int(v))
    return f"{v:.4g}"


def fill_report(docx: bytes, vals: dict) -> bytes:
    """Replace {{name}} in the report's paragraphs and table cells with the measured values ("—" when not measured)."""
    import io

    from docx import Document
    doc = Document(io.BytesIO(docx))

    def para(p):
        text = "".join(r.text for r in p.runs)
        if "{{" not in text:
            return
        new = PLACE.sub(lambda m_: fmt(vals[m_.group(1)]) if m_.group(1) in vals else "—", text)
        for i, run in enumerate(p.runs):
            run.text = new if i == 0 else ""

    for p in doc.paragraphs:
        para(p)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    para(p)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def sign(key: str, body: bytes) -> str:
    return hmac.new(key.encode(), body, hashlib.sha256).hexdigest()


# ---------------------------------------------------------------- JupyterLab file interface

async def upload(http, url: str, token: str, files: list[tuple[str, Any]]) -> None:
    """files: [(path, str | dict notebook)], paths relative to the JupyterLab root; folders are created first."""
    head = {"Authorization": f"token {token}"}
    base = url.rstrip("/") + "/api/contents/"
    dirs = sorted({"/".join(p.split("/")[:i]) for p, _ in files for i in range(1, p.count("/") + 1)})
    for d in dirs:
        await _req(http, "PUT", base + d, "JupyterLab", headers=head, json={"type": "directory"})
    for path, content in files:
        body = ({"type": "notebook", "format": "json", "content": content} if isinstance(content, dict)
                else {"type": "file", "format": "text", "content": content})
        await _req(http, "PUT", base + path, "JupyterLab", headers=head, json=body)


async def last_activity(http, url: str, token: str) -> int | None:
    """JupyterLab's /api/status last_activity (ISO time) as unix seconds; None if unknown."""
    try:
        data = await _req(http, "GET", url.rstrip("/") + "/api/status", "JupyterLab", headers={"Authorization": f"token {token}"})
    except EngineError:
        return None
    from datetime import datetime
    v = str(data.get("last_activity") or "")
    try:
        return int(datetime.fromisoformat(v.replace("Z", "+00:00")).timestamp())
    except ValueError:
        return None


# ---------------------------------------------------------------- API

class StartIn(BaseModel):
    book: str = Field(max_length=40)
    lab: str = Field(max_length=10)


class GrantIn(BaseModel):
    hours: float = Field(ge=0, le=1000)


class UserGrantIn(BaseModel):
    user_id: int = Field(ge=1)
    hours: float = Field(ge=0, le=1000)


class SubmitIn(BaseModel):
    sid: str = Field(max_length=64)
    values: dict[str, Any] = Field(default_factory=dict)
    env: dict[str, Any] = Field(default_factory=dict)


def public(s: dict) -> dict:
    out = {k: s[k] for k in ("id", "book", "lab", "tier", "provider", "state", "started_at", "ready_at", "ended_at", "end_reason")}
    out["cost"] = cost(s)
    if s["state"] == "ready":
        out["url"] = f"{s['jupyter_url'].rstrip('/')}/lab/tree/wq/{s['lab'].replace('.', '_')}?token={s['jupyter_token']}"
    return out


def register(app, m) -> None:
    current = m.current
    state = m.state
    _store: dict[str, Store] = {}
    starts: dict[str, asyncio.Lock] = {}

    def store() -> Store:
        k = state.settings.data_dir
        if k not in _store:
            _store[k] = Store(str(Path(k) / "gpulab"))
        return _store[k]

    def provider() -> Provider:
        p = getattr(state, "gpu_provider", None)
        return p if p is not None else provider_of(state.settings, state.http)

    def catalog(book: str) -> dict:
        if not BOOK.match(book):
            raise EngineError("not_found", "no such textbook", 404)
        p = Path(state.settings.textbook_dir) / book / "gpulab" / "index.json"
        if not p.exists():
            return {"labs": {}}
        return json.loads(p.read_text(encoding="utf-8"))

    def lab_of(book: str, lab: str) -> dict:
        if not LAB.match(lab):
            raise EngineError("not_found", "no such GPU lab", 404)
        spec = catalog(book)["labs"].get(lab)
        if not spec:
            raise EngineError("not_found", "no such GPU lab", 404)
        return spec

    async def quota(sess: Session) -> float:
        try:
            courses = await state.moodle.user_courses(sess.moodle_token, sess.user_id)
        except EngineError:
            courses = []
        return store().course_hours([int(c["id"]) for c in courses]) + store().user_hours(sess.user_id)

    async def finish(s: dict, reason: str) -> None:
        try:
            await provider().stop(s["pid"])
        except EngineError as e:
            log.warning("stopping %s (%s): %s", s["id"], s["pid"], e)
        store().update(s["id"], state="ended", ended_at=int(time.time()), end_reason=reason)

    async def advance(s: dict) -> dict:
        """Move a starting machine on: wait for JupyterLab, upload the lab, mark it ready."""
        if s["state"] not in ("starting", "provisioning"):
            return s
        st = await provider().status(s["pid"])
        if st["state"] in ("stopped", "failed"):
            store().update(s["id"], state="ended", ended_at=int(time.time()), end_reason=f"provider_{st['state']}")
            return store().get(s["id"])
        if st["state"] != "running" or not st.get("jupyter_url"):
            return s
        store().update(s["id"], state="provisioning", jupyter_url=st["jupyter_url"], jupyter_token=st.get("jupyter_token", ""),
                       price_hour=st.get("price_hour") or 0)
        spec = lab_of(s["book"], s["lab"])
        folder = Path(state.settings.textbook_dir) / s["book"] / "gpulab"
        files: list[tuple[str, Any]] = []
        for name in [spec["notebook"], *spec.get("files", [])]:
            if ".." in name or name.startswith("/"):
                continue
            data = (folder / name).read_text(encoding="utf-8")
            files.append((f"wq/{s['lab'].replace('.', '_')}/{Path(name).name}", json.loads(data) if name.endswith(".ipynb") else data))
        helper = folder / "wqgpu.py"                                   # the textbook build copies it next to the labs
        if not helper.exists():
            helper = Path(__file__).resolve().parents[3] / "deploy" / "gpu-lab" / "wqgpu.py"
        files.append(("wq/wqgpu.py", helper.read_text(encoding="utf-8")))
        cb = (state.settings.public_url or state.settings.app_url).rstrip("/") + "/api/v1/gpulab/submit"
        files.append(("wq/session.json", json.dumps({"sid": s["id"], "key": s["key"], "callback": cb, "book": s["book"],
                                                     "lab": s["lab"], "results": spec.get("results", [])}, ensure_ascii=False)))
        try:
            await upload(state.http, st["jupyter_url"], st.get("jupyter_token", ""), files)
        except EngineError as e:
            log.info("provisioning %s not yet: %s", s["id"], e)          # JupyterLab may still be starting; try again next poll
            return store().get(s["id"])
        store().update(s["id"], state="ready", ready_at=int(time.time()), last_active=int(time.time()))
        return store().get(s["id"])

    async def me_info(sess: Session) -> dict:
        q = await quota(sess)
        used = store().used_hours(sess.user_id)
        live = store().live(sess.user_id)
        return {"quota_hours": round(q, 2), "used_hours": round(used, 2), "left_hours": round(max(0.0, q - used), 2),
                "active": public(live[0]) if live else None, "provider": state.settings.gpu_provider or ""}

    @app.get("/api/v1/gpulab/me")
    async def me(sess: Annotated[Session, Depends(current)]):
        return await me_info(sess)

    @app.get("/api/v1/gpulab/labs/{book}")
    async def labs(book: str, sess: Annotated[Session, Depends(current)]):
        return {k: {"title": v.get("title"), "tier": v.get("tier"), "results": v.get("results", [])} for k, v in catalog(book)["labs"].items()}

    async def do_start(sess: Session, body: StartIn) -> dict:
        spec = lab_of(body.book, body.lab)
        lock = starts.setdefault(str(sess.user_id), asyncio.Lock())
        async with lock:
            live = store().live(sess.user_id)
            if live:
                raise EngineError("gpu_busy", "you already have a GPU machine running; stop it first", 409)
            q = await quota(sess)
            if q - store().used_hours(sess.user_id) < 0.05:
                raise EngineError("gpu_quota", "no GPU hours left; ask your teacher", 403)
            sid = secrets.token_hex(8)
            pid = await provider().launch(spec.get("tier", "basic"), f"wq-{sess.user_id}-{body.lab}-{sid[:6]}")
            store().new(id=sid, user_id=sess.user_id, book=body.book, lab=body.lab, tier=spec.get("tier", "basic"),
                        provider=provider().name, pid=pid, state="starting", key=secrets.token_hex(16), started_at=int(time.time()))
        return public(store().get(sid))

    @app.post("/api/v1/gpulab/sessions")
    async def start(body: StartIn, sess: Annotated[Session, Depends(current)]):
        return await do_start(sess, body)

    def own(sid: str, sess: Session) -> dict:
        s = store().get(sid)
        if not s or s["user_id"] != sess.user_id:
            raise EngineError("not_found", "no such GPU session", 404)
        return s

    @app.get("/api/v1/gpulab/sessions/{sid}")
    async def poll(sid: str, sess: Annotated[Session, Depends(current)]):
        return public(await advance(own(sid, sess)))

    async def do_stop(sess: Session, sid: str) -> dict:
        s = own(sid, sess)
        if s["state"] in LIVE:
            await finish(s, "student")
        return public(store().get(sid))

    @app.post("/api/v1/gpulab/sessions/{sid}/stop")
    async def stop(sid: str, sess: Annotated[Session, Depends(current)]):
        return await do_stop(sess, sid)

    @app.post("/api/v1/gpulab/submit")
    async def submit(request: Request):
        raw = await request.body()
        try:
            body = SubmitIn.model_validate_json(raw)
        except ValueError as e:
            raise EngineError("bad_request", "malformed results", 400) from e
        s = store().get(body.sid)
        if not s or not hmac.compare_digest(sign(s["key"], raw), request.headers.get("X-WQ-Signature", "")):
            raise EngineError("forbidden", "bad signature", 403)
        allowed = set(lab_of(s["book"], s["lab"]).get("results", []))
        vals = {k: v for k, v in body.values.items() if NAME.match(k) and (not allowed or k in allowed)
                and isinstance(v, (int, float, str, bool)) and len(str(v)) <= 200}
        env = {k: str(v)[:200] for k, v in list(body.env.items())[:30] if NAME.match(k)}
        store().add_result(s, vals, env)
        store().update(s["id"], last_active=int(time.time()))
        return {"ok": True, "kept": sorted(vals)}

    @app.get("/api/v1/gpulab/results/{book}/{lab}")
    async def my_results(book: str, lab: str, sess: Annotated[Session, Depends(current)]):
        lab_of(book, lab)
        return {"results": [{k: r[k] for k in ("at", "vals", "env")} for r in store().results(sess.user_id, book, lab)]}

    # teachers: the course's grant and its students' usage
    async def need_teacher(sess: Session, courseid: int) -> None:
        try:
            ok = await state.moodle.can_edit_course(sess.moodle_token, courseid)
        except EngineError:
            ok = False
        if not ok:
            raise EngineError("forbidden", "teachers of this course only", 403)

    @app.put("/api/v1/courses/{courseid}/gpulab/grant")
    async def set_grant(courseid: int, body: GrantIn, sess: Annotated[Session, Depends(current)]):
        await need_teacher(sess, courseid)
        store().grant_course(courseid, body.hours, sess.user_id)
        return {"course_id": courseid, "hours": body.hours}

    @app.get("/api/v1/courses/{courseid}/gpulab/usage")
    async def usage(courseid: int, sess: Annotated[Session, Depends(current)]):
        await need_teacher(sess, courseid)
        users = await state.moodle.call(sess.moodle_token, "core_enrol_get_enrolled_users", None, courseid=courseid)
        names = {int(u["id"]): u.get("fullname", "") for u in users or []}
        rows = store().of_users(list(names))
        per: dict[int, dict] = {}
        for s in rows:
            p = per.setdefault(s["user_id"], {"user_id": s["user_id"], "name": names.get(s["user_id"], ""), "hours": 0.0, "cost": 0.0, "sessions": 0})
            p["hours"] += ((s["ended_at"] or int(time.time())) - s["started_at"]) / 3600
            p["cost"] += cost(s)
            p["sessions"] += 1
        return {"grant_hours": store().course_grant(courseid),
                "students": [dict(p, hours=round(p["hours"], 2), cost=round(p["cost"], 2)) for p in per.values()]}

    # administrators: personal grants and the whole ledger
    async def need_admin(sess: Session) -> None:
        if not await m.accounts_is_admin(sess):
            raise EngineError("forbidden", "administrators only", 403)

    @app.put("/api/v1/gpulab/admin/grant")
    async def admin_grant(body: UserGrantIn, sess: Annotated[Session, Depends(current)]):
        await need_admin(sess)
        store().grant_user(body.user_id, body.hours, sess.user_id)
        return {"user_id": body.user_id, "hours": body.hours}

    @app.get("/api/v1/gpulab/admin/ledger")
    async def ledger(sess: Annotated[Session, Depends(current)]):
        await need_admin(sess)
        rows = store().all()
        return {"sessions": [dict(public(s), user_id=s["user_id"], pid=s["pid"], price_hour=s["price_hour"]) for s in rows],
                "total_cost": round(sum(cost(s) for s in rows), 2)}

    # ---- the lab's own page, opened from the textbook (a personal link signed by the gateway, valid 12 hours)
    def link_of(signed: str) -> tuple[Session, str, str]:
        try:
            d = json.loads(state.codec.fernet.decrypt(signed.encode(), ttl=LINK_TTL))
        except Exception as e:  # noqa: BLE001
            raise EngineError("link_expired", "link expired; open the lab again from the textbook", 410) from e
        sess = state.codec.read(d.get("s", ""))
        if not sess or not BOOK.match(d.get("b", "")) or not LAB.match(d.get("l", "")):
            raise EngineError("link_expired", "link expired; open the lab again from the textbook", 410)
        return sess, d["b"], d["l"]

    def link(sess: Session, book: str, lab: str) -> str:
        tok = state.codec.fernet.encrypt(json.dumps({"s": state.codec.issue(sess), "b": book, "l": lab}).encode()).decode()
        return f"{state.settings.public_url.rstrip('/')}/api/v1/textbook-gpulab/{tok}"

    m.gpulab_link = link
    m.gpulab_catalog = catalog

    @app.get("/api/v1/textbook-gpulab/{signed}")
    async def lab_page(signed: str):
        sess, book, lab = link_of(signed)
        spec = lab_of(book, lab)
        en = sess.lang == "en"
        title = (spec.get("title") or ["", ""])[1 if en else 0]
        page = PAGE.replace("__TITLE__", html_escape(f"{'GPU lab' if en else 'GPU 实验'} {lab}　{title}"))
        page = page.replace("__LANG__", "en" if en else "zh").replace("__BASE__", json.dumps(f"/api/v1/textbook-gpulab/{signed}"))
        page = page.replace("__RESULTS__", json.dumps(spec.get("results", [])))
        return HTMLResponse(page, headers={"Cache-Control": "no-store", "Referrer-Policy": "no-referrer"})

    @app.get("/api/v1/textbook-gpulab/{signed}/state")
    async def link_state(signed: str):
        sess, book, lab = link_of(signed)
        info = await me_info(sess)
        if info["active"]:
            info["active"] = public(await advance(store().get(info["active"]["id"])))
        info["results"] = [{k: r[k] for k in ("at", "vals", "env")} for r in store().results(sess.user_id, book, lab)]
        return info

    @app.post("/api/v1/textbook-gpulab/{signed}/start")
    async def link_start(signed: str):
        sess, book, lab = link_of(signed)
        return await do_start(sess, StartIn(book=book, lab=lab))

    @app.post("/api/v1/textbook-gpulab/{signed}/stop/{sid}")
    async def link_stop(signed: str, sid: str):
        sess, _, _ = link_of(signed)
        return await do_stop(sess, sid)

    @app.get("/api/v1/textbook-gpulab/{signed}/report")
    async def link_report(signed: str):
        """The lab report template with this student's latest measured values filled into the {{name}} cells."""
        sess, book, lab = link_of(signed)
        en = sess.lang == "en"
        stem = f"lab{lab.replace('.', '_')}-report{'.en' if en else ''}.docx"
        p = Path(state.settings.textbook_dir) / book / "gpulab" / stem
        if not p.exists():
            raise EngineError("not_found", "no report template for this lab", 404)
        res = store().results(sess.user_id, book, lab)
        vals = res[-1]["vals"] if res else {}
        data = fill_report(p.read_bytes(), vals)
        return Response(data, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        headers={"Content-Disposition": f'attachment; filename="{stem}"', "Cache-Control": "no-store"})

    m.gpulab_reap = lambda: reap(m)
    m.gpulab_store = store


async def reap(m) -> None:
    """Stop machines past the time limit or idle in JupyterLab; let starting machines move on."""
    state = m.state
    st = m.gpulab_store()
    now = int(time.time())
    cfg = state.settings
    prov = getattr(state, "gpu_provider", None) or provider_of(cfg, state.http)
    for s in st.live():
        reason = None
        if now - s["started_at"] > cfg.gpu_max_hours * 3600:
            reason = "time_limit"
        elif s["state"] == "ready":
            act = await last_activity(state.http, s["jupyter_url"], s["jupyter_token"])
            last = max(act or 0, s["last_active"] or 0)
            st.update(s["id"], last_active=last)
            if now - last > cfg.gpu_idle_minutes * 60:
                reason = "idle"
        elif now - s["started_at"] > 20 * 60:
            reason = "never_ready"
        if reason:
            try:
                await prov.stop(s["pid"])
            except EngineError as e:
                log.warning("reaper: stopping %s: %s", s["id"], e)
            st.update(s["id"], state="ended", ended_at=now, end_reason=reason)


async def reap_loop(m) -> None:
    await asyncio.sleep(30)
    while True:
        if m.state.settings.gpu_provider:
            try:
                await reap(m)
            except Exception:  # noqa: BLE001 - keep the loop alive
                log.exception("gpu lab reaper")
        await asyncio.sleep(60)


# The lab's control page (served by the gateway, so it needs nothing from outside). Texts in both languages.
PAGE = """<!doctype html><html lang="__LANG__"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>__TITLE__</title><style>
body{margin:0;background:#f6f7f6;color:#1d2327;font:15px/1.7 "Noto Sans CJK SC","PingFang SC",system-ui,sans-serif}
main{max-width:760px;margin:0 auto;padding:24px 16px}h1{font-size:20px;margin:0 0 6px}
.card{background:#fff;border:1px solid #d9dee2;border-radius:10px;padding:16px;margin:14px 0}
.muted{color:#5d6b73;font-size:13.5px}button,a.btn{display:inline-block;border:0;border-radius:8px;padding:10px 18px;font:inherit;cursor:pointer;text-decoration:none}
.go{background:#b8860b;color:#fff}.stop{background:#fff;color:#a12;border:1px solid #a12}.open{background:#1d2327;color:#fff}
button:disabled{opacity:.5;cursor:default}table{border-collapse:collapse;width:100%;font-size:14px}td,th{border:1px solid #d9dee2;padding:5px 8px;text-align:left}
.err{color:#a12}</style></head><body><main>
<h1>__TITLE__</h1><div class="muted" id="quota"></div>
<div class="card" id="box"><div id="msg"></div><p id="acts"></p></div>
<div class="card"><b id="rt"></b><div id="res" class="muted"></div><p><a class="btn open" id="rep" href="#">…</a></p></div>
<p class="muted" id="note"></p></main><script>
const BASE=__BASE__, RESULTS=__RESULTS__, EN=document.documentElement.lang==="en";
const T=(z,e)=>EN?e:z, $=id=>document.getElementById(id);
let timer=null, busy=false;
$("rt").textContent=T("我的测量结果","My results");$("rep").textContent=T("下载填好数据的实验报告（Word）","Download my lab report (Word)");$("rep").href=BASE+"/report";
$("note").textContent=T("机器按秒计费：做完请点“关机”。空闲 30 分钟或开机满 3 小时会自动关机，笔记本里的文件随之删除，请把需要的代码下载保存。","Machines are billed by the second: press Stop when done. They stop by themselves after 30 idle minutes or 3 hours, and their files are deleted, so download what you want to keep.");
const ERR={gpu_quota:T("没有剩余的 GPU 时长，请联系任课老师。","No GPU hours left; please ask your teacher."),gpu_busy:T("你已经有一台机器在运行，请先关机。","You already have a machine running; stop it first."),
 gpu_unavailable:T("本平台还没有接通 GPU，请联系管理员。","GPU labs are not set up on this platform yet."),gpu_provider:T("GPU 平台暂时没有响应或没有空闲的显卡，请稍后再试。","The GPU platform did not respond or has no free GPU; try again later."),
 link_expired:T("链接已过期，请回到教材重新打开本实验。","The link has expired; open the lab again from the textbook.")};
async function call(path,method){const r=await fetch(BASE+path,{method:method||"GET"});const j=await r.json().catch(()=>({}));if(!r.ok)throw new Error(ERR[j.error]||j.detail||r.status);return j}
function btn(text,cls,fn){const b=document.createElement("button");b.className=cls;b.textContent=text;b.onclick=async()=>{if(busy)return;busy=true;b.disabled=true;try{await fn()}catch(e){$("msg").innerHTML='<span class="err">'+e.message+'</span>'}busy=false;refresh()};return b}
function show(st){
 $("quota").textContent=T(`GPU 时长：已用 ${st.used_hours} / 额度 ${st.quota_hours} 小时`,`GPU hours: ${st.used_hours} used of ${st.quota_hours}`);
 const a=st.active, acts=$("acts");acts.innerHTML="";
 if(!a){$("msg").textContent=st.left_hours>0?T("还没有开机。开机后约一两分钟可以进入 JupyterLab。","No machine yet. JupyterLab is ready a minute or two after you start one."):T("没有剩余的 GPU 时长，请联系任课老师。","No GPU hours left; please ask your teacher.");
  if(st.left_hours>0)acts.appendChild(btn(T("开机","Start a GPU machine"),"go",()=>call("/start","POST")));return}
 const names={starting:T("正在开机……","Starting…"),provisioning:T("正在准备实验文件……","Preparing the lab files…"),ready:T("已就绪","Ready")};
 const cur=st.provider==="lambda"?"US$":"¥";$("msg").textContent=(names[a.state]||a.state)+T(`（已用 ${cur}${a.cost}）`,` (cost so far ${cur}${a.cost})`);
 if(a.state==="ready"){const o=document.createElement("a");o.className="btn open";o.href=a.url;o.target="_blank";o.rel="noopener";o.textContent=T("进入 JupyterLab","Open JupyterLab");acts.appendChild(o);acts.appendChild(document.createTextNode(" "))}
 acts.appendChild(btn(T("关机","Stop"),"stop",()=>call("/stop/"+a.id,"POST")));
}
function results(rs){if(!rs.length){$("res").textContent=T("还没有回传的结果：运行笔记本的最后一格后会出现在这里。","No results yet: they appear here after you run the notebook's last cell.");return}
 const last=rs[rs.length-1];let h="<table><tr><th>"+T("结果","Result")+"</th><th>"+T("数值","Value")+"</th></tr>";
 for(const k of (RESULTS.length?RESULTS:Object.keys(last.vals)))h+="<tr><td>"+k+"</td><td>"+(k in last.vals?last.vals[k]:"—")+"</td></tr>";
 h+="</table><p>"+T("测量环境：","Measured on: ")+[last.env.gpu,last.env.driver&&("driver "+last.env.driver),last.env.nvcc&&("CUDA "+last.env.nvcc)].filter(Boolean).join("，")+"</p>";$("res").innerHTML=h}
async function refresh(){clearTimeout(timer);try{const st=await call("/state");show(st);results(st.results);timer=setTimeout(refresh,st.active&&st.active.state!=="ready"?4000:20000)}catch(e){$("msg").innerHTML='<span class="err">'+e.message+'</span>';timer=setTimeout(refresh,20000)}}
refresh();
</script></body></html>"""
