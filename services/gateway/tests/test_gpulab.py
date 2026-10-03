"""Cloud GPU labs (《人工智能》第 14 轮附): quota from the course grant, start → JupyterLab provisioning → signed results →
stop; the reaper; the teacher's usage; the AutoDL and Lambda adapters against their documented interfaces."""
import asyncio
import importlib.util
import json
import time
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import httpx
import pytest
from fastapi.testclient import TestClient

from app import gpulab, main
from app.ai import ModelGateway
from app.moodle import EngineError, MoodleClient
from tests.test_materials import BASE, fake_moodle, login

ROOT = Path(__file__).resolve().parents[3]
spec = importlib.util.spec_from_file_location("wqgpu", ROOT / "deploy" / "gpu-lab" / "wqgpu.py")
wqgpu = importlib.util.module_from_spec(spec)
spec.loader.exec_module(wqgpu)

JUPYTER = {"files": {}, "last": "2026-10-03T00:00:00Z"}
NB = {"cells": [], "metadata": {}, "nbformat": 4, "nbformat_minor": 5}


def handler(request: httpx.Request) -> httpx.Response:
    url = urlparse(str(request.url))
    if url.netloc == "jupyter.mock":
        if request.headers.get("authorization", "").startswith("token tok-") is False:
            return httpx.Response(403, json={"message": "no"})
        if url.path.startswith("/api/contents/") and request.method == "PUT":
            JUPYTER["files"][url.path[len("/api/contents/"):]] = json.loads(request.content)
            return httpx.Response(201, json={"ok": True})
        if url.path == "/api/status":
            return httpx.Response(200, json={"last_activity": JUPYTER["last"]})
        return httpx.Response(404, json={})
    q = parse_qs(url.query)
    fn = (q.get("wsfunction") or [""])[0]
    tok = (q.get("wstoken") or [""])[0]
    if fn == "core_enrol_get_users_courses":
        return httpx.Response(200, json=[{"id": 7, "fullname": "人工智能"}])
    if fn == "core_course_get_user_administration_options":
        return httpx.Response(200, json={"courses": [{"id": 7, "options": [{"name": "update", "available": tok == "teacher"}]}]})
    if fn == "core_enrol_get_enrolled_users":
        return httpx.Response(200, json=[{"id": 6, "fullname": "学生甲"}, {"id": 5, "fullname": "老师"}])
    return fake_moodle(request)


@pytest.fixture
def client(tmp_path):
    lab = tmp_path / "books" / "ai" / "gpulab"
    lab.mkdir(parents=True)
    (lab / "index.json").write_text(json.dumps({"labs": {"4.1": {"title": ["第一个核函数", "First kernel"], "tier": "basic",
                                    "notebook": "lab4_1.ipynb", "files": ["vadd.cu"], "results": ["t_ms", "bw_GBs"]}}}), encoding="utf-8")
    (lab / "lab4_1.ipynb").write_text(json.dumps(NB), encoding="utf-8")
    (lab / "vadd.cu").write_text("__global__ void k() {}\n", encoding="utf-8")
    with TestClient(main.app) as c:
        st = main.state.settings
        saved = {k: getattr(st, k) for k in ("textbook_dir", "data_dir", "gpu_provider", "public_url")}
        st.textbook_dir, st.data_dir, st.gpu_provider, st.public_url = str(tmp_path / "books"), str(tmp_path / "data"), "mock", "https://learn.test"
        main.state.http = httpx.AsyncClient(transport=httpx.MockTransport(handler))
        main.state.moodle = MoodleClient(BASE, "moodle_mobile_app", main.state.http)
        main.state.ai = ModelGateway("fake", main.state.http)
        main.state.gpu_provider = gpulab.Mock()
        JUPYTER["files"].clear()
        JUPYTER["last"] = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
        yield c
        main.state.gpu_provider = None
        for k, v in saved.items():
            setattr(st, k, v)


def test_quota_start_provision_submit_stop(client):
    s, t = login(client, "s"), login(client, "t")
    me = client.get("/api/v1/gpulab/me", headers=s).json()
    assert me["quota_hours"] == 0 and me["active"] is None
    r = client.post("/api/v1/gpulab/sessions", headers=s, json={"book": "ai", "lab": "4.1"})
    assert r.status_code == 403                                             # no hours yet
    assert client.put("/api/v1/courses/7/gpulab/grant", headers=s, json={"hours": 5}).status_code == 403   # students cannot grant
    assert client.put("/api/v1/courses/7/gpulab/grant", headers=t, json={"hours": 2}).json()["hours"] == 2
    assert client.get("/api/v1/gpulab/me", headers=s).json()["quota_hours"] == 2
    r = client.post("/api/v1/gpulab/sessions", headers=s, json={"book": "ai", "lab": "4.1"}).json()
    sid = r["id"]
    assert r["state"] == "starting" and "url" not in r
    assert client.post("/api/v1/gpulab/sessions", headers=s, json={"book": "ai", "lab": "4.1"}).status_code == 409   # one at a time
    assert client.get(f"/api/v1/gpulab/sessions/{sid}", headers=s).json()["state"] == "starting"     # the mock boots on the 2nd call
    r = client.get(f"/api/v1/gpulab/sessions/{sid}", headers=s).json()
    assert r["state"] == "ready" and r["url"].startswith("http://jupyter.mock/lab/tree/wq/4_1?token=tok-")
    files = JUPYTER["files"]
    assert files["wq"] == {"type": "directory"} and files["wq/4_1"] == {"type": "directory"}
    assert files["wq/4_1/lab4_1.ipynb"]["type"] == "notebook" and files["wq/4_1/vadd.cu"]["type"] == "file"
    assert "def submit" in files["wq/wqgpu.py"]["content"]
    sess = json.loads(files["wq/session.json"]["content"])
    assert sess["sid"] == sid and sess["callback"] == "https://learn.test/api/v1/gpulab/submit" and sess["results"] == ["t_ms", "bw_GBs"]
    # signed results, exactly as the notebook helper sends them
    body, sig = wqgpu.payload({"t_ms": 1.25, "bw_GBs": 812.0, "junk": 1}, {"gpu": "NVIDIA GeForce RTX 4090"}, sess)
    r = client.post("/api/v1/gpulab/submit", content=body, headers={"X-WQ-Signature": sig, "Content-Type": "application/json"})
    assert r.json()["kept"] == ["bw_GBs", "t_ms"]                          # unregistered names are dropped
    assert client.post("/api/v1/gpulab/submit", content=body, headers={"X-WQ-Signature": "0" * 64}).status_code == 403
    res = client.get("/api/v1/gpulab/results/ai/4.1", headers=s).json()["results"]
    assert res[0]["vals"] == {"t_ms": 1.25, "bw_GBs": 812.0} and res[0]["env"]["gpu"].endswith("4090")
    assert client.get(f"/api/v1/gpulab/sessions/{sid}", headers=t).status_code == 404                # not the teacher's machine
    r = client.post(f"/api/v1/gpulab/sessions/{sid}/stop", headers=s).json()
    assert r["state"] == "ended" and r["end_reason"] == "student"
    assert main.state.gpu_provider.machines["mock-1"]["on"] is False
    u = client.get("/api/v1/courses/7/gpulab/usage", headers=t).json()
    assert u["grant_hours"] == 2 and u["students"][0]["name"] == "学生甲" and u["students"][0]["sessions"] == 1
    assert client.get("/api/v1/courses/7/gpulab/usage", headers=s).status_code == 403
    assert client.get("/api/v1/gpulab/admin/ledger", headers=t).status_code == 403                  # site administrators only


def test_reaper_stops_idle_and_overlong_machines(client):
    s, t = login(client, "s"), login(client, "t")
    client.put("/api/v1/courses/7/gpulab/grant", headers=t, json={"hours": 3})
    sid = client.post("/api/v1/gpulab/sessions", headers=s, json={"book": "ai", "lab": "4.1"}).json()["id"]
    client.get(f"/api/v1/gpulab/sessions/{sid}", headers=s)
    assert client.get(f"/api/v1/gpulab/sessions/{sid}", headers=s).json()["state"] == "ready"
    asyncio.run(gpulab.reap(main))
    assert main.gpulab_store().get(sid)["state"] == "ready"                 # active a moment ago
    JUPYTER["last"] = "2026-01-01T00:00:00Z"
    main.gpulab_store().update(sid, last_active=int(time.time()) - 3600)
    asyncio.run(gpulab.reap(main))
    got = main.gpulab_store().get(sid)
    assert got["state"] == "ended" and got["end_reason"] == "idle"
    sid = client.post("/api/v1/gpulab/sessions", headers=s, json={"book": "ai", "lab": "4.1"}).json()["id"]
    main.gpulab_store().update(sid, started_at=int(time.time()) - 4 * 3600)
    asyncio.run(gpulab.reap(main))
    assert main.gpulab_store().get(sid)["end_reason"] == "time_limit"
    me = client.get("/api/v1/gpulab/me", headers=s).json()
    assert me["used_hours"] >= 4 and me["left_hours"] == 0                  # the long one used up the quota
    assert client.post("/api/v1/gpulab/sessions", headers=s, json={"book": "ai", "lab": "4.1"}).status_code == 403


def test_unknown_lab_and_unset_platform(client):
    s = login(client, "s")
    assert client.post("/api/v1/gpulab/sessions", headers=s, json={"book": "ai", "lab": "9.9"}).status_code == 404
    assert client.get("/api/v1/gpulab/labs/ai", headers=s).json()["4.1"]["tier"] == "basic"
    with pytest.raises(EngineError):
        gpulab.provider_of(type("S", (), {"gpu_provider": ""})(), None)


def run(coro):
    return asyncio.run(coro)


def test_autodl_adapter_follows_its_interface():
    calls, state = [], {"st": "creating"}

    def h(req: httpx.Request) -> httpx.Response:
        u = urlparse(str(req.url))
        calls.append((req.method, u.path, req.headers.get("authorization"), req.content.decode() or u.query))
        ok = lambda d: httpx.Response(200, json={"code": "Success", "msg": "", "data": d})  # noqa: E731
        if u.path.endswith("/pro/create"):
            b = json.loads(req.content)
            assert b["gpu_spec_uuid"] == "v-48g" and b["image_uuid"] == "base-image-mbr2n4urrc" and b["cuda_v_from"] == 121
            return ok("pro-abc")
        if u.path.endswith("/pro/status"):
            return ok(state["st"])
        if u.path.endswith("/pro/snapshot"):
            return ok({"jupyter_domain": "u1-x.westc.seetacloud.com:8443", "jupyter_token": "jt", "payg_price": 2.68})
        if u.path.endswith("/pro/power_off"):
            state["st"] = "shutdown"
            return ok(None)
        if u.path.endswith("/pro/release"):
            return ok(None)
        return httpx.Response(404)

    http = httpx.AsyncClient(transport=httpx.MockTransport(h))
    a = gpulab.AutoDL(http, "T0K", "https://api.autodl.com", "basic=v-48g:base-image-mbr2n4urrc", 121)
    pid = run(a.launch("basic", "wq-6-4.1"))
    assert pid == "pro-abc" and calls[0][2] == "T0K"
    assert run(a.status(pid)) == {"state": "starting"}
    state["st"] = "running"
    assert run(a.status(pid)) == {"state": "running", "jupyter_url": "https://u1-x.westc.seetacloud.com:8443", "jupyter_token": "jt",
                                  "price_hour": 2.68}
    run(a.stop(pid))
    assert [c[1].rsplit("/", 1)[1] for c in calls][-3:] == ["power_off", "status", "release"]
    with pytest.raises(EngineError):
        run(a.launch("hopper", "x"))                                       # no machine configured for this tier


def test_lambda_adapter_follows_its_interface():
    calls = []

    def h(req: httpx.Request) -> httpx.Response:
        u = urlparse(str(req.url))
        calls.append((req.method, u.path, req.headers.get("authorization"), req.content.decode()))
        if u.path.endswith("/instance-operations/launch"):
            b = json.loads(req.content)
            assert b["instance_type_name"] == "gpu_1x_h100_sxm5" and b["ssh_key_names"] == ["wq"] and b["region_name"] == "us-east-1"
            return httpx.Response(200, json={"data": {"instance_ids": ["i-1"]}})
        if u.path.endswith("/instances/i-1"):
            return httpx.Response(200, json={"data": {"id": "i-1", "status": "active", "instance_type": {"name": "gpu_1x_h100_sxm5"},
                                                      "jupyter_url": "https://jupyter-i1.lambdaspaces.com/?token=jt", "jupyter_token": "jt"}})
        if u.path.endswith("/instance-operations/terminate"):
            assert json.loads(req.content) == {"instance_ids": ["i-1"]}
            return httpx.Response(200, json={"data": {"terminated_instances": []}})
        return httpx.Response(404)

    http = httpx.AsyncClient(transport=httpx.MockTransport(h))
    a = gpulab.Lambda(http, "K", "https://cloud.lambda.ai/api/v1", "hopper=gpu_1x_h100_sxm5", "us-east-1", "wq", "gpu_1x_h100_sxm5=3.99")
    assert run(a.launch("hopper", "x")) == "i-1" and calls[0][2] == "Bearer K"
    assert run(a.status("i-1")) == {"state": "running", "jupyter_url": "https://jupyter-i1.lambdaspaces.com", "jupyter_token": "jt",
                                    "price_hour": 3.99}
    run(a.stop("i-1"))
    assert calls[-1][1].endswith("/terminate")


def test_lab_page_link_and_filled_report(client, tmp_path):
    import io

    from docx import Document
    doc = Document()
    doc.add_paragraph("实验 4.1 报告")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "耗时（ms）", "{{t_ms}}"
    t.cell(1, 0).text, t.cell(1, 1).text = "带宽（GB/s）", "{{ bw_GBs }}"
    doc.save(str(tmp_path / "books" / "ai" / "gpulab" / "lab4_1-report.docx"))
    s, t_ = login(client, "s"), login(client, "t")
    client.put("/api/v1/courses/7/gpulab/grant", headers=t_, json={"hours": 2})
    sess = main.state.codec.read(s["Authorization"].split(" ", 1)[1])
    url = main.gpulab_link(sess, "ai", "4.1")
    base = urlparse(url).path
    page = client.get(base)
    assert page.status_code == 200 and "GPU 实验 4.1" in page.text and "第一个核函数" in page.text
    st = client.get(base + "/state").json()
    assert st["quota_hours"] == 2 and st["active"] is None and st["results"] == []
    sid = client.post(base + "/start").json()["id"]
    client.get(base + "/state")
    st = client.get(base + "/state").json()
    assert st["active"]["state"] == "ready" and st["active"]["url"].startswith("http://jupyter.mock/")
    sj = json.loads(JUPYTER["files"]["wq/session.json"]["content"])
    body, sig = wqgpu.payload({"t_ms": 0.123456, "bw_GBs": 905}, {"gpu": "RTX 4090"}, sj)
    client.post("/api/v1/gpulab/submit", content=body, headers={"X-WQ-Signature": sig})
    rep = client.get(base + "/report")
    cells = [c.text for row in Document(io.BytesIO(rep.content)).tables[0].rows for c in row.cells]
    assert cells == ["耗时（ms）", "0.1235", "带宽（GB/s）", "905"]
    assert client.post(base + f"/stop/{sid}").json()["state"] == "ended"
    assert client.get(base[:-5] + "xxxxx/state").status_code == 410          # a tampered link
