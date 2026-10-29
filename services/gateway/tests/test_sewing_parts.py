"""第 49 轮：本单位零件库第一期——同单位共享、别的单位看不见；手填、表格批量、AI 整理样本（待复核 → 人确认才可用）；
接进设计参数、CAE 电机选型和机壳模态；ERPNext 物料导入表；FreeCAD 钥匙。模型和数字工厂都用替身。"""
import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from app import main
from app.sewcore import partlib as PL
from tests.test_sewing_api import sc  # noqa: F401  (管理员 = 问渠；用户 s 开通给“甲缝纫机厂”)
from tests.test_sewing_api import UID
from tests.test_sewing_cae import cae  # noqa: F401
from tests.test_course_pack import client  # noqa: F401
from tests.test_materials import login

try:
    import build123d  # noqa: F401
    HAS_BD = True
except Exception:  # noqa: BLE001
    HAS_BD = False

MOTOR = {"power_w": 400, "rated_nm": 1.27, "peak_nm": 4.46, "rated_rpm": 3000, "max_rpm": 6000, "inertia": 0.42, "flange": 60, "length": 90, "shaft_d": 14}


def users(c):
    admin, user = login(c), login(c, "s")
    UID["s"] = c.get("/api/v1/me", headers=user).json()["id"]
    c.post("/api/v1/admin/sewing/licenses", json={"username": "s", "org": "甲缝纫机厂"}, headers=admin)
    return admin, user


def test_access_and_isolation(sc):  # noqa: F811
    c = sc
    assert len(c.get("/api/v1/sewing/parts/cats").json()["cats"]) == len(PL.CATS)
    assert c.get("/api/v1/sewing/parts").status_code == 401
    assert c.get("/api/v1/sewing/parts", headers=login(c, "s")).status_code == 403          # 还没开通
    admin, user = users(c)
    r = c.post("/api/v1/sewing/parts", json={"cat": "motor", "code": "M-400", "name": "400 W 伺服", "erp": "MOT-0400", "fields": MOTOR}, headers=user)
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["status"] == "ok" and p["source"] == "manual" and p["has_step"] == HAS_BD
    assert c.post("/api/v1/sewing/parts", json={"cat": "motor", "code": "M-400", "name": "x", "fields": MOTOR}, headers=user).status_code == 409
    bad = c.post("/api/v1/sewing/parts", json={"cat": "motor", "code": "M-1", "name": "x", "fields": {"rated_nm": -1}}, headers=user)
    assert bad.status_code == 400 and "额定功率没填" in bad.json()["detail"]
    assert c.get("/api/v1/sewing/parts", headers=user).json()["total"] == 1
    assert c.get("/api/v1/sewing/parts", headers=admin).json()["total"] == 0                  # 问渠看不见甲厂的
    assert c.get(f"/api/v1/sewing/parts/{p['id']}", headers=admin).status_code == 404
    assert c.delete(f"/api/v1/sewing/parts/{p['id']}", headers=admin).status_code == 404
    # 用在设计里：花样机 X / Y 轴电机
    use = c.get("/api/v1/sewing/parts/for?machine=pattern", headers=user).json()["items"]
    assert {u["label"] for u in use} == {"X 轴电机", "Y 轴电机"}
    assert next(u for u in use if u["label"] == "X 轴电机")["params"] == {"TxR": 1.27, "TxP": 4.46, "Jx": 0.42, "Nmot": 3000}
    assert c.get("/api/v1/sewing/parts/for?machine=lockstitch", headers=user).json()["items"] == []
    # FreeCAD 钥匙也能读
    k = c.post("/api/v1/sewing/keys", headers=user).json()["key"]
    assert c.get("/api/v1/sewing/parts?q=M-4", headers={"X-WQ-Key": k}).json()["parts"][0]["code"] == "M-400"
    if HAS_BD:
        s = c.get(f"/api/v1/sewing/parts/{p['id']}/step", headers={"X-WQ-Key": k})
        assert s.status_code == 200 and s.content.startswith(b"ISO-10303-21")
        mesh = c.get(f"/api/v1/sewing/parts/{p['id']}/mesh", headers=user).json()
        assert len(mesh["f"]) > 30 and mesh["size"][0] == pytest.approx(60, abs=0.5)


def test_table_import(sc):  # noqa: F811
    c = sc
    _, user = users(c)
    t = c.get("/api/v1/sewing/parts/template.xlsx?cat=link", headers=user)
    wb = load_workbook(io.BytesIO(t.content))
    ws = wb.active
    keys = [x.value for x in ws[2]]
    assert keys[:3] == ["code", "name", "erp"] and "L" in keys
    ws.delete_rows(3)
    col = {k: i + 1 for i, k in enumerate(keys)}
    for i, (code, L) in enumerate([("CR-15", 15.5), ("CR-16", 16), ("CR-16", 17), ("CR-X", 999)]):
        ws.cell(3 + i, col["code"], code)
        ws.cell(3 + i, col["name"], "针杆曲柄")
        ws.cell(3 + i, col["L"], L)
        ws.cell(3 + i, col["t"], 8)
    b = io.BytesIO()
    wb.save(b)
    files = {"file": ("t.xlsx", b.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = c.post("/api/v1/sewing/parts/import", files=files, data={"cat": "link", "commit": "true"}, headers=user).json()
    assert r["imported"] == 0 and r["bad"] == 2                     # 有错就一个都不入库
    assert "重复" in r["rows"][2]["errors"][0] and "不在" in r["rows"][3]["errors"][0]
    ws.delete_rows(5, 2)
    b = io.BytesIO()
    wb.save(b)
    files = {"file": ("t.xlsx", b.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    assert c.post("/api/v1/sewing/parts/import", files=files, data={"cat": "link"}, headers=user).json()["imported"] == 0   # 只预览
    r = c.post("/api/v1/sewing/parts/import", files=files, data={"cat": "link", "commit": "true"}, headers=user).json()
    assert r["imported"] == 2
    ps = c.get("/api/v1/sewing/parts?cat=link", headers=user).json()["parts"]
    assert {p["code"] for p in ps} == {"CR-15", "CR-16"} and all(p["source"] == "table" and p["has_step"] == HAS_BD for p in ps)
    use = c.get("/api/v1/sewing/parts/for?machine=lockstitch", headers=user).json()["items"]
    got = {(u["code"], u["label"], json.dumps(u["params"])) for u in use}
    assert ("CR-15", "刺布曲柄", '{"r": 15.5}') in got and ("CR-16", "刺布曲柄", '{"r": 16}') in got
    assert not any(u["label"] == "针杆连杆" for u in use)            # 孔距 16 不在连杆长 35–90 的范围里，不给
    # 拿别的类别的模板来导不行
    other = c.get("/api/v1/sewing/parts/template.xlsx?cat=motor", headers=user).content
    r = c.post("/api/v1/sewing/parts/import", files={"file": ("m.xlsx", other, "application/x")}, data={"cat": "link"}, headers=user)
    assert r.status_code == 400
    # ERPNext 物料导入表
    x = load_workbook(io.BytesIO(c.get("/api/v1/sewing/parts/erp.xlsx", headers=user).content)).active
    rows = list(x.iter_rows(values_only=True))
    assert rows[0][:4] == ("Item Code", "Item Name", "Item Group", "Stock UOM") and {r[0] for r in rows[1:]} == {"CR-15", "CR-16"}
    assert "孔距 15.5 mm" in next(r[5] for r in rows[1:] if r[0] == "CR-15")


def pdf_with_text(lines):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    b = io.BytesIO()
    with plt.rc_context({"pdf.fonttype": 42, "font.family": "DejaVu Sans"}):      # 别的测试可能改过字体设置
        fig = plt.figure(figsize=(8, 4))
        for i, s in enumerate(lines):
            fig.text(0.05, 0.9 - i * 0.1, s, fontsize=10)
        fig.savefig(b, format="pdf")
        plt.close(fig)
    return b.getvalue()


def test_ai_reads_a_catalogue_into_review(sc, monkeypatch):  # noqa: F811
    c = sc
    _, user = users(c)
    calls = []

    async def fake_json(**kw):
        calls.append(kw)
        return {"rows": [{"code": "SV-400", "name": "400 W 伺服",
                          "fields": {"power_w": 400, "rated_nm": 1.27, "peak_nm": 4.46, "rated_rpm": 3000, "max_rpm": 6000, "inertia": 0.42,
                                     "flange": 60, "length": 95, "shaft_d": None},
                          "quotes": {"power_w": "400 W", "rated_nm": "1.27 N·m", "peak_nm": "4.46", "rated_rpm": "3000", "max_rpm": "6000",
                                     "inertia": "0.42", "flange": "60", "length": "95", "shaft_d": None}}], "notes": []}

    monkeypatch.setattr(main.state.ai, "json", fake_json)
    monkeypatch.setattr(main.state.ai, "provider", "fake")
    pdf = pdf_with_text(["Servo SV-400: rated power 400 W, rated torque 1.27 N m, peak 4.46 N m",
                         "rated speed 3000 r/min, max speed 6000 r/min, rotor inertia 0.42 x 1e-4 kg m2, flange 60 mm"])
    r = c.post("/api/v1/sewing/parts/extract", files={"file": ("cat.pdf", pdf, "application/pdf")}, data={"cat": "motor"}, headers=user)
    assert r.status_code == 200, r.text
    p = r.json()["parts"][0]
    assert "SV-400" in calls[0]["prompt"] and p["status"] == "review" and p["source"] == "ai"
    assert p["evidence"]["rated_nm"]["found"] is True and p["evidence"]["length"]["found"] is False      # 95 在样本里找不到
    # 待复核的不进设计、不进选型
    assert c.get("/api/v1/sewing/parts/for?machine=pattern", headers=user).json()["items"] == []
    assert c.get("/api/v1/sewing/parts?status=review", headers=user).json()["review"] == 1
    body = {"cat": "motor", "code": p["code"], "name": p["name"], "fields": dict(p["fields"], length=90), "confirm": True}
    q = c.put(f"/api/v1/sewing/parts/{p['id']}", json=body, headers=user).json()
    assert q["status"] == "ok" and q["fields"]["length"] == 90
    assert len(c.get("/api/v1/sewing/parts/for?machine=pattern", headers=user).json()["items"]) == 2
    # 不是 PDF、没有文字层
    assert c.post("/api/v1/sewing/parts/extract", files={"file": ("a.pdf", b"hello", "application/pdf")}, data={"cat": "motor"},
                  headers=user).status_code == 400
    monkeypatch.setattr(main.state.ai, "provider", "none")
    assert c.post("/api/v1/sewing/parts/extract", files={"file": ("cat.pdf", pdf, "application/pdf")}, data={"cat": "motor"},
                  headers=user).status_code == 503


def test_parts_in_cae(cae):  # noqa: F811
    c, user, fk = cae
    big = dict(MOTOR, power_w=600, rated_nm=6.5, peak_nm=19, inertia=4.0)
    c.post("/api/v1/sewing/parts", json={"cat": "motor", "code": "SP-600", "name": "主轴伺服", "fields": big}, headers=user)
    m = c.post("/api/v1/sewing/cae/motor", json={"machine": "lockstitch", "params": {"n": 5000}}, headers=user).json()
    assert m["pick"]["name"] == "本单位 · SP-600 主轴伺服" and m["pick"]["org"] is True        # 本单位的更小又够用，排在前面
    x = c.post("/api/v1/sewing/cae/motor", json={"machine": "pattern", "params": {}}, headers=user).json()
    assert any(r.get("org") for r in x["axes"][0]["rows"])
    # 机壳直接从库里选
    f = c.post("/api/v1/sewing/parts", json={"cat": "frame", "code": "FR-1", "name": "机壳", "fields": {"material": "HT200"}}, headers=user).json()
    assert f["has_step"] is False
    r = c.post("/api/v1/sewing/cae/frame", data={"machine": "lockstitch", "part_id": str(f["id"])}, headers=user)
    assert r.status_code == 404                                       # 还没上传 STEP
    up = c.post(f"/api/v1/sewing/parts/{f['id']}/step", files={"file": ("frame.step", b"ISO-10303-21;\nHEADER;", "model/step")}, headers=user)
    assert up.json()["has_step"] and up.json()["step_name"] == "frame.step"
    assert c.post(f"/api/v1/sewing/parts/{f['id']}/step", files={"file": ("x.txt", b"hello", "text/plain")}, headers=user).status_code == 400
    fk.desc = None
    r = c.post("/api/v1/sewing/cae/frame", data={"machine": "lockstitch", "part_id": str(f["id"])}, headers=user)
    assert r.status_code == 200 and fk.posted[-1]["kind"] == "fea"
    assert c.post("/api/v1/sewing/cae/frame", data={"machine": "lockstitch"}, headers=user).status_code == 400


def test_partlib_checks():
    f, e = PL.check_fields("link", {"L": "15.5", "t": 8, "d1": "abc"})
    assert f == {"L": 15.5, "t": 8} and e == ["大孔径不是数：abc"]
    assert PL.param_values("pulley", {"r": 15}, "embroid") == [{"label": "绷架带轮", "params": {"rf": 15}}]
    assert PL.cae_motor({"id": 1, "code": "A", "name": "B", "fields": MOTOR})["inertia_kgm2"] == pytest.approx(0.42e-4)
    wb = Workbook()
    assert PL.read_xlsx("link", PL.template_xlsx("link")) == []
    del wb
