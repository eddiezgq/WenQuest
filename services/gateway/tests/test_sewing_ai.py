"""第 49 轮：缝纫机设计软件的 AI 助手——AI 起草、计算核心把关、人拍板；规则版；FreeCAD 钥匙；计算书带设计说明。
模型用替身（不联网）。"""
import io
import sqlite3
from pathlib import Path

import docx

from app import main
from app.sewcore import aicheck
from app.sewcore import design as D
from tests.test_course_pack import client  # noqa: F401
from tests.test_materials import login


def fake(monkeypatch, answers):
    """按顺序给出模型的回答；记下每次的提示词"""
    calls = []

    async def fake_json(**kw):
        calls.append(kw)
        a = answers[min(len(calls) - 1, len(answers) - 1)]
        return a(kw) if callable(a) else a

    monkeypatch.setattr(main.state.ai, "json", fake_json)
    monkeypatch.setattr(main.state.ai, "provider", "fake")
    return calls


def test_needs_an_activated_account(client, monkeypatch):  # noqa: F811
    fake(monkeypatch, [{"changes": [], "assumptions": [], "questions": []}])
    body = {"machine": "lockstitch", "params": {}, "sentence": "转速提到 6000"}
    assert client.post("/api/v1/sewing/ai/draft", json=body).status_code == 401
    assert client.post("/api/v1/sewing/ai/draft", json=body, headers=login(client, "s")).status_code == 403
    assert client.post("/api/v1/sewing/ai/draft", json=body, headers={"X-WQ-Key": "wqs_nonsense"}).status_code == 401


def test_draft_rejects_new_red_lights_then_accepts_the_fix(client, monkeypatch):  # noqa: F811
    calls = fake(monkeypatch, [
        {"changes": [{"key": "n", "value": 6000, "why": "句子要求"}], "assumptions": [], "questions": []},
        {"changes": [{"key": "n", "value": 6000, "why": "句子要求"}, {"key": "mr", "value": 0.1, "why": "减小针杆惯性力"},
                     {"key": "fn", "value": 240, "why": "避开二阶共振"}], "assumptions": ["机壳加筋"], "questions": []},
    ])
    h = login(client)                                          # 管理员 = 已开通（问渠）
    r = client.post("/api/v1/sewing/ai/draft", json={"params": {}, "sentence": "转速提到 6000，全部合格"}, headers=h).json()
    assert r["status"] == "ok" and r["rounds"] == 2
    assert {c["key"]: c["new"] for c in r["changes"]} == {"n": 6000, "mr": 0.1, "fn": 240}
    assert r["params"]["n"] == 6000 and r["after"]["bad"] == 0 and r["before"]["bad"] == 0
    assert "rejected" in calls[1]["prompt"] and "needle_force" in calls[1]["prompt"]     # 第二轮带上了新增的红灯
    assert calls[0]["schema"]["properties"]["changes"]["items"]["properties"]["key"]["enum"][0] == "n"
    # 两轮都不行就退回，不给参数
    fake(monkeypatch, [{"changes": [{"key": "n", "value": 6500, "why": ""}], "assumptions": [], "questions": []}])
    r = client.post("/api/v1/sewing/ai/draft", json={"params": {}, "sentence": "转速 6500"}, headers=h).json()
    assert r["status"] == "rejected" and r["params"] is None and {x["key"] for x in r["new_bad"]} == {"needle_force", "frame_res"}
    # 超范围的值夹到边上并说明；不认识的键丢掉
    fake(monkeypatch, [{"changes": [{"key": "s", "value": 9, "why": ""}, {"key": "nope", "value": 1, "why": ""}], "assumptions": [], "questions": []}])
    r = client.post("/api/v1/sewing/ai/draft", json={"params": {}, "sentence": "针距 9"}, headers=h).json()
    assert [(c["key"], c["new"]) for c in r["changes"]] == [("s", 5)] and "夹到边上" in r["changes"][0]["note"]


def test_explain_tries_every_fix(client, monkeypatch):  # noqa: F811
    fake(monkeypatch, [{"items": [{"check": "hook", "why": "旋梭来晚了", "fixes": [
        {"key": "psi", "value": 0, "how": "旋梭定时回到 0°"}, {"key": "Td", "value": 120, "how": "无关的改法"}]}]}])
    h = login(client)
    r = client.post("/api/v1/sewing/ai/explain", json={"params": {"psi": -20}}, headers=h).json()
    it = next(x for x in r["items"] if x["check"] == "hook")
    assert it["why"] == "旋梭来晚了"
    assert it["fixes"][0]["clears"] is True and it["fixes"][0]["params"]["psi"] == 0
    assert it["fixes"][1]["clears"] is False
    # 全部合格不调模型
    assert client.post("/api/v1/sewing/ai/explain", json={"params": {}}, headers=h).json()["all_ok"] is True


def test_describe_drops_sentences_with_invented_numbers(client, monkeypatch):  # noqa: F811
    fake(monkeypatch, [{"sentences": ["本设计为平缝机，设计转速 5000 r/min。", "针杆惯性力 654 N，低于轴套允许的 800 N。",
                                      "整机寿命 12345 小时。"]}])
    r = client.post("/api/v1/sewing/ai/describe", json={"params": {}}, headers=login(client)).json()
    assert r["sentences"][:2] == ["本设计为平缝机，设计转速 5000 r/min。", "针杆惯性力 654 N，低于轴套允许的 800 N。"]
    assert r["dropped"][0]["missing"] == [12345]


def test_rules_when_no_model_is_configured(client, monkeypatch):  # noqa: F811
    monkeypatch.setattr(main.state.ai, "provider", "none")
    h = login(client)
    r = client.post("/api/v1/sewing/ai/draft", json={"params": {}, "sentence": "转速改成 4800，针距 3.5"}, headers=h).json()
    assert r["model"] == "rules" and {c["key"]: c["new"] for c in r["changes"]} == {"n": 4800, "s": 3.5} and r["status"] == "ok"
    r = client.post("/api/v1/sewing/ai/draft", json={"params": {}, "sentence": "做得好看一点"}, headers=h).json()
    assert r["changes"] == [] and r["questions"]
    r = client.post("/api/v1/sewing/ai/explain", json={"params": {"psi": -20}}, headers=h).json()
    assert r["model"] == "rules" and any("旋梭" in x["why"] for x in r["items"])
    r = client.post("/api/v1/sewing/ai/describe", json={"params": {"n": 6000}}, headers=h).json()
    assert r["sentences"][0].startswith("本设计为平缝机") and not r["dropped"]


def test_quota_and_log(client, monkeypatch):  # noqa: F811
    fake(monkeypatch, [{"changes": [], "assumptions": [], "questions": []}])
    monkeypatch.setattr(main.state.settings, "sewing_ai_daily", 2)
    h = login(client)
    body = {"params": {}, "sentence": "转速 5000"}
    assert client.post("/api/v1/sewing/ai/draft", json=body, headers=h).json()["left_today"] == 1
    client.post("/api/v1/sewing/ai/draft", json=body, headers=h)
    assert client.post("/api/v1/sewing/ai/draft", json=body, headers=h).status_code == 429
    with sqlite3.connect(str(Path(main.state.settings.data_dir) / "sewing_ai.db")) as c:
        rows = c.execute("select kind, input, org, passed from runs order by id").fetchall()
    assert rows[-1] == ("draft", "转速 5000", "问渠", 1)


def test_freecad_key(client, monkeypatch):  # noqa: F811
    fake(monkeypatch, [{"changes": [{"key": "n", "value": 4500, "why": ""}], "assumptions": [], "questions": []}])
    h = login(client)
    assert client.get("/api/v1/sewing/keys", headers=h).json()["has_key"] is False
    k = client.post("/api/v1/sewing/keys", headers=h).json()["key"]
    assert k.startswith("wqs_") and client.get("/api/v1/sewing/keys", headers=h).json()["has_key"] is True
    r = client.post("/api/v1/sewing/ai/draft", json={"params": {}, "sentence": "转速 4500"}, headers={"X-WQ-Key": k})
    assert r.status_code == 200 and r.json()["changes"][0]["new"] == 4500
    k2 = client.post("/api/v1/sewing/keys", headers=h).json()["key"]          # 重发：旧钥匙作废
    assert client.post("/api/v1/sewing/ai/draft", json={"params": {}, "sentence": "x 1"}, headers={"X-WQ-Key": k}).status_code == 401
    client.delete("/api/v1/sewing/keys", headers=h)
    assert client.post("/api/v1/sewing/ai/draft", json={"params": {}, "sentence": "x 1"}, headers={"X-WQ-Key": k2}).status_code == 401
    assert client.post("/api/v1/sewing/keys", headers=login(client, "s")).status_code == 403


def test_report_carries_the_checked_description(client):  # noqa: F811
    h = login(client)
    body = {"params": {}, "desc": ["本设计为平缝机，设计转速 5000 r/min。", "寿命 99999 小时。"]}
    r = client.post("/api/v1/sewing/export/report", json=body, headers=h)
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
    assert "设计说明" in text and "设计转速 5000 r/min" in text and "99999" not in text and "AI 起草" in text
    r = client.post("/api/v1/sewing/export/report", json={"params": {}}, headers=h)
    assert "设计说明" not in "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)


def test_number_check_allows_rounding_only():
    R = D.analyze({})
    bag = aicheck.allowed_numbers(R)
    assert aicheck.found(5000, 0, bag) and aicheck.found(654, 0, bag)
    keep, drop = aicheck.verify(["转速 5,000 r/min。", "惯性力 655 N。"], R)
    assert keep == ["转速 5,000 r/min。"] and drop[0]["missing"] == [655]
