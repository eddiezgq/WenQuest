"""A3 virtual labs: the lab kit page, the safety check, the lab engineer's retries, publishing one lab page per chapter."""
import json
import math
import re
from io import BytesIO
from urllib.parse import parse_qs, urlparse

import httpx
import pytest
from PIL import Image

from app import main
from app.production import labs
from tests.test_materials import fake_moodle, login
from tests.test_studio import CALLS, _write_first_lesson, client, publish, settle  # noqa: F401 - the fixture

COURSE, CHAPTER = ["大学物理A（上）", "University Physics A (I)"], ["第2章 牛顿运动定律", "Chapter 2 Newton's laws"]


def png() -> bytes:
    b = BytesIO()
    Image.new("RGB", (160, 100), "#f8faf9").save(b, "PNG")
    return b.getvalue()


def test_the_example_lab_passes_the_static_check_and_builds_a_safe_page():
    code = labs.example_code()
    assert labs.static_problems(code) == []
    page = labs.page([("2.1", code), ("2.2", code)], course=COURSE, chapter=CHAPTER)
    assert "Content-Security-Policy" in page and "connect-src 'none'" in page
    assert page.count('WQ.begin("2-1")') == 1 and page.count('WQ.begin("2-2")') == 1
    assert "fonts.googleapis" not in page and "WQ.start()" in page
    assert "第2章 牛顿运动定律 虚拟实验" in page


def test_unsafe_labs_are_rejected_before_they_run():
    ok = "WQ.lab({title: ['a', 'b'], scenes: [], params: [], tasks: [], draw() {}});"
    bad = {
        "fetch('/api/v1/me')": "network", "new XMLHttpRequest()": "network", "new WebSocket('x')": "network",
        "import('x.js')": "loading other code", "eval('1')": "eval", "new Function('x')": "eval",
        "document.cookie": "storage", "localStorage.setItem('a', 1)": "storage", "location.href = 'x'": "navigation",
        "window.open('x')": "navigation", "parent.postMessage(1, '*')": "reaching the page", "el.innerHTML = '<b>'": "writing HTML",
        "'</script><script>alert(1)'": "script tags", "const u = 'https://evil.example/x'": "web addresses",
        "setInterval(f, 10)": "own timers", "WQ.demo('2-1', 'x')": "kit internals",
    }
    for snippet, what in bad.items():
        problems = labs.static_problems(ok + "\n" + snippet)
        assert any(what in p for p in problems), (snippet, problems)
    assert labs.static_problems(ok) == []
    assert "exactly once" in labs.static_problems(ok + ok)[0]
    assert labs.static_problems("") == ["the code is empty"]


def test_lab_is_written_checked_fixed_and_published_once_per_chapter(client, monkeypatch):
    """The checker's problems go back to the lab engineer; the passing lab is on the slides, in the lesson's
    outputs (sandboxed) and in the course as one chapter lab page, which a later publish replaces."""
    runs: list[dict] = []

    async def trial_run(url, page_html, lab, timeout=150.0):
        runs.append({"lab": lab, "html": page_html})
        if len(runs) == 1:
            return {"ok": False, "problems": ["task 'safe': after its demo the task was not ticked"], "tasks": {"safe": False},
                    "screenshot": b"", "seconds": 5.0}
        return {"ok": True, "problems": [], "tasks": {"slide": True, "safe": True, "strap": True}, "screenshot": png(), "seconds": 7.0}

    prompts: list[str] = []
    real_json = main.state.ai.json if hasattr(main.state, "ai") else None

    monkeypatch.setattr(labs, "trial_run", trial_run)
    monkeypatch.setattr(main.state.settings, "labcheck_url", "http://labcheck.test")
    edits: list[dict] = []

    def moodle(request: httpx.Request) -> httpx.Response:
        fn = (parse_qs(urlparse(str(request.url)).query).get("wsfunction") or [""])[0]
        if not fn:  # file uploads, sign-in
            return fake_moodle(request)
        body = {k: v[0] for k, v in parse_qs(request.content.decode()).items()}
        if fn == "local_wenquest_add_activities":  # one course module per activity, in order
            CALLS.append((fn, body))
            n = len({k.split("]")[0] for k in body if k.startswith("activities[")})
            return httpx.Response(200, json={"courseid": 7, "sectionid": 30, "cmids": list(range(200, 200 + n))})
        if fn == "local_wenquest_edit_course":
            edits.append(body)
            return httpx.Response(200, json={"ok": True, "sectionid": 0, "cmid": int(body["cmid"])})
        if fn == "local_wenquest_create_course":
            CALLS.append((fn, body))
        return fake_moodle(request)
    main.state.http = httpx.AsyncClient(transport=httpx.MockTransport(moodle))
    main.state.moodle.http = main.state.http
    h = login(client)
    ai = main.state.ai
    orig = ai.json

    async def spy(**kw):
        if "实验师" in kw.get("system", ""):
            prompts.append(kw["prompt"])
        return await orig(**kw)
    monkeypatch.setattr(ai, "json", spy)

    pid, p = _write_first_lesson(client, h)
    les = p["outline"]["chapters"][0]["lessons"][0]
    assert len(runs) == 2 and runs[0]["lab"] == runs[1]["lab"] and re.fullmatch(r"\d+-\d+", runs[0]["lab"])
    assert "after its demo the task was not ticked" in prompts[1] and "Previous code" in prompts[1]
    lab_file = next(f for f in les["files"] if f["kind"] == "lab")
    r = client.get(lab_file["url"])
    assert r.status_code == 200 and "sandbox allow-scripts" in r.headers["content-security-policy"]
    assert "allow-same-origin" not in r.headers["content-security-policy"] and b"WQ.lab(" in r.content
    assert "虚拟实验" in p["messages"][-1]["text"]
    # the slides show the trial run's picture instead of the placeholder
    from pptx import Presentation
    studio = main._studio()
    d = studio.lesson_dir(studio.projects.load(pid), les)
    deck = Presentation(str(next(d.glob("*课件.pptx"))))
    assert any(s.shape_type == 13 for s in deck.slides[4].shapes)  # a picture on the lab slide

    # publish: the chapter lab page goes in once
    publish(client, h, pid, les["id"])
    acts = [c for f, c in CALLS if f == "local_wenquest_add_activities"][-1]
    names = [v for k, v in acts.items() if re.fullmatch(r"activities\[\d+\]\[name\]", k)]
    assert sum("虚拟实验" in n for n in names) == 1
    proj = studio.projects.load(pid)
    lab_cm = list(proj["course"]["labs"].values())[0]
    lab_pos = names.index(next(n for n in names if "虚拟实验" in n))
    assert lab_cm == 200 + lab_pos
    # publishing the next lesson of the chapter replaces that file instead of adding another
    nxt = p["outline"]["chapters"][0]["lessons"][1]
    client.post(f"/api/v1/studio/projects/{pid}/lessons/next", headers=h)
    settle(client, h, pid)
    publish(client, h, pid, nxt["id"])
    acts = [c for f, c in CALLS if f == "local_wenquest_add_activities"][-1]
    names = [v for k, v in acts.items() if re.fullmatch(r"activities\[\d+\]\[name\]", k)]
    assert not any("虚拟实验" in n for n in names)
    assert edits[-1]["action"] == "replacefile" and int(edits[-1]["cmid"]) == lab_cm


def test_a_lab_that_keeps_failing_is_left_out_and_can_be_redone(client, monkeypatch):
    calls = {"n": 0, "fail": True}

    async def trial_run(url, page_html, lab, timeout=150.0):
        calls["n"] += 1
        if calls["fail"]:
            return {"ok": False, "problems": ["scene 'robot': the picture is blank"], "tasks": {}, "screenshot": b"", "seconds": 3.0}
        return {"ok": True, "problems": [], "tasks": {}, "screenshot": png(), "seconds": 3.0}

    monkeypatch.setattr(labs, "trial_run", trial_run)
    monkeypatch.setattr(main.state.settings, "labcheck_url", "http://labcheck.test")
    h = login(client)
    pid, p = _write_first_lesson(client, h)
    les = p["outline"]["chapters"][0]["lessons"][0]
    assert calls["n"] == 3 and les["status"] == "awaiting"
    assert "lab" not in [f["kind"] for f in les["files"]]
    assert any("没有通过试运行" in i["text"] for i in les["review"]["issues"])
    # 重做实验
    calls["fail"] = False
    r = client.post(f"/api/v1/studio/projects/{pid}/lessons/{les['id']}/lab", headers=h)
    assert r.status_code == 200
    p = settle(client, h, pid)
    les = p["outline"]["chapters"][0]["lessons"][0]
    assert "lab" in [f["kind"] for f in les["files"]] and not any("没有通过试运行" in i["text"] for i in les["review"]["issues"])
    assert "重做好了" in p["messages"][-1]["text"]


def test_no_checker_no_lab(client, monkeypatch):
    async def down(url, page_html, lab, timeout=150.0):
        raise labs.CheckError("lab checker unreachable: ConnectError", "service")

    monkeypatch.setattr(labs, "trial_run", down)
    monkeypatch.setattr(main.state.settings, "labcheck_url", "http://labcheck.test")
    h = login(client)
    _, p = _write_first_lesson(client, h)
    les = p["outline"]["chapters"][0]["lessons"][0]
    assert les["status"] == "awaiting" and "lab" not in [f["kind"] for f in les["files"]]
    assert any("实验检查服务暂时不可用" in i["text"] for i in les["review"]["issues"])


def test_3d_labs_carry_the_engine_and_only_the_models_they_name(tmp_path):
    from app.production import library_sample as L, scene3d as S
    models = S.load_models(L.build(tmp_path) / L.VERSION, ["B-ARM-6R-S", "C-LNK-4BAR-S"])
    code = labs.example_code3d()
    assert labs.static_problems(code) == [] and labs.models_used(code, models) == ["B-ARM-6R-S"]
    page = labs.page([("1.1", code)], course=COURSE, chapter=CHAPTER, models=models)
    assert "window.WQ_MODELS" in page and "B-ARM-6R-S" in page and "C-LNK-4BAR-S" not in page
    assert "connect-src 'none'" in page                       # still no network
    plain = labs.page([("1.1", labs.example_code())], course=COURSE, chapter=CHAPTER, models=models)
    assert "window.WQ_MODELS =" not in plain and len(plain) < 200_000   # 2D labs stay small (no engine)


def test_circuit_labs_carry_the_simulator_without_network():
    """第 7 轮第 6 步：view "circuit" labs embed CircuitJS (GPL-2.0, built from source) once, with its source notice; others do not."""
    code = labs.example_code().replace("WQ.lab({", 'WQ.lab({\n  view: "circuit",\n  circuit: "$ 1 0.000005 10 50 5 50\\n",', 1)
    assert labs.static_problems(code) == []
    page = labs.page([("2.1", code)], course=COURSE, chapter=CHAPTER)
    assert page.count("window.WQ_CIRCUITJS =") == 1 and "CircuitJSEmbedded" in page and "GPL-2.0" in page
    assert "connect-src 'none'" in page and "<!--WQ-CONFIG-->" in page
    assert not re.search(r"""<script[^>]+src=|<link[^>]+href=["']?(?!data:)""", page)   # nothing is fetched
    plain = labs.page([("2.1", labs.example_code())], course=COURSE, chapter=CHAPTER)
    assert "WQ_CIRCUITJS" not in plain.split("<script>\n", 1)[0] and "window.WQ_CIRCUITJS =" not in plain


def test_quantum_bench_finds_known_levels_and_keeps_probability():
    """《大学物理》第 9 轮 2.6: api.qm — infinite well and oscillator levels, orthogonal degenerate states, Crank–Nicolson keeps ∫|ψ|² = 1."""
    import shutil
    import subprocess
    if not shutil.which("node"):
        pytest.skip("node is not installed")
    kit = (labs.KIT / "kit.js").read_text(encoding="utf-8")
    qm = kit[kit.index("const QM = {"):kit.index("function graph(ctx, o)")]
    script = qm + """
let N = 999, dx = 1 / (N + 1), V = new Float64Array(N);
const w = QM.levels(V, dx, 3);
N = 1999; const a = -5; dx = 10 / (N + 1); V = new Float64Array(N);
for (let i = 0; i < N; i++) { const x = a + (i + 1) * dx; V[i] = 0.01 * x * x / (4 * QM.C); }
const ho = QM.levels(V, dx, 3);
N = 1500; dx = 75 / (N + 1); V = new Float64Array(N); const x = new Float64Array(N);
for (let i = 0; i < N; i++) { x[i] = (i + 1) * dx; V[i] = x[i] > 40 && x[i] < 41 ? 0.3 : 0; }
const p = QM.packet(x, 20, 3, QM.k(0.2));
for (let k = 0; k < 500; k++) QM.step(p, V, dx, 0.5);
const norm = QM.prob(p, dx);
N = 240; dx = 10 / (N + 1); V = new Float64Array(N);          // two identical deep wells: degenerate levels
for (let i = 0; i < N; i++) { const y = (i + 1) * dx; V[i] = (y > 1 && y < 3) || (y > 7 && y < 9) ? 0 : 2; }
const dw = QM.levels(V, dx, 2);
let ov = 0; for (let i = 0; i < N; i++) ov += dw.psi[0][i] * dw.psi[1][i] * dx;
console.log(JSON.stringify({ well: w.E, ho: ho.E, norm, dE: dw.E[1] - dw.E[0], overlap: ov }));
"""
    out = json.loads(subprocess.run(["node", "-e", script], capture_output=True, text=True, check=True).stdout)
    for n, e in enumerate(out["well"], 1):          # E_n = (ħ²/2m)(nπ/L)², L = 1 nm
        assert abs(e - 0.0380998212 * (n * math.pi) ** 2) < 1e-4 * e
    for n, e in enumerate(out["ho"]):                # E_n = (n + 1/2) ħω, ħω = 0.1 eV
        assert abs(e - (n + 0.5) * 0.1) < 1e-4
    assert abs(out["norm"] - 1) < 1e-9
    assert abs(out["dE"]) < 1e-9 and abs(out["overlap"]) < 1e-9   # degenerate states come out orthogonal


def test_gpu_teaching_model_hand_cases():
    """《人工智能》第 14 轮 2.5: api.gpu — divergence efficiency and Little's-law latency hiding on cases worked by hand."""
    import shutil
    import subprocess
    if not shutil.which("node"):
        pytest.skip("node is not installed")
    kit = (labs.KIT / "kit.js").read_text(encoding="utf-8")
    gpu = kit[kit.index("const GPU = {"):kit.index("// ---------- tasks & progress")]
    script = gpu + """
const d = [GPU.diverge({ cond: "odd", pre: 0, post: 0 }).eff, GPU.diverge({ cond: "warp", warps: 4 }).eff,
           GPU.diverge({ cond: "odd" }).issued];
const u = [[1, 4, 600], [16, 4, 600], [16, 40, 600], [16, 8, 100]].map(([n, k, L]) =>
  [GPU.schedule({ warps: n, k, latency: L }).util, GPU.little(n, k, L)]);
console.log(JSON.stringify({ d, u }));
"""
    out = json.loads(subprocess.run(["node", "-e", script], capture_output=True, text=True, check=True).stdout)
    assert out["d"] == [0.5, 1.0, 12]
    for got, want in out["u"]:
        assert abs(got - want) < 0.03 * want + 1e-3
