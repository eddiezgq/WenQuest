"""AI professor team course building (round 3, D30), end to end against a fake Moodle and the fake model."""
import asyncio
import re
import time
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import httpx
import pytest
from fastapi.testclient import TestClient

from app import main
from app import materials as mt
from app import studio as st
from app.ai import ModelGateway
from app.moodle import MoodleClient
from tests.test_materials import FOLDER, fake_moodle, login, upload

CALLS: list[tuple[str, dict]] = []


def moodle(request: httpx.Request) -> httpx.Response:
    url = urlparse(str(request.url))
    q = parse_qs(url.query)
    fn = (q.get("wsfunction") or [""])[0]
    if fn in ("local_wenquest_create_course", "local_wenquest_add_activities"):
        CALLS.append((fn, {k: v[0] for k, v in parse_qs(request.content.decode()).items()}))
        if fn == "local_wenquest_add_activities":
            return httpx.Response(200, json={"courseid": 7, "sectionid": 30, "cmids": [101, 102, 103]})
    return fake_moodle(request)


@pytest.fixture
def client(tmp_path):
    with TestClient(main.app) as c:
        main.state.http = httpx.AsyncClient(transport=httpx.MockTransport(moodle))
        main.state.moodle = MoodleClient("http://moodle.test", "moodle_mobile_app", main.state.http)
        main.state.ai = ModelGateway("fake", main.state.http)
        main.state.store = mt.Store(str(tmp_path / "imports"))
        main.state.projects_dir = str(tmp_path / "projects")
        CALLS.clear()
        yield c


def settle(c, h, pid, timeout=10.0):
    end = time.time() + timeout
    while time.time() < end:
        p = c.get(f"/api/v1/studio/projects/{pid}", headers=h).json()
        if not p["busy"]:
            return p
        time.sleep(0.05)
    raise AssertionError("the team did not finish")


def confirm_chapters(c, h, pid):
    """本章确认 for every chapter (round 5: a chapter is written only after the teacher confirms it)."""
    p = c.get(f"/api/v1/studio/projects/{pid}", headers=h).json()
    for ch in p["outline"]["chapters"]:
        c.post(f"/api/v1/studio/projects/{pid}/chapters/{ch['no']}/confirm", headers=h)


def make_project(c, h):
    p = c.post("/api/v1/studio/projects", headers=h, json={"description": "大学物理A（上），给大一工科学生"}).json()
    for path, data in FOLDER.items():
        r = c.post(f"/api/v1/studio/projects/{p['id']}/files", headers=h, data={"path": path},
                   files={"file": (path.rsplit("/", 1)[-1], data, "application/octet-stream")})
        assert r.status_code == 200, r.text
    return p["id"]


def test_whole_flow_from_materials_to_a_published_lesson(client):
    h = login(client)
    pid = make_project(client, h)
    p = client.get(f"/api/v1/studio/projects/{pid}", headers=h).json()
    assert p["stage"] == "intake" and len(p["files"]) == len(FOLDER) and p["messages"][0]["role"] == "lead"

    # 1. the librarian and the lead: materials list and at most five questions
    client.post(f"/api/v1/studio/projects/{pid}/start", headers=h)
    p = settle(client, h, pid)
    assert p["stage"] == "materials"
    roles = {f["path"]: f["role"] for f in p["files"]}
    assert roles["课程/大学物理教学大纲.docx"] == "syllabus" and roles["课程/第1章/第1章教案.docx"] == "lesson_plan"
    assert roles["课程/测验/期中参考答案.docx"] == "answer_key"
    open_q = [q for q in p["questions"] if q["status"] == "open"]
    assert 1 <= len(open_q) <= 5 and p["messages"][-1]["kind"] == "report"

    # the teacher answers a question and corrects one file
    lang_q = next(q for q in open_q if "语言" in q["text"])
    p = client.post(f"/api/v1/studio/projects/{pid}/questions/{lang_q['id']}", headers=h, json={"answer": "中文"}).json()
    assert p["requirements"]["language"] == "zh"
    txt = next(f for f in p["files"] if f["path"].endswith("新建文本文档.txt"))
    p = client.put(f"/api/v1/studio/projects/{pid}/materials", headers=h,
                   json={"files": [{"id": txt["id"], "role": "reference", "chapters": []}]}).json()
    assert next(f for f in p["files"] if f["id"] == txt["id"])["by"] == "teacher"

    # 2. the designer: outline with real names, quality-checked
    client.post(f"/api/v1/studio/projects/{pid}/approve-materials", headers=h)
    p = settle(client, h, pid)
    assert p["stage"] == "outline"
    titles = [c["title"]["zh"] for c in p["outline"]["chapters"]]
    assert "第1章 质点运动学" in titles and "第2章 牛顿运动定律" in titles
    assert st.check_outline(p["outline"]) == []

    # the teacher renames a lesson
    o = p["outline"]
    o["chapters"][0]["lessons"][0]["title"]["zh"] = "1.1 参考系、坐标系与质点"
    p = client.put(f"/api/v1/studio/projects/{pid}/outline", headers=h, json=o).json()
    assert p["outline"]["chapters"][0]["lessons"][0]["title"]["zh"] == "1.1 参考系、坐标系与质点"

    # 3. outline settled: the course is created with its chapters, nothing published yet
    p = client.post(f"/api/v1/studio/projects/{pid}/approve-outline", headers=h).json()
    confirm_chapters(client, h, pid)
    assert p["stage"] == "lessons" and p["course"]["id"] == 7
    created = dict(CALLS)["local_wenquest_create_course"]
    assert created["sections[1][name]"] == "第1章 质点运动学"

    # 4. write the next lesson: author, assessor, reviewer
    client.post(f"/api/v1/studio/projects/{pid}/lessons/next", headers=h)
    p = settle(client, h, pid)
    les = p["outline"]["chapters"][0]["lessons"][0]
    assert les["status"] == "awaiting" and les["review"]["verdict"] == "pass"
    assert "机器人问题" in les["content"]["zh"] and "虚拟实验" in les["content"]["zh"]
    assert les["exercises"]["zh"] and les["answers"]["zh"]
    kinds = {f["kind"]: f for f in les["files"]}
    assert set(kinds) == {"slides", "guide", "report", "plan", "figure"} and kinds["plan"]["teacher_only"]
    # 插图: drawn, checked, placed in the notes (preview links are signed) and on their own slides
    figs = [f for f in les["files"] if f["kind"] == "figure"]
    assert len(figs) == 2 and all(f["url"] in les["content"]["zh"] for f in figs)
    assert "wqfig/" not in les["content"]["zh"] and "图 1.1" in les["content"]["zh"]
    assert les["content"]["zh"].index("图 1.1") < les["content"]["zh"].index("图 1.2")
    assert client.get(figs[0]["url"]).content[:4] == b"\x89PNG"
    assert next(c for c in les["checklist"] if c["key"] == "figures")["ok"] is True
    assert p["progress"]["awaiting"] == 1
    # the teacher can download each deliverable to check it
    r = client.get(kinds["slides"]["url"])
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert client.get(kinds["plan"]["url"][:-4] + "abcd").status_code == 410

    # 5. the teacher approves: lesson, practice and a hidden answer key, plus the chapter's files
    p = client.post(f"/api/v1/studio/projects/{pid}/lessons/{les['id']}/approve", headers=h).json()
    les = p["outline"]["chapters"][0]["lessons"][0]
    assert les["status"] == "published" and les["cmids"]
    sent = [c for f, c in CALLS if f == "local_wenquest_add_activities"][-1]
    assert sent["section"] == "2"
    order = [(sent[f"activities[{i}][type]"], sent[f"activities[{i}][name]"], sent.get(f"activities[{i}][visible]", "1"))
             for i in range(20) if f"activities[{i}][type]" in sent]
    names = [n for _, n, _ in order]
    # benchmark order: notes, slides, practice, lab guide, report template, then teacher-only answers and plan
    assert order[0][0] == "page" and "课件" in order[1][1] and order[1][0] == "resource"
    # the notes page carries its figures: one draft area, content points at the page's own files
    assert "@@PLUGINFILE@@/fig-1.png" in sent["activities[0][content]"] and int(sent["activities[0][draftitemid]"]) > 0
    assert "练习" in order[2][1] and "实验指导书" in order[3][1] and "实验报告模板" in order[4][1]
    assert "答案" in order[5][1] and order[5][2] == "0" and "教案" in order[6][1] and order[6][2] == "0"
    assert "第1章教案" in names and "讲义" in names


def test_rewrite_with_a_note_and_stop(client):
    h = login(client)
    pid = make_project(client, h)
    client.post(f"/api/v1/studio/projects/{pid}/start", headers=h)
    settle(client, h, pid)
    client.post(f"/api/v1/studio/projects/{pid}/approve-materials", headers=h)
    settle(client, h, pid)
    client.post(f"/api/v1/studio/projects/{pid}/approve-outline", headers=h)
    confirm_chapters(client, h, pid)
    main.state.ai.fake_delay = 1.0
    client.post(f"/api/v1/studio/projects/{pid}/lessons/next", headers=h)
    time.sleep(0.2)
    assert client.get(f"/api/v1/studio/projects/{pid}", headers=h).json()["busy"]
    client.post(f"/api/v1/studio/projects/{pid}/stop", headers=h)
    p = settle(client, h, pid)
    assert p["outline"]["chapters"][0]["lessons"][0]["status"] == "planned"
    assert "已停止" in p["messages"][-1]["text"]
    main.state.ai.fake_delay = 0
    lid = p["outline"]["chapters"][0]["lessons"][0]["id"]
    client.post(f"/api/v1/studio/projects/{pid}/lessons/{lid}/write", headers=h, json={"note": "加一个 AGV 转弯的例子"})
    p = settle(client, h, pid)
    les = p["outline"]["chapters"][0]["lessons"][0]
    assert les["status"] == "awaiting" and les["notes"] == "加一个 AGV 转弯的例子"


def test_daily_pace_writes_one_lesson_a_day_and_waits_for_review(client):
    h = login(client)
    pid = make_project(client, h)
    client.post(f"/api/v1/studio/projects/{pid}/start", headers=h)
    settle(client, h, pid)
    client.post(f"/api/v1/studio/projects/{pid}/approve-materials", headers=h)
    settle(client, h, pid)
    client.post(f"/api/v1/studio/projects/{pid}/approve-outline", headers=h)
    confirm_chapters(client, h, pid)
    client.put(f"/api/v1/studio/projects/{pid}/pace", headers=h, json={"mode": "daily", "hour": 0, "tz": "UTC"})
    studio = main._studio()

    def tick():
        client.portal.call(studio.tick)  # run in the app's event loop, like the background loop does

    tick()
    p = settle(client, h, pid)
    assert p["progress"]["awaiting"] == 1 and p["pace"]["last_auto"]
    p["pace"]["last_auto"] = ""  # even on a new day, an unreviewed lesson blocks the next one
    proj = studio.projects.load(pid)
    proj["pace"]["last_auto"] = ""
    studio.projects.save(proj)
    tick()
    p = settle(client, h, pid)
    assert p["progress"]["awaiting"] == 1 and p["progress"]["planned"] == p["progress"]["total"] - 1


def test_projects_are_private_and_need_permission(client):
    h = login(client)
    pid = make_project(client, h)
    student = login(client, "s")
    assert client.get(f"/api/v1/studio/projects/{pid}", headers=student).status_code == 403
    assert client.post("/api/v1/studio/projects", headers=student, json={}).status_code == 403
    assert client.get("/api/v1/studio/projects", headers=student).json() == {"projects": []}
    mine = client.get("/api/v1/studio/projects", headers=h).json()["projects"]
    assert [x["id"] for x in mine] == [pid]


def test_quality_check_rejects_equations_and_empty_names():
    o = {"title": {"zh": "新课程"}, "chapters": [{"no": 1, "title": {"zh": "第1章"}, "lessons": [
        {"title": {"zh": "0.1 mm = 0.0001 m"}}, {"title": {"zh": "1.0598 T = 110 N"}}, {"title": {"zh": "1.2 Units and Standards"}}]}]}
    problems = st.check_outline(o)
    assert any("新课程" in x for x in problems) and any("第1章" in x for x in problems)
    assert sum("像算式" in x for x in problems) == 2


def test_textbook_contents_and_page_ranges():
    pages = {1: "University Physics", 2: "Contents\nChapter 1 Units and Measurement 1\n1.1 The Scope and Scale of Physics 2\n"
                                         "1.2 Units and Standards 4\nChapter 2 Vectors 9\n2.1 Scalars and Vectors 10\n2.2 Coordinate Systems 12",
             3: "Preface"}
    for n in range(4, 20):
        pages[n] = f"body page {n}"
    pages[5] += "\nCHAPTER 1 UNITS AND MEASUREMENT"
    pages[6] += "\n1.1 The Scope and Scale of Physics"
    pages[8] += "\n1.2 Units and Standards"
    pages[13] += "\nChapter 2 Vectors\n2.1 Scalars and Vectors"
    pages[15] += "\n2.2 Coordinate Systems"
    toc_pages = st.find_toc_pages(pages)
    assert toc_pages[0] == 2
    chapters = st.fake_toc(pages, toc_pages)["chapters"]
    assert [c["title"] for c in chapters] == ["Units and Measurement", "Vectors"]
    for c in chapters:
        c["sections"] = [{**s} for s in c["sections"]]
    st.locate(chapters, pages, 3)
    s11, s12 = chapters[0]["sections"]
    assert (s11["start"], s11["end"], s12["start"]) == (6, 7, 8)
    assert chapters[1]["sections"][0]["start"] == 13
    assert st.pages_of("[第1页]\nA\n\n[第2页]\nB") == {1: "A", 2: "B"}


EDDIE = Path(__file__).resolve().parents[3] / "samples" / "Eddie物理资料"


@pytest.mark.skipif(not EDDIE.exists(), reason="Eddie's physics materials are not in this checkout")
def test_eddies_materials_two_courses_one_chosen(client, monkeypatch):
    """Eddie's real files: university slides (Ch 1-10) and high-school slides (Units 5-14), no textbook.
    The model is scripted the way a good librarian answers; the gateway must turn that into a clean
    university outline with real chapter names and no high-school material."""
    from app import team
    h = login(client)
    p = client.post("/api/v1/studio/projects", headers=h, json={"description": ""}).json()
    pid = p["id"]
    for f in sorted(EDDIE.iterdir()):
        r = client.post(f"/api/v1/studio/projects/{pid}/files", headers=h, data={"path": f.name},
                        files={"file": (f.name, f.read_bytes(), "application/octet-stream")})
        assert r.status_code == 200, r.text
    real_json = main.state.ai.json

    async def scripted(*, system, prompt, schema, max_tokens=8000, fake=None):
        if system == team.LIBRARIAN:
            ids = re.findall(r"### id=(\w+) \| ([^|]+)", prompt)
            return {"course": {"title": "Physics for Scientists and Engineers I", "subject": "physics", "language": "en"},
                    "main_textbook": "", "summary": "两套课件：PHYS 2425（大学，第1-10章）和 High School Physics（第5-14单元）；主教材 OpenStax University Physics Vol 1 未包含。",
                    "files": [{"id": i, "role": "slides", "chapters": [int(re.search(r"(?:Chapter|Physics)_0*(\d+)", n).group(1))],
                               "title": n, "language": "en", "confidence": "high"} for i, n in ids],
                    "questions": [{"text": "资料里有两套课件：大学物理（PHYS 2425）和高中物理。这门课建哪一套？", "options": ["大学物理 PHYS 2425", "高中物理", "两套合并"]}]}
        if system == team.LIBRARIAN_UPDATE:
            ids = re.findall(r"- (\w+) \| ([^|]*HSPhysics[^|]*) \|", prompt)
            return {"files": [{"id": i, "role": "other", "chapters": [], "why": "高中课程，不用于本课"} for i, _ in ids], "note": ""}
        return await real_json(system=system, prompt=prompt, schema=schema, max_tokens=max_tokens, fake=fake)

    monkeypatch.setattr(main.state.ai, "json", scripted)
    client.post(f"/api/v1/studio/projects/{pid}/start", headers=h)
    p = settle(client, h, pid)
    assert p["requirements"]["course_title"] == "Physics for Scientists and Engineers I"
    which = next(q for q in p["questions"] if "哪一套" in q["text"])
    client.post(f"/api/v1/studio/projects/{pid}/questions/{which['id']}", headers=h, json={"answer": "大学物理 PHYS 2425"})
    client.post(f"/api/v1/studio/projects/{pid}/approve-materials", headers=h)
    p = settle(client, h, pid)
    assert any("调整了 10 份资料" in m["text"] for m in p["messages"])
    names = [c["title"]["zh"] for c in p["outline"]["chapters"]]
    assert len(names) == 10 and names[0] == "Chapter 1 Units and Measurement" and names[6] == "Chapter 7 Work and Kinetic Energy"
    ch7 = p["outline"]["chapters"][6]["lessons"]
    assert [l["title"]["zh"] for l in ch7] == ["Work", "Work Done by Forces that Vary", "Kinetic Energy", "Work-Energy Theorem", "Power"]
    assert st.check_outline(p["outline"]) == []
    assert not any("Keplers" in l["title"]["zh"] or "Magnetism" in c["title"]["zh"] for c in p["outline"]["chapters"] for l in c["lessons"])


def _write_first_lesson(client, h):
    pid = make_project(client, h)
    client.post(f"/api/v1/studio/projects/{pid}/start", headers=h)
    settle(client, h, pid)
    client.post(f"/api/v1/studio/projects/{pid}/approve-materials", headers=h)
    settle(client, h, pid)
    client.post(f"/api/v1/studio/projects/{pid}/approve-outline", headers=h)
    confirm_chapters(client, h, pid)
    client.post(f"/api/v1/studio/projects/{pid}/lessons/next", headers=h)
    return pid, settle(client, h, pid)


def test_animation_is_rendered_fixed_and_published(client, monkeypatch):
    """A2: the animator's scene is rendered; a failed render goes back to the animator with the error."""
    from app.production import anim
    sent: list[str] = []

    async def render(url, code, timeout=600.0):
        sent.append(code)
        if len(sent) == 1:
            raise anim.RenderError("NameError: name 'Arow' is not defined")
        from io import BytesIO
        from PIL import Image
        png = BytesIO()
        Image.new("RGB", (64, 36), "#0f1419").save(png, "PNG")
        return b"\x00\x00\x00\x18ftypmp42video", png.getvalue(), 42.0

    monkeypatch.setattr(anim, "render", render)
    monkeypatch.setattr(main.state.settings, "animator_url", "http://animator.test")
    h = login(client)
    pid, p = _write_first_lesson(client, h)
    les = p["outline"]["chapters"][0]["lessons"][0]
    assert len(sent) == 2 and "class Lesson(Base)" in sent[1]
    kinds = [f["kind"] for f in les["files"]]
    assert kinds[0] == "animation" and les["files"][0]["seconds"] == 42.0
    assert client.get(les["files"][0]["url"]).content.startswith(b"\x00\x00\x00\x18ftyp")
    assert "动画" in p["messages"][-1]["text"]
    # the slides carry the video on the animation slide
    from pptx import Presentation
    d = main._studio().lesson_dir(main._studio().projects.load(pid), les)
    deck = Presentation(str(next(d.glob("*课件.pptx"))))
    assert any(s.shape_type == 16 or "movie" in s.name.lower() or s.shape_type == 13
               for s in deck.slides[3].shapes if hasattr(s, "shape_type"))
    # published right after the lecture notes
    client.post(f"/api/v1/studio/projects/{pid}/lessons/{les['id']}/approve", headers=h)
    sent_acts = [c for f, c in CALLS if f == "local_wenquest_add_activities"][-1]
    assert sent_acts["activities[0][type]"] == "page"
    assert sent_acts["activities[1][type]"] == "resource" and "动画" in sent_acts["activities[1][name]"]


def test_animation_falls_back_to_the_storyboard(client, monkeypatch):
    from app.production import anim
    tries: list[str] = []

    async def render(url, code, timeout=600.0):
        tries.append(code)
        if len(tries) <= 3:  # the animator's three tries fail; the storyboard version renders
            raise anim.RenderError("render took longer than 420 s")
        return b"video", b"", 20.0

    monkeypatch.setattr(anim, "render", render)
    monkeypatch.setattr(main.state.settings, "animator_url", "http://animator.test")
    h = login(client)
    _, p = _write_first_lesson(client, h)
    les = p["outline"]["chapters"][0]["lessons"][0]
    assert len(tries) == 4  # the animator's code three times, then the storyboard version
    assert "agv(" not in tries[-1] and "cargo(" not in tries[-1]  # the plain version draws no stock robot
    assert les["files"][0]["kind"] == "animation"
    # honest: the lesson needs the teacher, the message says so, the checklist crosses the animation
    assert [x["kind"] for x in les["attention"]] == ["animation"] and "简版" in les["attention"][0]["text"]
    assert "需要你处理" in p["messages"][-1]["text"] and "审稿通过" not in p["messages"][-1]["text"]
    anim_item = next(c for c in les["checklist"] if c["key"] == "animation")
    assert anim_item["ok"] is False


def test_no_renderer_no_video(client, monkeypatch):
    from app.production import anim

    async def down(url, code, timeout=600.0):
        raise anim.RenderError("renderer unreachable: ConnectError", "service")

    monkeypatch.setattr(anim, "render", down)
    monkeypatch.setattr(main.state.settings, "animator_url", "http://animator.test")
    h = login(client)
    _, p = _write_first_lesson(client, h)
    les = p["outline"]["chapters"][0]["lessons"][0]
    assert les["status"] == "awaiting" and "animation" not in [f["kind"] for f in les["files"]]
    assert any("动画渲染服务暂时不可用" in i["text"] for i in les["review"]["issues"])


def test_materials_can_be_added_downloaded_and_deleted_at_any_stage(client):
    """Materials list: add files after the outline (sorted at once, written lessons named for a rewrite),
    download one or all (zip, original folders), delete one."""
    import io
    import zipfile
    from tests.test_materials import docx_bytes
    h = login(client)
    pid, p = _write_first_lesson(client, h)
    written = [f"{c['no']}.{i + 1}" for c in p["outline"]["chapters"] for i, les in enumerate(c["lessons"]) if les["status"] == "awaiting"]
    before = len(p["files"])
    r = client.post(f"/api/v1/studio/projects/{pid}/files", headers=h, data={"path": "新资料/第1章/补充例题.docx"},
                    files={"file": ("补充例题.docx", docx_bytes(["第1章 补充例题", "例1 AGV 转弯"]), "application/octet-stream")})
    assert r.status_code == 200
    client.post(f"/api/v1/studio/projects/{pid}/files/done", headers=h)
    p = settle(client, h, pid)
    new = next(f for f in p["files"] if f["name"] == "补充例题.docx")
    assert len(p["files"]) == before + 1 and new["role"] and new["url"]
    msg = p["messages"][-1]["text"]
    assert "新加了 1 份资料" in msg and "补充例题.docx" in msg and "还没写的课会自动用上" in msg
    assert not written or (written[0] in msg and "重写" in msg)
    # download one, then all
    one = client.get(new["url"])
    assert one.status_code == 200 and one.content[:2] == b"PK" and "attachment" in one.headers["content-disposition"]
    z = zipfile.ZipFile(io.BytesIO(client.get(p["zip_url"]).content))
    assert "新资料/第1章/补充例题.docx" in z.namelist() and len(z.namelist()) == before + 1
    # someone else's link does not open, a forged one neither
    assert client.get("/api/v1/studio/materials/not-a-token").status_code == 410
    # delete
    p = client.delete(f"/api/v1/studio/projects/{pid}/files/{new['id']}", headers=h).json()
    assert len(p["files"]) == before and all(f["id"] != new["id"] for f in p["files"])
    assert client.delete(f"/api/v1/studio/projects/{pid}/files/{new['id']}", headers=h).status_code == 404
    s = login(client, "s")
    assert client.delete(f"/api/v1/studio/projects/{pid}/files/{p['files'][0]['id']}", headers=s).status_code == 403
