"""Slides presented in the browser (round 3, step 3): parsing, conversion, API and permissions."""
import io
import shutil
import time

import httpx
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app import main
from app import slides as sl
from app.moodle import MoodleClient
from tests.test_api import BASE, fake_moodle, login

DECK_URL = f"{BASE}/webservice/pluginfile.php/9/mod_resource/content/0/ch1.pptx"
TEACHER = {"on": False}


def make_deck() -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "第一章 质点运动学"
    s1.notes_slide.notes_text_frame.text = "先问学生：车厢里的人看到什么？"
    s2 = prs.slides.add_slide(blank)
    s2.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1)).text_frame.text = "1.3  虚拟实验  Virtual lab"
    button = s2.shapes.add_textbox(Inches(8), Inches(5), Inches(4), Inches(1))
    run = button.text_frame.paragraphs[0].add_run()
    run.text = "打开虚拟实验"
    run.hyperlink.address = "https://example.test/labs/ch1.html#lab-1-3"
    s2.shapes.add_movie(io.BytesIO(b"\x00\x00\x00\x18ftypmp42fake-video"), Inches(1), Inches(2), Inches(6),
                        Inches(3.375), mime_type="video/mp4")
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


DECK = make_deck()


def fake(request: httpx.Request) -> httpx.Response:
    url = str(request.url)
    if url.startswith(DECK_URL):
        return httpx.Response(200, content=DECK, headers={"content-type": "application/vnd.openxmlformats-officedocument.presentationml.presentation"})
    body = request.content.decode()
    if "core_course_get_course_module" in url and "cmid=9" in body:
        return httpx.Response(200, json={"cm": {"id": 9, "course": 2, "instance": 3, "modname": "resource", "name": "第一章课件"}})
    if "core_course_get_contents" in url:
        return httpx.Response(200, json=[{"id": 11, "section": 1, "name": "第1章", "summary": "", "modules": [
            {"id": 9, "modname": "resource", "name": "第一章课件", "uservisible": True, "visible": 1,
             "contents": [{"type": "file", "filename": "ch1.pptx", "filesize": len(DECK), "timemodified": 1,
                           "mimetype": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                           "fileurl": DECK_URL}]}]}])
    if "core_course_get_user_administration_options" in url:
        return httpx.Response(200, json={"courses": [{"id": 2, "options": [{"name": "update", "available": TEACHER["on"]}]}]})
    return fake_moodle(request)


@pytest.fixture
def client(tmp_path):
    with TestClient(main.app) as c:
        main.state.http = httpx.AsyncClient(transport=httpx.MockTransport(fake))
        main.state.moodle = MoodleClient(BASE, "moodle_mobile_app", main.state.http)
        main.state.slides = sl.SlideStore(str(tmp_path / "slides"))
        main.state.slide_settings = sl.Settings(str(tmp_path / "settings.json"))
        TEACHER["on"] = False
        yield c


# --- parsing (no LibreOffice needed) -----------------------------------------------------------

def test_parse_finds_labs_links_videos_and_notes():
    info = sl.parse_pptx(DECK)
    s1, s2 = info["slides"]
    assert s1["notes"] == "先问学生：车厢里的人看到什么？" and s1["labs"] == []
    assert s2["labs"] == ["1.3"]
    link = s2["links"][0]
    assert link["lab"] == "1.3" and link["href"].endswith("#lab-1-3")
    assert 0.55 < link["x"] < 0.65 and 0.6 < link["y"] < 0.7
    video = s2["videos"][0]
    assert video["src"] == "v1.mp4" and 0.07 < video["x"] < 0.08 and 0.44 < video["w"] < 0.46
    assert list(info["media"].values()) == ["v1.mp4"]


def test_lab_references_in_text():
    assert sl.lab_refs("1.3  虚拟实验  Virtual lab") == ["1.3"]
    assert sl.lab_refs("课后做实验1.4，再做虚拟实验 2.1") == ["1.4", "2.1"]
    assert sl.lab_refs("Lab 3.2 and lab 3-2") == ["3.2"]
    assert sl.lab_refs("第1.3节 抛体运动") == []


def test_slide_files_are_whitelisted(tmp_path):
    store = sl.SlideStore(str(tmp_path))
    assert store.file("../etc", "passwd") is None
    assert store.file("a" * 64, "../../x.webp") is None
    assert store.file("a" * 64, "p001.webp") is None  # well-formed but missing


# --- API ----------------------------------------------------------------------------------------

def test_slides_unavailable_without_converter(client):
    main.state.slides.available = False
    h = login(client)
    r = client.get("/api/v1/slides/9", headers=h).json()
    assert r["status"] == "unavailable" and r["teacher"] is False
    assert r["download_url"] is None  # students do not download by default


def test_download_setting_is_teacher_only(client):
    main.state.slides.available = False
    h = login(client)
    assert client.put("/api/v1/slides/9/settings", headers=h, json={"allow_download": True}).status_code == 403
    TEACHER["on"] = True
    assert client.put("/api/v1/slides/9/settings", headers=h, json={"allow_download": True}).json() == {"allow_download": True}
    TEACHER["on"] = False
    r = client.get("/api/v1/slides/9", headers=h).json()
    assert r["allow_download"] is True and r["download_url"].startswith("/api/v1/files/")


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice is not installed here")
def test_convert_and_present(client):
    h = login(client)
    first = client.get("/api/v1/slides/9", headers=h).json()
    assert first["status"] in ("converting", "ready")
    for _ in range(300):
        r = client.get("/api/v1/slides/9", headers=h).json()
        if r["status"] != "converting":
            break
        time.sleep(0.2)
    assert r["status"] == "ready", r
    assert r["pages"] == 2 and r["width"] == 1600 and len(r["slides"]) == 2
    s2 = r["slides"][1]
    assert s2["labs"] == ["1.3"] and s2["videos"][0]["src"].endswith("/v1.mp4")
    assert "notes" not in r["slides"][0]  # students never receive speaker notes
    img = client.get(s2["image"])
    assert img.status_code == 200 and img.headers["content-type"] == "image/webp"
    assert "immutable" in img.headers["cache-control"]
    video = client.get(s2["videos"][0]["src"], headers={"Range": "bytes=0-3"})
    assert video.status_code == 206 and video.content == b"\x00\x00\x00\x18"
    TEACHER["on"] = True
    t = client.get("/api/v1/slides/9", headers=h).json()
    assert t["slides"][0]["notes"] == "先问学生：车厢里的人看到什么？" and t["download_url"]


def test_a_deck_someone_opens_jumps_the_queue(tmp_path, monkeypatch):
    """Decks converted ahead of time wait; the deck a person is looking at goes next (Eddie, 2026-09-28)."""
    import asyncio
    order = []

    def fake_convert(data, folder, ext=".pptx"):
        time.sleep(0.05)
        order.append(data.decode())
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "manifest.json").write_text(f'{{"version": {sl.VERSION}, "slides": []}}')

    monkeypatch.setattr(sl, "convert", fake_convert)

    async def scenario():
        store = sl.SlideStore(str(tmp_path))
        store.available = True

        def fetch(name):
            async def f():
                return name.encode()
            return f

        for n in ("a", "b", "c", "d"):
            store.start(n, fetch(n))            # ahead of time (priority 1)
        await asyncio.sleep(0.01)                # "a" is converting now
        assert store.progress("d")["queue"] == 3
        store.start("d", fetch("d"), priority=0)  # someone opens "d"
        assert store.progress("d")["queue"] == 1  # only the deck already converting is ahead
        while store.pending or store.current:
            await asyncio.sleep(0.02)
        return store

    store = asyncio.run(scenario())
    assert order == ["a", "d", "b", "c"]
    assert store.overview()["converted"] == 4 and store.overview()["waiting"] == 0
