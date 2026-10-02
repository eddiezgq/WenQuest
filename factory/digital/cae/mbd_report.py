# -*- coding: utf-8 -*-
"""运动与动力分析的 Word 报告（第 12 轮 D5）：模型、设置、驱动汇总、关节受力峰值、动画截图与曲线、电机选型、AI 解释。"""
import base64
import datetime as dt
import io

KIND = {"speed": "匀速", "move": "点到点（梯形速度）", "sine": "正弦", "hold": "保持", "torque": "给定力矩 / 力", "coupled": "机构带动"}


def _model_name(job):
    m = job.get("model") or {}
    if m.get("source") == "mech":
        from cae import mech_mjcf as MC
        return "零件库机构 {} {}".format(m.get("id"), MC.NAMES.get(m.get("id"), ""))
    if m.get("source") == "library":
        from cae import mbd_models as MM
        return "零件库机械臂 {} {}".format(m.get("id"), MM.ROBOTS.get(m.get("id"), ("", ""))[1])
    return "上传的模型"


def _drive_text(d):
    k = d.get("kind")
    if k == "speed":
        return "匀速 {:g}（rad/s 或 m/s）".format(d.get("value", 0))
    if k == "move":
        return "点到点：到 {:g}，{:g}–{:g} 秒".format(d.get("to", 0), d.get("t0", 0), d.get("t1", 0))
    if k == "sine":
        return "正弦：幅值 {:g}，{:g} Hz".format(d.get("amp", 0), d.get("freq", 0))
    if k == "torque":
        return "给定力矩 / 力 {:g}".format(d.get("value", 0))
    return KIND.get(k, k)


def build(job, peaks, images=(), motors=None, ai_text=None):
    """peaks：[(构件, 最大反力 N, 时刻 s, 最大反力矩 N·m, 时刻 s)]，由调用方从曲线里取"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    st, setup = job["stats"], job["setup"]
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for name in ("Title", "Heading 1"):
        d.styles[name].element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)

    def table(rows):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for k, v in enumerate(r):
                t.cell(i, k).text = str(v)
                if i == 0:
                    for run in t.cell(i, k).paragraphs[0].runs:
                        run.bold = True
        d.add_paragraph()

    d.add_heading("运动与动力分析报告", 0)
    table([["项目", "内容"], ["计算名称", job.get("title") or "—"], ["模型", _model_name(job)],
           ["计算人", job.get("owner_name") or "—"], ["时间", dt.datetime.fromtimestamp(job["created"]).strftime("%Y-%m-%d %H:%M")],
           ["任务编号", job["id"]]])

    d.add_heading("1  计算设置", 1)
    table([["项目", "内容"],
           ["仿真时长", "{:g} 秒（步长 {:g} 秒）".format(st["duration_s"], st["timestep_s"])],
           ["重力", "计入（−Z，9.81 m/s²）" if setup.get("gravity", True) else "不计"],
           ["初始位置", "；".join("{} = {:.4g}".format(k, v) for k, v in (setup.get("initial") or {}).items()) or "模型默认"],
           ["末端负载", "；".join("{} 上 {:g} kg".format(p["body"], p["mass"]) for p in setup.get("payloads") or []) or "无"],
           ["外力", "；".join("{} 上 ({:g}, {:g}, {:g}) N".format(f["body"], *f["force"]) for f in setup.get("forces") or []) or "无"]])
    table([["关节", "驱动方式"]] + [[x["joint"], _drive_text(x)] for x in setup.get("drives") or []])
    p = d.add_paragraph("求解：MuJoCo 多体动力学（隐式积分），给定运动的驱动以约束实现，驱动力矩为该约束的约束力；不计碰撞。")
    p.runs[0].font.size = Pt(9)

    d.add_heading("2  驱动力矩与功率", 1)
    table([["关节", "方式", "峰值", "均方根", "最高速度", "峰值功率 W", "平均功率 W"]] +
          [[x["joint"], KIND.get(x["kind"], x["kind"]), "{:.3g}".format(x["peak"]), "{:.3g}".format(x["rms"]),
            "{:.3g}".format(x["speed_max"]), "{:.3g}".format(x["power_peak"]), "{:.3g}".format(x["power_mean"])] for x in st["drives"]])
    d.add_paragraph("力矩单位：转动关节 N·m、移动关节 N；速度：rad/s 或 m/s。")

    d.add_heading("3  关节受力峰值", 1)
    if peaks:
        table([["构件（受父构件的力）", "最大反力 N", "时刻 s", "最大反力矩 N·m", "时刻 s"]] +
              [[b, "{:.4g}".format(f), "{:.3f}".format(tf), "{:.4g}".format(m), "{:.3f}".format(tm)] for b, f, tf, m, tm in peaks])

    for img in images:
        try:
            raw = base64.b64decode(img["data"].split(",", 1)[-1])
        except Exception:  # noqa: BLE001
            continue
        d.add_picture(io.BytesIO(raw), width=Cm(15.5))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = d.add_paragraph(img.get("caption") or "")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    n = 4
    if motors and motors.get("joints"):
        d.add_heading("{}  电机选型".format(n), 1)
        n += 1
        table([["关节", "峰值 N·m", "均方根 N·m", "最高 r/min", "推荐", "裕量", "是否够用"]] +
              [[m["joint"], "{:.1f}".format(m["need"]["peak_Nm"]), "{:.1f}".format(m["need"]["rms_Nm"]), "{:.1f}".format(m["need"]["speed_rpm"]),
                "{} + {}，i = {}".format(m["best"]["motor"], m["best"]["gear"], m["best"]["ratio"]),
                "很大" if m["best"]["margin_pct"] is None else "{}%".format(m["best"]["margin_pct"]), "够用" if m["ok"] else "不够"]
               for m in motors["joints"]])
        p = d.add_paragraph("安全系数 {}；参数：{}。".format(motors["joints"][0]["safety"], motors["table"]["source"]))
        p.runs[0].font.size = Pt(9)
    if ai_text:
        d.add_heading("{}  AI 分析与改进建议".format(n), 1)
        for para in str(ai_text).split("\n"):
            if para.strip():
                d.add_paragraph(para.strip())
        p = d.add_paragraph("（AI 根据上面的设置和结果写出，仅供参考。）")
        p.runs[0].font.size = Pt(9)
    p = d.add_paragraph("问渠数字工厂 · 运动与动力分析 生成。刚体模型，不计构件弹性、间隙和摩擦（设置了关节阻尼、摩擦的除外）；用于教学和方案比较。")
    p.runs[0].font.size = Pt(8)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
