# -*- coding: utf-8 -*-
"""计算报告（第 11 轮 F4）：Word 文件，带设置、网格、结果、云图、结论。云图由浏览器截好传过来。"""
import base64
import datetime as dt
import io

from cae import materials as M

DOF_NAME = {"radial": "径向", "tangential": "转动", "axial": "轴向"}
SF_LINE = 1.5          # 判断线：一般机械零件静强度常用 1.3–1.5，具体以设计规范或任务书为准


def describe_load(l):
    f = "面 " + "、".join(str(x) for x in l["faces"])
    t = l["type"]
    if t == "fixed":
        return "固定", f, "三个方向都不能动"
    if t == "cyl_support":
        dofs = l.get("dofs") or ["radial"]
        kind = "限制转动（联轴器 / 键连接）" if dofs == ["tangential"] else "轴承支承"
        return kind, f, "限制" + "、".join(DOF_NAME[d] for d in dofs)
    if t == "force":
        v = l["vector_n"]
        return "力", f, "Fx {:g}，Fy {:g}，Fz {:g} N（均匀分布）".format(*v)
    if t == "pressure":
        return "压力", f, "{:g} MPa".format(l["value_mpa"])
    if t == "torque":
        return "扭矩", f, "{:g} N·m，绕轴线 {}".format(l["value_nmm"] / 1000, _vec(l["axis"]["dir"]))
    return t, f, ""


def _vec(v):
    return "(" + ", ".join("{:g}".format(round(x, 3)) for x in v) + ")"


def conclusion(stats, mat):
    sf = stats.get("safety_factor")
    if sf is None:
        return "没有强度数据，无法判断。"
    kind = "屈服强度" if mat.get("yield_mpa") else "抗拉强度"
    s = "按{} {} MPa 计，安全系数 {:.2f}。".format(kind, mat["strength_mpa"], sf)
    if sf < 1:
        s += "最大应力已超过材料强度，零件在这个工况下会{}，必须修改设计（加大尺寸、加圆角、换更强的材料）或减小载荷。".format(
            "屈服（产生永久变形）" if mat.get("yield_mpa") else "断裂")
    elif sf < SF_LINE:
        s += "低于常用判断线 {}，裕量不足，建议改进；若最大应力出现在尖角处，先加圆角再算，或按规范的应力集中系数校核。".format(SF_LINE)
    else:
        s += "高于常用判断线 {}，静强度满足要求（交变载荷下还要做疲劳校核）。".format(SF_LINE)
    return s


def build(job, images=(), ai_text=None):
    if (job.get("setup") or {}).get("analysis") in ("thermal", "thermo_mech"):
        return build_thermal(job, images, ai_text)
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    st = job["stats"]
    setup = job["setup"]
    mat = M.get(setup["material_id"])
    srcs = [M.SRC[s] for s in M.BY_ID[mat["id"]]["src"]]
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for name in ("Title", "Heading 1", "Heading 2"):
        d.styles[name].element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)

    def table(rows, widths=None, head=True):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                c = t.cell(i, j)
                c.text = str(v)
                if head and i == 0:
                    for run in c.paragraphs[0].runs:
                        run.bold = True
                if widths:
                    c.width = Cm(widths[j])
        d.add_paragraph()
        return t

    d.add_heading("有限元强度计算报告", 0)
    created = dt.datetime.fromtimestamp(job["created"]).strftime("%Y-%m-%d %H:%M")
    table([["项目", "内容"],
           ["计算名称", job.get("title") or "—"],
           ["零件", job.get("item") or "上传的零件"],
           ["计算人", job.get("owner_name") or "—"],
           ["时间", created],
           ["任务编号", job["id"]]], [4, 12])

    d.add_heading("1  材料", 1)
    table([["材料", "弹性模量 E", "泊松比 ν", "屈服强度", "抗拉强度", "疲劳极限 σ₋₁"],
           [mat["name"], "{:g} MPa".format(mat["E_mpa"]), "{:g}".format(mat["nu"]),
            "{} MPa".format(mat["yield_mpa"]) if mat["yield_mpa"] else "—（脆性）", "{} MPa".format(mat["ultimate_mpa"]),
            "{} MPa".format(mat["sigma_1"])]])
    p = d.add_paragraph("说明：" + mat["note"] + "。出处：" + "；".join(srcs) + "。")
    p.runs[0].font.size = Pt(9)

    d.add_heading("2  约束与载荷", 1)
    table([["类型", "作用面", "大小 / 方式"]] + [list(describe_load(l)) for l in setup["loads"]], [5, 3.5, 7.5])

    d.add_heading("3  网格与求解", 1)
    table([["项目", "内容"],
           ["单元", "二阶四面体（10 节点，CalculiX C3D10）"],
           ["单元尺寸", "{} mm（载荷面附近加密到约 1/3）".format(st.get("mesh_size_mm"))],
           ["规模", "{} 个单元，{} 个节点".format(st["elements"], st["nodes"])],
           ["求解", "线弹性静力，CalculiX；网格 Gmsh；计算用时 {} 秒".format(st["seconds"])]], [4, 12])

    d.add_heading("4  结果", 1)
    rows = [["项目", "数值", "位置（mm）"],
            ["最大 Von Mises 应力（评估值）", "{:.1f} MPa".format(st["vm_max_mpa"]),
             "{}，面 {}".format(_vec(st["vm_max_at"]), "、".join(str(x) for x in st.get("vm_max_faces") or []) or "内部")],
            ["最大位移", "{:.4g} mm".format(st["u_max_mm"]), _vec(st["u_max_at"])]]
    if st.get("safety_factor") is not None:
        rows.append(["安全系数（强度 ÷ 最大应力）", "{:.2f}".format(st["safety_factor"]), ""])
    table(rows, [6, 4, 6])
    if st.get("vm_peak_all_mpa", 0) > st["vm_max_mpa"] * 1.01:
        d.add_paragraph("约束面附近最高 {:.1f} MPa，是约束方式造成的局部值（实际支承没有那么“死”），评估时避开了约束面 {:.1f} mm 以内的点。".format(
            st["vm_peak_all_mpa"], 1.5 * (st.get("mesh_size_mm") or 0)))
    d.add_paragraph("注意：没有圆角的内角（尖角）处，应力理论上没有上限，网格越细数值越大。最大值出现在尖角时，应加圆角后重算，或按规范的应力集中系数校核。")

    for img in images:
        try:
            raw = base64.b64decode(img["data"].split(",", 1)[-1])
        except Exception:  # noqa: BLE001
            continue
        d.add_picture(io.BytesIO(raw), width=Cm(15.5))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = d.add_paragraph(img.get("caption") or "")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs and setattr(cap.runs[0].font, "size", Pt(9))

    fat = job.get("fatigue")
    n = 5
    if fat:
        d.add_heading("5  疲劳寿命", 1)
        table([["项目", "内容"],
               ["载荷谱", fat.get("label") or "—"],
               ["计算工况载荷", "{} {}".format(fat.get("ref_load"), fat.get("ref_unit") or "")],
               ["每块循环数（雨流计数，pyLife 四点法）", str(fat["cycles_per_block"])],
               ["S-N 曲线", "σ₋₁ = {} MPa × 表面 β {}（{}）× 尺寸 ε {} ÷ Kf {} = S_D {} MPa；N_D = {:g}，k = {:g}{}".format(
                   mat["sigma_1"], fat["beta"], fat["surface"], fat["size_factor"], fat["kf"], fat["S_D"], fat["N_D"], fat["k"],
                   "，低于 S_D 按 Haibach 斜率 {:g}".format(fat["k2"]) if fat.get("haibach") else "，低于 S_D 不计损伤")],
               ["平均应力修正", "Goodman：σa,eq = σa / (1 − |σm| / σb)，σb = {} MPa".format(mat["ultimate_mpa"])],
               ["累积损伤", "Miner 线性累积"],
               ["最危险点", "{}（计算工况下 {} MPa）".format(_vec(fat["hot_at"]), fat["hot_vm_ref"])],
               ["寿命", "无限寿命（应力幅都低于疲劳极限）" if fat.get("infinite") else "{:.3g} 块 ≈ {:.3g} 次循环 ≈ {:.3g} 小时".format(
                   fat["life_blocks"], fat["life_blocks"] * fat["cycles_per_block"], fat["life_hours"])]], [5, 11])
        if fat.get("hot_cycles"):
            table([["载荷幅", "载荷均值", "σa MPa", "σm MPa", "σa,eq MPa", "该幅值下的寿命 N"]] + [
                [c["load_amp"], c["load_mean"], c["sigma_a"], c["sigma_m"], c["sigma_a_eq"],
                 "∞" if c["N"] is None else "{:.3g}".format(c["N"])] for c in fat["hot_cycles"]])
        n = 6
    d.add_heading("{}  结论".format(n), 1)
    d.add_paragraph(conclusion(st, mat))
    if fat:
        d.add_paragraph("疲劳：" + ("在这个载荷谱下为无限寿命。" if fat.get("infinite") else
                                    "按这个载荷谱连续运行约 {:.3g} 小时达到疲劳损伤 1（出现裂纹）。".format(fat["life_hours"])))
    if ai_text:
        d.add_heading("{}  AI 分析与改进建议".format(n + 1), 1)
        for para in str(ai_text).split("\n"):
            if para.strip():
                d.add_paragraph(para.strip())
        p = d.add_paragraph("（AI 根据上面的设置和结果写出，仅供参考，请结合手算和规范复核。）")
        p.runs[0].font.size = Pt(9)

    p = d.add_paragraph("问渠数字工厂 · 仿真与分析 生成。计算模型为线弹性小变形，没有考虑接触、塑性、残余应力和表面状态；用于教学和方案比较，正式设计请按相关标准复核。")
    p.runs[0].font.size = Pt(8)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


TH_NAME = {"temperature": "固定温度", "convection": "对流散热", "heat_flux": "面发热", "heat_body": "整体发热"}


def describe_thermal(l):
    f = "其余所有面" if l.get("faces") == "rest" else ("整个零件" if l["type"] == "heat_body" else "面 " + "、".join(str(x) for x in l["faces"]))
    if l["type"] == "temperature":
        return TH_NAME[l["type"]], f, "{:g} ℃".format(l["value_c"])
    if l["type"] == "convection":
        return TH_NAME[l["type"]], f, "散热系数 {:g} W/(m²·K)，环境 {:g} ℃".format(l["h_w_m2k"], l["t_inf_c"])
    return TH_NAME.get(l["type"], l["type"]), f, "{:g} W".format(l["power_w"])


def build_thermal(job, images=(), ai_text=None):
    """热分析 / 热—结构耦合的报告（第 14 轮）"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    st, setup = job["stats"], job["setup"]
    mat = M.get(setup["material_id"])
    coupled = setup.get("analysis") == "thermo_mech"
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for name in ("Title", "Heading 1", "Heading 2"):
        d.styles[name].element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)

    def table(rows, widths=None):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                c = t.cell(i, j)
                c.text = str(v)
                if i == 0:
                    for run in c.paragraphs[0].runs:
                        run.bold = True
                if widths:
                    c.width = Cm(widths[j])
        d.add_paragraph()

    d.add_heading("热—结构耦合计算报告" if coupled else "温度场计算报告", 0)
    table([["项目", "内容"], ["计算名称", job.get("title") or "—"], ["零件", job.get("item") or "上传的零件"],
           ["计算人", job.get("owner_name") or "—"], ["时间", dt.datetime.fromtimestamp(job["created"]).strftime("%Y-%m-%d %H:%M")],
           ["任务编号", job["id"]]], [4, 12])
    d.add_heading("1  材料", 1)
    rows = [["材料", "导热系数 k", "比热 c", "密度", "线膨胀系数 α"],
            [mat["name"], "{:g} W/(m·K)".format(mat["k_w_mk"]), "{:g} J/(kg·K)".format(mat["c_j_kgk"]),
             "{:g} g/cm³".format(mat["density"]), "{:g} ×10⁻⁶/K".format(mat["alpha_1e6"])]]
    table(rows)
    p = d.add_paragraph("出处：" + mat["thermal_src"] + "。")
    p.runs[0].font.size = Pt(9)
    d.add_heading("2  热边界与载荷", 1)
    table([["类型", "作用面", "大小"]] + [list(describe_thermal(l)) for l in setup.get("thermal") or []], [4, 4, 8])
    if setup.get("transient"):
        tr = setup["transient"]
        d.add_paragraph("瞬态：初始温度 {:g} ℃，计算 {:g} 秒。".format(tr.get("t0_c", 20), tr["duration_s"]))
    if coupled:
        table([["类型", "作用面", "大小 / 方式"]] + [list(describe_load(l)) for l in setup.get("loads") or []], [5, 3.5, 7.5])
        d.add_paragraph("无应力参考温度 {:g} ℃（零件在这个温度下没有热应力）。".format(setup.get("ref_temp_c", 20)))
    d.add_heading("3  网格与求解", 1)
    table([["项目", "内容"], ["单元", "二阶四面体（10 节点，CalculiX C3D10）"],
           ["规模", "{} 个单元，{} 个节点，单元尺寸 {} mm".format(st["elements"], st["nodes"], st.get("mesh_size_mm"))],
           ["求解", ("稳态热—结构耦合" if coupled else "瞬态导热" if setup.get("transient") else "稳态导热") + "，CalculiX；用时 {} 秒".format(st["seconds"])]],
          [4, 12])
    d.add_heading("4  结果", 1)
    rows = [["项目", "数值", "位置（mm）"],
            ["最高温度", "{:.1f} ℃".format(st["t_max_c"]), _vec(st["t_max_at"])],
            ["最低温度", "{:.1f} ℃".format(st["t_min_c"]), _vec(st["t_min_at"])],
            ["表面平均温度", "{:.1f} ℃".format(st["t_surface_mean_c"]), ""],
            ["输入热量 / 对流散走", "{:.1f} W / {:.1f} W".format(st["heat_in_w"], st["heat_out_convection_w"]), "两者应相等（稳态）"]]
    for g in st.get("film_groups") or []:
        rows.append(["对流面组 {} 平均温度".format(g["load"] + 1), "{:.1f} ℃".format(g["mean_c"]),
                     "面积 {:.4f} m²，散热 {:.1f} W".format(g["area_m2"], g["heat_w"])])
    if coupled:
        rows += [["最大 Von Mises 应力（评估值）", "{:.1f} MPa".format(st["vm_max_mpa"]), _vec(st["vm_max_at"])],
                 ["最大热变形", "{:.4g} mm".format(st["u_max_mm"]), _vec(st["u_max_at"])]]
        if st.get("safety_factor") is not None:
            rows.append(["安全系数", "{:.2f}".format(st["safety_factor"]), ""])
    table(rows, [5, 4, 7])
    f = setup.get("formula")
    n = 5
    if f:
        d.add_heading("5  与教材公式对比", 1)
        table([["项目", "内容"]] + [[k, str(v)] for k, v in f.items() if k not in ("name", "note")] + [["公式", f.get("name", "")],
                                                                                                    ["说明", f.get("note", "")]], [5, 11])
        n = 6
    for img in images:
        try:
            raw = base64.b64decode(img["data"].split(",", 1)[-1])
        except Exception:  # noqa: BLE001
            continue
        d.add_picture(io.BytesIO(raw), width=Cm(15.5))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = d.add_paragraph(img.get("caption") or "")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if ai_text:
        d.add_heading("{}  AI 分析与改进建议".format(n), 1)
        for para in str(ai_text).split("\n"):
            if para.strip():
                d.add_paragraph(para.strip())
    p = d.add_paragraph("问渠数字工厂 · 仿真与分析 生成。表面散热按给定的散热系数（对流与辐射合计），不考虑接触热阻、辐射细节和内部油液流动；用于教学和方案比较。")
    p.runs[0].font.size = Pt(8)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
