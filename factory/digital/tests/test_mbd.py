# -*- coding: utf-8 -*-
"""第 12 轮 运动与动力分析：标准题（单摆周期、机械臂重力矩、动态力矩与逆动力学互核）、计算服务的动力学任务。
机器人模型要 Menagerie（WQ_MENAGERIE，CI 用 deploy/fetch_menagerie.sh 取）；没有就跳过机器人的几项。"""
import math
import os
import re
import time
from pathlib import Path

import numpy as np
import pytest

mujoco = pytest.importorskip("mujoco")

from cae import mbd  # noqa: E402
from cae import mbd_models as MM  # noqa: E402

REPO = Path(__file__).resolve().parents[3]
HAS_MEN = os.path.exists(os.path.join(MM.MENAGERIE, MM.ROBOTS["B-ARM-UR5E"][0]))
need_men = pytest.mark.skipif(not HAS_MEN, reason="没有 Menagerie 模型（WQ_MENAGERIE）")

PENDULUM = """<mujoco><worldbody><body name="rod"><joint name="pivot" type="hinge" axis="0 1 0"/>
<inertial pos="0 0 -0.5" mass="1" diaginertia="0.08333333 0.08333333 1e-6"/>
<geom type="capsule" fromto="0 0 0 0 0 -1" size="0.01" mass="0"/></body></worldbody></mujoco>"""
POSE = {"shoulder_pan_joint": 0.3, "shoulder_lift_joint": -1.2, "elbow_joint": 1.0, "wrist_1_joint": -0.6,
        "wrist_2_joint": 0.4, "wrist_3_joint": 0.0}


def test_pendulum_period():
    amp = 0.05
    s, ch, an = mbd.simulate(mbd.load_spec(xml=PENDULUM), {"duration_s": 10, "initial": {"pivot": amp}, "sample_hz": 200, "dt": 1e-4})
    q, t = ch["q.pivot"], ch["t"]
    z = [t[i] - q[i] * (t[i + 1] - t[i]) / (q[i + 1] - q[i]) for i in range(len(q) - 1) if q[i] > 0 >= q[i + 1]]
    T = float(np.mean(np.diff(z)))
    theory = 2 * math.pi * math.sqrt((1 / 3) / (9.81 * 0.5)) * (1 + amp ** 2 / 16)
    assert abs(T - theory) / theory < 0.01
    assert an["pos"].shape[1] == 1 and len(an["t"]) > 400
    # 铰点反力：静止下垂时 ≈ m·g（摆动很小）
    assert abs(ch["rf.rod.abs"][0] - 9.81) / 9.81 < 0.01


def test_setup_errors():
    with pytest.raises(mbd.SetupError, match="没有关节"):
        mbd.simulate(mbd.load_spec(xml=PENDULUM), {"drives": [{"joint": "nope", "kind": "hold"}]})
    with pytest.raises(mbd.SetupError, match="最多"):
        mbd.simulate(mbd.load_spec(xml=PENDULUM), {"duration_s": 120})
    with pytest.raises(mbd.SetupError, match="两个驱动"):
        mbd.simulate(mbd.load_spec(xml=PENDULUM), {"drives": [{"joint": "pivot", "kind": "hold"}, {"joint": "pivot", "kind": "hold"}]})


def test_menagerie_pin_matches_library():
    pin = (REPO / "library" / "vendor" / "menagerie.yaml").read_text(encoding="utf-8")
    assert MM.MENAGERIE_COMMIT in pin
    sh = (REPO / "factory" / "deploy" / "fetch_menagerie.sh").read_text(encoding="utf-8")
    assert MM.MENAGERIE_COMMIT in sh
    assert sorted(re.search(r'DIRS="([^"]+)"', sh).group(1).split()) == MM.MENAGERIE_DIRS


@need_men
def test_ur5e_gravity_torque_matches_hand_calc():
    pay = [{"body": "wrist_3_link", "mass": 3, "pos": [0, 0.1, 0]}]
    s, ch, _ = mbd.simulate(mbd.load_spec(path=MM.robot_path("B-ARM-UR5E")),
                            {"duration_s": 1.0, "initial": POSE, "payloads": pay, "drives": [{"joint": j, "kind": "hold"} for j in POSE]})
    model, _ = mbd.build(mbd.load_spec(path=MM.robot_path("B-ARM-UR5E")), {"payloads": pay})
    qp = np.zeros(model.nq)
    for j, v in POSE.items():
        qp[model.jnt_qposadr[mujoco.mj_name2id(model, mujoco.mjtObj.mjOBJ_JOINT, j)]] = v
    g = mbd.gravity_torques(model, qp)
    for j in ("shoulder_lift_joint", "elbow_joint", "wrist_1_joint"):
        assert abs(ch["drive." + j][-1] + g[j]) / abs(g[j]) < 0.02
    assert s["tracking_error_max"] < math.radians(0.05)


@need_men
def test_ur5e_move_torque_matches_inverse_dynamics():
    """点到点搬运：驱动力矩与按实际运动做逆动力学（mj_inverse）的结果一致"""
    pay = [{"body": "wrist_3_link", "mass": 3, "pos": [0, 0.1, 0]}]
    drives = [{"joint": "shoulder_pan_joint", "kind": "move", "to": 1.8, "t0": 0.2, "t1": 1.7},
              {"joint": "shoulder_lift_joint", "kind": "move", "to": -0.6, "t0": 0.2, "t1": 1.7}] + \
             [{"joint": j, "kind": "hold"} for j in ("elbow_joint", "wrist_1_joint", "wrist_2_joint", "wrist_3_joint")]
    setup = {"duration_s": 2.0, "initial": POSE, "payloads": pay, "drives": drives, "sample_hz": 200}
    s, ch, an = mbd.simulate(mbd.load_spec(path=MM.robot_path("B-ARM-UR5E")), setup)
    model, _ = mbd.build(mbd.load_spec(path=MM.robot_path("B-ARM-UR5E")), {"payloads": pay})   # 不带驱动约束的模型
    data = mujoco.MjData(model)
    names = list(POSE)
    jid = [mujoco.mj_name2id(model, mujoco.mjtObj.mjOBJ_JOINT, j) for j in names]
    worst = 0.0
    peak = max(abs(ch["drive.shoulder_lift_joint"]))
    for k in range(5, len(ch["t"]) - 5, 7):
        for j, n in zip(jid, names):
            data.qpos[model.jnt_qposadr[j]] = ch["q." + n][k]
            data.qvel[model.jnt_dofadr[j]] = ch["qd." + n][k]
            data.qacc[model.jnt_dofadr[j]] = ch["qdd." + n][k]
        mujoco.mj_inverse(model, data)
        for j, n in zip(jid, names):
            worst = max(worst, abs(data.qfrc_inverse[model.jnt_dofadr[j]] - ch["drive." + n][k]))
    assert worst / peak < 0.02
    d = {x["joint"]: x for x in s["drives"]}
    assert d["shoulder_pan_joint"]["speed_max"] > 1.0 and d["shoulder_lift_joint"]["peak"] > 40
    assert abs(ch["q.shoulder_pan_joint"][-1] - 1.8) < 1e-3


@need_men
def test_service_runs_dynamics_job(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from cae import service
    monkeypatch.setattr(service, "DATA", str(tmp_path))
    with TestClient(service.app) as c:
        assert "B-ARM-UR5E" in [r["id"] for r in c.get("/mbd/robots").json()["robots"]]
        info = c.post("/mbd/models/load", json={"source": "library", "id": "B-ARM-UR5E"}).json()
        assert len(info["joints"]) == 6 and c.get("/mbd/models/{}/model.glb".format(info["key"])).content[:4] == b"glTF"
        assert c.post("/mbd/jobs", json={"model": {"source": "library", "id": "B-ARM-UR5E"},
                                         "setup": {"drives": [{"joint": "x", "kind": "hold"}]}}).status_code == 400
        j = c.post("/mbd/jobs", json={"model": {"source": "library", "id": "B-ARM-UR5E"}, "owner": "7", "factory": "wq_test",
                                      "setup": {"duration_s": 0.5, "initial": POSE, "drives": [{"joint": n, "kind": "hold"} for n in POSE]}}).json()
        for _ in range(240):
            j = c.get("/jobs/" + j["id"]).json()
            if j["status"] in ("done", "failed"):
                break
            time.sleep(0.5)
        assert j["status"] == "done", j.get("error")
        ser = service.read_series(os.path.join(str(tmp_path), "jobs", j["id"], "series.bin"))
        assert "drive.elbow_joint" in ser and len(ser["t"]) >= 50
        assert c.get("/jobs/{}/anim.bin".format(j["id"])).content[:4] == b"WQA1"
        assert [x["id"] for x in c.get("/jobs", params={"kind": "mbd"}).json()["jobs"]] == [j["id"]]
        assert c.get("/jobs", params={"kind": "fea"}).json()["jobs"] == []
        # 上传：单个 MJCF
        up = c.post("/mbd/models/upload", params={"name": "pend.xml"}, content=PENDULUM.encode()).json()
        assert up["joints"][0]["name"] == "pivot"
        assert c.post("/mbd/models/upload", params={"name": "x.txt"}, content=b"hello").status_code == 422


# ---------------------------------------------------------------- 第 2 步：零件库机构
import sys  # noqa: E402

from cae import mech_mjcf as MC  # noqa: E402

sys.path.insert(0, str(REPO / "library"))


def _mech(mid, setup, given=None):
    s = MC.prepare_setup(mid, given or {}, setup)
    return mbd.simulate(mbd.load_spec(xml=MC.mjcf(mid, given)), s)


def test_mechanism_defaults_match_library():
    import yaml
    for mid, d in MC.DEFAULTS.items():
        e = yaml.safe_load((REPO / "library" / "catalog" / "C" / mid / "entry.yaml").read_text(encoding="utf-8"))
        assert {k: float(v) for k, v in e["defaults"].items()} == {k: float(v) for k, v in d.items()}, mid
    assert set(MC.DEFAULTS) == {p.name for p in (REPO / "library" / "catalog" / "C").iterdir() if p.is_dir()}


@pytest.mark.parametrize("mid", sorted(MC.DEFAULTS))
def test_every_mechanism_runs_and_stays_closed(mid):
    s, ch, an = _mech(mid, {"duration_s": 1.5, "drives": [{"joint": MC.DRIVER[mid], "kind": "speed", "value": 2 * math.pi}]})
    assert s["tracking_error_max"] < 0.01
    th = ch["q." + MC.DRIVER[mid]]
    assert abs(th[-1] - th[0] - 2 * math.pi * 1.5) < 0.02
    for k in (0, len(th) // 3, len(th) - 1):           # 各关节始终满足闭环（与按公式算的位置一致）
        want = MC.initial(mid, {}, th[k])
        for j, v in want.items():
            tol = 0.02 if mid in ("C-GNV-GENEVA", "C-RAT-RATCHET") else 2e-3   # 槽轮、棘轮啮合瞬间有冲击
            assert abs(ch["q." + j][k] - v) < tol, (mid, j, ch["q." + j][k], v)


def test_kinematics_match_library_generators():
    from generators import c_mech as L
    p = MC.params("C-LNK-4BAR")
    for th in np.linspace(0, 2 * math.pi, 37):
        a, b = MC.four_bar_solve(p, th), L.four_bar_solve(p, th)
        assert abs(a[2] - b[2]) < 1e-12 and abs(a[3] - b[3]) < 1e-12
    for mid, f1, f2 in (("C-LNK-SLIDER", MC.slider_crank_x, L.slider_crank_x),):
        p = MC.params(mid)
        assert all(abs(f1(p, t) - f2(p, t)) < 1e-12 for t in np.linspace(0, 6.3, 50))
    p = MC.params("C-CAM-DISC")
    assert all(MC.cam_lift(p, d) == L.cam_lift(p, d) for d in range(0, 720, 7))
    p = MC.params("C-GNV-GENEVA")
    assert all(MC.geneva_wheel_angle(p, d) == L.geneva_wheel_angle(p, d) for d in range(0, 720, 7))
    p = MC.params("C-RAT-RATCHET")
    assert all(MC.ratchet_state(p, d) == L.ratchet_state(p, d) for d in range(0, 1440, 7))


def test_slider_crank_torque_matches_virtual_work():
    """准静态（慢转、不计重力），滑块受 200 N 阻力：驱动力矩 τ = F·dx/dθ"""
    p = MC.params("C-LNK-SLIDER")
    s, ch, _ = _mech("C-LNK-SLIDER", {"duration_s": 10, "gravity": False, "sample_hz": 50,
                                      "drives": [{"joint": "crank", "kind": "speed", "value": 0.2 * math.pi}],
                                      "forces": [{"body": "slider", "force": [-200, 0, 0]}]})
    th, tau = ch["q.crank"], ch["drive.crank"]
    ref = np.array([200 * (MC.slider_crank_x(p, t + 1e-6) - MC.slider_crank_x(p, t - 1e-6)) / 2e-6 for t in th])
    k = slice(10, -10)
    assert np.max(np.abs(tau[k] - ref[k])) / np.max(np.abs(ref)) < 0.01
    # 连杆受力：两端反力大小相近（连杆质量小），约等于 200 N / cos(连杆倾角)
    assert 190 < np.median(ch["rf.rod.abs"][k]) < 230


def test_wqr105_input_torque_is_output_over_ratio():
    """WQR-105 传动链：输出轴加 350 N·m 负载，慢速匀速 → 输入力矩 = 350 / 10.5（理想传动）"""
    i = MC.ratio("C-RED-WQR105")
    s, ch, _ = _mech("C-RED-WQR105", {"duration_s": 2, "gravity": False,
                                      "drives": [{"joint": "shaft1", "kind": "speed", "value": 1.0},
                                                 {"joint": "shaft3", "kind": "torque", "value": 350.0}]})
    tau = ch["drive.shaft1"][len(ch["t"]) // 2:]
    assert abs(i - 10.5) < 1e-12
    assert abs(abs(np.mean(tau)) - 350 / i) / (350 / i) < 0.01


def test_link_step_for_fea():
    pytest.importorskip("build123d")
    pytest.importorskip("gmsh")
    from cae import geometry as G
    faces, _, solid = G.faces(MC.member_step("C-LNK-SLIDER", "rod"))
    holes = sorted(f["center"][0] for f in faces if f["kind"] == "cylinder" and abs(f["radius_mm"] - 3.0) < 1e-6)
    assert holes == pytest.approx([0.0, 140.0], abs=1e-6)
    assert solid["bbox_mm"][1] == pytest.approx(-3.0) and solid["bbox_mm"][4] == pytest.approx(3.0)   # 厚度沿 Y
    with pytest.raises(ValueError, match="不是杆件"):
        MC.member_step("C-GER-TRAIN", "shaft1")


# ---------------------------------------------------------------- 第 4 步：送到有限元、疲劳；电机选型
def test_slider_crank_rod_to_fea_and_fatigue(tmp_path, monkeypatch):
    pytest.importorskip("gmsh")
    import shutil
    if not shutil.which(os.environ.get("WQ_CCX", "ccx")):
        pytest.skip("没有 CalculiX")
    from fastapi.testclient import TestClient
    from cae import service
    monkeypatch.setattr(service, "DATA", str(tmp_path))
    with TestClient(service.app) as c:
        ref = {"source": "mech", "id": "C-LNK-SLIDER", "params": {}}
        j = c.post("/mbd/jobs", json={"model": ref, "factory": "t", "setup": {
            "duration_s": 1.0, "gravity": False, "drives": [{"joint": "crank", "kind": "speed", "value": 2 * math.pi}],
            "forces": [{"body": "slider", "force": [-2000, 0, 0]}]}}).json()
        for _ in range(200):
            j = c.get("/jobs/" + j["id"]).json()
            if j["status"] in ("done", "failed"):
                break
            time.sleep(0.3)
        assert j["status"] == "done", j.get("error")
        assert c.post("/mbd/jobs/{}/to-fea".format(j["id"]), json={"member": "slider"}).status_code == 400
        r = c.post("/mbd/jobs/{}/to-fea".format(j["id"]), json={"member": "rod"}).json()
        F = r["force_local_N"]
        # 二力杆：力基本沿杆长方向（构件 X），大小 ≈ 阻力 / cos(连杆倾角)，在 2000–2200 N 之间
        assert abs(F[0]) / math.hypot(*F) > 0.95 and 1990 < r["peak_N"] < 2250
        assert [x["kind"] for x in r["rows"]] == ["fixed", "force"]
        loads = [{"type": "fixed", "faces": r["rows"][0]["faces"]},
                 {"type": "force", "faces": r["rows"][1]["faces"], "vector_n": [r["rows"][1][k] for k in ("fx", "fy", "fz")]}]
        fj = c.post("/jobs", json={"step_sha": r["geometry"]["sha"], "factory": "t",
                                   "setup": {"material_id": "45-QT", "mesh": {"size_mm": 3}, "loads": loads, "source": r["source"]}}).json()
        for _ in range(300):
            fj = c.get("/jobs/" + fj["id"]).json()
            if fj["status"] in ("done", "failed"):
                break
            time.sleep(0.5)
        assert fj["status"] == "done", fj.get("error")
        # 杆身名义应力 F/A（12×6 截面，扣孔前）≈ 29 MPa；评估最大值在孔边，应大于名义值
        nominal = r["peak_N"] / (12 * 6)
        assert fj["stats"]["vm_max_mpa"] > nominal
        ms = c.get("/mbd/jobs/{}/member-series".format(j["id"]), params={"member": "rod",
                   "direction": ",".join(str(x) for x in loads[1]["vector_n"])}).json()
        assert abs(max(abs(x) for x in ms["values"]) - r["peak_N"]) / r["peak_N"] < 0.01
        fat = c.post("/jobs/{}/fatigue".format(fj["id"]), json={"ref_load": r["peak_N"], "series": ms["values"],
                                                                "block_seconds": ms["duration_s"]}).json()
        assert fat["summary"]["cycles_per_block"] >= 1


def test_motor_sizing_hand_check():
    from cae import motors as MT
    t = np.linspace(0, 1, 101)
    tau, w, a = np.full_like(t, 30.0), np.full_like(t, 1.0), np.zeros_like(t)     # 恒定 30 N·m、1 rad/s
    r = MT.check(tau, w, a, MT.MOTORS[1], MT.HARMONIC[2], 100)
    assert r["motor_peak"] == pytest.approx(30 / (100 * MT.ETA)) and r["motor_rms"] == pytest.approx(0.4)
    assert r["motor_speed_rpm"] == pytest.approx(100 * 60 / (2 * math.pi))
    # 200 W 电机连续 0.4 / 0.637；20 号谐波额定 30 / 34 —— 都够，最紧的是减速器额定力矩
    assert r["ok"] is True and r["worst"] == "gear_rms" and r["utilization"]["gear_rms"] == pytest.approx(30 / 34, abs=1e-3)
    s = MT.select(tau, w, a, safety=1.2)
    assert s["ok"] and s["best"]["utilization"][s["best"]["worst"]] <= 1
    big = MT.select(np.full_like(t, 2000.0), w, a)                  # 2000 N·m：表里没有够用的
    assert not big["ok"] and "没有够用的组合" in big["note"]


def test_no_startup_acceleration_spike():
    s, ch, _ = _mech("C-LNK-SLIDER", {"duration_s": 0.5, "drives": [{"joint": "crank", "kind": "speed", "value": 2 * math.pi}]})
    assert np.max(np.abs(ch["qdd.crank"])) < 5          # 匀速：曲柄角加速度≈0


# ---------------------------------------------------------------- 第 5 步：一句话设置、AI 解释
MECH_MODEL = {"kind": "mech", "joints": ["crank", "rod", "slider"], "driver": "crank", "labels": {"crank": "曲柄", "rod": "连杆", "slider": "滑块"},
              "followers": ["rod", "slider"], "bodies": ["crank", "rod", "slider"], "end_body": "slider"}
ARM_MODEL = {"kind": "robot", "joints": ["shoulder_pan_joint", "shoulder_lift_joint", "elbow_joint", "wrist_1_joint", "wrist_2_joint", "wrist_3_joint"],
             "labels": {}, "end_body": "wrist_3_link", "bodies": ["wrist_3_link"]}


def test_one_sentence_rules():
    from hub import mbd_ai as A
    r = A.setup(None, "曲柄 60 rpm 匀速转，滑块上有 200 N 阻力，仿真 3 秒", MECH_MODEL)
    assert r["engine"] == "rules" and r["drives"] == {"crank": {"kind": "speed", "speed": 60.0}}
    assert r["forces"] == [{"body": "slider", "fx": -200.0, "fz": 0}] and r["duration"] == 3 and not r["unmatched"]
    gear = dict(MECH_MODEL, joints=["shaft1", "shaft2", "shaft3"], driver="shaft1", followers=["shaft2", "shaft3"], labels={"shaft3": "输出轴"})
    r = A.rules_setup("输入轴 1450 rpm，输出轴负载 350 N·m", gear)
    assert r["drives"]["shaft1"]["speed"] == 1450 and r["loads"] == {"shaft3": -350.0}
    r = A.rules_setup("底座转 90°，大臂抬 30°，J3 不动，带 5 kg，1.5 秒", ARM_MODEL)
    assert r["drives"]["shoulder_pan_joint"] == {"kind": "move", "rel": True, "to": 90.0}
    assert r["drives"]["elbow_joint"] == {"kind": "hold"} and r["payloads"][0]["mass"] == 5 and r["duration"] == 1.5
    r = A.rules_setup("UR5e 两秒内把 3 kg 工件搬过去", ARM_MODEL)
    assert set(r["drives"]) == {"shoulder_pan_joint", "shoulder_lift_joint"} and r["duration"] == 2


def test_explain_rules_and_retime_action():
    from hub import mbd_ai as A
    job = {"model": {"source": "library", "id": "B-ARM-UR5E"},
           "setup": {"drives": [{"joint": "j1", "kind": "move", "to": 1, "t0": 0, "t1": 1}]},
           "stats": {"drives": [{"joint": "j1", "kind": "move", "peak": 80.0, "rms": 30.0, "speed_max": 2, "power_peak": 100, "power_mean": -20}],
                     "peaks": {"drive.j1": {"max_abs": 80, "at_s": 0.25, "rms": 30}, "rf.link1.abs": {"max_abs": 400, "at_s": 0.3, "rms": 200}}}}
    r = A.explain(None, job)
    assert "惯性" in r["text"] and "发电" in r["text"] and "link1" in r["text"]
    assert r["actions"] == [{"kind": "retime", "factor": 1.5, "label": "运动时间放长到 1.5 倍重算"}]


def test_explain_fixed_base_is_mounting_load():
    from hub import mbd_ai as A
    job = {"model": {"source": "library"}, "setup": {"drives": []},
           "stats": {"drives": [{"joint": "j1", "kind": "hold", "peak": 50.0, "rms": 49.0, "speed_max": 0, "power_peak": 0, "power_mean": 0},
                                {"joint": "j6", "kind": "hold", "peak": 1.0, "rms": 0.2, "speed_max": 0, "power_peak": 0, "power_mean": 0}],
                     "peaks": {"rf.base.abs": {"max_abs": 270, "at_s": 1}, "rm.base.abs": {"max_abs": 60, "at_s": 1},
                               "rf.link1.abs": {"max_abs": 200, "at_s": 1}}}}
    r = A.rules_explain(job, {}, {"base"})
    assert "base的安装处" in r["text"] and "link1与上一个构件" in r["text"] and "j6的峰值" not in r["text"]


def test_lab9_documents():
    import io
    import docx
    from cae import labdoc
    g = docx.Document(io.BytesIO(labdoc.guide_docx("lab9")))
    assert "实验 9" in "\n".join(p.text for p in g.paragraphs) and len(g.tables) >= 1
    t = docx.Document(io.BytesIO(labdoc.report_template_docx("lab9")))
    cells = " ".join(c.text for tb in t.tables for r in tb.rows for c in r.cells)
    assert "J2 大臂" in cells and "虚功原理" in cells
