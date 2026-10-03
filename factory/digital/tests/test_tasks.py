# -*- coding: utf-8 -*-
"""第 11 轮《机械设计》2.7（5）：工程任务单——AI 设计评审员的规则；下达—提交—评审—批准（连带设计发布生效）—评分—
成绩回传记录；任务单钥匙；学生看不到别人的提交。数据库部分需要本机 PostgreSQL（WQ_TEST_DB）。"""
import io
import os
import tempfile

import pytest

import wqbus

DSN = os.environ.get("WQ_TEST_DB", "postgresql://postgres@localhost:5433/wq_test").rsplit("/", 1)[0] + "/wq_task_test"
B_SEGS = [(35, 37), (40, 53), (48, 10), (42, 15), (35, 57), (30, 68)]

SPEC = {
    "编号": "TS-33-1", "book": "mechdesign", "no": "33.1", "学时": 8,
    "标题": ["SH-301 改版", "SH-301 redesign"], "角色": ["设计工程师", "design engineer"],
    "背景": ["改版", "Redesign"], "工位": ["CAD", "PLM"],
    "交付物": [
        {"名称": ["设计计算书", "Design calculation sheet"], "验收": ["齐全", "Complete"], "类型": "文件", "格式": ["docx"]},
        {"名称": ["三维模型（STEP）", "3D model"], "验收": ["有圆角", "Fillets"], "类型": "设计发布", "零件": "SH-301", "轴承": ["6207", "6207"]},
        {"名称": ["工程更改申请", "Change request"], "验收": ["列全", "All listed"], "类型": "更改单", "必列": ["GR-302", "KEY|键"]},
    ],
    "步骤": [["做", "Do"]], "评审要点": [["查", "Check"]],
    "评分": [{"项": ["计算书", "Sheet"], "分": 60, "标准": ["对", "Right"]}, {"项": ["模型", "Model"], "分": 40, "标准": ["对", "Right"]}],
}


# ---------------------------------------------------------------- 测试数据
def calc_docx(rows=None, checks=None):
    """CalcSheet 导出的格式：计算表（项目、公式、代入、结果、出处）、校核表（校核项、要求、计算值、结论）"""
    from docx import Document
    d = Document()
    t = d.add_table(rows=1, cols=5)
    for i, h in enumerate(["项目", "公式", "代入", "结果", "出处"]):
        t.rows[0].cells[i].text = h
    for r in rows if rows is not None else [["圆周力", "F_t = 2T/d", "2×350/0.21", "F_t = 3333 N", "式 (33.4.1)"],
                                            ["理论应力集中系数", "查表", "", "α = 3.0", "[SHI] 首轮估计"]]:
        c = t.add_row().cells
        for i, v in enumerate(r):
            c[i].text = v
    t = d.add_table(rows=1, cols=4)
    for i, h in enumerate(["校核项", "要求", "计算值", "结论"]):
        t.rows[0].cells[i].text = h
    for r in checks if checks is not None else [["疲劳强度", "S_ca,min >= 1.5（33.5.5 节）", "1.71", "合格"]]:
        c = t.add_row().cells
        for i, v in enumerate(r):
            c[i].text = v
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def shaft_step(segs=B_SEGS, fillet=1.0, keys=((60, 40, 12, 5, 0), (200, 30, 8, 4, 0))):
    import build123d as bd
    z, s, steps = 0, None, []
    for d, L in segs:
        c = bd.Pos(0, 0, z) * bd.Cylinder(d / 2, L, align=(bd.Align.CENTER, bd.Align.CENTER, bd.Align.MIN))
        s = c if s is None else s + c
        z += L
        steps.append(z)
    s = s.clean()
    if fillet:
        sel = []
        for zz in steps[:-1]:
            es = [e for e in s.edges() if e.geom_type == bd.GeomType.CIRCLE and abs(e.center().Z - zz) < 1e-6]
            if len(es) == 2:
                sel.append(min(es, key=lambda e: e.radius))
        s = s.fillet(fillet, sel)
    for zc, d, b, t, ang in keys:
        s = s - bd.Pos(0, 0, zc) * bd.Rot(0, 0, ang) * bd.Pos(0, d / 2, 0) * bd.Box(b, 2 * t, 30)
    p = os.path.join(tempfile.mkdtemp(), "s.step")
    bd.export_step(s, p)
    return open(p, "rb").read()


def levels(f, level="error"):
    return [x["text"] for x in f if x["level"] == level]


# ---------------------------------------------------------------- 规则
def test_calc_sheet_rules():
    from hub import task_review as R
    out = []
    R.check_calc(calc_docx(), 0, out)
    assert levels(out) == [] and levels(out, "warning") == []
    out = []
    R.check_calc(calc_docx(rows=[["圆周力", "F_t = 2T/d", "", "F_t = ", "式 (33.4.1)"], ["α", "查表", "", "α = 3.0", ""]],
                           checks=[["疲劳强度", "S_ca,min >= 1.5", "1.38", "合格"], ["静强度", "S >= 1.5", "", ""]]), 0, out)
    e = " ".join(levels(out))
    assert "没有结果" in e and "没有写出处" in e and "不满足" in e and "没有填计算值" in e
    out = []
    from docx import Document
    b = io.BytesIO()
    Document().save(b)
    R.check_calc(b.getvalue(), 0, out)
    assert "不是任务单发的计算书格式" in levels(out)[0]


def test_model_rules():
    pytest.importorskip("build123d")
    from hub import design as D
    from hub import task_review as R
    out = []
    R.check_model(D.step_bytes(D.normalize(D.defaults())), 1, out, ["6207", "6207"])      # A 版：没有圆角，左轴承位短
    e = " ".join(levels(out))
    assert e.count("没有过渡圆角") == 4 and "轴承装不下" in e
    out = []
    R.check_model(shaft_step(), 1, out, ["6207", "6207"])                                # B 版样子：都对
    assert out == []
    out = []
    R.check_model(shaft_step(keys=((60, 40, 12, 5, 0), (200, 30, 8, 4, 90))), 1, out, ["6207"])
    assert "不在同一条母线上" in levels(out, "warning")[0]


def test_change_and_score_rules():
    from hub import task_review as R
    out = []
    R.check_change([{"item": "GR-302", "change": "带毂锻件", "reason": ""}], ["GR-302", "KEY|键"], 2, out)
    assert levels(out) == ["受影响的物料里没有“KEY / 键”。"] and "没有写更改理由" in levels(out, "warning")[0]
    f = [{"level": "error"}, {"level": "warning"}, {"level": "info"}]
    assert R.suggested_score(f) == 87
    assert R.rubric_split(SPEC, 87) == [52, 35] and sum(R.rubric_split(SPEC, 87)) == 87


# ---------------------------------------------------------------- 整条流程（老师、学生两个账号；学习平台用钥匙）
@pytest.fixture()
def db():
    psycopg = pytest.importorskip("psycopg")
    pytest.importorskip("build123d")
    try:
        with psycopg.connect(DSN.rsplit("/", 1)[0] + "/postgres", autocommit=True) as c:
            c.execute("drop database if exists wq_task_test with (force)")
            c.execute("create database wq_task_test")
    except psycopg.OperationalError:
        pytest.skip("没有可用的 PostgreSQL")
    from hub.db import DB
    d = DB(DSN)
    d.init()
    return d


def test_issue_submit_review_approve_grade(db, monkeypatch):
    import importlib
    monkeypatch.setenv("WQ_HUB_NO_START", "1")
    monkeypatch.setenv("WQ_AUTH", "wenquest")
    monkeypatch.setenv("WQ_SSO_URL", "https://learn.example.com/api/v1/auth/sso")
    monkeypatch.setenv("WQ_SECRET", "test-secret")
    monkeypatch.setenv("WQ_FACTORY_TASK_KEY", "task-key-123")
    import httpx
    from test_auth import ACCOUNTS, _Resp, _Teach, _client
    accounts = dict(ACCOUNTS, **{"stu2-cookie": {"id": 102, "fullname": "张三", "username": "zs", "can_create_courses": False}})

    def fake_get(url, cookies=None, timeout=None):
        u = accounts.get((cookies or {}).get("wq_sso", ""))
        return _Resp(200, {"token": "x", "user": u}) if u else _Resp(401, {"code": "not_logged_in"})
    monkeypatch.setattr(httpx, "get", fake_get)
    import hub.app as A
    A = importlib.reload(A)
    A.H.teach = _Teach()
    A.H.db = db
    sent = []

    class _Bus:
        def publish_msg(self, tp, msg):
            sent.append((tp, msg))
            db.insert_message(tp, msg)

    class _Hist:
        def handle(self, tp, msg):
            pass
    A.H.bus, A.H.hist = _Bus(), _Hist()
    try:
        gw = _client(A)
        key = {"Authorization": "Bearer task-key-123"}
        assert gw.post("/api/tasks/course/issue", json={"spec": SPEC, "course_id": 5}).status_code == 401
        tid = gw.post("/api/tasks/course/issue", json={"spec": SPEC, "course_id": 5, "course_name": "机械设计", "cmid": 77,
                                                       "due": "2026-10-20T23:59:00Z", "issued_by": "李老师"}, headers=key).json()["id"]
        assert gw.post("/api/tasks/course/issue", json={"spec": SPEC, "course_id": 5, "cmid": 77}, headers=key).json()["id"] == tid

        stu, stu2, tea = _client(A, "stu-cookie"), _client(A, "stu2-cookie"), _client(A, "tea-cookie")
        hs = {"x-wq-token": stu.post("/api/login", json={"role": "engineer", "mode": "teach"}).json()["token"]}
        h2 = {"x-wq-token": stu2.post("/api/login", json={"role": "engineer", "mode": "teach"}).json()["token"]}
        ht = {"x-wq-token": tea.post("/api/login", json={"role": "engineer", "mode": "teach"}).json()["token"]}
        lst = stu.get("/api/tasks", headers=hs).json()
        assert lst[0]["code"] == "TS-33-1" and lst[0]["mine"] is None

        # 交付物：格式不对被拒；计算书、模型（教学模式也进待审）、更改单
        r = stu.post(f"/api/tasks/{tid}/deliverables/0/file", files={"file": ("calc.pdf", b"%PDF-1.4", "application/pdf")}, headers=hs)
        assert r.status_code == 422 and "docx" in r.json()["detail"]
        bad_calc = calc_docx(checks=[["疲劳强度", "S_ca,min >= 1.5", "1.38", "合格"]])
        stu.post(f"/api/tasks/{tid}/deliverables/0/file", files={"file": ("calc.docx", bad_calc, "application/octet-stream")}, headers=hs)
        r = stu.post(f"/api/tasks/{tid}/deliverables/1/design", files={"step": ("sh301.step", shaft_step(), "application/step")}, headers=hs)
        assert r.status_code == 200, r.text
        plm_sid = r.json()["deliverables"]["1"]["plm"]
        assert db.one("select status from design_submission where id=%s", (plm_sid,))["status"] == "pending"
        assert not any(m["type"] == "design.release" for _, m in sent)
        stu.post(f"/api/tasks/{tid}/deliverables/2/change", json={"rows": [{"item": "GR-302", "change": "带毂锻件", "reason": "毂宽 55"}]},
                 headers=hs)

        pre = stu.post(f"/api/tasks/{tid}/precheck", headers=hs).json()
        errs = levels(pre["findings"])
        assert any("不满足" in x for x in errs) and any("KEY" in x for x in errs) and pre["suggested_score"] == 80
        sub = stu.post(f"/api/tasks/{tid}/submit", headers=hs).json()
        assert sub["status"] == "submitted" and sum(c["level"] == "error" for c in sub["comments"]) == 2
        assert stu.post(f"/api/tasks/{tid}/deliverables/2/change", json={"rows": []}, headers=hs).status_code == 422   # 提交后不能改
        sid = sub["id"]

        # 别人看不到、学生不能批准；老师看列表
        assert stu2.get(f"/api/task-submissions/{sid}", headers=h2).status_code == 403
        assert stu.post(f"/api/task-submissions/{sid}/decision", json={"decision": "approve"}, headers=hs).status_code == 403
        rows = tea.get(f"/api/tasks/{tid}/submissions", headers=ht).json()["submissions"]
        assert rows[0]["name"].startswith("王小明") and rows[0]["open_errors"] == 2 and rows[0]["suggested"] == 80

        # 退回 → 学生改 → 再提交；意见回复、老师关闭；有“必改”没关闭不能批准
        assert tea.post(f"/api/task-submissions/{sid}/decision", json={"decision": "return"}, headers=ht).status_code == 422
        tea.post(f"/api/task-submissions/{sid}/decision", json={"decision": "return", "note": "计算书结论与数值矛盾"}, headers=ht)
        stu.post(f"/api/tasks/{tid}/deliverables/0/file", files={"file": ("calc.docx", calc_docx(), "application/octet-stream")}, headers=hs)
        sub = stu.post(f"/api/tasks/{tid}/submit", headers=hs).json()
        assert sub["rounds"] == 2 and sum(c["level"] == "error" and not c["resolved"] for c in sub["comments"]) == 1
        c = next(c for c in sub["comments"] if c["level"] == "error")
        stu.post(f"/api/task-submissions/{sid}/comments", json={"reply_to": c["id"], "body": "键 8×7×63 漏列，已在更改单补上"}, headers=hs)
        assert tea.post(f"/api/task-submissions/{sid}/decision", json={"decision": "approve"}, headers=ht).status_code == 422
        assert stu.post(f"/api/task-submissions/{sid}/comments/{c['id']}", json={}, headers=hs).status_code == 403
        tea.post(f"/api/task-submissions/{sid}/comments/{c['id']}", json={}, headers=ht)
        assert tea.post(f"/api/task-submissions/{sid}/grade", json={"items": [50, 30]}, headers=ht).status_code == 422   # 批准前不能评分
        r = tea.post(f"/api/task-submissions/{sid}/decision", json={"decision": "approve", "note": "同意"}, headers=ht).json()
        assert r["status"] == "approved" and "SH-301 设计发布生效" in r["decision"]
        assert db.one("select status from design_submission where id=%s", (plm_sid,))["status"] == "approved"
        assert any(m["type"] == "design.release" and m["data"]["item"] == "SH-301" for _, m in sent)

        # 评分；学习平台读分数、记已回传
        assert tea.post(f"/api/task-submissions/{sid}/grade", json={"items": [70, 30]}, headers=ht).status_code == 422   # 超过分值
        assert tea.post(f"/api/task-submissions/{sid}/grade", json={"items": [52, 36], "note": "好"}, headers=ht).json()["score"]["total"] == 88
        assert gw.get("/api/tasks/course/5").status_code == 401
        course = gw.get("/api/tasks/course/5", headers=key).json()
        s0 = course[0]["submissions"][0]
        assert course[0]["cmid"] == 77 and s0["author_uid"] == "101" and s0["total"] == 88 and s0["pushed_at"] is None
        assert gw.post("/api/tasks/course/5/pushed", json={"submissions": [sid]}, headers=key).json()["marked"] == 1
        assert gw.post("/api/tasks/course/6/pushed", json={"submissions": [sid]}, headers=key).json()["marked"] == 0
        assert gw.get("/api/tasks/course/5", headers=key).json()[0]["submissions"][0]["pushed_at"]
    finally:
        for k in ("WQ_AUTH", "WQ_SSO_URL", "WQ_FACTORY_TASK_KEY"):
            os.environ.pop(k, None)
        importlib.reload(A)
