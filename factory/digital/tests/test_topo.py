# -*- coding: utf-8 -*-
"""第 14 轮 H6：平面拓扑优化（SIMP）。验收：经典 MBB 梁得到公认的桁架形状，柔度随迭代下降并收敛；结果能拉伸成板件做有限元校核。"""
import time

import numpy as np
import pytest

from cae import topo as T


def test_mbb_classic():
    """Andreassen 等 2011（88 行程序）的标准算例：60×20、体积比 0.5、p = 3、rmin = 1.5，柔度约 203"""
    r = T.run("mbb")
    c = [h[1] for h in r["history"]]
    assert r["converged"] and r["iterations"] < 200
    assert abs(r["compliance"] - 203) / 203 < 0.02, r["compliance"]
    assert c[-1] < 0.25 * c[0]                                    # 从均匀灰色开始，柔度大幅下降
    assert max(c[10:]) - min(c[10:]) < 0.05 * c[-1] + 1e-9 or c[-1] <= min(c[10:]) * 1.01   # 后段基本不再变（收敛）
    assert abs(r["history"][-1][2] - 0.5) < 0.005                  # 体积比守住
    d = np.array(r["density"]) >= 0.5
    ny, nx = d.shape
    assert d[0, : nx // 2].all() and d[-1].all()                 # 上弦（受压，左半段贴着对称面）和下弦（受拉）连续
    assert d[0, 0] and d[-1, -1]                                   # 载荷点、支座处有材料
    holes = (~d[1:-1, 1:-1]).sum()
    assert holes > 0.3 * (ny - 2) * (nx - 2)                      # 中间挖空，只剩斜杆——桁架
    assert r["grey"] < 0.25
    # 斜杆数：中间一行从左到右数“有材料”的段数（桁架的腹杆）
    mid = d[ny // 2]
    segs = int(mid[0]) + int(np.sum(mid[1:] & ~mid[:-1]))
    assert 3 <= segs <= 8, segs


def test_presets_and_errors():
    r = T.run("cantilever", nelx=32, nely=16, volfrac=0.4)
    d = np.array(r["density"]) >= 0.5
    assert d[0, 0] and d[-1, 0]                                     # 固定端上下两角有材料（弯矩最大处）
    assert r["history"][-1][1] < r["history"][0][1]
    with pytest.raises(ValueError):
        T.run("nope")
    with pytest.raises(ValueError):
        T.run("mbb", nelx=200, nely=100)
    with pytest.raises(ValueError):
        T.run("mbb", volfrac=0.95)


def test_fea_part_and_rows():
    pytest.importorskip("build123d")
    from cae import geometry as G
    r = T.run("mbb")
    step, pads, note, info = T.fea_part("mbb", r["density"], 300, 10)
    faces, _, solid = G.faces(step)
    assert solid["solids"] == 1 and info["width_mm"] == 300 and "镜像" in note
    rows = T.fea_rows("mbb", pads, faces, 2000)
    assert [x["kind"] for x in rows] == ["fixed", "fixed", "force"] and rows[2]["fy"] == -2000
    assert all(len(x["faces"]) == 1 for x in rows)


def test_service_topo_job_and_to_fea(tmp_path, monkeypatch):
    pytest.importorskip("build123d")
    from fastapi.testclient import TestClient
    from cae import service
    monkeypatch.setattr(service, "DATA", str(tmp_path))
    with TestClient(service.app) as c:
        assert {p["key"] for p in c.get("/topo/presets").json()["presets"]} == {"mbb", "cantilever", "bridge", "bracket"}
        assert c.post("/topo/jobs", json={"spec": {"preset": "mbb", "volfrac": 0.95}}).status_code == 400
        assert c.post("/topo/jobs", json={"spec": {"preset": "mbb", "material_id": "X"}}).status_code == 400
        j = c.post("/topo/jobs", json={"spec": {"preset": "bracket", "nelx": 30, "nely": 20, "length_mm": 120, "thickness_mm": 8, "force_n": 1500},
                                       "factory": "wq_test", "owner": "7"}).json()
        assert j["kind"] == "topo" and j["status"] == "queued"
        for _ in range(300):
            j = c.get("/jobs/" + j["id"]).json()
            if j["status"] in ("done", "failed"):
                break
            time.sleep(0.3)
        assert j["status"] == "done", j.get("error")
        r = c.get("/topo/jobs/{}/result".format(j["id"])).json()
        assert len(r["density"]) == 20 and len(r["density"][0]) == 30 and r["frames"]
        f = c.post("/topo/jobs/{}/to-fea".format(j["id"]), json={}).json()
        assert f["material_id"] == "6061-T6" and f["geometry"]["solid"]["solids"] == 1
        assert f["rows"][0]["kind"] == "fixed" and f["rows"][-1] == dict(f["rows"][-1], kind="force", fy=-1500.0)
        assert c.get("/jobs", params={"kind": "topo"}).json()["jobs"][0]["id"] == j["id"]


def test_topo_explain():
    from hub import opt_ai as A
    r = T.run("mbb", nelx=30, nely=10)
    r["volfrac"] = 0.5
    t = A.topo_explain({}, r)
    assert "桁架" in t and "柔度" in t and "有限元校核" in t


def test_lab11_documents():
    import io
    import docx
    from cae import labdoc
    g = docx.Document(io.BytesIO(labdoc.guide_docx("lab11")))
    text = "\n".join(p.text for p in g.paragraphs)
    assert "实验 11" in text and "热平衡" in text and "帕累托" in text and len(g.tables) >= 1
    t = docx.Document(io.BytesIO(labdoc.report_template_docx("lab11")))
    cells = " ".join(c.text for tb in t.tables for r in tb.rows for c in r.cells)
    assert "482.6" in cells and "自由热伸长" in cells and "卡住最优解" in cells
