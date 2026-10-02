# -*- coding: utf-8 -*-
"""第 11 轮 有限元：悬臂梁标准题（与理论解误差 < 3%）、SH-301 扭转、计算服务排队、材料库。
需要 gmsh、CalculiX（ccx）、build123d；没有就跳过（CI 的镜像里都有）。"""
import math
import os
import shutil
import time

import pytest

pytest.importorskip("gmsh")
pytest.importorskip("build123d")
if not shutil.which(os.environ.get("WQ_CCX", "ccx")):
    pytest.skip("没有 CalculiX（ccx）", allow_module_level=True)

from cae import geometry as G  # noqa: E402
from cae import materials as M  # noqa: E402
from cae import solve as S  # noqa: E402

AX = {"origin": [0, 0, 0], "dir": [0, 0, 1]}


def beam_step(L=100.0, b=10.0):
    import build123d as bd
    import tempfile
    part = bd.Pos(L / 2, 0, 0) * bd.Box(L, b, b)
    with tempfile.TemporaryDirectory() as t:
        p = os.path.join(t, "b.step")
        bd.export_step(part, p)
        return open(p, "rb").read()


def face_at(faces, axis, value):
    return next(f["id"] for f in faces if f["kind"] == "plane" and abs(f["center"][axis] - value) < 1e-6)


def test_cantilever_matches_beam_theory():
    step = beam_step()
    faces, glb, solid = G.faces(step)
    assert solid["faces"] == 6 and glb[:4] == b"glTF"
    fix, tip = face_at(faces, 0, 0.0), face_at(faces, 0, 100.0)
    mat = M.get("45-QT")
    st, surf, _ = S.solve(step, {"material": mat, "mesh": {"size_mm": 4},
                                 "loads": [{"type": "fixed", "faces": [fix]},
                                           {"type": "force", "faces": [tip], "vector_n": [0, -100, 0]}]})
    theory = 100 * 100 ** 3 / (3 * mat["E_mpa"] * (10 * 10 ** 3 / 12))
    assert abs(st["u_max_mm"] - theory) / theory < 0.03
    assert abs(st["u_max_at"][0] - 100) < 1e-6
    # 根部弯曲应力 M/W = 100·100/(10·10²/6) = 60 MPa；评估值避开固定面附近
    assert 50 < st["vm_max_mpa"] < 75
    assert st["applied_force_n"] == [0, -100, 0]
    assert st["safety_factor"] == round(mat["yield_mpa"] / st["vm_max_mpa"], 3)
    assert len(surf["vm"]) == len(surf["positions"]) and surf["triangles"].max() < len(surf["positions"])


def test_sh301_torsion_nominal_and_keyway_hotspot():
    from hub import design as D
    step = D.step_bytes(D.normalize(D.defaults()))
    faces, _, _ = G.faces(step)
    cyl = {f["radius_mm"]: f for f in faces if f["kind"] == "cylinder" and f.get("radius_mm") and abs(f["axis"][2]) > 0.99}
    brg = sorted([f["id"] for f in faces if f.get("radius_mm") == 17.5])
    out = max((f for f in faces if f.get("radius_mm") == 15.0), key=lambda f: f["center"][2])["id"]
    walls = [f for f in faces if f["kind"] == "plane" and f.get("normal") and abs(abs(f["normal"][0]) - 1) < 1e-3]
    assert len(walls) == 2 and 20.0 in cyl
    loads = [{"type": "cyl_support", "faces": [brg[0]], "axis": AX, "dofs": ["radial", "axial"]},
             {"type": "cyl_support", "faces": [brg[1]], "axis": AX, "dofs": ["radial"]},
             {"type": "cyl_support", "faces": [out], "axis": AX, "dofs": ["tangential"]},
             {"type": "torque", "faces": [walls[0]["id"]], "value_nmm": 350e3, "axis": AX}]
    st, _, full = S.solve(step, {"material": M.get("45-QT"), "loads": loads, "mesh": {"size_mm": 5}})
    # Ø35 光滑段中部：纯扭转 τ = 16T/πd³，Von Mises = √3·τ
    seg = next(f for f in faces if f.get("radius_mm") == 17.5 and 112 < f["center"][2] < 137)
    vals = [full["vm"][n] for n, p in full["nodes"].items()
            if abs(p[2] - seg["center"][2]) < 3 and abs(math.hypot(p[0], p[1]) - 17.5) < 0.05]
    theory = math.sqrt(3) * 16 * 350e3 / (math.pi * 35 ** 3)
    assert vals and abs(sum(vals) / len(vals) - theory) / theory < 0.03
    # 最危险的点在键槽根部（槽底面与侧壁相交处）
    key_faces = {f["id"] for f in faces if f["kind"] == "plane" and abs(f["center"][1] - 15) < 0.01} | {w["id"] for w in walls}
    assert set(st["vm_max_faces"]) & key_faces
    assert abs(st["vm_max_at"][1] - 15) < 0.5 and 52 < st["vm_max_at"][2] < 112


def test_setup_errors_are_plain():
    step = beam_step()
    faces, _, _ = G.faces(step)
    mat = M.get("Q235")
    with pytest.raises(ValueError, match="固定"):
        S.solve(step, {"material": mat, "loads": [{"type": "force", "faces": [1], "vector_n": [1, 0, 0]}]})
    with pytest.raises(ValueError, match="载荷"):
        S.solve(step, {"material": mat, "loads": [{"type": "fixed", "faces": [1]}]})
    with pytest.raises(ValueError, match="没有这些面"):
        S.solve(step, {"material": mat, "loads": [{"type": "fixed", "faces": [99]},
                                                  {"type": "force", "faces": [1], "vector_n": [1, 0, 0]}]})


def test_materials_have_sources_and_fields():
    pub = M.public()
    assert len(pub) == 12
    for m in pub:
        assert m["sources"] and all(m["sources"])
        assert m["E_mpa"] > 0 and 0 < m["nu"] < 0.5 and m["density"] > 0 and m["sigma_1"] > 0
        assert m["ultimate_mpa"] >= (m["yield_mpa"] or 0)
        g = M.get(m["id"])
        assert g["strength_mpa"] > 0
    assert M.get("HT200")["strength_kind"] == "抗拉强度"
    with pytest.raises(ValueError):
        M.get("unobtainium")


def test_service_queue_end_to_end(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from cae import service
    monkeypatch.setattr(service, "DATA", str(tmp_path))
    with TestClient(service.app) as c:
        step = beam_step()
        g = c.post("/geometry", content=step).json()
        assert len(g["faces"]) == 6 and c.get("/geometry/{}/model.glb".format(g["sha"])).content[:4] == b"glTF"
        assert c.post("/geometry", content=b"not a step").status_code == 422
        fix, tip = face_at(g["faces"], 0, 0.0), face_at(g["faces"], 0, 100.0)
        bad = c.post("/jobs", json={"step_sha": g["sha"], "setup": {"material_id": "45-QT",
                                                                    "loads": [{"type": "force", "faces": [tip], "vector_n": [0, 1, 0]}]}})
        assert bad.status_code == 400 and "固定" in bad.json()["detail"]
        assert c.post("/jobs", json={"step_sha": g["sha"], "setup": {"material_id": "x", "loads": []}}).status_code == 400
        j = c.post("/jobs", json={"step_sha": g["sha"], "owner": "7", "factory": "wq_test", "title": "悬臂梁",
                                  "setup": {"material_id": "45-QT", "mesh": {"size_mm": 5},
                                            "loads": [{"type": "fixed", "faces": [fix]},
                                                      {"type": "force", "faces": [tip], "vector_n": [0, -100, 0]}]}}).json()
        assert j["status"] == "queued"
        for _ in range(240):
            j = c.get("/jobs/" + j["id"]).json()
            if j["status"] in ("done", "failed"):
                break
            time.sleep(0.5)
        assert j["status"] == "done", j.get("error")
        assert abs(j["stats"]["u_max_mm"] - 0.1905) / 0.1905 < 0.03
        surf = service.read_surface(os.path.join(str(tmp_path), "jobs", j["id"], "surface.bin"))
        assert len(surf["vm"]) == len(surf["positions"]) and surf["face_of_triangle"].max() <= 6
        assert c.get("/jobs/{}/surface.bin".format(j["id"])).content[:4] == b"WQS1"
        assert [x["id"] for x in c.get("/jobs", params={"factory": "wq_test", "owner": "7"}).json()["jobs"]] == [j["id"]]
        assert c.get("/jobs", params={"factory": "other"}).json()["jobs"] == []
        # 疲劳：恒幅对称循环（载荷 ±1 倍计算工况）→ 最危险点寿命与手算一致
        from cae import fatigue as FT
        fat = c.post("/jobs/{}/fatigue".format(j["id"]), json={"ref_load": 1, "series": [5, -5], "block_seconds": 1,
                                                               "surface": "ground", "haibach": False}).json()
        sm = fat["summary"]
        hand = FT.hand_check(5 * sm["hot_vm_ref"], 0, M.get("45-QT"))
        assert sm["cycles_per_block"] == 1 and abs(sm["life_blocks"] - hand) / hand < 1e-3
        assert c.post("/jobs/{}/fatigue".format(j["id"]), json={"ref_load": 1, "series": [1]}).status_code == 400
        # Word 报告：设置、结果、截图、结论都在
        png = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
        rep = c.post("/jobs/{}/report".format(j["id"]), json={"images": [{"data": png, "caption": "应力云图"}], "ai_text": "最危险在根部。"})
        assert rep.status_code == 200 and rep.content[:2] == b"PK"
        import docx
        import io
        text = "\n".join(p.text for p in docx.Document(io.BytesIO(rep.content)).paragraphs)
        cells = " ".join(c.text for t in docx.Document(io.BytesIO(rep.content)).tables for r in t.rows for c in r.cells)
        assert "结论" in text and "安全系数" in text and "最危险在根部" in text and "应力云图" in text and "疲劳寿命" in text
        assert "45 钢" in cells and "固定" in cells and "C3D10" in cells


def test_rainflow_closes_repeating_blocks():
    from cae import fatigue as FT
    assert FT.rainflow([350, -350]).tolist() == [[350, 0]]
    assert FT.rainflow([350, 0]).tolist() == [[175, 175]]
    # 一块里一个大循环（-50..350）、一个小循环（100..300）、一个 0..350
    assert sorted(FT.rainflow([0, 350, 100, 300, -50, 350, 0]).tolist()) == [[100, 200], [175, 175], [200, 150]]
    m = M.get("45-QT")
    _, s = FT.compute([300.0, 10.0], 350, [350, -350], 2.0, m, haibach=False)
    assert abs(s["life_blocks"] - FT.hand_check(300, 0, m)) < 1 and abs(s["life_hours"] - s["life_blocks"] * 2 / 3600) < 1e-6
    _, s = FT.compute([100.0], 350, [350, -350], 1.0, m, haibach=False)
    assert s["infinite"]
    _, s = FT.compute([100.0], 350, [350, -350], 1.0, m, haibach=True)
    assert not s["infinite"] and s["life_blocks"] > 1e9


def test_one_sentence_setup_rules_pick_faces():
    """一句话设置（没配模型时的规则兜底）：示范题的话能选出端面、键槽侧面；轴承、轴伸、面编号、方向、材料"""
    from hub import cae_ai as A
    from hub import design as D
    step = D.step_bytes(D.normalize(D.defaults()))
    faces, _, solid = G.faces(step)
    tags, _, _ = A.tag_faces(faces, solid)
    by = lambda t: sorted(i for i, v in tags.items() if t in v)  # noqa: E731
    assert len(by("端面")) == 2 and len(by("键槽侧面")) == 2 and len(by("键槽底面")) == 1 and len(by("外圆")) == 5
    left_end = by("左端面")[0]
    r = A.setup(None, "左端面固定，键槽侧面加 350 N·m 扭矩，45 钢调质", faces, solid)
    assert r["engine"] == "rules" and r["material_id"] == "45-QT" and not r["unmatched"]
    assert r["rows"][0] == {"kind": "fixed", "faces": [left_end]}
    assert r["rows"][1]["kind"] == "torque" and r["rows"][1]["value"] == 350 and r["rows"][1]["faces"][0] in by("键槽侧面")
    r = A.rules_setup("两个 Ø35 轴承位支承（左边那个止推），右端 Ø30 轴伸限制转动，键槽侧面加 0.35 kN·m 扭矩，40Cr", faces, solid)
    kinds = [(x["kind"], x.get("thrust")) for x in r["rows"]]
    assert kinds == [("bearing", True), ("bearing", False), ("coupling", None), ("torque", None)]
    assert r["rows"][2]["faces"] == by("右端轴伸") and r["rows"][3]["value"] == 350 and r["material_id"] == "40Cr-QT"
    r = A.rules_setup("面 18 固定，右端面加 2 kN 向下的力，看不懂的话", faces, solid)
    assert r["rows"][1]["faces"] == by("右端面") and r["rows"][1]["fy"] == -2000 and r["unmatched"] == ["看不懂的话"]


def test_explain_rules_and_suggestions():
    from hub import cae_ai as A
    from hub import design as D
    job = {"item": "SH-301", "setup": {"material_id": "45-QT", "loads": []},
           "stats": {"vm_max_mpa": 260.0, "vm_peak_all_mpa": 260.0, "vm_max_faces": [9, 14], "safety_factor": 1.37}}
    faces, _, _ = G.faces(D.step_bytes(D.normalize(D.defaults())))
    r = A.explain(None, job, faces, D.normalize(D.defaults()))
    assert "键槽" in r["text"] and "尖角" in r["text"]
    kinds = {a["kind"]: a for a in r["actions"]}
    assert kinds["material"]["material_id"] == "40Cr-QT"
    p = kinds["design"]["params"]
    assert p["segments"][2][0] == 44 and (p["keyway"]["b"], p["keyway"]["t"]) == (12.0, 5.0)
    assert D.check(D.normalize(p))["ok"]


def test_llm_paths_validate_and_fall_back():
    """模型给的面不存在时丢掉；模型出错时退回规则"""
    from hub import cae_ai as A
    step = beam_step()
    faces, _, solid = G.faces(step)

    class Fake:
        name = "fake"

        def __init__(self, args=None, boom=False):
            self.args, self.boom = args, boom

        def available(self):
            return True

        def run(self, system, messages, tools, call_tool, max_turns=6):
            if self.boom:
                raise RuntimeError("network")
            assert "faces" in messages[0]["content"] or "stats" in messages[0]["content"]
            call_tool(tools[0]["name"], self.args)
            return "", []
    r = A.setup(Fake({"rows": [{"kind": "fixed", "faces": [1]}, {"kind": "force", "faces": [999], "fy": -5}],
                      "material_id": "Q235"}), "随便", faces, solid)
    assert r["engine"] == "fake" and r["rows"] == [{"kind": "fixed", "faces": [1]}] and r["material_id"] == "Q235"
    r = A.setup(Fake(boom=True), "左端面固定，右端面加 100 N 向下的力", faces, solid)
    assert r["engine"] == "rules" and "模型暂时不可用" in r["note"] and [x["kind"] for x in r["rows"]] == ["fixed", "force"]
    job = {"item": None, "setup": {"material_id": "6061-T6", "loads": []},
           "stats": {"vm_max_mpa": 300.0, "vm_peak_all_mpa": 300.0, "vm_max_faces": [1], "safety_factor": 0.92}}
    r = A.explain(Fake({"text": "根部最危险。", "material_id": "7075-T6"}), job, faces)
    assert r["text"] == "根部最危险。" and r["actions"][0]["material_id"] == "7075-T6" and r["engine"] == "fake"


def test_lab8_documents():
    import io
    import docx
    from cae import labdoc
    g = docx.Document(io.BytesIO(labdoc.guide_docx()))
    text = "\n".join(p.text for p in g.paragraphs)
    assert "实验 8" in text and "任务 3" in text and len(g.tables) >= 2
    t = docx.Document(io.BytesIO(labdoc.report_template_docx()))
    cells = " ".join(c.text for tb in t.tables for r in tb.rows for c in r.cells)
    assert "Haibach" in cells and "方案 B" in cells and "安全系数" in cells


def test_size_limit_gives_plain_message(monkeypatch, tmp_path):
    """超过教学版规模：直接提示调大网格，不卡死；计算服务里同样把原因带回任务"""
    monkeypatch.setattr(S, "MAX_ELEMENTS", 200)
    step = beam_step()
    faces, _, _ = G.faces(step)
    fix, tip = face_at(faces, 0, 0.0), face_at(faces, 0, 100.0)
    with pytest.raises(S.TooBig, match="教学版上限"):
        S.solve(step, {"material": M.get("Q235"), "mesh": {"size_mm": 2},
                       "loads": [{"type": "fixed", "faces": [fix]}, {"type": "force", "faces": [tip], "vector_n": [0, -1, 0]}]})
