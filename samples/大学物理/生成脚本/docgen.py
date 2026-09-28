"""Small helpers to write teacher-style Word documents (python-docx)."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_CN, HEAD_CN, LATIN = "宋体", "黑体", "Times New Roman"


def _font(run, cn=BODY_CN, size=None, bold=None, color=None, italic=None):
    run.font.name = LATIN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


class Doc:
    def __init__(self):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        st = self.d.styles["Normal"]
        st.font.name = LATIN
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)
        st.font.size = Pt(11)
        st.paragraph_format.line_spacing = 1.35
        st.paragraph_format.space_after = Pt(4)
        for lvl, size in ((1, 16), (2, 13.5), (3, 12)):
            h = self.d.styles[f"Heading {lvl}"]
            h.font.name = LATIN
            h.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_CN)
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = RGBColor(0x1F, 0x2D, 0x3A)
            h.paragraph_format.space_before = Pt(10 if lvl > 1 else 14)
            h.paragraph_format.space_after = Pt(6)

    def title(self, text, sub=None):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(text), HEAD_CN, 18, True)
        if sub:
            q = self.d.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _font(q.add_run(sub), BODY_CN, 10.5, color="555555")

    def h(self, text, level=1):
        self.d.add_heading(text, level)

    def p(self, text, bold_prefix=None, indent=True):
        para = self.d.add_paragraph()
        if indent:
            para.paragraph_format.first_line_indent = Pt(22)
        if bold_prefix:
            _font(para.add_run(bold_prefix), HEAD_CN, bold=True)
        _font(para.add_run(text))
        return para

    def kv(self, key, value):
        para = self.d.add_paragraph()
        _font(para.add_run(key + "："), HEAD_CN, bold=True)
        _font(para.add_run(value))

    def eq(self, text):
        para = self.d.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(para.add_run(text), BODY_CN, 11.5, italic=False)

    def bullets(self, items, numbered=False):
        style = "List Number" if numbered else "List Bullet"
        for it in items:
            para = self.d.add_paragraph(style=style)
            if isinstance(it, tuple):
                _font(para.add_run(it[0]), HEAD_CN, bold=True)
                _font(para.add_run(it[1]))
            else:
                _font(para.add_run(it))

    def table(self, header, rows, widths=None, font_size=10):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, htext in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            _font(cell.paragraphs[0].add_run(htext), HEAD_CN, font_size, True)
            shade = cell._element.get_or_add_tcPr()
            from docx.oxml import OxmlElement
            sh = OxmlElement("w:shd")
            sh.set(qn("w:val"), "clear")
            sh.set(qn("w:color"), "auto")
            sh.set(qn("w:fill"), "E8EEF0")
            shade.append(sh)
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = ""
                _font(cells[i].paragraphs[0].add_run(str(v)), BODY_CN, font_size)
        if widths:
            t.autofit = False
            for i, w in enumerate(widths):
                t.columns[i].width = Cm(w)  # grid width (LibreOffice, WPS)
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)  # cell width (Word)
        self.d.add_paragraph()
        return t

    def image(self, path, width_cm=12, caption=None):
        self.d.add_picture(path, width=Cm(width_cm))
        self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            c = self.d.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _font(c.add_run(caption), BODY_CN, 9.5, color="555555")

    def example(self, title, question, solution_lines):
        para = self.d.add_paragraph()
        _font(para.add_run(title + "  "), HEAD_CN, bold=True, color="1F6F8B")
        _font(para.add_run(question))
        s = self.d.add_paragraph()
        _font(s.add_run("解："), HEAD_CN, bold=True)
        _font(s.add_run(solution_lines[0]))
        for line in solution_lines[1:]:
            q = self.d.add_paragraph()
            q.paragraph_format.left_indent = Pt(22)
            _font(q.add_run(line))

    # ---- bilingual (中英对照) helpers ----
    def en(self, text, size=9.5, indent=False):
        para = self.d.add_paragraph()
        if indent:
            para.paragraph_format.first_line_indent = Pt(22)
        para.paragraph_format.space_after = Pt(6)
        r = para.add_run(text)
        r.font.name = LATIN
        r._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string("5B6B75")
        return para

    def bh(self, zh, en, level=1):
        h = self.d.add_heading(zh, level)
        r = h.add_run("  " + en)
        r.font.size = Pt({1: 12, 2: 11, 3: 10.5}[level])
        r.font.color.rgb = RGBColor.from_string("5B6B75")
        r.bold = False

    def bp(self, zh, en, bold_prefix=None):
        self.p(zh, bold_prefix=bold_prefix)
        self.en(en, indent=True)

    def blist(self, items, numbered=False):
        # numbered lists restart at 1: write the number, not a shared Word numbering
        for i, (zh, en) in enumerate(items, 1):
            if numbered:
                para = self.d.add_paragraph()
                para.paragraph_format.left_indent = Pt(22)
                para.paragraph_format.first_line_indent = Pt(-16)
                _font(para.add_run(f"{i}. " + zh))
            else:
                para = self.d.add_paragraph(style="List Bullet")
                _font(para.add_run(zh))
            para.add_run().add_break()
            r = para.add_run(en)
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor.from_string("5B6B75")

    def lines(self, n=4):
        for _ in range(n):
            para = self.d.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            pPr = para._p.get_or_add_pPr()
            from docx.oxml import OxmlElement
            bdr = OxmlElement("w:pBdr")
            b = OxmlElement("w:bottom")
            b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:space"), "1"); b.set(qn("w:color"), "B8C4C8")
            bdr.append(b); pPr.append(bdr)
            para.paragraph_format.line_spacing = 1.9

    def page_break(self):
        self.d.add_page_break()

    def save(self, path):
        self.d.save(path)
