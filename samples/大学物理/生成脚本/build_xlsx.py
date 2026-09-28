from pathlib import Path
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path("out/大学物理（上）课程资料")
F = "Arial"
HEAD = PatternFill("solid", fgColor="E8EEF0")
INPUT = PatternFill("solid", fgColor="FFF7D6")
thin = Side(style="thin", color="B8C4C8")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name=F, bold=True)
        c.fill = HEAD
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w


def body(ws, r, values, wrap_cols=()):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=F)
        c.border = BOX
        c.alignment = Alignment(vertical="center", wrap_text=i in wrap_cols)


# ---------------------------------------------------------------- 教学日历
wb = Workbook(); ws = wb.active; ws.title = "教学日历"
ws["A1"] = "大学物理A（上）教学日历  2026 年秋季学期（第 1–8 周）"; ws["A1"].font = Font(name=F, bold=True, size=13)
ws["A2"] = "每周 2 次课，每次 2 学时；周二 8:00–9:40（理科楼 201），周四 10:00–11:40（理科楼 201）"; ws["A2"].font = Font(name=F, color="555555")
header(ws, 4, ["周次", "日期", "章节", "教学内容", "学时", "作业 / 实验 / 测验"], [7, 13, 22, 46, 7, 30])
rows = [
    [1, "9/8–9/12", "第1章 质点运动学", "课程介绍；参考系与质点；位矢、位移、速度、加速度", 4, "预习：讲义 1.1–1.2"],
    [2, "9/15–9/19", "第1章 质点运动学", "抛体运动；积分求运动方程；圆周运动；相对运动", 2, "习题1（第3周周二提交）"],
    [2, "9/15–9/19", "第2章 牛顿定律", "牛顿三定律；常见的力", 2, ""],
    [3, "9/22–9/26", "第2章 牛顿定律", "隔离法；斜面问题；连接体", 4, "实验一：单摆法测重力加速度"],
    [4, "9/29–10/3", "—", "国庆假期", 0, ""],
    [5, "10/6–10/10", "第2章 牛顿定律", "变力问题；非惯性系；习题课", 2, "第二章作业（第5周周二提交）"],
    [5, "10/6–10/10", "第3章 动量与能量", "冲量与动量定理；动量守恒", 2, ""],
    [6, "10/13–10/17", "第3章 动量与能量", "质心；功与动能定理", 4, ""],
    [7, "10/20–10/24", "第3章 动量与能量", "保守力与势能；机械能守恒；碰撞", 4, "作业3（第8周周二提交）"],
    [8, "10/27–10/31", "复习", "第1–3章综合习题课；单元测验", 4, "第1–3章单元测验（周四课上）"],
]
for i, r in enumerate(rows, 5):
    body(ws, i, r, wrap_cols=(4, 6))
last = 4 + len(rows)
ws.cell(row=last + 1, column=4, value="合计学时").font = Font(name=F, bold=True)
ws.cell(row=last + 1, column=5, value=f"=SUM(E5:E{last})").font = Font(name=F, bold=True)
ws.freeze_panes = "A5"
wb.save(ROOT / "00_课程说明" / "教学日历_2026秋.xlsx")

# ---------------------------------------------------------------- 评分标准
wb = Workbook(); ws = wb.active; ws.title = "成绩构成"
ws["A1"] = "大学物理A（上）成绩构成"; ws["A1"].font = Font(name=F, bold=True, size=13)
header(ws, 3, ["考核环节", "权重", "说明"], [18, 10, 50])
comp = [["平时作业", 0.15, "每章一次，按“作业评分量规”评分，取平均"],
        ["课堂表现与测验", 0.10, "课堂练习 + 单元测验"],
        ["实验", 0.15, "4 个实验报告的平均分"],
        ["期中考试", 0.20, "闭卷，第 1–3 章"],
        ["期末考试", 0.40, "闭卷，全部内容"]]
for i, r in enumerate(comp, 4):
    body(ws, i, r, wrap_cols=(3,))
    ws.cell(row=i, column=2).number_format = "0%"
ws.cell(row=9, column=1, value="合计").font = Font(name=F, bold=True)
ws.cell(row=9, column=2, value="=SUM(B4:B8)").number_format = "0%"
ws.cell(row=9, column=2).font = Font(name=F, bold=True)
ws.cell(row=10, column=1, value="权重来源：课程教学大纲第四部分。合计应为 100%。").font = Font(name=F, color="777777", italic=True)

ws2 = wb.create_sheet("作业评分量规")
ws2["A1"] = "作业评分量规（每道计算题按 100 分折算）"; ws2["A1"].font = Font(name=F, bold=True, size=13)
header(ws2, 3, ["评分维度", "满分", "优秀（满分）", "合格（约一半）", "不合格（0 分）"], [16, 8, 38, 34, 30])
rubric = [["物理模型与受力分析", 30, "研究对象、模型选择正确；受力图完整无多画漏画", "模型基本正确，受力图有个别遗漏", "模型错误或没有受力分析"],
          ["方程与推导", 30, "规律选用正确，方程完整，推导步骤清楚", "规律正确但推导跳步或有小错误", "规律用错或只有结果"],
          ["计算结果与单位", 20, "结果正确，单位和有效数字规范", "结果有计算错误或缺单位", "无结果"],
          ["规范与讨论", 20, "书写规范，对结果做合理性检验或讨论", "书写基本规范，无讨论", "字迹潦草、无法辨认"]]
for i, r in enumerate(rubric, 4):
    body(ws2, i, r, wrap_cols=(3, 4, 5))
    ws2.row_dimensions[i].height = 34
ws2.cell(row=8, column=1, value="合计").font = Font(name=F, bold=True)
ws2.cell(row=8, column=2, value="=SUM(B4:B7)").font = Font(name=F, bold=True)
ws2.cell(row=9, column=1, value="用于问渠 AI 批改：AI 按本量规给出评分建议，教师审核后发布。").font = Font(name=F, color="777777", italic=True)
wb.save(ROOT / "00_课程说明" / "课程评分标准.xlsx")

# ---------------------------------------------------------------- 单摆数据记录表
wb = Workbook(); ws = wb.active; ws.title = "数据记录"
ws["A1"] = "实验一  单摆法测定重力加速度  数据记录表"; ws["A1"].font = Font(name=F, bold=True, size=13)
ws["A2"] = "姓名：________   学号：________   日期：________"; ws["A2"].font = Font(name=F)
ws["A3"] = "黄色格子填测量值；其余格子自动计算。第 5 行是示例数据，做实验时请覆盖。"; ws["A3"].font = Font(name=F, color="777777", italic=True)
header(ws, 4, ["序号", "线长 l / cm", "小球直径 d / cm", "摆长 L / m", "30T 第1次 / s", "30T 第2次 / s", "30T 第3次 / s", "周期 T / s", "T² / s²", "g / (m/s²)"],
       [6, 12, 14, 11, 13, 13, 13, 11, 10, 12])
sample = [[1, 48.8, 2.4, 42.55, 42.45, 42.51], [2, 58.8, 2.4, 46.5, 46.58, 46.63], [3, 68.8, 2.4, 50.32, 50.22, 50.29], [4, 78.8, 2.4, 53.72, 53.82, 53.73], [5, 88.8, 2.4, 57.06, 57.03, 56.94], [6, 98.8, 2.4, 60.05, 60.15, 60.12]]
for i, (n, l, dd, t1, t2, t3) in enumerate(sample, 5):
    vals = [n, l, dd, f"=(B{i}+C{i}/2)/100", t1, t2, t3, f"=AVERAGE(E{i}:G{i})/30", f"=H{i}^2", f"=4*PI()^2*D{i}/I{i}"]
    body(ws, i, vals)
    for col in (2, 3, 5, 6, 7):
        ws.cell(row=i, column=col).fill = INPUT
    ws.cell(row=i, column=4).number_format = "0.0000"
    ws.cell(row=i, column=8).number_format = "0.0000"
    ws.cell(row=i, column=9).number_format = "0.0000"
    ws.cell(row=i, column=10).number_format = "0.000"
ws["I12"] = "平均 g"; ws["I12"].font = Font(name=F, bold=True)
ws["J12"] = "=AVERAGE(J5:J10)"; ws["J12"].number_format = "0.000"; ws["J12"].font = Font(name=F, bold=True)
ws["I13"] = "作图法 g"; ws["I13"].font = Font(name=F, bold=True)
ws["J13"] = "=4*PI()^2/SLOPE(I5:I10,D5:D10)"; ws["J13"].number_format = "0.000"; ws["J13"].font = Font(name=F, bold=True)
ws["J13"].comment = Comment("T²–L 直线的斜率 k = 4π²/g，所以 g = 4π²/k。", "教师")
ws["I14"] = "参考值"; ws["J14"] = 9.79; ws["J14"].number_format = "0.00"
ws["K14"] = "当地参考值，由实验室提供"; ws["K14"].font = Font(name=F, color="777777", italic=True)
ws["I15"] = "相对误差"; ws["J15"] = "=ABS(J13-J14)/J14"; ws["J15"].number_format = "0.0%"
wb.save(ROOT / "实验" / "单摆实验数据记录表.xlsx")
print("xlsx done")
