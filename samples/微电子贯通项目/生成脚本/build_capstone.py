"""贯通项目四步的实验指导书（R7）、实验报告模板（R8）与评分量规（第 17 轮附二 3.3），中英对照。
排版沿用数字集成电路设计第 1 章（samples/数字集成电路设计/生成脚本/docgen.py）；内容见 capstone_labs.py。
运行：python build_capstone.py
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1] / "数字集成电路设计" / "生成脚本"))
from docgen import Doc  # noqa: E402
from capstone_labs import H, HE, LABS, LINK, PROJECT, S4  # noqa: E402

BASE = HERE.parent
GUIDES = BASE / "指导书"
REPORTS = BASE / "实验报告模板"
GUIDES.mkdir(parents=True, exist_ok=True)
REPORTS.mkdir(parents=True, exist_ok=True)

# 第 17 轮附二 3.3：总分 100，流片就绪度 15（由检查结果自动给分）
RUBRIC = [
    ("原理 Principles", 15, "公式与原理正确，说明每个量的意义和模型的适用条件\nCorrect principles and formulas, meaning and conditions explained",
     "基本正确，缺少适用条件\nMostly correct, conditions missing", "错误或缺失\nWrong or missing"),
    ("数据 Data", 20, "表格完整，单位、有效数字规范；附检查结果（result.json）与至少一张图或版图截图\nComplete tables, units and figures; check results attached; at least one graph or layout screenshot",
     "有缺项或单位不规范\nGaps or unit problems", "无数据\nNo data"),
    ("分析 Analysis", 20, "参考值与自己的结果逐项比较，解释差异；与标准单元、经验公式等对照\nCompare with the reference item by item and explain differences; compare with standard cells or rules of thumb",
     "有比较但缺解释\nComparison without explanation", "无分析\nNo analysis"),
    ("实际问题 Real problem", 20, "五步完整：问题、模型（写明假设）、求解、检验、模型局限与改进\nAll five steps: problem, model with assumptions, solution, check, limits and improvements",
     "有求解但缺假设或局限\nSolved but assumptions or limits missing", "未完成\nNot attempted"),
    ("规范 Presentation", 10, "书写清楚，图表规范，代码有注释，AI 使用如实声明\nClear writing, tidy graphs, commented code, honest AI statement",
     "基本清楚\nMostly clear", "难以辨认或未声明 AI 使用\nIllegible or no AI statement"),
    ("流片就绪度 Tapeout readiness", 15, "按检查结果自动给分（见“流片就绪度”页）\nScored automatically from the check results (see the Readiness sheet)",
     "部分检查通过\nSome checks pass", "检查未运行或全部不通过\nNot run or all fail"),
]
# 各步“流片就绪度”15 分的自动给分项：(检查项, 分值, result.json 里的条件)
READINESS = {
    "一": [("检查命令运行成功，参数表完整 Check ran, table complete", 5, "params.json 有 VT、IDSAT、IOFF、S"),
           ("VT 与参考值相差 ≤ 0.02 V VT within 0.02 V of reference", 5, "|vt − 参考| ≤ 0.02"),
           ("IDSAT、S 与参考值相差 ≤ 3% IDSAT and S within 3%", 5, "相对误差 ≤ 3%")],
    "二": [("DRC 错误为 0 DRC errors 0", 5, "drc_errors = 0"),
           ("LVS 一致（含尺寸）LVS match incl. sizes", 7, "lvs_match = true"),
           ("与标准单元比较完成 Comparison with the standard cell", 3, "area_ratio_vs_stdcell 有值")],
    "三": [("自己的测试全部通过 Own tests all pass", 5, "wqlab test：all_passed"),
           ("参考测试全部通过 Reference tests all pass", 7, "wqlab test --ref：6/6"),
           ("闭环检验差 ≤ 2 Closed loop within 2 counts", 3, "test_motor 通过")],
    "四": [("时序：setup、hold 余量不为负 Timing: setup and hold slack ≥ 0", 4, "setup_ws_ns ≥ 0 且 hold_ws_ns ≥ 0"),
           ("DRC、LVS 为 0 DRC and LVS 0", 5, "magic_drc_errors = 0 且 lvs_errors = 0"),
           ("天线违例为 0 Antenna 0", 2, "antenna_violations = 0"),
           ("Tiny Tapeout 格式（名称、接口、1 块）TT format", 4, "tt_ready = true")],
}


def env_table(d, lab):
    d.table(["项目 Item", "内容 Details"], [[f"{a} {b}", f"{c}\n{e}" if e else c] for a, b, c, e in lab["env"]], widths=[4.0, 12.0], font_size=9.5)


def guide(lab):
    d = Doc()
    d.title(f"贯通项目 步骤{lab['no']}  {lab['zh']}", f"Capstone Step {lab['noe']}  {lab['en']}")
    d.en(f"{H}    {HE}", size=9)
    d.bp(*PROJECT)
    d.bh("一、实验目的", "1  Objectives")
    d.blist(lab["goal"], numbered=True)
    d.bh("二、实验原理", "2  Principles")
    for zh, en in lab["theory"]:
        d.bp(zh, en)
    d.bh("三、实验环境与起始文件", "3  Environment and starting files")
    d.bp("本步在问渠微电子实验镜像中完成（网页远程桌面或命令行）。", "This step runs in the WenQuest microelectronics lab image (web desktop or command line).")
    d.en("入口 Access: " + LINK)
    env_table(d, lab)
    d.bh("四、实验步骤", "4  Procedure")
    d.blist(lab["steps"], numbered=True)
    d.bh("五、数据记录", "5  Data record")
    d.bp("以下表格同时出现在实验报告模板中，请在报告里填写，并附上检查命令生成的 result.json。", "The same tables are in the report template; fill them in there and attach the result.json from the check.")
    for cap, head, rows, widths in lab["tables"]:
        d.p(cap, indent=False)
        d.table(head, rows, widths=widths, font_size=9.5)
    d.bh("六、注意事项", "6  Notes")
    d.blist(lab["cautions"])
    d.bh("七、思考题", "7  Questions")
    d.blist(lab["questions"], numbered=True)
    t, te, q, qe = lab["problem"]
    d.bh("八、与实际问题的联系", "8  Connection to a real problem")
    d.p(f"机器人问题：{t}", indent=False).runs[0].bold = True
    d.en(f"Robot problem: {te}")
    d.bp(q, qe)
    d.bp("请在实验报告“实际问题的建模与求解”一栏中按五步完成：① 实际问题 ② 建立模型（写出简化假设）③ 求解 ④ 用实验或仿真检验 ⑤ 指出模型在哪里失效、怎样修正。",
         "In the report section “Modelling a real problem”, work through five steps: ① the problem ② the model (state your assumptions) ③ the solution ④ a check by experiment or simulation ⑤ where the model fails and how to improve it.")
    d.p("生活中的例子：" + lab["everyday"][0], indent=False)
    d.en("Everyday example: " + lab["everyday"][1])
    d.bh("九、评分", "9  Grading")
    d.bp("按贯通项目评分量规：原理 15、数据 20、分析 20、实际问题 20、规范 10、流片就绪度 15。本步的流片就绪度按下表由检查结果自动给分：",
         "Graded with the capstone rubric: principles 15, data 20, analysis 20, real problem 20, presentation 10, tapeout readiness 15. This step's readiness points come from the check results:")
    d.table(["检查项 Check", "分值 Points", "条件 Condition"], [[a, str(b), c] for a, b, c in READINESS[lab["no"]]], widths=[8.0, 2.0, 6.0], font_size=9.5)
    d.save(GUIDES / f"{lab['file']} 实验指导书.docx")


SECTIONS = [
    ("一、实验目的", "1  Objectives", 3), ("二、实验原理（写出主要公式）", "2  Principles (key formulas)", 5),
    ("三、实验步骤与设计说明（简要）", "3  Procedure and design notes (brief)", 5), ("四、数据记录与处理", "4  Data and processing", 0),
    ("五、结果与分析", "5  Results and analysis", 6), ("六、实际问题的建模与求解", "6  Modelling a real problem", 0),
    ("七、结论", "7  Conclusion", 3), ("八、思考题", "8  Questions", 6),
]


def report(lab):
    d = Doc()
    title = "设计报告" if lab["no"] == "四" else "实验报告"
    d.title(f"{title}  贯通项目 步骤{lab['no']}  {lab['zh']}", f"{'Design' if lab['no'] == '四' else 'Lab'} Report  Capstone Step {lab['noe']}  {lab['en']}")
    d.table(["姓名 Name", "", "学号 ID", "", "日期 Date", ""], [["班级 Class", "", "同组 Partner", "", "成绩 Score", ""]],
            widths=[2.4, 3.2, 2.2, 2.8, 2.2, 3.2], font_size=10)
    d.en("提交：本报告、检查命令生成的 result.json，以及本步要求的文件（代码、网表、版图或 GDS）一起在问渠“作业”中上传。  Submit this report, the check's result.json and the step's files (code, netlist, layout or GDS) in WenQuest Assignments.", size=9)
    for zh, en, n in SECTIONS:
        d.bh(zh, en)
        if zh.startswith("四"):
            for cap, head, rows, widths in lab["tables"]:
                d.p(cap, indent=False)
                d.table(head, rows, widths=widths, font_size=9.5)
            d.bp("附检查结果：把 result.json 的主要项目抄在这里或附在报告后；至少一张图或截图。", "Attach the check results: copy the main items of result.json here or append it; include at least one graph or screenshot.")
            d.lines(4)
        elif zh.startswith("六"):
            t, te, q, qe = lab["problem"]
            d.p(f"问题：{t}", indent=False).runs[0].bold = True
            d.bp(q, qe)
            for step, en in [("① 实际问题（用自己的话描述，给出已知数据）", "① The problem (in your own words, with the given data)"),
                             ("② 建立模型（写出简化假设）", "② The model (state your assumptions)"),
                             ("③ 求解（公式与数值）", "③ Solution (formulas and numbers)"),
                             ("④ 用实验或仿真检验（写出设置与结果）", "④ Check by experiment or simulation (settings and results)"),
                             ("⑤ 模型在哪里失效？怎样修正？", "⑤ Where does the model fail? How would you improve it?")]:
                d.p(step, indent=False)
                d.en(en)
                d.lines(3)
        elif zh.startswith("八"):
            for zhq, enq in lab["questions"]:
                d.p(zhq, indent=False); d.en(enq); d.lines(3)
        else:
            d.lines(n)
    d.bh("九、AI 使用声明", "9  AI use statement")
    d.p("□ 未使用 AI　　□ 用 AI 解释概念　　□ 用 AI 检查代码或计算　　□ 其他：____________", indent=False)
    d.en("□ No AI used   □ AI to explain concepts   □ AI to check code or calculations   □ Other: ____________")
    d.p("如使用了 AI，请写明用在哪里：", indent=False); d.en("If you used AI, say where and how:"); d.lines(2)
    d.bh("评分量规", "Rubric")
    d.table(["项目 Criterion", "分值 Points", "要求 What earns full marks"], [[a, str(b), c.replace("\n", " / ")] for a, b, c, _, _ in RUBRIC],
            widths=[4.2, 2.0, 9.8], font_size=9.5)
    d.en("AI 按本量规给出评分建议，教师复核后发布；流片就绪度由检查结果自动给分。  AI suggests a score against this rubric and the teacher reviews it; tapeout readiness is scored from the check results.", size=9)
    d.save(REPORTS / f"贯通项目步骤{lab['no']} {title}模板.docx")


def rubric():
    F = "Arial"; HEAD = PatternFill("solid", fgColor="E8EEF0"); INPUT = PatternFill("solid", fgColor="FFF7D6")
    thin = Side(style="thin", color="B8C4C8"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook(); ws = wb.active; ws.title = "评分量规 Rubric"
    ws["A1"] = "微电子贯通项目评分量规  Microelectronics Capstone Rubric"; ws["A1"].font = Font(name=F, bold=True, size=13)
    ws["A2"] = "四步各按本量规评分；AI 给出建议，教师复核后发布。 Each step is scored with this rubric; AI suggests, the teacher reviews."; ws["A2"].font = Font(name=F, color="555555", italic=True)
    hdr = ["评分项目 Criterion", "满分 Points", "优秀（满分） Excellent", "合格（约一半） Adequate", "不合格（0 分） Missing"]
    for i, v in enumerate(hdr, 1):
        c = ws.cell(row=4, column=i, value=v); c.font = Font(name=F, bold=True); c.fill = HEAD; c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(RUBRIC, 5):
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v); c.font = Font(name=F); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 52
    last = 4 + len(RUBRIC)
    ws.cell(row=last + 1, column=1, value="合计 Total").font = Font(name=F, bold=True)
    ws.cell(row=last + 1, column=2, value=f"=SUM(B5:B{last})").font = Font(name=F, bold=True)
    for col, w in zip("ABCDE", [24, 11, 52, 30, 28]):
        ws.column_dimensions[col].width = w

    # 流片就绪度：每步的检查项与分值
    w3 = wb.create_sheet("流片就绪度 Readiness")
    w3["A1"] = "流片就绪度 15 分的自动给分项（由 wqlab 检查结果判定） Readiness points from the wqlab check results"; w3["A1"].font = Font(name=F, bold=True, size=12)
    for i, v in enumerate(["步骤 Step", "检查项 Check", "分值 Points", "条件 Condition"], 1):
        c = w3.cell(row=3, column=i, value=v); c.font = Font(name=F, bold=True); c.fill = HEAD; c.border = BOX
    r = 4
    for step, items in READINESS.items():
        for a, b, cnd in items:
            for i, v in enumerate([f"步骤{step}", a, b, cnd], 1):
                c = w3.cell(row=r, column=i, value=v); c.font = Font(name=F); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="center")
            r += 1
    w3.cell(row=r + 1, column=1, value=f"参考设计（步骤四）：setup 余量 {S4['setup_ws_ns']:.1f} ns，hold 余量 {S4['hold_ws_ns']:.2f} ns，DRC/LVS/天线 0 → 15 分。").font = Font(name=F, color="555555")
    for col, w in zip("ABCD", [10, 46, 10, 40]):
        w3.column_dimensions[col].width = w

    ws2 = wb.create_sheet("成绩登记 Scores")
    ws2["A1"] = "贯通项目成绩登记（示例数据，请替换） Capstone scores (sample rows — replace)"; ws2["A1"].font = Font(name=F, bold=True, size=12)
    h2 = ["学号 ID", "姓名 Name", "步骤 Step", "原理 /15", "数据 /20", "分析 /20", "实际问题 /20", "规范 /10", "流片就绪度 /15", "总分 Total", "等级 Grade"]
    for i, v in enumerate(h2, 1):
        c = ws2.cell(row=3, column=i, value=v); c.font = Font(name=F, bold=True); c.fill = HEAD; c.border = BOX; c.alignment = Alignment(wrap_text=True, horizontal="center")
    sample = [["2026001", "示例学生 A", "一", 14, 18, 17, 16, 9, 15], ["2026002", "示例学生 B", "二", 12, 16, 13, 14, 8, 12], ["2026003", "示例学生 C", "四", 15, 19, 18, 18, 10, 15]]
    dv = DataValidation(type="list", formula1='"一,二,三,四"', allow_blank=True); ws2.add_data_validation(dv)
    caps = [15, 20, 20, 20, 10, 15]
    for r in range(4, 64):
        vals = sample[r - 4] if r - 4 < len(sample) else [None] * 9
        for i, v in enumerate(vals, 1):
            c = ws2.cell(row=r, column=i, value=v); c.font = Font(name=F, color="0000FF" if i >= 4 else "000000"); c.border = BOX
            if i >= 4:
                c.fill = INPUT
        dv.add(ws2.cell(row=r, column=3))
        for k, cap in enumerate(caps):
            col = 4 + k
            dvn = DataValidation(type="decimal", operator="between", formula1="0", formula2=str(cap), allow_blank=True)
            ws2.add_data_validation(dvn); dvn.add(ws2.cell(row=r, column=col))
        ws2.cell(row=r, column=10, value=f'=IF(COUNT(D{r}:I{r})=0,"",SUM(D{r}:I{r}))').border = BOX
        ws2.cell(row=r, column=11, value=f'=IF(J{r}="","",IF(J{r}>=90,"优秀 A",IF(J{r}>=80,"良好 B",IF(J{r}>=70,"中等 C",IF(J{r}>=60,"及格 D","不及格 F")))))').border = BOX
        for i in (10, 11):
            ws2.cell(row=r, column=i).font = Font(name=F)
    ws2["M3"] = "班级平均 Class average"; ws2["M3"].font = Font(name=F, bold=True)
    ws2["N3"] = '=IFERROR(AVERAGE(J4:J63),"")'; ws2["N3"].number_format = "0.0"
    ws2["M4"] = "流片就绪度平均 /15"; ws2["N4"] = '=IFERROR(AVERAGE(I4:I63),"")'; ws2["N4"].number_format = "0.0"
    for col, w in zip("ABCDEFGHIJKLMN", [11, 14, 8, 9, 9, 9, 11, 9, 12, 10, 12, 2, 22, 10]):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A4"
    out = BASE / "贯通项目评分量规.xlsx"
    wb.save(out)
    return out


if __name__ == "__main__":
    for lab in LABS:
        guide(lab)
        report(lab)
    print("guides and reports written;", "rubric", rubric())
