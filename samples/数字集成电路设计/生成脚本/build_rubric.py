"""第 1 章实验报告评分量规（R8），沿用大学物理第 1 章的量规（原理 20 · 数据 25 · 分析 25 · 实际问题 20 · 规范 10）。"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from pathlib import Path
F = "Arial"; HEAD = PatternFill("solid", fgColor="E8EEF0"); INPUT = PatternFill("solid", fgColor="FFF7D6")
thin = Side(style="thin", color="B8C4C8"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
out = Path(__file__).resolve().parents[1] / "课程资料" / "第1章 CMOS反相器" / "虚拟实验" / "实验报告评分量规.xlsx"
wb = Workbook(); ws = wb.active; ws.title = "评分量规 Rubric"
ws["A1"] = "虚拟实验报告评分量规  Virtual Lab Report Rubric"; ws["A1"].font = Font(name=F, bold=True, size=13)
ws["A2"] = "AI 按本量规给出评分建议，教师复核后发布。 AI suggests scores against this rubric; the teacher reviews before release."; ws["A2"].font = Font(name=F, color="555555", italic=True)
hdr = ["评分项目 Criterion", "满分 Points", "优秀（满分） Excellent", "合格（约一半） Adequate", "不合格（0 分） Missing"]
rows = [
 ["原理 Principles", 20, "公式正确，说明每个量的意义和模型的适用条件\nCorrect formulas; meaning and model conditions explained", "公式基本正确，缺少适用条件\nMostly correct; conditions missing", "公式错误或缺失\nWrong or missing"],
 ["数据 Data", 25, "数据表填写完整，单位、有效数字规范，有至少一张图\nComplete tables, correct units and significant figures, at least one graph", "数据有缺项或单位不规范\nGaps or unit problems", "无数据\nNo data"],
 ["分析 Analysis", 25, "逐项比较测量值与理论值，给出相对误差并分析来源\nMeasured vs theory item by item, relative errors and their sources", "有比较但缺误差分析\nComparison without error analysis", "无分析\nNo analysis"],
 ["实际问题 Real problem", 20, "五步完整：问题、模型（写明假设）、求解、实验检验、模型局限与改进\nAll five steps: problem, model with assumptions, solution, lab check, limits and improvements", "有求解但缺假设或局限\nSolved but assumptions or limits missing", "未完成\nNot attempted"],
 ["规范 Presentation", 10, "书写清楚，图表规范，AI 使用如实声明\nClear writing and graphs, honest AI statement", "基本清楚\nMostly clear", "难以辨认或未声明 AI 使用\nIllegible or no AI statement"],
]
for i, v in enumerate(hdr, 1):
    c = ws.cell(row=4, column=i, value=v); c.font = Font(name=F, bold=True); c.fill = HEAD; c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="center")
for r, row in enumerate(rows, 5):
    for i, v in enumerate(row, 1):
        c = ws.cell(row=r, column=i, value=v); c.font = Font(name=F); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 48
ws.cell(row=10, column=1, value="合计 Total").font = Font(name=F, bold=True)
ws.cell(row=10, column=2, value="=SUM(B5:B9)").font = Font(name=F, bold=True)
for col, w in zip("ABCDE", [20, 11, 52, 34, 28]): ws.column_dimensions[col].width = w

ws2 = wb.create_sheet("成绩登记 Scores")
ws2["A1"] = "第 1 章实验报告成绩登记（示例数据，请替换） Chapter 1 lab report scores (sample rows—replace)"; ws2["A1"].font = Font(name=F, bold=True, size=12)
h2 = ["学号 ID", "姓名 Name", "实验 Lab", "原理 /20", "数据 /25", "分析 /25", "实际问题 /20", "规范 /10", "总分 Total", "等级 Grade"]
for i, v in enumerate(h2, 1):
    c = ws2.cell(row=3, column=i, value=v); c.font = Font(name=F, bold=True); c.fill = HEAD; c.border = BOX; c.alignment = Alignment(wrap_text=True, horizontal="center")
sample = [["2026001", "示例学生 A", "1.3", 18, 22, 20, 17, 9], ["2026002", "示例学生 B", "1.4", 15, 20, 14, 12, 8], ["2026003", "示例学生 C", "1.5", 20, 24, 23, 19, 10]]
dv = DataValidation(type="list", formula1='"1.1,1.2,1.3,1.4,1.5"', allow_blank=True); ws2.add_data_validation(dv)
for r in range(4, 44):
    vals = sample[r - 4] if r - 4 < len(sample) else [None] * 8
    for i, v in enumerate(vals, 1):
        c = ws2.cell(row=r, column=i, value=v); c.font = Font(name=F, color="0000FF" if i >= 4 else "000000"); c.border = BOX
        if i >= 4: c.fill = INPUT
    dv.add(ws2.cell(row=r, column=3))
    ws2.cell(row=r, column=9, value=f'=IF(COUNT(D{r}:H{r})=0,"",SUM(D{r}:H{r}))').border = BOX
    ws2.cell(row=r, column=10, value=f'=IF(I{r}="","",IF(I{r}>=90,"优秀 A",IF(I{r}>=80,"良好 B",IF(I{r}>=70,"中等 C",IF(I{r}>=60,"及格 D","不及格 F")))))').border = BOX
    for i in (9, 10): ws2.cell(row=r, column=i).font = Font(name=F)
ws2["L3"] = "班级平均 Class average"; ws2["L3"].font = Font(name=F, bold=True)
ws2["M3"] = "=IFERROR(AVERAGE(I4:I43),\"\")"; ws2["M3"].number_format = "0.0"
ws2["L4"] = "实际问题平均 /20"; ws2["M4"] = "=IFERROR(AVERAGE(G4:G43),\"\")"; ws2["M4"].number_format = "0.0"
for col, w in zip("ABCDEFGHIJKLM", [11, 14, 8, 9, 9, 9, 11, 9, 10, 12, 2, 22, 10]): ws2.column_dimensions[col].width = w
ws2.freeze_panes = "A4"
wb.save(out); print("rubric saved")
